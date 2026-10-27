from datetime import date
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import Company, Role, User
from apps.estimates.models import EstimateCatalog
from apps.planning.models import CalendarDay, LoadProfile, LoadProfileItem, ProductionCalendar
from apps.projects.models import ConstructionObject, Project, Section
from apps.resources.models import Brigade, BrigadeGroup, BrigadeMacroGroup, Employee, Equipment, EquipmentCategory, EquipmentType
from apps.resources.position_catalog import position_for
from apps.works.models import ComplexObject, MeasurementUnit, ProjectWork, ProjectWorkItem, WorkGroup


class CatalogueExcelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.company = Company.objects.create(name='Catalogue Excel')
        cls.foreign = Company.objects.create(name='Foreign catalogue Excel')
        cls.user = User.objects.create_user(username='catalogue-excel-admin', company=cls.company, role=Role.objects.get(code='ADMIN'))
        cls.project = Project.objects.create(company=cls.company, name='Проект', start_date=date(2026, 1, 1))
        cls.obj = ConstructionObject.objects.create(company=cls.company, project=cls.project, name='КТП')
        cls.section = Section.objects.create(company=cls.company, construction_object=cls.obj, name='УЗА')
        cls.complex = ComplexObject.objects.create(company=cls.company, construction_object=cls.obj, section=cls.section, name='УЗА 33')
        cls.work_group = WorkGroup.objects.create(company=cls.company, name='СМР')
        cls.profile = LoadProfile.objects.create(company=cls.company, name='Равномерный')
        LoadProfileItem.objects.create(company=cls.company, profile=cls.profile, workday_number=1, percentage=Decimal('37.5'))
        cls.work = ProjectWork.objects.create(company=cls.company, section=cls.section, complex_object=cls.complex,
            name='Монтаж', kind='COMPOSITE', unit='компл', work_group=cls.work_group, unit_price=Decimal('123.45'))
        cls.item = ProjectWorkItem.objects.create(company=cls.company, project_work=cls.work, name='Кабель',
            unit='м', quantity_per_unit=Decimal('12.5'), weight=Decimal('100'), load_profile=cls.profile)
        cls.unit, _ = MeasurementUnit.objects.get_or_create(company=cls.company, symbol='м', defaults={'name': 'Метр'})
        cls.group = BrigadeGroup.objects.create(company=cls.company, name='ИТР')
        cls.macro = BrigadeMacroGroup.objects.create(company=cls.company, name='Персонал')
        cls.brigade = Brigade.objects.create(company=cls.company, name='Инженер', group=cls.group)
        cls.position = position_for(cls.brigade)
        cls.employee = Employee.objects.create(company=cls.company, position=cls.position, first_name='Иван', last_name='Иванов', employee_number='000123')
        cls.category = EquipmentCategory.objects.create(company=cls.company, name='Специализированная техника')
        cls.equipment_type = EquipmentType.objects.create(company=cls.company, name='Экскаватор', category=cls.category)
        cls.machine = Equipment.objects.create(company=cls.company, type=cls.equipment_type, plate_number='00123', name='Экскаватор', manufacture_year=2020)
        cls.calendar = ProductionCalendar.objects.create(company=cls.company, name='2026', year=2026)
        CalendarDay.objects.create(company=cls.company, calendar=cls.calendar, date=date(2026, 1, 1), is_working=False, is_holiday=True)
        EstimateCatalog.objects.create(company=cls.company, kind='work', identity='work-1', code='01', name='Монтаж', unit='м')
        cls.routes = [
            ('resources:employee_list', {}), ('resources:equipment_category_list', {}),
            ('resources:brigade_group_list', {}), ('resources:brigade_macro_group_list', {}),
            ('resources:equipment_unit_list', {}), ('works:work_list', {}), ('works:group_list', {}),
            ('works:unit_list', {}), ('works:complex_object_list', {}), ('projects:project_list', {}),
            ('projects:object_list', {'project_pk': cls.project.pk}), ('projects:section_list', {'pk': cls.obj.pk}),
            ('planning:profile_list', {}), ('planning:calendar_list', {}), ('estimates:catalog_list', {}),
        ]

    def setUp(self):
        self.client.force_login(self.user)

    def workbook(self, route, params=None, **kwargs):
        response = self.client.get(reverse(route, kwargs=kwargs), {'export': 'xlsx', **(params or {})})
        self.assertEqual(response.status_code, 200, response.content[:300])
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('.xlsx', response['Content-Disposition'])
        return load_workbook(BytesIO(response.content))

    def test_every_catalogue_has_button_and_a_readable_export(self):
        for route, kwargs in self.routes:
            with self.subTest(route=route):
                self.assertContains(self.client.get(reverse(route, kwargs=kwargs)), 'Экспорт в Excel')
                book = self.workbook(route, **kwargs)
                self.assertGreater(book.active.max_row, 1)
                self.assertEqual(book.active.freeze_panes, 'A2')
                self.assertTrue(book.active.auto_filter.ref)
                book.close()

    def test_work_hierarchy_and_numeric_component_values(self):
        book = self.workbook('works:work_list', {'construction_object': self.obj.pk})
        main = list(book.active.values)
        values = dict(zip(main[0], main[1]))
        self.assertEqual([values[key] for key in ['Проект', 'Объект', 'Раздел', 'Комплексный объект', 'Название']], ['Проект', 'КТП', 'УЗА', 'УЗА 33', 'Монтаж'])
        self.assertEqual(values['Цена за единицу'], 123.45)
        rows = list(book['Работы в комплексах'].values)
        values = dict(zip(rows[0], rows[1]))
        self.assertEqual(values['ID комплекса'], self.work.pk)
        self.assertEqual(values['Работа в комплексе'], 'Кабель')
        self.assertEqual(values['Норматив на единицу комплекса'], 12.5)
        book.close()

    def test_profiles_and_calendars_include_their_day_values(self):
        book = self.workbook('planning:profile_list')
        self.assertEqual(list(book['Дни профилей'].values)[1][-2:], (1, 37.5))
        book.close()
        book = self.workbook('planning:calendar_list')
        row = list(book['Дни календарей'].values)[1]
        self.assertEqual(row[2].date(), date(2026, 1, 1))
        self.assertEqual(row[3:6], ('Нет', 'Да', 'Нет'))
        book.close()

    def test_employee_filters_export_all_pages_preserve_identifiers_and_exclude_private_data(self):
        Employee.objects.bulk_create([Employee(company=self.company, first_name='Иван', last_name=f'Иванов {i}', position=self.position) for i in range(25)])
        Employee.objects.create(company=self.company, first_name='Иван', last_name='Уволен', position=self.position, is_active=False)
        Employee.objects.create(company=self.foreign, first_name='Иван', last_name='Чужой')
        book = self.workbook('resources:employee_list', {'q': 'иван', 'page': 2, 'catalogue_position': self.position.pk})
        rows = list(book.active.values)
        self.assertEqual(len(rows), 27)
        self.assertIn('000123', [row[0] for row in rows[1:]])
        self.assertEqual(book.active.cell(2, 1).data_type, 's')
        self.assertNotIn('zup_card', rows[0])
        self.assertNotIn('Уволен', str(rows)); self.assertNotIn('Чужой', str(rows))
        book.close()
        book = self.workbook('resources:employee_list', {'include_archived': 'on'})
        self.assertEqual(book.active.max_row, 28)
        book.close()
        response = self.client.get(reverse('resources:employee_list'), {'q': 'Иван', 'page': 2})
        self.assertContains(response, 'q=%D0%98%D0%B2%D0%B0%D0%BD&amp;export=xlsx')

    def test_invalid_filters_are_rejected_and_empty_results_still_have_headers(self):
        for route, params in [('resources:employee_list', {'catalogue_position': 'bad'}), ('works:work_list', {'section': 'bad'})]:
            self.assertEqual(self.client.get(reverse(route), {'export': 'xlsx', **params}).status_code, 400)
        book = self.workbook('works:work_list', {'q': 'Нет такой работы'})
        self.assertEqual(book.active.max_row, 1)
        self.assertEqual(book['Работы в комплексах'].max_row, 1)
        book.close()

    def test_scoped_objects_and_their_work_components_only(self):
        hidden = ConstructionObject.objects.create(company=self.company, project=self.project, name='Скрытый объект')
        section = Section.objects.create(company=self.company, construction_object=hidden, name='Скрытый раздел')
        work = ProjectWork.objects.create(company=self.company, section=section, name='Скрытый комплекс', kind='COMPOSITE')
        ProjectWorkItem.objects.create(company=self.company, project_work=work, name='Скрытая работа', unit='шт', load_profile=self.profile)
        planner = User.objects.create_user(username='catalogue-excel-planner', company=self.company, role=Role.objects.get(code='PLANNER'))
        planner.assigned_objects.add(self.obj)
        self.client.force_login(planner)
        book = self.workbook('works:work_list')
        self.assertNotIn('Скрытый', str(list(book.active.values)))
        self.assertEqual(book['Работы в комплексах'].max_row, 2)
        book.close()
        book = self.workbook('projects:object_list', project_pk=self.project.pk)
        self.assertEqual(book.active.max_row, 2)
        book.close()
        self.assertEqual(self.client.get(reverse('projects:section_list', args=[hidden.pk]), {'export': 'xlsx'}).status_code, 404)

    def test_text_is_not_a_formula_and_other_company_catalogues_are_excluded(self):
        self.category.name = '=1+1\x01'
        # Update deliberately bypasses form validation to cover legacy imported values.
        EquipmentCategory.objects.filter(pk=self.category.pk).update(name=self.category.name)
        EquipmentCategory.objects.create(company=self.foreign, name='Чужая категория')
        book = self.workbook('resources:equipment_category_list')
        self.assertEqual(book.active.max_row, 2)
        self.assertEqual(book.active['A2'].value, '=1+1')
        self.assertEqual(book.active['A2'].data_type, 's')
        book.close()

    def test_read_access_required_and_estimate_kind_and_search_preserved(self):
        EstimateCatalog.objects.create(company=self.company, kind='resource', identity='resource-1', name='Ресурс')
        EstimateCatalog.objects.create(company=self.foreign, kind='work', identity='foreign-1', name='Монтаж чужой')
        book = self.workbook('estimates:catalog_list', {'kind': 'resource', 'q': 'Ресурс'})
        self.assertEqual(book.active.max_row, 2)
        self.assertIn('Ресурс', list(book.active.values)[1])
        book.close()
        self.user.role = Role.objects.get(code='FOREMAN'); self.user.save(update_fields=['role'])
        for route, kwargs in self.routes:
            self.assertEqual(self.client.get(reverse(route, kwargs=kwargs), {'export': 'xlsx'}).status_code, 403, route)
        self.client.logout()
        self.assertEqual(self.client.get(reverse('resources:employee_list'), {'export': 'xlsx'}).status_code, 302)
