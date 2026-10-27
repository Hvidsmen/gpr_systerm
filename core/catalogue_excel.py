"""Excel exports of catalogue querysets, with the list's filters and access scope."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .permissions import READ_ROLES, require_roles


def xlsx_response(filename, sheets):
    workbook = Workbook(write_only=True)
    for title, headers, rows in sheets:
        sheet = workbook.create_sheet(title)
        sheet.freeze_panes = 'A2'
        from openpyxl.cell import WriteOnlyCell
        heading = []
        for index, label in enumerate(headers, 1):
            cell = WriteOnlyCell(sheet, label)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='087F9B')
            heading.append(cell)
            sheet.column_dimensions[get_column_letter(index)].width = 28 if index > 1 else 18
        sheet.append(heading)
        count = 1
        for values in rows:
            cells = []
            for value in values:
                if isinstance(value, bool):
                    value = 'Да' if value else 'Нет'
                if value is not None and not isinstance(value, (int, float, Decimal, date, datetime)):
                    value = ILLEGAL_CHARACTERS_RE.sub('', str(value))
                cell = WriteOnlyCell(sheet, value)
                if isinstance(value, str):
                    # Catalogue names and identifiers are always text, even '=...'.
                    cell.data_type = 's'
                elif isinstance(value, date):
                    cell.number_format = 'DD.MM.YYYY'
                cells.append(cell)
            sheet.append(cells)
            count += 1
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{count}'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response


def value_at(row, path):
    value = row
    for part in path.split('.'):
        if value is None:
            return None
        value = getattr(value, part)
        if callable(value):
            value = value()
    return value


def sheet_rows(queryset, columns):
    for row in queryset.iterator(chunk_size=500):
        yield [value_at(row, path) for label, path in columns]


# Explicit public columns: no raw integration payloads or personal employee cards.
SCHEMAS = {
    'resources.Employee': ('employees', 'Сотрудники', [
        ('Табельный номер', 'employee_number'), ('ФИО', 'display_name'),
        ('Должность', 'display_position'), ('Организация', 'zup_organisation'),
        ('Подразделение', 'display_department'), ('Активен', 'is_active')], ['position']),
    'resources.EquipmentCategory': ('equipment-categories', 'Категории техники', [('Название', 'name')], []),
    'resources.BrigadeGroup': ('position-groups', 'Группы должностей', [('Название', 'name')], []),
    'resources.BrigadeMacroGroup': ('position-macro-groups', 'Макрогруппы должностей', [('Название', 'name')], []),
    'resources.Equipment': ('equipment-units', 'Конкретная техника', [
        ('Госномер', 'plate_number'), ('Название', 'name'), ('Вид или модель', 'type.name'),
        ('Родительский вид', 'type.parent.name'), ('Марка / модель', 'brand_model'),
        ('Год выпуска', 'manufacture_year'), ('Инвентарный номер', 'inventory_number'),
        ('VIN / заводской номер', 'vin'), ('Участок', 'site_name'),
        ('Должность машиниста', 'operator_position.name'), ('Активна', 'is_active')],
        ['type__parent', 'operator_position']),
    'projects.Project': ('projects', 'Проекты', [
        ('ID', 'pk'), ('Проект', 'name'), ('Описание', 'description'), ('Статус', 'get_status_display'),
        ('Начало', 'start_date'), ('Окончание', 'end_date')], []),
    'projects.ConstructionObject': ('objects', 'Строительные объекты', [
        ('ID', 'pk'), ('Проект', 'project.name'), ('Родительский объект', 'parent.name'),
        ('Объект', 'name')], ['project', 'parent']),
    'projects.Section': ('sections', 'Разделы', [
        ('ID', 'pk'), ('Проект', 'construction_object.project.name'),
        ('Объект', 'construction_object.name'), ('Раздел', 'name'), ('Описание', 'description')],
        ['construction_object__project']),
    'works.ComplexObject': ('complex-objects', 'Комплексные объекты', [
        ('ID', 'pk'), ('Проект', 'construction_object.project.name'),
        ('Объект', 'construction_object.name'), ('Раздел', 'section.name'),
        ('Комплексный объект', 'name'), ('Количество, шт.', 'quantity'), ('Описание', 'description')], ['construction_object__project', 'section']),
    'works.WorkGroup': ('work-groups', 'Группы работ', [('Название', 'name')], []),
    'works.MeasurementUnit': ('measurement-units', 'Единицы измерения', [
        ('Обозначение', 'symbol'), ('Название', 'name')], []),
    'works.ProjectWork': ('works', 'Работы и комплексы', [
        ('ID', 'pk'), ('Проект', 'section.construction_object.project.name'),
        ('Объект', 'section.construction_object.name'), ('Раздел', 'section.name'),
        ('Комплексный объект', 'complex_object.name'), ('Вид работы', 'get_kind_display'),
        ('Название', 'name'), ('Единица измерения', 'unit'), ('Цена за единицу', 'unit_price'),
        ('Группа работ', 'work_group.name'), ('Профиль нагрузки', 'load_profile.name'),
        ('Статус', 'get_status_display'), ('Раздел РД', 'rd_section'), ('Пункт в РДЦ', 'rdc_point')],
        ['section__construction_object__project', 'complex_object', 'work_group', 'load_profile']),
    'planning.LoadProfile': ('load-profiles', 'Профили нагрузки', [
        ('ID', 'pk'), ('Название', 'name'), ('Описание', 'description')], []),
    'planning.ProductionCalendar': ('calendars', 'Календари', [
        ('ID', 'pk'), ('Название', 'name'), ('Год', 'year'), ('По умолчанию', 'is_default')], []),
    'estimates.EstimateCatalog': ('estimate-catalogue', 'Справочник смет', [
        ('Тип', 'get_kind_display'), ('Обоснование', 'code'), ('Наименование', 'name'), ('Единица измерения', 'unit')], []),
}


def catalogue_response(request, queryset):
    require_roles(request.user, READ_ROLES)
    filename, title, columns, related = SCHEMAS[queryset.model._meta.label]
    queryset = queryset.filter(company=request.user.company)
    if related:
        queryset = queryset.select_related(*related)
    sheets = [(title, [label for label, path in columns], sheet_rows(queryset, columns))]
    label = queryset.model._meta.label
    if label == 'works.ProjectWork':
        from apps.works.models import ProjectWorkItem
        children = ProjectWorkItem.objects.filter(company=request.user.company,
            project_work_id__in=queryset.values('pk')).select_related(
                'project_work__section__construction_object__project', 'project_work__complex_object', 'load_profile', 'parent')
        columns = [
            ('ID', 'pk'), ('ID комплекса', 'project_work_id'), ('Проект', 'project_work.section.construction_object.project.name'),
            ('Объект', 'project_work.section.construction_object.name'), ('Раздел', 'project_work.section.name'),
            ('Комплексный объект', 'project_work.complex_object.name'), ('Комплекс работ', 'project_work.name'),
            ('Родительская работа', 'parent.name'), ('Работа в комплексе', 'name'), ('Порядок', 'sequence'),
            ('Единица измерения', 'unit'), ('Норматив на единицу комплекса', 'quantity_per_unit'),
            ('Доля стоимости, %', 'weight'), ('Плановый объём', 'planned_quantity'), ('Профиль нагрузки', 'load_profile.name')]
        sheets.append(('Работы в комплексах', [label for label, path in columns], sheet_rows(children, columns)))
    elif label == 'planning.LoadProfile':
        from apps.planning.models import LoadProfileItem
        children = LoadProfileItem.objects.filter(company=request.user.company,
            profile_id__in=queryset.values('pk')).select_related('profile').order_by('profile_id', 'workday_number')
        columns = [('ID профиля', 'profile_id'), ('Профиль', 'profile.name'),
            ('Рабочий день', 'workday_number'), ('Процент', 'percentage')]
        sheets.append(('Дни профилей', [label for label, path in columns], sheet_rows(children, columns)))
    elif label == 'planning.ProductionCalendar':
        from apps.planning.models import CalendarDay
        children = CalendarDay.objects.filter(company=request.user.company,
            calendar_id__in=queryset.values('pk')).select_related('calendar').order_by('calendar_id', 'date')
        columns = [('ID календаря', 'calendar_id'), ('Календарь', 'calendar.name'), ('Дата', 'date'),
            ('Рабочий', 'is_working'), ('Праздник', 'is_holiday'), ('Сокращённый', 'is_shortened'), ('Примечание', 'note')]
        sheets.append(('Дни календарей', [label for label, path in columns], sheet_rows(children, columns)))
    return xlsx_response(filename, sheets)


class CatalogueExcelMixin:
    def get(self, request, *args, **kwargs):
        if request.GET.get('export') == 'xlsx':
            queryset = self.get_queryset()
            form = getattr(self, 'filter_form', None)
            if form is not None and not form.is_valid():
                return HttpResponse('Некорректные фильтры. Исправьте значения и повторите экспорт.',
                    status=400, content_type='text/plain; charset=utf-8')
            return catalogue_response(request, queryset)
        return super().get(request, *args, **kwargs)
