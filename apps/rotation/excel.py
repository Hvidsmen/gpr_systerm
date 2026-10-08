"""Editable rotation matrices with Excel dropdowns and recalculated coverage."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.workbook.properties import CalcProperties
from .services import matrix


def export_matrix(plan, start, end):
    days, positions = matrix(plan, start, end)
    book = Workbook()
    book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True)
    sheet = book.active
    sheet.title = 'Перевахта'
    sheet.append([plan.title])
    sheet.append([f'{plan.source.construction_object} · {start:%d.%m.%Y} — {end:%d.%m.%Y}'])
    sheet.append(['Редактируйте статусы людей: В = вахта, О = отдых. Итоги пересчитываются в Excel. Изменения файла не загружаются в систему автоматически.'])
    sheet.append(['Должность / человек', 'График', 'Начало вахты', *days])
    sheet.freeze_panes = 'D5'
    sheet.column_dimensions['A'].width = 36
    sheet.column_dimensions['B'].width = 14
    sheet.column_dimensions['C'].width = 16
    validation = DataValidation(type='list', formula1='"В,О"', allow_blank=False)
    validation.errorTitle = 'Статус перевахты'
    validation.error = 'Выберите В (вахта) или О (отдых).'
    validation.showErrorMessage = True
    validation.errorStyle = 'stop'
    sheet.add_data_validation(validation)
    green, gray = PatternFill('solid', fgColor='E6F5EE'), PatternFill('solid', fgColor='F1F5F9')
    red, amber = PatternFill('solid', fgColor='FEE2E2'), PatternFill('solid', fgColor='FFF4D8')
    for position in positions:
        sheet.append([position['position'].brigade.name])
        sheet.cell(sheet.max_row,1).font = Font(bold=True, color='174EA6')
        sheet.append(['Потребность', '', '', *[v if v is not None else '—' for v in position['needed']]])
        demand_row = sheet.max_row
        first_person = sheet.max_row+1
        for item in position['people']:
            person = item['person']
            sheet.append([person.name, f'{person.on_days}/{person.off_days}', person.anchor, *['В' if cell['present'] else 'О' for cell in item['statuses']]])
            sheet.cell(sheet.max_row,3).number_format = 'dd.mm.yyyy'
            status_range = f'D{sheet.max_row}:{get_column_letter(3+len(days))}{sheet.max_row}'
            validation.add(status_range)
            sheet.conditional_formatting.add(status_range, FormulaRule(formula=[f'D{sheet.max_row}="В"'], fill=green))
            sheet.conditional_formatting.add(status_range, FormulaRule(formula=[f'D{sheet.max_row}="О"'], fill=gray))
            for offset, cell in enumerate(item['statuses'],4):
                if cell['manual']:
                    sheet.cell(sheet.max_row,offset).comment = Comment('Ручной статус в системе', 'Перевахта')
        last_person = sheet.max_row
        present_row = last_person+1
        sheet.append(['На вахте', '', '', *[f'=COUNTIF({get_column_letter(c)}{first_person}:{get_column_letter(c)}{last_person},"В")' if position['people'] else '=0' for c in range(4,4+len(days))]])
        sheet.append(['Нехватка', '', '', *[f'=IF(ISNUMBER({get_column_letter(c)}{demand_row}),MAX({get_column_letter(c)}{demand_row}-{get_column_letter(c)}{present_row},0),"—")' for c in range(4,4+len(days))]])
        sheet.conditional_formatting.add(f'D{sheet.max_row}:{get_column_letter(3+len(days))}{sheet.max_row}', CellIsRule(operator='greaterThan', formula=['0'], fill=red))
        sheet.append(['Избыток', '', '', *[f'=IF(ISNUMBER({get_column_letter(c)}{demand_row}),MAX({get_column_letter(c)}{present_row}-{get_column_letter(c)}{demand_row},0),"—")' for c in range(4,4+len(days))]])
        sheet.conditional_formatting.add(f'D{sheet.max_row}:{get_column_letter(3+len(days))}{sheet.max_row}', CellIsRule(operator='greaterThan', formula=['0'], fill=amber))
        sheet.append([])
    for column in range(4,4+len(days)):
        sheet.column_dimensions[get_column_letter(column)].width = 7
        sheet.cell(4,column).number_format = 'dd.mm'
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical='center', horizontal='left' if cell.column <= 3 else 'center')
            # User-supplied names are always text, including names starting with '='.
            if cell.column <= 3 and isinstance(cell.value,str): cell.data_type = 's'
            if cell.row == 4:
                cell.fill = PatternFill('solid', fgColor='EAF1FF')
                cell.font = Font(bold=True)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:4'
    output = BytesIO()
    book.save(output)
    return output.getvalue()
