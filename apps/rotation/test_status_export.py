from io import BytesIO
from datetime import timedelta
from django.test import TestCase
from django.core.exceptions import ValidationError
from django.urls import reverse
from openpyxl import load_workbook
from apps.accounts.models import Role
from .tests import RotationTests
from .models import RotationStatus
from .services import generate_people, matrix, refresh_demand
from .excel import export_matrix


class StatusExportTests(TestCase):
    setUpTestData = classmethod(RotationTests.setUpTestData.__func__)

    def setUp(self):
        RotationTests.setUp(self)
        generate_people(self.position)
        self.person = self.position.people.first()
        self.url = reverse('rotation:status_update',args=[self.person.pk])
        self.export = reverse('rotation:plan_export',args=[self.plan.pk])

    def test_manual_status_recalculates_coverage_for_one_day_and_reset_restores_cycle(self):
        response = self.client.post(self.url, {'day':self.start,'status':'OFF'})
        self.assertEqual(response.status_code,302)
        row=matrix(self.plan,self.start,self.start+timedelta(days=1))[1][0]
        self.assertEqual(row['present'],[0,1])
        self.assertEqual(row['shortage'],[1,0])
        self.assertTrue(row['people'][0]['statuses'][0]['manual'])
        refresh_demand(self.plan)
        self.assertEqual(RotationStatus.objects.count(),1)
        self.client.post(self.url,{'day':self.start,'status':'AUTO'})
        self.assertEqual(RotationStatus.objects.count(),0)
        self.assertEqual(matrix(self.plan,self.start,self.start)[1][0]['present'],[1])

    def test_validation_and_isolation(self):
        for day,status in [('bad','ON'),(self.end+timedelta(days=1),'ON'),(self.start,'BAD')]:
            response=self.client.post(self.url,{'day':day,'status':status})
            self.assertEqual(response.status_code,200)
            self.assertTrue(response.context['form'].errors)
        self.assertFalse(RotationStatus.objects.exists())
        with self.assertRaises(ValidationError):
            RotationStatus.objects.create(company=self.foreign,person=self.person,day=self.start,status='ON')
        self.client.force_login(self.other_user)
        self.assertEqual(self.client.get(self.url).status_code,404)
        self.assertEqual(self.client.post(self.url,{'day':self.start,'status':'OFF'}).status_code,404)
        self.assertEqual(self.client.get(self.export).status_code,404)

    def test_export_contains_matrix_editable_statuses_formulas_and_manual_marker(self):
        self.client.post(self.url,{'day':self.start,'status':'OFF'})
        response=self.client.get(self.export)
        self.assertEqual(response.status_code,200)
        self.assertIn('.xlsx',response['Content-Disposition'])
        sheet=load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.freeze_panes,'D5')
        self.assertEqual(sheet.cell(4,4).value.date(),self.start)
        self.assertEqual(sheet.max_column,184)
        self.assertEqual(sheet['D7'].value,'О')
        self.assertIn('Ручной',sheet['D7'].comment.text)
        self.assertEqual(sheet['D9'].value,'=COUNTIF(D7:D8,"В")')
        self.assertIn('MAX(D6-D9,0)',sheet['D10'].value)
        validation=list(sheet.data_validations.dataValidation)[0]
        self.assertEqual(validation.formula1,'"В,О"')
        self.assertTrue(validation.showErrorMessage)
        self.assertIn('D7',validation.sqref)

    def test_month_export_and_safe_names_empty_roster(self):
        self.person.name='=HYPERLINK("https://example.com")';self.person.save()
        response=self.client.get(self.export,{'month':'2026-02'})
        sheet=load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_column,31)
        self.assertEqual(sheet['A7'].data_type,'s')
        self.assertEqual(sheet['D4'].value.month,2)
        self.assertEqual(self.client.get(self.export,{'month':'bad'}).status_code,400)
        self.assertEqual(self.client.get(self.export,{'month':'2025-01'}).status_code,400)
        self.position.people.all().delete()
        sheet=load_workbook(BytesIO(export_matrix(self.plan,self.start,self.start))).active
        self.assertEqual(sheet['D7'].value,'=0')

    def test_reader_can_export_but_cannot_modify_status(self):
        self.user.role=Role.objects.get(code='MANAGER');self.user.save(update_fields=['role'])
        self.assertEqual(self.client.get(self.export,{'month':'2026-01'}).status_code,200)
        self.assertEqual(self.client.get(self.url).status_code,403)
        self.assertEqual(self.client.post(self.url,{'day':self.start,'status':'OFF'}).status_code,403)
        response=self.client.get(reverse('rotation:plan_detail',args=[self.plan.pk]))
        self.assertNotContains(response,'class="rotation-status"')
