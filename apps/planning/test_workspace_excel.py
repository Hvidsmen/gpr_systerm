from datetime import date
from decimal import Decimal
from io import BytesIO
from copy import deepcopy
from django.core.exceptions import ValidationError, PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from apps.accounts.models import User, Role, Company
from apps.resources.models import Brigade
from apps.production.models import FuelFact, LaborFact
from apps.works.models import ProjectWork, ProjectWorkItem
from . import test_workspace as workspace_tests
from .test_workspace import JAN, FEB, MAR
from .workspace_excel import make_workbook, import_workbook
from .workspace_services import WorkspaceService
from .models import WorkMonthAllocation, ResourceMonthAllocation
from .workspace_forms import WorkAllocationForm

D=Decimal


class WorkspaceExcelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        workspace_tests.WorkspaceTests.setUpTestData.__func__(cls)

    create=workspace_tests.WorkspaceTests.create
    add_resources=workspace_tests.WorkspaceTests.add_resources
    approve=workspace_tests.WorkspaceTests.approve

    def setUp(self):
        self.workspace=self.create([self.simple,self.composite])
        self.version=self.workspace.baseline_version
        self.client.force_login(self.planner)

    def workbook(self, months=None):
        return load_workbook(BytesIO(make_workbook(self.version,months or [JAN,FEB,MAR],True)))

    def upload(self, workbook):
        stream=BytesIO();workbook.save(stream);workbook.close()
        return SimpleUploadedFile('plan.xlsx',stream.getvalue())

    def set_value(self,workbook,sheet,name,value,month=JAN,column=None):
        ws=workbook[sheet]
        col=column or month.strftime('%m.%Y')
        headers={cell.value:cell.column for cell in ws[1]}
        name_col={'Работы':'Подработа' if name in ['A','B'] else 'Работа','Люди':'Бригада','Техника':'Вид техники','ГСМ':'Вид ГСМ'}[sheet]
        for row in ws.iter_rows(min_row=2):
            if row[headers[name_col]-1].value==name:
                ws.cell(row[0].row,headers[col]).value=value;return
        raise AssertionError(name)

    def fill_composite(self,workbook,a=20,b=24,month=JAN):
        self.set_value(workbook,'Работы','A',a,month)
        self.set_value(workbook,'Работы','B',b,month)

    def row(self,work=None,month=JAN):
        return self.version.work_allocations.get(work=work or self.composite,month=month)

    def test_period_and_month_templates_have_four_sheets_and_no_codes(self):
        for months in [[JAN,FEB,MAR],[FEB]]:
            wb=self.workbook(months)
            self.assertEqual(wb.sheetnames,['Работы','Люди','Техника','ГСМ'])
            self.assertEqual([cell.value for cell in wb['Работы'][1]][8:8+len(months)],[m.strftime('%m.%Y') for m in months])
            self.assertTrue(wb['Работы'].column_dimensions['A'].hidden)
            for ws in wb:
                self.assertNotIn('Код',[cell.value for cell in ws[1]])
            self.assertEqual(wb['ГСМ'].max_column,2+2*len(months))
            wb.close()

    def test_import_derives_minimum_and_keeps_both_subwork_totals_in_daily_plan(self):
        wb=self.workbook();self.fill_composite(wb)
        self.assertEqual(import_workbook(self.planner,self.version,self.upload(wb),[JAN,FEB,MAR]),(1,0))
        row=self.row()
        self.assertEqual(row.quantity,D(8))
        self.assertEqual(row.item_quantities,{str(self.a.pk):'20',str(self.b.pk):'24'})
        version=WorkspaceService.refresh(self.version,self.planner)
        spec=next(spec for spec in version.snapshot['works'] if spec['id']==self.composite.pk)
        for item,total in [(self.a,20),(self.b,24)]:
            self.assertEqual(sum(D(plan['quantity']) for plan in spec['plans'] if plan['item_id']==item.pk and plan['date'].startswith('2026-01')),D(total))
        export=load_workbook(BytesIO(make_workbook(version,[JAN])))
        self.assertEqual(import_workbook(self.planner,version,self.upload(export),[JAN]),(2,0))
        self.assertEqual(self.row().item_quantities,row.item_quantities)

    def test_fractional_setting_and_zero_main_work_with_partial_subwork_plan(self):
        wb=self.workbook([JAN]);self.fill_composite(wb,20.5,24.9)
        import_workbook(self.planner,self.version,self.upload(wb),[JAN])
        self.assertEqual(self.row().quantity,D('8.300'))
        self.composite.allow_fractional=False;self.composite.save()
        wb=self.workbook([JAN]);self.fill_composite(wb,20.5,24.9)
        import_workbook(self.planner,self.version,self.upload(wb),[JAN])
        self.assertEqual(self.row().quantity,D(8))
        wb=self.workbook([JAN]);self.fill_composite(wb,20,0)
        import_workbook(self.planner,self.version,self.upload(wb),[JAN])
        version=WorkspaceService.refresh(self.version,self.planner)
        spec=next(s for s in version.snapshot['works'] if s['id']==self.composite.pk)
        self.assertEqual(sum(D(p['quantity']) for p in spec['plans'] if p['item_id']==self.a.pk and p['date'].startswith('2026-01')),D(20))
        self.assertEqual(self.row().quantity,0)

    def test_month_import_all_resources_repeats_counts_and_balance_and_distributes_expense(self):
        wb=self.workbook([JAN])
        self.set_value(wb,'Работы','Simple',12)
        self.set_value(wb,'Люди','Brigade',10)
        self.set_value(wb,'Техника','Excavator',3)
        self.set_value(wb,'ГСМ','Дизельное топливо',700,column='Остаток 01.2026')
        self.set_value(wb,'ГСМ','Дизельное топливо',3100,column='Расход 01.2026')
        self.assertEqual(import_workbook(self.planner,self.version,self.upload(wb),[JAN]),(1,3))
        self.assertEqual(self.row(self.simple,JAN).quantity,D(12))
        self.assertEqual(self.row(self.simple,FEB).quantity,D(200))
        version=WorkspaceService.refresh(self.version,self.planner)
        self.assertEqual({row['planned_workers'] for row in version.snapshot['resources']['labor']},{10})
        self.assertEqual({row['planned_count'] for row in version.snapshot['resources']['equipment']},{3})
        self.assertTrue(all(D(row['planned_balance'])==700 and D(row['planned_liters'])==100 for row in version.snapshot['resources']['fuel']))
        self.assertEqual(len(version.snapshot['resources']['fuel']),31)

    def test_blank_cells_preserve_old_values_and_zero_overwrites(self):
        self.add_resources(self.workspace)
        wb=load_workbook(BytesIO(make_workbook(self.version,[JAN,FEB,MAR])))
        self.set_value(wb,'Работы','Simple',None,JAN)
        self.set_value(wb,'Работы','Simple',0,FEB)
        import_workbook(self.planner,self.version,self.upload(wb),[JAN,FEB,MAR])
        self.assertEqual(self.row(self.simple,JAN).quantity,300)
        self.assertEqual(self.row(self.simple,FEB).quantity,0)
        self.assertEqual(self.row(self.simple,MAR).quantity,500)

    def test_errors_anywhere_prevent_all_changes(self):
        wb=self.workbook([JAN]);self.set_value(wb,'Работы','Simple',12)
        self.set_value(wb,'ГСМ','Дизельное топливо',-1,column='Остаток 01.2026')
        self.set_value(wb,'ГСМ','Дизельное топливо',100,column='Расход 01.2026')
        with self.assertRaises(ValidationError):import_workbook(self.planner,self.version,self.upload(wb),[JAN])
        self.assertEqual(self.row(self.simple).quantity,300)
        self.assertFalse(self.version.resource_allocations.exists())

    def test_missing_child_duplicate_formula_nan_and_partial_fuel_are_rejected(self):
        for mode in ['missing','duplicate','formula','nan','partial_fuel']:
            with self.subTest(mode=mode):
                wb=self.workbook([JAN])
                self.fill_composite(wb)
                if mode=='missing':self.set_value(wb,'Работы','B',None)
                if mode=='duplicate':
                    ws=wb['Работы'];row=next(row for row in ws.iter_rows(min_row=2) if row[5].value=='A');ws.append([cell.value for cell in row])
                if mode=='formula':self.set_value(wb,'Работы','A','=10*2')
                if mode=='nan':self.set_value(wb,'Работы','A','NaN')
                if mode=='partial_fuel':self.set_value(wb,'ГСМ','Дизельное топливо',1,column='Остаток 01.2026')
                with self.assertRaises(ValidationError):import_workbook(self.planner,self.version,self.upload(wb),[JAN])
                self.assertEqual(self.row().quantity,300)

    def test_foreign_ids_and_unknown_names_rejected_but_unique_names_without_ids_work(self):
        wb=self.workbook([JAN]);self.set_value(wb,'Работы','Simple',12)
        row=next(row for row in wb['Работы'].iter_rows(min_row=2) if row[1].value=='Simple');row[0].value=999999
        with self.assertRaises(ValidationError):import_workbook(self.planner,self.version,self.upload(wb),[JAN])
        wb=self.workbook([JAN]);self.set_value(wb,'Работы','Simple',12)
        row=next(row for row in wb['Работы'].iter_rows(min_row=2) if row[1].value=='Simple');row[0].value=None
        import_workbook(self.planner,self.version,self.upload(wb),[JAN])
        self.assertEqual(self.row(self.simple).quantity,12)

    def test_manual_main_quantity_resets_imported_children_but_profile_change_preserves_them(self):
        wb=self.workbook([JAN]);self.fill_composite(wb);import_workbook(self.planner,self.version,self.upload(wb),[JAN])
        row=self.row()
        form=WorkAllocationForm(data={'month':JAN,'work':self.composite.pk,'quantity':'8','load_profile':''},instance=row,user=self.planner,version=self.version,month=JAN)
        self.assertTrue(form.is_valid(),form.errors);form.save()
        self.assertTrue(self.row().item_quantities)
        form=WorkAllocationForm(data={'month':JAN,'work':self.composite.pk,'quantity':'9','load_profile':''},instance=self.row(),user=self.planner,version=self.version,month=JAN)
        self.assertTrue(form.is_valid(),form.errors);form.save()
        self.assertEqual(self.row().item_quantities,{})
        self.assertEqual(self.row().quantity,9)

    def test_subwork_values_survive_baseline_approval_forecast_copy_and_export(self):
        wb=self.workbook();self.fill_composite(wb);import_workbook(self.planner,self.version,self.upload(wb),[JAN,FEB,MAR])
        self.approve(self.version)
        forecast=WorkspaceService.forecast(self.workspace,self.planner,JAN,'BASELINE')
        self.assertEqual(forecast.work_allocations.get(work=self.composite).item_quantities,self.row().item_quantities)
        WorkspaceService.refresh(forecast,self.planner)
        content=make_workbook(forecast,[JAN,FEB,MAR])
        self.assertTrue(content.startswith(b'PK'))
        with self.assertRaises(PermissionDenied):import_workbook(self.planner,self.version,SimpleUploadedFile('x.xlsx',content),[JAN])

    def test_http_export_import_rights_scope_and_invalid_uploads(self):
        export=reverse('planning:workspace_excel_export',args=[self.version.pk]);imp=reverse('planning:workspace_excel_import',args=[self.version.pk])
        response=self.client.get(export,{'scope':'month','month':JAN,'template':1})
        self.assertEqual(response.status_code,200)
        self.assertEqual(load_workbook(BytesIO(response.content)).sheetnames,['Работы','Люди','Техника','ГСМ'])
        wb=self.workbook([JAN]);self.fill_composite(wb)
        response=self.client.post(imp,{'scope':'month','month':JAN,'file':self.upload(wb)})
        self.assertEqual(response.status_code,302)
        self.assertEqual(self.row().quantity,8)
        response=self.client.post(imp,{'scope':'month','month':JAN,'file':SimpleUploadedFile('bad.xlsx',b'bad')})
        self.assertEqual(response.status_code,400)
        self.client.force_login(self.manager)
        self.assertEqual(self.client.get(export).status_code,200)
        self.assertEqual(self.client.get(imp).status_code,403)
        other=Company.objects.create(name='Excel foreign')
        user=User.objects.create_user(username='excel-foreign',company=other,role=Role.objects.get(code='ADMIN'))
        self.client.force_login(user)
        self.assertEqual(self.client.get(export).status_code,404)
        self.assertEqual(self.client.get(imp).status_code,404)

    def test_frozen_export_uses_snapshot_not_mutated_catalog_profile_and_formula_names_are_text(self):
        self.approve(self.version)
        content=make_workbook(self.version,[JAN])
        self.assertTrue(content.startswith(b'PK'))
        self.simple.name='=Dangerous()';self.simple.save()
        workbook=load_workbook(BytesIO(make_workbook(self.version,[JAN],True)))
        cell=next(row[1] for row in workbook['Работы'].iter_rows(min_row=2) if row[1].value=='=Dangerous()')
        self.assertEqual(cell.data_type,'s')
        workbook.close()

    def test_year_template_includes_all_twelve_months_and_month_template_only_one(self):
        self.version.end_date=date(2026,12,31)
        self.workspace.end_date=self.version.end_date
        self.workspace.save();self.version.save()
        url=reverse('planning:workspace_excel_export',args=[self.version.pk])
        response=self.client.get(url,{'template':1,'scope':'period'})
        wb=load_workbook(BytesIO(response.content))
        self.assertEqual([c.value for c in wb['Работы'][1]][8:20],[f'{m:02}.2026' for m in range(1,13)])
        self.assertEqual(wb['ГСМ'].max_column,26)
        self.assertTrue(all(c.data_type=='f' for row in wb['Работы'].iter_rows(min_row=2,min_col=21) for c in row))
        wb.close()

    def test_large_subwork_amounts_round_trip_without_float_precision_loss(self):
        wb=self.workbook([JAN]);self.fill_composite(wb,'123456789012.123456','123456789012.123456')
        import_workbook(self.planner,self.version,self.upload(wb),[JAN])
        expected=deepcopy(self.row().item_quantities)
        exported=load_workbook(BytesIO(make_workbook(self.version,[JAN])))
        import_workbook(self.planner,self.version,self.upload(exported),[JAN])
        self.assertEqual(self.row().item_quantities,expected)

    def test_remaining_forecast_import_skips_automatically_calculated_future(self):
        self.add_resources(self.workspace)
        self.approve(self.version)
        forecast=WorkspaceService.forecast(self.workspace,self.planner,FEB,'REMAINING')
        content=make_workbook(forecast,[JAN,FEB,MAR])
        wb=load_workbook(BytesIO(content))
        self.set_value(wb,'Работы','Simple',210,FEB)
        self.set_value(wb,'Работы','Simple',999,MAR)
        self.set_value(wb,'ГСМ','Дизельное топливо',999,column='Расход 03.2026')
        import_workbook(self.planner,forecast,self.upload(wb),[FEB,MAR])
        self.assertEqual(forecast.work_allocations.get(work=self.simple,month=FEB).quantity,210)
        self.assertFalse(forecast.work_allocations.filter(month=MAR).exists())
        self.assertFalse(forecast.resource_allocations.filter(month=MAR).exists())
        refreshed=WorkspaceService.refresh(forecast,self.planner)
        spec=next(row for row in refreshed.snapshot['works'] if row['id']==self.simple.pk)
        self.assertEqual(sum(D(r['quantity']) for r in spec['daily'] if r['date'].startswith('2026-03')),790)

    def test_stock_and_counts_in_fact_months_export_with_explanatory_comments(self):
        self.add_resources(self.workspace)
        self.approve(self.version)
        for day,count,stock in [(1,5,100),(2,7,90)]:
            LaborFact.objects.create(company=self.company,construction_object=self.obj,date=JAN.replace(day=day),brigade=self.brigade,actual_workers=count,actual_hours=1)
            FuelFact.objects.create(company=self.company,construction_object=self.obj,date=JAN.replace(day=day),fuel_type='DIESEL',equipment_ref='A1',actual_balance=stock,actual_liters=10)
        forecast=WorkspaceService.forecast(self.workspace,self.planner,FEB,'REMAINING')
        wb=load_workbook(BytesIO(make_workbook(forecast,[JAN,FEB,MAR])))
        self.assertEqual(wb['Люди'].cell(2,6).value,6)
        self.assertIsNotNone(wb['Люди'].cell(2,6).comment)
        self.assertEqual(wb['ГСМ'].cell(2,3).value,90)
        self.assertIsNotNone(wb['ГСМ'].cell(2,3).comment)
        wb.close()

    def test_explicit_reset_flag_restores_normatives_even_when_quantity_is_unchanged(self):
        wb=self.workbook([JAN]);self.fill_composite(wb);import_workbook(self.planner,self.version,self.upload(wb),[JAN])
        form=WorkAllocationForm(data={'month':JAN,'work':self.composite.pk,'quantity':'8','load_profile':'','reset_item_quantities':'on'},instance=self.row(),user=self.planner,version=self.version,month=JAN)
        self.assertTrue(form.is_valid(),form.errors);form.save()
        self.assertEqual(self.row().item_quantities,{})

    def test_extra_foreign_subwork_and_fractional_people_are_rejected(self):
        for mode in ['foreign_child','fractional_people']:
            with self.subTest(mode=mode):
                wb=self.workbook([JAN]);self.fill_composite(wb)
                if mode=='foreign_child':
                    row=next(row for row in wb['Работы'].iter_rows(min_row=2) if row[5].value=='A');row[4].value=99999
                else:
                    self.set_value(wb,'Люди','Brigade',1.5)
                with self.assertRaises(ValidationError):import_workbook(self.planner,self.version,self.upload(wb),[JAN])
                self.assertFalse(self.row().item_quantities)

    def test_year_scope_in_multi_year_plan_imports_only_that_year(self):
        self.version.end_date=date(2027,12,31)
        self.workspace.end_date=self.version.end_date
        self.workspace.save();self.version.save()
        export=reverse('planning:workspace_excel_export',args=[self.version.pk])
        response=self.client.get(export,{'scope':'year','year':2027,'template':1})
        self.assertEqual(response.status_code,200)
        wb=load_workbook(BytesIO(response.content))
        self.assertEqual([cell.value for cell in wb['Работы'][1]][8:20],[f'{m:02}.2027' for m in range(1,13)])
        self.set_value(wb,'Работы','Simple',10,date(2027,1,1))
        response=self.client.post(reverse('planning:workspace_excel_import',args=[self.version.pk]),{'scope':'year','year':2027,'file':self.upload(wb)})
        self.assertEqual(response.status_code,302)
        self.assertEqual(self.version.work_allocations.get(work=self.simple,month=date(2027,1,1)).quantity,10)
        self.assertEqual(self.row(self.simple,JAN).quantity,300)

    def test_derived_work_quantities_are_calculated_columns_not_trusted_input(self):
        wb=self.workbook([JAN]);self.fill_composite(wb)
        self.assertEqual(wb['Работы'].cell(1,10).value,'План работы 01.2026')
        for row in wb['Работы'].iter_rows(min_row=2):row[9].value=999999
        import_workbook(self.planner,self.version,self.upload(wb),[JAN])
        self.assertEqual(self.row().quantity,8)

    def test_forecast_rejects_imported_subworks_not_in_frozen_baseline(self):
        wb=self.workbook([JAN]);self.fill_composite(wb)
        import_workbook(self.planner,self.version,self.upload(wb),[JAN])
        self.approve(self.version)
        forecast=WorkspaceService.forecast(self.workspace,self.planner,JAN,'BASELINE')
        third=ProjectWorkItem.objects.create(company=self.company,project_work=self.composite,name='New part',unit='m',quantity_per_unit=1,weight=0)
        wb=load_workbook(BytesIO(make_workbook(forecast,[JAN],True)))
        self.fill_composite(wb)
        row=next(row for row in wb['Работы'].iter_rows(min_row=2) if row[5].value==third.name)
        row[8].value=3
        with self.assertRaises(ValidationError):import_workbook(self.planner,forecast,self.upload(wb),[JAN])
        self.assertEqual(forecast.work_allocations.get(work=self.composite).item_quantities,self.row().item_quantities)
