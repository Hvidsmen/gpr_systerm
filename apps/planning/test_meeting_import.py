from datetime import date
from decimal import Decimal
from io import BytesIO

from django.core.exceptions import ValidationError, PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from apps.projects.models import ConstructionObject, Section, Project
from apps.resources.models import Brigade, EquipmentType
from apps.works.models import ProjectWork
from .meeting_import import parse_meeting_workbook, resolve_sheet, apply_meeting_import
from .models import (
    PlanningWorkspace,
    WorkMonthAllocation,
    ResourceMonthAllocation,
    LoadProfileItem,
)
from . import test_workspace

JAN = date(2026, 1, 1)
FEB = date(2026, 2, 1)


def upload(name="o", work_names=None, variable=False, invalid=False, sections=None):
    book = Workbook()
    sheet = book.active
    sheet.title = name
    sheet.append(
        [
            "№",
            "Наименование работ",
            "Ед. изм.",
            None,
            None,
            None,
            None,
            None,
            "План на месяц",
            None,
            None,
            None,
            "План/Факт",
        ]
    )
    sheet.append([None] * 13 + [date(2025, 1, 1), JAN, date(2026, 1, 2), FEB])
    sheet.append([None, "s"])
    for index, (work, unit, volume) in enumerate(work_names or [("Imported work", "м", 2)]):
        if sections:
            sheet.append([None, sections[index]])
        sheet.append(
            [
                "1",
                work,
                unit,
                None,
                None,
                None,
                None,
                None,
                9999,
                None,
                None,
                None,
                "план",
                "#REF!",
                volume,
                "#REF!" if invalid else 3,
                4,
            ]
        )
        sheet.append([None] * 12 + ["факт", 999, 999, 999, 999])
    sheet.append([None] * 6 + ["Людские ресурсы"])
    sheet.append(
        [None] * 5
        + [
            "1",
            "Imported brigade",
            None,
            None,
            None,
            None,
            "чел.",
            "план",
            "#REF!",
            2,
            3 if variable else 2,
            5,
        ]
    )
    sheet.append([None] * 12 + ["факт", 99, 99, 99, 99])
    sheet.append([None] * 6 + ["Технические ресурсы"])
    sheet.append(
        [None] * 5
        + [
            "1",
            "Imported equipment",
            None,
            None,
            None,
            None,
            "ед.",
            "план",
            "#REF!",
            1,
            1,
            2,
        ]
    )
    sheet.append([None] * 12 + ["факт", 99, 99, 99, 99])
    content = BytesIO()
    book.save(content)
    return SimpleUploadedFile("meeting.xlsx", content.getvalue())


class MeetingImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        test_workspace.WorkspaceTests.setUpTestData.__func__(cls)

    def parse(self, **kwargs):
        return parse_meeting_workbook(upload(**kwargs), JAN, date(2026, 2, 1))

    def test_exact_period_daily_plan_only_and_monthly_totals_ignored(self):
        sheet = self.parse()[0]
        self.assertEqual(sheet["errors"], [])
        works = [r for r in sheet["entries"] if r["kind"] == "work"]
        self.assertEqual(
            [(r["month"], Decimal(r["quantity"])) for r in works],
            [("2026-01-01", Decimal(5)), ("2026-02-01", Decimal(4))],
        )
        resources = [r for r in sheet["entries"] if r["kind"] == "labor"]
        self.assertEqual(
            [Decimal(r["quantity"]) for r in resources], [Decimal(2), Decimal(5)]
        )
        self.assertFalse(any(Decimal(r["quantity"]) == 999 for r in sheet["entries"]))

    def test_partial_period_excludes_other_dates_even_errors(self):
        sheet = parse_meeting_workbook(upload(invalid=True), JAN, JAN)[0]
        self.assertFalse(sheet["errors"])
        self.assertEqual(sheet["entries"][0]["quantity"], "2.000000")

    def test_errors_in_selected_dates_and_variable_count_rule(self):
        self.assertTrue(self.parse(invalid=True)[0]["errors"])
        sheet = self.parse(variable=True)[0]
        self.assertEqual(
            next(r for r in sheet["entries"] if r["kind"] == "labor")["quantity"], "3"
        )
        self.assertTrue(
            any("количество по дням меняется" in w for w in sheet["warnings"])
        )
        first = parse_meeting_workbook(
            upload(variable=True), JAN, date(2026, 1, 31), "first"
        )[0]
        self.assertEqual(
            next(r for r in first["entries"] if r["kind"] == "labor")["quantity"], "2"
        )

    def test_existing_objects_only_and_no_cross_company_or_project(self):
        sheet = self.parse(name="missing")[0]
        count = ConstructionObject.objects.count()
        with self.assertRaises(ValidationError):
            apply_meeting_import(
                self.planner, self.obj.project_id, [sheet], {"0"}, JAN, date(2026, 2, 1)
            )
        self.assertEqual(ConstructionObject.objects.count(), count)
        wrong_project = Project.objects.create(
            company=self.company, name="Other project"
        )
        with self.assertRaises(ValidationError):
            resolve_sheet(self.company, wrong_project, self.parse()[0])
        with self.assertRaises(Project.DoesNotExist):
            apply_meeting_import(self.planner, 99999, self.parse(), {"0"}, JAN, FEB)
        with self.assertRaises(PermissionDenied):
            apply_meeting_import(
                self.manager, self.obj.project_id, self.parse(), {"0"}, JAN, FEB
            )

    def test_preview_no_writes_confirm_new_draft_no_duplicate_catalogs_by_month(self):
        self.client.force_login(self.planner)
        response = self.client.post(
            reverse("planning:meeting_import"),
            {
                "project": self.obj.project_id,
                "start": JAN,
                "end": FEB,
                "resource_rule": "maximum",
                "file": upload(),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(PlanningWorkspace.objects.exists())
        self.assertFalse(ProjectWork.objects.filter(name="Imported work").exists())
        token = response.context["preview"]
        self.assertContains(response, "Imported work")
        confirmed = self.client.post(
            reverse("planning:meeting_import"),
            {"action": "confirm", "preview": token, "sheets": ["0"]},
        )
        self.assertEqual(confirmed.status_code, 302)
        workspace = PlanningWorkspace.objects.get()
        self.assertEqual((workspace.start_date, workspace.end_date), (JAN, FEB))
        self.assertEqual(workspace.baseline_version.status, "DRAFT")
        self.assertEqual(ProjectWork.objects.filter(name="Imported work").count(), 1)
        self.assertEqual(Brigade.objects.filter(name="Imported brigade").count(), 1)
        self.assertEqual(
            EquipmentType.objects.filter(name="Imported equipment").count(), 1
        )
        self.assertEqual(WorkMonthAllocation.objects.count(), 2)
        self.assertEqual(ResourceMonthAllocation.objects.count(), 4)
        work = ProjectWork.objects.get(name="Imported work")
        self.assertEqual(
            LoadProfileItem.objects.get(profile=work.load_profile).percentage, 100
        )
        self.client.post(
            reverse("planning:meeting_import"),
            {"action": "confirm", "preview": token, "sheets": ["0"]},
        )
        self.assertEqual(PlanningWorkspace.objects.count(), 1)

    def test_confirmation_tampering_and_permissions(self):
        self.client.force_login(self.planner)
        self.client.post(
            reverse("planning:meeting_import"),
            {"action": "confirm", "preview": "tampered", "sheets": ["0"]},
        )
        self.assertFalse(PlanningWorkspace.objects.exists())
        self.client.force_login(self.manager)
        self.assertEqual(
            self.client.get(reverse("planning:meeting_import")).status_code, 403
        )
        self.assertEqual(
            self.client.post(reverse("planning:meeting_import"), {}).status_code, 403
        )

    def test_composite_matches_items_and_calculates_minimum(self):
        sheets = self.parse(work_names=[("A", "m", 10), ("B", "m3", 12)])
        plans = apply_meeting_import(
            self.planner, self.obj.project_id, sheets, {"0"}, JAN, FEB
        )
        row = WorkMonthAllocation.objects.get(
            version=plans[0].baseline_version, month=JAN
        )
        self.assertEqual(row.work, self.composite)
        self.assertEqual(
            row.item_quantities,
            {str(self.a.pk): "13.000000", str(self.b.pk): "15.000000"},
        )
        self.assertEqual(row.quantity, Decimal(5))
        with self.assertRaises(ValidationError):
            apply_meeting_import(
                self.planner,
                self.obj.project_id,
                self.parse(work_names=[("A", "m", 10)]),
                {"0"},
                JAN,
                FEB,
            )

    def test_atomic_rollback_on_later_sheet_error(self):
        sheets = self.parse() + self.parse(name="not-in-database")
        with self.assertRaises(ValidationError):
            apply_meeting_import(
                self.planner, self.obj.project_id, sheets, {"0", "1"}, JAN, FEB
            )
        self.assertFalse(PlanningWorkspace.objects.exists())
        self.assertFalse(ProjectWork.objects.filter(name="Imported work").exists())

    def test_duplicate_dates_and_resource_subtotals(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "o"
        sheet.append(
            [
                "№",
                "Наименование работ",
                "Ед. изм.",
                None,
                None,
                None,
                None,
                None,
                "План на месяц",
                None,
                None,
                None,
                "План/Факт",
            ]
        )
        sheet.append(
            [None] * 13
            + [
                date(2025, 1, 1),
                JAN,
                JAN,
                FEB,
                date(2026, 2, 2),
                date(2026, 2, 3),
                date(2026, 2, 4),
            ]
        )
        sheet.append([None] * 6 + ["Людские ресурсы"])
        sheet.append(
            [None] * 5
            + ["1", "ИТР", None, None, None, None, "чел.", "план", None, 3, 3]
        )
        sheet.append(
            [None] * 5
            + ["1.1", "Инженер", None, None, None, None, "чел.", "план", None, 3, 3]
        )
        content = BytesIO()
        book.save(content)
        result = parse_meeting_workbook(
            SimpleUploadedFile("meeting.xlsx", content.getvalue()), JAN, FEB
        )[0]
        self.assertTrue(any("повторные дневные даты" in e for e in result["errors"]))
        self.assertEqual([r["name"] for r in result["entries"]], ["Инженер"])

    def test_bad_extension_rejected(self):
        with self.assertRaises(ValidationError):
            parse_meeting_workbook(SimpleUploadedFile("bad.txt", b"x"), JAN, FEB)

    def test_work_subtotals_become_sections_without_double_counting(self):
        from openpyxl import load_workbook

        book = load_workbook(upload())
        sheet = book.active
        sheet.insert_rows(4)
        sheet.cell(4, 1, "1")
        sheet.cell(4, 2, "Монтаж, в том числе:")
        sheet.cell(4, 3, "м")
        sheet.cell(4, 13, "план")
        sheet.cell(4, 15, 999)
        content = BytesIO()
        book.save(content)
        result = parse_meeting_workbook(
            SimpleUploadedFile("meeting.xlsx", content.getvalue()), JAN, JAN
        )[0]
        works = [r for r in result["entries"] if r["kind"] == "work"]
        self.assertEqual(len(works), 1)
        self.assertEqual(works[0]["section"], "Монтаж, в том числе")
        self.assertEqual(Decimal(works[0]["quantity"]), 2)

    def test_shifted_resource_heading_and_same_named_brigades_in_groups(self):
        from openpyxl import load_workbook

        book = load_workbook(upload())
        sheet = book.active
        sheet.cell(6, 7, None)
        sheet["G6"] = None
        sheet["F6"] = "Людские ресурсы"
        content = BytesIO()
        book.save(content)
        result = parse_meeting_workbook(
            SimpleUploadedFile("meeting.xlsx", content.getvalue()), JAN, JAN
        )[0]
        labor = next(r for r in result["entries"] if r["kind"] == "labor")
        self.assertEqual(labor["name"], "Imported brigade")
        sheets = self.parse()
        jan_labor = next(r for r in sheets[0]["entries"] if r["kind"] == "labor")
        jan_labor["section"] = "Группа 1"
        other = dict(jan_labor, section="Группа 2", row=999)
        sheets[0]["entries"].append(other)
        plans = apply_meeting_import(
            self.planner, self.obj.project_id, sheets, {"0"}, JAN, FEB
        )
        self.assertEqual(
            ResourceMonthAllocation.objects.filter(
                version=plans[0].baseline_version, kind="labor", month=JAN
            ).count(),
            2,
        )
        self.assertEqual(
            Brigade.objects.filter(
                name="Imported brigade", group__isnull=False
            ).count(),
            2,
        )
