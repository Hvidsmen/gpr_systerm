"""Four-sheet, month-based planning interchange. Validate everything before writing."""
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
from zipfile import ZipFile, BadZipFile

from django import forms
from django.contrib import messages
from django.core.exceptions import ValidationError, PermissionDenied
from django.db import transaction
from django.http import HttpResponse, Http404
from django.shortcuts import render, redirect
from django.urls import reverse
from django.views import View
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

from .approval_workflow import baseline_snapshot
from core.permissions import require_roles, PLAN_ROLES, READ_ROLES
from apps.resources.models import Brigade, EquipmentType
from apps.works.models import ProjectWork
from apps.works.progress import work_specification, quantity_from_totals
from .models import GlobalPlanVersion, WorkMonthAllocation, ResourceMonthAllocation
from .workspace_views import version_for
from .workspace_services import months_between, build_workspace_snapshot
from .bulk_add import editable_months

SHEETS = ['Работы', 'Люди', 'Техника', 'ГСМ']
BASE_HEADERS = {
    'Работы': ['ID работы', 'Работа', 'Раздел', 'Группа работ', 'ID подработы', 'Подработа', 'Ед. изм.', 'Норматив'],
    'Люди': ['ID бригады', 'Бригада', 'Макрогруппа', 'Группа', 'Ед. изм.'],
    'Техника': ['ID техники', 'Вид техники', 'Категория', 'Номер машины', 'Ед. изм.'],
    'ГСМ': ['Служебный ключ', 'Вид ГСМ'],
}
MAX_SIZE = 8 * 1024 * 1024
MAX_ROWS = 10000


def label(month):
    return month.strftime('%m.%Y')


def headers(sheet, months):
    return BASE_HEADERS[sheet] + ([title + ' ' + label(month) for month in months for title in ['Остаток', 'Расход']] if sheet == 'ГСМ' else [label(month) for month in months])


def selected_months(version, scope, month=None, year=None):
    all_months = months_between(version.start_date, version.end_date)
    if scope == 'period':
        return all_months
    if scope == 'year':
        try:
            chosen=int(year or (version.planning_month or all_months[0]).year)
        except (ValueError,TypeError):
            raise Http404('Неизвестный год.')
        result=[m for m in all_months if m.year==chosen]
        if not result:raise Http404('Год за пределами плана.')
        return result
    if scope != 'month':
        raise Http404('Неизвестный период Excel.')
    try:
        selected = date.fromisoformat(month or (version.planning_month or all_months[0]).isoformat())
    except (ValueError, TypeError):
        raise Http404('Неизвестный месяц.')
    if selected not in all_months:
        raise Http404('Месяц за пределами плана.')
    return [selected]


def numeric(value):
    amount=Decimal(str(value))
    return format(amount,'f') if len(amount.normalize().as_tuple().digits)>15 else float(amount)


def make_workbook(version, months, template=False):
    workbook = Workbook()
    workbook.remove(workbook.active)
    worksheets = {}
    cell_notes = []
    for sheet in SHEETS:
        ws = workbook.create_sheet(sheet)
        ws.append(headers(sheet, months) + (['План работы '+label(month) for month in months] if sheet=='Работы' else []))
        ws.freeze_panes = 'I2' if sheet == 'Работы' else ('F2' if sheet in ['Люди', 'Техника'] else 'C2')
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='6C757D' if cell.column>len(headers(sheet,months)) else '1263D6')
        worksheets[sheet] = ws
    frozen = version.status in ["SUBMITTED", "APPROVED", "COMPLETED"]
    from apps.resources.equipment_merge import equipment_aliases, normalize_equipment_snapshot
    from apps.resources.brigade_merge import brigade_aliases, normalize_resource_snapshot
    snapshot = normalize_resource_snapshot(version.snapshot, version.company) if frozen else {}
    if not template and not frozen and version.version_kind == 'FORECAST':
        snapshot = build_workspace_snapshot(version)
    allocations = {(row.work_id, row.month): row for row in version.work_allocations.select_related('work')}
    specs = {spec['id']: spec for spec in snapshot.get('works', [])}
    works = ProjectWork.objects.filter(merged_source__isnull=True,company=version.company, section__construction_object=version.construction_object).select_related('section', 'work_group').prefetch_related('items').order_by('work_group__name', 'name', 'pk')
    for work in works:
        relevant = any(key[0] == work.pk for key in allocations) or work.pk in specs
        if not template and not relevant:
            continue
        spec = specs.get(work.pk) or work_specification(work)
        items = spec['items'] if spec['kind'] == 'COMPOSITE' else [{'id':None, 'name':'', 'unit':work.unit, 'norm':'1'}]
        first_row=worksheets['Работы'].max_row+1
        for item in items:
            quantities = []
            for month in months:
                row = allocations.get((work.pk, month))
                if template:
                    quantities.append(None)
                elif work.pk in specs:
                    quantities.append(numeric(sum((Decimal(plan['quantity']) for plan in spec.get('plans', []) if plan['date'][:7] == month.isoformat()[:7] and plan['item_id'] == item['id']), Decimal(0))))
                elif row:
                    if row.item_quantities and set(row.item_quantities)!={str(child['id']) for child in spec['items']}:
                        raise ValidationError(f'{work.name}: состав подработ изменён. Обновите импортированные объёмы или задайте объём основной работы.')
                    quantities.append(numeric(row.item_quantities[str(item['id'])] if row.item_quantities else row.quantity * Decimal(item['norm'])))
                else:
                    quantities.append(None)
            worksheets['Работы'].append([work.pk, work.name, work.section.name, work.work_group.name if work.work_group_id else '', item['id'], item['name'], item['unit'], numeric(item['norm']), *quantities])
        last_row=worksheets['Работы'].max_row
        for index,month in enumerate(months):
            quantity_column=get_column_letter(len(BASE_HEADERS['Работы'])+index+1)
            ratios=','.join(f'{quantity_column}{row}/$H{row}' for row in range(first_row,last_row+1))
            if not ratios:continue
            formula=f'=IF(COUNTA({quantity_column}{first_row}:{quantity_column}{last_row})={len(items)},ROUNDDOWN(MIN({ratios}),{3 if spec["allow_fractional"] or spec["kind"]=="SIMPLE" else 0}),"")'
            for row in range(first_row,last_row+1):
                cell=worksheets['Работы'].cell(row,len(BASE_HEADERS['Работы'])+len(months)+index+1,formula)
                cell.comment=Comment('Расчёт основной работы по минимуму нормативов. Этот столбец вычисляется, при импорте не считывается.','ГПР')
    resource_inputs = list(version.resource_allocations.select_related('brigade', 'equipment_type'))
    aliases = equipment_aliases(version.company)
    labor_aliases = brigade_aliases(version.company)
    for row in resource_inputs:
        if row.brigade_id in labor_aliases:
            row.brigade_id = labor_aliases[row.brigade_id]
        if row.equipment_type_id in aliases:
            row.equipment_type_id = aliases[row.equipment_type_id]
    for kind, sheet, catalog, fk in [('labor','Люди',Brigade,'brigade_id'),('equipment','Техника',EquipmentType,'equipment_type_id')]:
        input_map = {(getattr(row,fk), row.equipment_number if kind=='equipment' else '',row.month):row for row in resource_inputs if row.kind==kind}
        snapshot_map = defaultdict(list)
        for row in snapshot.get('resources', {}).get(kind, []):
            snapshot_map[(row[fk], row.get('equipment_number','') if kind=='equipment' else '', date.fromisoformat(row['date']).replace(day=1))].append(row)
        identities = {(key[0],key[1]) for key in input_map} | {(key[0],key[1]) for key in snapshot_map}
        if template:
            identities = {(item.pk,'') for item in catalog.objects.filter(company=version.company,is_active=True)}
        catalog_rows = {row.pk:row for row in catalog.objects.filter(company=version.company).select_related(*(['group','macro_group'] if kind=='labor' else ['category']))}
        for pk, number in sorted(identities, key=lambda key:(catalog_rows[key[0]].name,key[1])):
            item = catalog_rows[pk]
            values=[]
            for month in months:
                daily = snapshot_map.get((pk,number,month), [])
                row = input_map.get((pk,number,month))
                if template:
                    values.append(None)
                elif daily:
                    counts = {int(day['planned_workers' if kind=='labor' else 'planned_count']) for day in daily}
                    if len(counts)!=1:
                        values.append(numeric(sum(Decimal(day['planned_workers' if kind=='labor' else 'planned_count']) for day in daily) / len(daily)))
                        cell_notes.append((sheet, worksheets[sheet].max_row+1,len(BASE_HEADERS[sheet])+len(values),'Прошлый месяц по факту: среднее дневное количество по датам с зарегистрированными значениями. Этот месяц не импортируется в месячную версию.'))
                    else:
                        values.append(counts.pop())
                else:
                    values.append(row.count if row else None)
            info = [pk,item.name,item.macro_group.name if item.macro_group_id else '',item.group.name if item.group_id else '',item.unit] if kind=='labor' else [pk,item.name,item.category.name if item.category_id else '',number,item.unit]
            worksheets[sheet].append(info+values)
    fuel_names = dict(ResourceMonthAllocation._meta.get_field('fuel_type').choices)
    fuel_inputs = {(row.fuel_type,row.equipment_ref,row.month):row for row in resource_inputs if row.kind=='fuel'}
    daily_fuel = defaultdict(list)
    for row in snapshot.get('resources', {}).get('fuel', []):
        daily_fuel[(row['fuel_type'],row.get('equipment_ref',''),date.fromisoformat(row['date']).replace(day=1))].append(row)
    identities = {(k[0],k[1]) for k in fuel_inputs} | {(k[0],k[1]) for k in daily_fuel}
    if template:
        identities = {(key,'') for key in fuel_names}
    for fuel, ref in sorted(identities):
        values=[]
        for month in months:
            row=fuel_inputs.get((fuel,ref,month));daily=daily_fuel.get((fuel,ref,month),[])
            if template:
                values.extend([None,None])
            elif daily:
                balances={Decimal(str(day.get('planned_balance') or 0)) for day in daily}
                if len(balances)!=1:
                    stock=Decimal(str(max(daily,key=lambda r:r['date']).get('planned_balance') or 0))
                    cell_notes.append(('ГСМ',worksheets['ГСМ'].max_row+1,len(BASE_HEADERS['ГСМ'])+len(values)+1,'Прошлый месяц по факту: остаток на последнюю зарегистрированную дату. Этот месяц не импортируется в месячную версию.'))
                else:
                    stock=balances.pop()
                values.extend([numeric(stock),numeric(sum(Decimal(day['planned_liters']) for day in daily))])
            elif row:
                values.extend([numeric(row.balance),numeric(row.liters)])
            else:
                values.extend([None,None])
        worksheets['ГСМ'].append([json.dumps([fuel,ref],ensure_ascii=False),fuel_names[fuel],*values])
    if any(ws.max_row>MAX_ROWS+1 for ws in worksheets.values()) or sum(ws.max_row*ws.max_column for ws in worksheets.values())>500000:
        raise ValidationError('Для одного файла слишком много строк или месяцев. Используйте экспорт по отдельным месяцам.')
    for sheet,ws in worksheets.items():
        ws.auto_filter.ref=ws.dimensions
        for index in range(1,ws.max_column+1):
            letter=ws.cell(1,index).column_letter
            ws.column_dimensions[letter].width=18 if index>len(BASE_HEADERS[sheet]) else 26
        hidden={'Работы':['A','E'],'Люди':['A'],'Техника':['A'],'ГСМ':['A']}[sheet]
        for letter in hidden:
            ws.column_dimensions[letter].hidden=True
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,str) and cell.column<=len(headers(sheet,months)):
                    cell.value=ILLEGAL_CHARACTERS_RE.sub('',cell.value)
                    cell.data_type='s'
                if cell.column>len(BASE_HEADERS[sheet]):
                    cell.number_format='0.######'
    for sheet,row,column,note in cell_notes:
        worksheets[sheet].cell(row,column).comment=Comment(note,'ГПР')
    if version.planning_month:
        allowed=set(editable_months(version))
        for sheet,ws in worksheets.items():
            for index,month in enumerate(months):
                if month not in allowed:
                    for offset in range(2 if sheet=='ГСМ' else 1):
                        column=len(BASE_HEADERS[sheet])+index*(2 if sheet=='ГСМ' else 1)+offset+1
                        ws.cell(1,column).fill=PatternFill('solid',fgColor='6C757D')
                        ws.cell(1,column).comment=Comment('Месяц формируется автоматически. Импорт этот столбец не меняет.','ГПР')
    stream=BytesIO();workbook.save(stream);workbook.close()
    if stream.tell()>MAX_SIZE:raise ValidationError('Файл слишком большой. Используйте экспорт по месяцам.')
    return stream.getvalue()


def text(value):
    return str(value or '').strip()


def number(value, decimals=3, maximum=Decimal('999999999999.999')):
    try:
        result=Decimal(str(value).strip().replace(' ','').replace(',','.'))
        if not result.is_finite() or result<0 or result>maximum or result!=result.quantize(Decimal(1).scaleb(-decimals)):
            raise ValueError
        return result
    except (InvalidOperation, ValueError, TypeError):
        raise ValidationError(f'Нужно неотрицательное число, до {decimals} знаков после запятой (максимум {maximum}).')


def resolve(rows, identifier, name, title):
    if identifier not in (None,''):
        try:
            pk=int(number(identifier,0))
        except ValidationError:
            raise ValidationError(f'{title}: неверный служебный ID.')
        row=rows.get(pk)
        if row is None:
            raise ValidationError(f'{title}: позиция не принадлежит этому объекту или компании.')
        if text(name).casefold()!=row.name.strip().casefold():
            raise ValidationError(f'{title}: название не совпадает с ID. Скачайте актуальный шаблон или очистите служебный ID.')
        return row
    matches=[row for row in rows.values() if row.name.strip().casefold()==text(name).casefold()]
    if len(matches)!=1:
        raise ValidationError(f'{title}: название не найдено или неоднозначно. Используйте строку из шаблона.')
    return matches[0]


def automatic_future(version, kind, identity, month):
    if not version.planning_month or version.scenario != 'REMAINING' or month <= version.planning_month:
        return False
    if not hasattr(version, '_excel_automatic_inputs'):
        inputs=baseline_snapshot(version).get('monthly_inputs', {})
        identities={'work':set(),'labor':set(),'equipment':set(),'fuel':set()}
        for row in inputs.get('works', []):
            if date.fromisoformat(row['month'])>version.planning_month and Decimal(row['quantity'])>0:
                identities['work'].add(row['work_id'])
        for row in inputs.get('resources', []):
            field='liters' if row['kind']=='fuel' else 'hours'
            if date.fromisoformat(row['month'])>version.planning_month and Decimal(row[field])>0:
                keys={'labor':['brigade_id'],'equipment':['equipment_type_id','equipment_number'],'fuel':['fuel_type','equipment_ref']}[row['kind']]
                identities[row['kind']].add(tuple(row.get(key) for key in keys))
        version._excel_automatic_inputs=identities
    return identity in version._excel_automatic_inputs[kind]


def parse_workbook(upload, version, months):
    if upload.size>MAX_SIZE or not upload.name.lower().endswith('.xlsx'):
        raise ValidationError('Нужен файл .xlsx размером до 8 МБ.')
    try:
        with ZipFile(upload) as archive:
            if len(archive.infolist())>3000 or sum(entry.file_size for entry in archive.infolist())>80*1024*1024:
                raise ValidationError('Распакованный Excel слишком большой.')
        upload.seek(0)
        workbook=load_workbook(upload,read_only=True,data_only=False)
    except Exception:
        raise ValidationError('Не удалось прочитать файл Excel.')
    errors=[];work_values=defaultdict(dict);resource_values=[];seen=set();total_cells=0
    works={w.pk:w for w in ProjectWork.objects.filter(merged_source__isnull=True,company=version.company,section__construction_object=version.construction_object).select_related('section').prefetch_related('items')}
    brigades={r.pk:r for r in Brigade.objects.filter(company=version.company)}
    equipment={r.pk:r for r in EquipmentType.objects.filter(company=version.company)}
    fuel_names=dict(ResourceMonthAllocation._meta.get_field('fuel_type').choices)
    permitted_refs={(r.fuel_type,r.equipment_ref) for r in version.resource_allocations.filter(kind='fuel')}
    try:
        if set(workbook.sheetnames)!=set(SHEETS):
            raise ValidationError('Файл должен содержать ровно четыре листа: Работы, Люди, Техника, ГСМ.')
        for sheet in SHEETS:
            ws=workbook[sheet]
            if ws.max_row is None or ws.max_column is None:
                raise ValidationError(f'{sheet}: сохраните лист в Excel, чтобы восстановить его размеры.')
            if ws.max_row>MAX_ROWS+1 or ws.max_column>300:
                raise ValidationError(f'{sheet}: не более {MAX_ROWS} строк и 300 столбцов.')
            total_cells+=ws.max_row*ws.max_column
            if total_cells>500000:
                raise ValidationError('В файле слишком много ячеек.')
            iterator=ws.iter_rows();first=next(iterator,())
            titles=[text(cell.value) for cell in first]
            if len(set(titles))!=len(titles):
                errors.append(f'{sheet}: повторяются заголовки столбцов.');continue
            columns={title:i for i,title in enumerate(titles)}
            required=headers(sheet,months)
            # IDs may be removed; visible names then resolve within the company/object.
            missing=[name for name in required if name not in columns and name not in ['ID работы','ID подработы','ID бригады','ID техники','Служебный ключ']]
            if missing:
                errors.append(f'{sheet}: нет столбцов {", ".join(missing)}.');continue
            for index,cells in enumerate(iterator,2):
                if not any(cell.value not in (None,'') for cell in cells):continue
                try:
                    if any(cell.data_type in ('f','e') for col,cell in enumerate(cells) if col<len(titles) and titles[col] in required):
                        raise ValidationError('Замените формулы и ошибки Excel обычными значениями.')
                    def value(name):
                        col=columns.get(name)
                        return cells[col].value if col is not None and col<len(cells) else None
                    entered=[name for name in required[len(BASE_HEADERS[sheet]):] if value(name) not in (None,'')]
                    if not entered:continue
                    if sheet=='Работы':
                        candidates=works
                        if value('ID работы') in (None,'') and text(value('Раздел')):
                            candidates={pk:w for pk,w in works.items() if w.section.name.strip().casefold()==text(value('Раздел')).casefold()}
                        work=resolve(candidates,value('ID работы'),value('Работа'),'Работа')
                        if work.kind=='COMPOSITE':
                            items={item.pk:item for item in work.items.all() if item.company_id==version.company_id}
                            item=resolve(items,value('ID подработы'),value('Подработа'),'Подработа')
                            if item.quantity_per_unit<=0:raise ValidationError('Норматив подработы должен быть положительным.')
                            item_id=item.pk
                        else:
                            if value('ID подработы') not in (None,'') or text(value('Подработа')):
                                raise ValidationError('У простой работы не должно быть подработ.')
                            item_id=None
                        for month in months:
                            if automatic_future(version,'work',work.pk,month):continue
                            raw=value(label(month))
                            if raw in (None,''):continue
                            key=('work',work.pk,item_id,month)
                            if key in seen:raise ValidationError('Повторная строка работы/подработы в том же месяце.')
                            seen.add(key)
                            work_values[(work.pk,month)][item_id]=number(raw,6 if item_id else 3,Decimal('999999999999.999999') if item_id else Decimal('999999999999.999'))
                    else:
                        kind={'Люди':'labor','Техника':'equipment','ГСМ':'fuel'}[sheet]
                        if kind=='labor':
                            item=resolve(brigades,value('ID бригады'),value('Бригада'),'Бригада');identity={'brigade':item};token=(item.pk,)
                        elif kind=='equipment':
                            item=resolve(equipment,value('ID техники'),value('Вид техники'),'Техника');machine=text(value('Номер машины'))
                            if len(machine)>50:raise ValidationError('Номер машины: не более 50 символов.')
                            identity={'equipment_type':item,'equipment_number':machine};token=(item.pk,machine)
                        else:
                            fuel=next((key for key,name in fuel_names.items() if text(value('Вид ГСМ')).casefold() in [key.casefold(),str(name).casefold()]),None)
                            if not fuel:raise ValidationError('Неизвестный вид ГСМ.')
                            ref=''
                            if value('Служебный ключ') not in (None,''):
                                try:stored_fuel,ref=json.loads(str(value('Служебный ключ')))
                                except (ValueError,TypeError):raise ValidationError('Неверный служебный ключ ГСМ.')
                                if stored_fuel!=fuel or not isinstance(ref,str) or ref and (fuel,ref) not in permitted_refs:
                                    raise ValidationError('Служебный ключ ГСМ не принадлежит этому плану.')
                            identity={'fuel_type':fuel,'equipment_ref':ref};token=(fuel,ref)
                        for month in months:
                            if automatic_future(version,kind,token,month):continue
                            if kind=='fuel':
                                stock=value('Остаток '+label(month));expense=value('Расход '+label(month))
                                if stock in (None,'') and expense in (None,''):continue
                                if stock in (None,'') or expense in (None,''):raise ValidationError(f'{label(month)}: заполните и остаток, и расход.')
                                quantities={'balance':number(stock,2,Decimal('9999999999999.99')),'liters':number(expense,2,Decimal('9999999999999.99'))}
                            else:
                                raw=value(label(month))
                                if raw in (None,''):continue
                                quantities={'count':int(number(raw,0,Decimal('2147483647')))}
                            key=(kind,*token,month)
                            if key in seen:raise ValidationError('Повторная позиция ресурса в том же месяце.')
                            seen.add(key);resource_values.append((kind,month,identity,quantities))
                except ValidationError as error:
                    errors.append(f'{sheet}, строка {index}: {"; ".join(error.messages)}')
                    if len(errors)>=50:raise ValidationError(errors)
        work_rows=[]
        for (pk,month),totals in work_values.items():
            work=works[pk]
            spec=work_specification(work)
            if work.kind=='COMPOSITE' and set(totals)!={item['id'] for item in spec['items']}:
                errors.append(f'Работы, «{work.name}», {label(month)}: заполните все подработы, включая нулевые объёмы.');continue
            quantity=quantity_from_totals(spec,totals)
            work_rows.append((work,month,quantity,{str(key):str(value) for key,value in totals.items()} if work.kind=='COMPOSITE' else {}))
        if errors:raise ValidationError(errors)
        if not work_rows and not resource_values:raise ValidationError('Нет заполненных объёмов в выбранных месяцах.')
        return work_rows,resource_values
    finally:
        workbook.close()


@transaction.atomic
def import_workbook(user, version, upload, months):
    require_roles(user,PLAN_ROLES)
    locked=GlobalPlanVersion.objects.select_for_update().select_related('workspace').get(pk=version.pk,company=user.company)
    if locked.status not in ['DRAFT','REJECTED']:
        raise PermissionDenied('Импорт доступен только для черновика или отклонённой версии.')
    allowed=set(editable_months(locked))
    if not set(months)<=allowed:raise ValidationError('Некоторые месяцы формируются автоматически и недоступны для импорта.')
    works,resources=parse_workbook(upload,locked,months)
    for work,month,quantity,items in works:
        row,_=WorkMonthAllocation.objects.get_or_create(company=user.company,version=locked,work=work,month=month)
        row.quantity=quantity;row.item_quantities=items;row.save()
    for kind,month,identity,quantities in resources:
        row,_=ResourceMonthAllocation.objects.get_or_create(company=user.company,version=locked,kind=kind,month=month,**identity)
        for name,value in quantities.items():setattr(row,name,value)
        row.save()
    locked.snapshot={};locked.save(update_fields=['snapshot'])
    return len(works),len(resources)


class ExcelImportForm(forms.Form):
    scope=forms.ChoiceField(label='Загрузить',choices=[('period','Весь период'),('year','Один год'),('month','Один месяц')])
    year=forms.ChoiceField(label='Год',required=False)
    month=forms.ChoiceField(label='Месяц',required=False)
    file=forms.FileField(label='Файл Excel (.xlsx)')

    def __init__(self,*args,version,**kwargs):
        super().__init__(*args,**kwargs)
        self.fields['year'].choices=[(str(year),str(year)) for year in sorted({m.year for m in editable_months(version)})]
        self.fields['month'].choices=[(m.isoformat(),label(m)) for m in editable_months(version)]
        for field in self.fields.values():field.widget.attrs['class']='form-control'


class WorkspaceExcelExport(View):
    def get(self,request,pk):
        require_roles(request.user,READ_ROLES)
        version=version_for(request,pk)
        scope=request.GET.get('scope','period');months=selected_months(version,scope,request.GET.get('month'),request.GET.get('year'))
        template=request.GET.get('template')=='1'
        try:
            content=make_workbook(version,months,template)
        except ValidationError as error:
            return HttpResponse('; '.join(error.messages),status=400,content_type='text/plain; charset=utf-8')
        response=HttpResponse(content,content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        suffix=months[0].strftime('%Y-%m') if scope=='month' else str(months[0].year) if scope=='year' else 'period'
        response['Content-Disposition']=f'attachment; filename="plan-v{version.version_number}-{suffix}{"-template" if template else ""}.xlsx"'
        return response


class WorkspaceExcelImport(View):
    def version(self,request,pk):
        require_roles(request.user,PLAN_ROLES)
        version=version_for(request,pk)
        if version.status not in ['DRAFT','REJECTED']:
            raise PermissionDenied('Импорт доступен только для черновика или отклонённой версии.')
        return version

    def get(self,request,pk):
        version=self.version(request,pk)
        form=ExcelImportForm(version=version,initial={'scope':request.GET.get('scope','period'),'year':request.GET.get('year') or str((version.planning_month or editable_months(version)[0]).year),'month':request.GET.get('month') or (version.planning_month or editable_months(version)[0]).isoformat()})
        return self.display(request,version,form)

    def display(self,request,version,form,status=200):
        return render(request,'planning/workspace_excel_import.html',{'version':version,'form':form,'return_month':(version.planning_month or editable_months(version)[0]).isoformat()},status=status)

    def post(self,request,pk):
        version=self.version(request,pk)
        form=ExcelImportForm(request.POST,request.FILES,version=version)
        if form.is_valid():
            try:
                months=editable_months(version) if form.cleaned_data['scope']=='period' else selected_months(version,form.cleaned_data['scope'],form.cleaned_data['month'],form.cleaned_data['year'])
                if form.cleaned_data['scope']=='year':months=[m for m in months if m in editable_months(version)]
                if not months:raise ValidationError('В выбранном году нет месяцев, доступных для редактирования.')
                works,resources=import_workbook(request.user,version,form.cleaned_data['file'],months)
                messages.success(request,f'Excel загружен. Обновлено месячных строк работ: {works}, ресурсов: {resources}.')
                return redirect(reverse('planning:workspace_edit',args=[version.pk])+'?month='+months[0].isoformat())
            except ValidationError as error:
                form.add_error(None,error)
        return self.display(request,version,form,400)
