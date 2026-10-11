"""Styled plan/fact workbook with native Excel row outlines."""
from io import BytesIO
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse


def matrix_workbook(report):
    book = Workbook()
    sheet = book.active
    fact_only, plan_only = report.get('fact_only'), report.get('plan_only')
    titles = ['Факт'] if fact_only else ['План'] if plan_only else ['План', 'Факт', 'Отклонение']
    stride = len(titles)
    sheet.title = 'Журнал фактов' if fact_only else 'Журнал планов' if plan_only else 'План-факт'
    columns = report['columns']
    width = 2 + stride * (len(columns) + 1)
    sheet.sheet_properties.outlinePr.summaryBelow = False
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'C5'
    sheet.append([sheet.title])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    sheet.append(['Период: ' + str(report.get('start', '')) + ' — ' + str(report.get('end', ''))])
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    sheet.append(['Объект / раздел / показатель', 'Ед. изм.'])
    sheet.append(['', ''])
    labels = [column['label'] for column in columns]
    labels = ['За период / среднее', *labels] if report.get('total_first') else [*labels, 'За период / среднее']
    for index, label in enumerate(labels):
        first = 3 + index*stride
        sheet.cell(3, first, label)
        sheet.merge_cells(start_row=3, start_column=first, end_row=3, end_column=first+stride-1)
        for offset, title in enumerate(titles):
            sheet.cell(4, first+offset, title)
    edge = Side(style='hair', color='DCE4EE')

    def style(number, depth=0, category='data'):
        color = {'object': 'DCE9F8', 'section': 'EAF1FA', 'group': 'F3F6FA', 'total': 'EDF6F3'}.get(category, 'FFFFFF')
        for cell in sheet[number]:
            if cell.column < 3 and isinstance(cell.value, str):
                cell.data_type = "s"
            cell.fill = PatternFill('solid', fgColor=color)
            cell.font = Font(name='Calibri', size=11, color='263C55', bold=category!='data')
            cell.border = Border(bottom=edge)
            cell.alignment = Alignment(vertical='center', horizontal='left' if cell.column<3 else 'right', wrap_text=cell.column==1,
                                       indent=min(depth, 7) if cell.column==1 else 0)
        sheet.row_dimensions[number].outlineLevel = min(depth, 7)
        sheet.row_dimensions[number].height = 30

    def header(label, depth, category):
        sheet.append([label, ''])
        style(sheet.max_row, depth, category)

    def values(row, depth, kind, category='data', money=False):
        result = [row['label'], '₽' if money else row.get('unit', '')]
        cells = [row['total'], *row['cells']] if report.get('total_first') else [*row['cells'], row['total']]
        for cell in cells:
            keys = ('amount', 'fact_amount', 'delta_amount') if money else ('plan', 'fact', 'delta')
            keys = (keys[1],) if fact_only else (keys[0],) if plan_only else keys
            result.extend(cell.get(key) if cell.get(key) is not None else '—' for key in keys)
        sheet.append(result)
        number = sheet.max_row
        style(number, depth, category)
        for cell in list(sheet[number])[2:]:
            cell.number_format = '#,##0' if money or kind in ('labor', 'equipment', 'money') else '#,##0.###'
            if isinstance(cell.value, (int, float, Decimal)) and cell.value < 0:
                cell.font = Font(name='Calibri', size=11, color='B54747', bold=category!='data')

    def money_total(rows, depth, label):
        priced = [row for row in rows if row.get('has_amount')]
        if not priced:
            return
        cells = []
        for index in range(len(columns)+1):
            parts = [row['cells'][index] if index<len(columns) else row['total'] for row in priced]
            cell = {}
            for key in ('amount', 'fact_amount', 'delta_amount'):
                known = [part[key] for part in parts if part.get(key) is not None]
                cell[key] = sum(known, Decimal(0)) if known else None
            cells.append(cell)
        if not any(cell.get('amount') or cell.get('fact_amount') for cell in cells):
            return
        values({'label': label, 'cells': cells[:-1], 'total': cells[-1]}, depth, 'works', 'total', True)

    def group(node, depth, kind):
        header(node['label'], depth, 'group')
        rows = []
        for child in node.get('children', []):
            rows.extend(group(child, depth+1, kind))
        for row in node.get('rows', []):
            values(row, depth+1, kind)
            if row.get('has_amount'):
                values({**row, 'label': 'Стоимость: ' + row['label']}, depth+2, kind, money=True)
            for child in row.get('children', []):
                values(child, depth+2, kind)
            rows.append(row)
        for subtotal in node.get('subtotals', []):
            values(subtotal, depth, kind, 'total')
        if kind=='works':
            money_total(rows, depth, 'Итого: ' + node['label'])
        return rows

    for obj in report['objects']:
        header(obj['name'] if obj.get('overview_total') else str(obj['object'].project) + ' · ' + obj['name'], 0, 'object')
        for section in obj['sections']:
            header(section['label'], 1, 'section')
            rows = []
            for node in section['groups']:
                rows.extend(group(node, 2, section['kind']))
            for subtotal in section.get('subtotals', []):
                values(subtotal, 1, section['kind'], 'total')
            if section['kind']=='works':
                money_total(rows, 1, 'Итого: работы')
    for row in sheet.iter_rows(min_row=1, max_row=4, max_col=width):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='24558A')
            cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    sheet.column_dimensions['A'].width = 65
    sheet.column_dimensions['B'].width = 12
    for index in range(3, width+1):
        sheet.column_dimensions[get_column_letter(index)].width = 16
    sheet.print_title_rows = '1:4'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_format.outlineLevelRow = max((dim.outlineLevel for dim in sheet.row_dimensions.values()), default=0)
    sheet.sheet_view.showOutlineSymbols = True
    stream = BytesIO()
    book.save(stream)
    response = HttpResponse(stream.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="plan-fact.xlsx"'
    return response
