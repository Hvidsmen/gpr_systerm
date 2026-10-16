"""Additional importer for monthly contractor GPR templates (.xls / .xlsx)."""
from calendar import monthrange
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import re
from zipfile import ZipFile
from openpyxl import load_workbook
from django.core.exceptions import ValidationError

MONTHS = ['январ', 'феврал', 'март', 'апрел', 'май', 'июн', 'июл', 'август', 'сентябр', 'октябр', 'ноябр', 'декабр']
MAX_ROWS, MAX_COLUMNS, MAX_SIZE = 2500, 128, 8*1024*1024


def text(value):
    return ' '.join(str(value or '').split())


def norm(value):
    return text(value).casefold().strip(' .:')


def number(value, count=False):
    if value is None or text(value) in ('', '-', '—', '–'):
        return None
    try:
        value = Decimal(str(value).replace('\xa0', '').replace(' ', '').replace(',', '.'))
    except InvalidOperation:
        raise ValidationError('Значение не является числом.')
    if not value.is_finite() or value < 0 or value > Decimal('999999999999.999'):
        raise ValidationError('Объём должен быть конечным неотрицательным числом допустимого размера.')
    if count and (value != value.to_integral_value() or value > 2147483647):
        raise ValidationError('Численность людей и техники должна быть целым числом.')
    return value


def header(rows):
    for index, row in enumerate(rows[:40]):
        dates = [(col, value.replace(day=1).date() if isinstance(value, datetime) else value.replace(day=1)) for col, value in enumerate(row) if isinstance(value, (datetime, date))]
        names = [col for col, value in enumerate(row) if norm(value) in ('наименование работ', 'наименование')]
        if names and len(dates) >= 2:
            name = names[0]
            return {'row': index, 'name': name, 'unit': name+1, 'type': None, 'dates': dates, 'layout': 'classic'}
        if index and len([(col, value) for col, value in enumerate(row) if any(norm(value).startswith(month) for month in MONTHS)]) >= 2:
            dates, year = [], None
            above = rows[index-1]
            for col, value in enumerate(row):
                year_cell = above[col] if col < len(above) else None
                if (isinstance(year_cell,(int,float)) or text(year_cell).isdigit()) and 1900 <= float(year_cell) <= 2100:
                    year = int(float(year_cell))
                month = next((i+1 for i, prefix in enumerate(MONTHS) if norm(value).startswith(prefix)), None)
                if month and year:
                    dates.append((col, date(year, month, 1)))
            names = [col for col, value in enumerate(above) if norm(value) == 'наименование']
            if len(dates) >= 2 and names:
                name = names[0]
                kind_col = next((col for col,value in enumerate(above) if norm(value)=='вид ресурса'), None)
                return {'row': index, 'name': name, 'unit': next((col for col,value in enumerate(above) if norm(value) in ('ед. изм', 'ед.изм', 'единица измерения', 'ед. измерения')), name+2 if kind_col is not None else name+1), 'type': kind_col, 'dates': dates, 'layout': 'mixed'}
    return None


def workbook_sheets(content, filename):
    if len(content) > MAX_SIZE:
        raise ValidationError('Размер файла превышает 8 МБ.')
    if filename.lower().endswith('.xls'):
        try:
            import xlrd
        except ImportError as exc:
            raise ValidationError('Для XLS установите xlrd: python -m pip install -r requirements.txt') from exc
        try:
            book = xlrd.open_workbook(file_contents=content)
        except Exception as exc:
            raise ValidationError('Не удалось прочитать XLS.') from exc
        result = []
        for sheet in book.sheets():
            rows = []
            for i in range(min(sheet.nrows, MAX_ROWS+1)):
                row = []
                for j in range(min(sheet.ncols, MAX_COLUMNS)):
                    cell = sheet.cell(i,j)
                    value = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode) if cell.ctype==xlrd.XL_CELL_DATE else '#ERROR' if cell.ctype==xlrd.XL_CELL_ERROR else cell.value
                    row.append(value)
                rows.append(row)
            result.append((sheet.name, rows))
        return result
    if not filename.lower().endswith('.xlsx'):
        raise ValidationError('Выберите XLS или XLSX.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 64*1024*1024:
                raise ValidationError('Распакованный файл слишком большой.')
        book = load_workbook(BytesIO(content), data_only=True, read_only=True)
        result = []
        for sheet in book:
            first = list(sheet.iter_rows(max_row=min(sheet.max_row or 0, 40), max_col=min(sheet.max_column or 0, MAX_COLUMNS), values_only=True))
            if header(first):
                rows = list(sheet.iter_rows(max_row=min(sheet.max_row or 0, MAX_ROWS+1), max_col=min(sheet.max_column or 0, MAX_COLUMNS), values_only=True))
                result.append((sheet.title, rows))
            else:
                result.append((sheet.title, first))
        book.close()
        return result
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError('Не удалось прочитать XLSX. Проверьте файл.') from exc


def parse_gpr(content, filename, sheet_name='', year=None):
    sheets = workbook_sheets(content, filename)
    candidates = [(name, rows, header(rows)) for name, rows in sheets if header(rows)]
    if sheet_name:
        candidates = [candidate for candidate in candidates if norm(candidate[0])==norm(sheet_name)]
    if not candidates:
        raise ValidationError('Лист с месячным ГПР не найден. Укажите название листа; нужны месячные столбцы.')
    selected = next((candidate for candidate in candidates if norm(candidate[0])=='гпр'), candidates[0])
    sheet, rows, config = selected
    if len(rows) > MAX_ROWS:
        raise ValidationError('На выбранном листе более 2500 строк.')
    raw_dates = config['dates']
    dates = list(raw_dates)
    warnings = []
    if len(candidates)>1:
        warnings.append(f'Выбран лист «{sheet}». Другие варианты ГПР не объединяются: '+', '.join(candidate[0] for candidate in candidates if candidate[0]!=sheet))
    if year:
        january = next((index for index, (_,day) in enumerate(dates) if day.month==1), None)
        current_year = year-1 if january is not None and january>0 else year
        changed, previous_month = [], None
        for col, day in dates:
            if previous_month is not None and day.month < previous_month:
                current_year += 1
            changed.append((col, date(current_year, day.month, 1)))
            previous_month = day.month
        dates = changed
        if dates!=raw_dates:
            warnings.append('Годы в заголовках Excel заменены выбранным плановым годом. Проверьте даты каждого столбца перед сохранением.')
    if len({day for _,day in dates})!=len(dates):
        warnings.append('В заголовках повторяются месяцы. Исправьте даты столбцов в предварительном просмотре; сохранить одинаковые месяцы нельзя.')
    first_month_col = min(col for col,_ in dates)
    footer_starts = [i for i,row in enumerate(rows) if i>config['row'] and any('технические ресурсы' in norm(value) for value in row[:first_month_col])]
    footer = footer_starts[-1] if footer_starts else next((i for i,row in enumerate(rows) if i>config['row'] and any('специализированная техника' in norm(value) or re.fullmatch(r'(?:\d+(?:\.\d+)*[ .]*)?(?:гсм|горюче-смазочные материалы):?', norm(value)) for value in row[:first_month_col])), None)
    entries, section, resource_group, mode = [], sheet, '', None
    major_section, resource_col, number_col = '', config['name'], None
    ignored_work = 0
    aggregate_section = False
    for index, row in enumerate(rows[config['row']+1:], config['row']+2):
        def value(col): return row[col] if col is not None and col<len(row) else None
        prefix = row[:first_month_col]
        strings = [text(cell) for cell in prefix if isinstance(cell,str) and text(cell)]
        line = ' '.join(strings).lower()
        label = text(value(config['name']))
        unit = text(value(config['unit']))
        resource_type = norm(value(config['type']))
        if 'количество рабочих дней' in line:
            continue
        for col, cell in enumerate(prefix):
            if norm(cell).replace(' ', '') in ('гос.номер', 'госномер', 'гос.номер/инв.№'):
                number_col = col
        if 'технические ресурсы' in line:
            resource_col = next(col for col,cell in enumerate(prefix) if 'технические ресурсы' in norm(cell))
            mode, resource_group = 'equipment', 'Техника'; continue
        if 'людские ресурсы' in line or line.strip()=='люди':
            resource_col = next(col for col,cell in enumerate(prefix) if 'людские ресурсы' in norm(cell) or norm(cell)=='люди')
            mode, resource_group = 'labor', 'Люди'; continue
        if any(re.fullmatch(r'(?:\d+(?:\.\d+)*[ .]*)?(?:гсм|горюче-смазочные материалы):?', norm(cell)) for cell in prefix) and not unit:
            resource_col = next(col for col,cell in enumerate(prefix) if re.fullmatch(r'(?:\d+(?:\.\d+)*[ .]*)?(?:гсм|горюче-смазочные материалы):?', norm(cell)))
            mode, resource_group = 'fuel', 'ГСМ'; continue
        is_work = label and isinstance(value(config['unit']),str) and unit and not re.fullmatch(r'[+-]?\d+(?:[.,]\d+)?', unit) and resource_type not in ('техника','персонал','люди') and mode is None
        if is_work and aggregate_section and text(value(0)):
            section = sheet
            aggregate_section = False
        if is_work and config['layout'] == 'mixed' and text(value(0)) and index < len(rows):
            following = rows[index]
            child_unit = text(following[config['unit']]) if config['unit'] < len(following) else ''
            child_type = norm(following[config['type']]) if config['type'] is not None and config['type'] < len(following) else ''
            if not text(following[0]) and child_unit == unit and child_type not in ('техника','персонал','люди'):
                section = label
                aggregate_section = True
                warnings.append(f'Строка {index}: общая работа «{label}» используется как раздел; импортируются детальные строки ниже без повторного сложения.')
                continue
        if is_work:
            kind, group, row_name = 'work', section, label
        else:
            if footer is None or index-1 < footer:
                if label and not unit and mode is None and resource_type not in ('техника','персонал'):
                    if norm(label) not in ('техника','персонал'):
                        code = text(value(config['name']-1)) if config['name'] else ''
                        if re.fullmatch(r'\d+\.', code) or 'договор' in norm(label):
                            major_section, section = label, label
                        else:
                            section = major_section+' / '+label if major_section else label
                continue
            if mode is None and any('специализированная техника' in norm(cell) for cell in prefix):
                mode = 'equipment'
                resource_col = next(col for col,cell in enumerate(prefix) if 'специализированная техника' in norm(cell))
            kind = 'equipment' if resource_type=='техника' else 'labor' if resource_type in ('персонал','люди') else mode
            row_name = text(value(resource_col)) or (label if label else max(strings, key=len, default=''))
            if not kind or not row_name:
                continue
            if 'итого' in norm(row_name) or norm(row_name) in ('техника','персонал'):
                continue
            if row_name.endswith(':') or re.match(r'^\d+\.\d+\s', row_name) or any(marker in norm(row_name) for marker in ('специализированная техника','вспомогательная техника','итр. штаб','итр. производство','водители, машинисты','основной рабочий персонал','основной рабочий состав','обслуживание городка','персонал рмм','итр состав')):
                resource_group = re.sub(r'^\d+(?:\.\d+)*\s*', '', row_name).rstrip(': ')
                continue
            group = resource_group or ('Люди' if kind=='labor' else 'Техника')
            unit = 'чел.' if kind=='labor' else 'ед.' if kind=='equipment' else 'л'
        found = False
        for column_index, (col, month) in enumerate(dates):
            try:
                quantity = number(value(col), count=kind in ('labor','equipment'))
            except ValidationError as exc:
                if is_work and isinstance(value(col), str) and re.search(r'[А-Яа-яA-Za-z]', value(col)) and not re.match(r'^[+-]?(?:nan|inf)', value(col).strip(), re.I):
                    warnings.append(f'Строка {index}, столбец {col+1}: текстовая пометка «{text(value(col))}» пропущена; объем не вычислялся.')
                    continue
                raise ValidationError(f'Лист «{sheet}», строка {index}, столбец {col+1}: {"; ".join(exc.messages)}') from exc
            if quantity is None or quantity==0:
                continue
            precision = Decimal('.01') if kind=='fuel' else Decimal('.001')
            rounded = quantity.quantize(precision, rounding=ROUND_HALF_UP)
            if rounded != quantity:
                warning = f'Строка {index}: объёмы округлены до {str(precision).replace(".", ",")} {unit}.'
                if warning not in warnings: warnings.append(warning)
            entries.append({'kind': kind, 'name': row_name, 'section': group, 'unit': unit, 'quantity': str(rounded), 'month': month.isoformat(), 'column': column_index, 'row': index, 'equipment_number': text(value(number_col)) if kind=='equipment' and number_col is not None else ''})
            found = True
        if is_work and not found:
            ignored_work += 1
    if ignored_work:
        warnings.append(f'Работы без ненулевых месячных объёмов пропущены: {ignored_work}. Объёмы не вычисляются из цветных полос, сроков или суточного темпа.')
    if not entries:
        raise ValidationError('На выбранном листе нет ненулевых месячных объёмов или ресурсов.')
    if len(entries)>10000:
        raise ValidationError('В файле более 10000 месячных записей. Разделите ГПР.')
    combined = {}
    for row in entries:
        identity = row['kind'], norm(row['name']), norm(row['section']), norm(row['unit']), row['column'], row.get('equipment_number', '')
        if identity in combined:
            combined[identity]['quantity'] = str(Decimal(combined[identity]['quantity'])+Decimal(row['quantity']))
        else:
            combined[identity] = row
    for row in combined.values():
        value = number(row['quantity'], count=row['kind'] in ('labor','equipment'))
        if row['kind']=='fuel' and value > Decimal('9999999999.99'):
            raise ValidationError('Объём ГСМ превышает допустимый размер.')
    return {'sheet': sheet, 'filename': filename, 'months': [{'column': i, 'excel_column': col+1, 'month': day.isoformat(), 'original': raw_dates[i][1].isoformat()} for i,(col,day) in enumerate(dates)], 'entries': list(combined.values()), 'warnings': warnings, 'errors': []}
