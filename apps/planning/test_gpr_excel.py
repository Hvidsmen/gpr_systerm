from datetime import datetime, date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from uuid import uuid4
import base64
from django.test import TestCase, SimpleTestCase
from django.core.exceptions import ValidationError
from django.urls import reverse
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook
from .gpr_excel import parse_gpr
from .gpr_excel_views import apply_gpr
from .models import PlanningWorkspace, WorkMonthAllocation, ResourceMonthAllocation
from . import test_workspace as fixtures
from apps.projects.models import ConstructionObject
from apps.accounts.models import Company, User, Role


def classic(name='Монтаж', unit='м3', section='Секция', first=10):
    book=Workbook();sheet=book.active;sheet.title='ГПР'
    for row in [['№','Наименование работ','Ед. изм.',None,None,None,datetime(2026,1,1),datetime(2026,2,1)],['1.',section],['1.1',name,unit,None,None,None,first,20],[None,None,None,None,'Технические ресурсы'],[None,None,None,None,'Специализированная техника:'],[None,None,None,None,'Экскаватор',None,2,3],[None,None,None,None,'Людские ресурсы'],[None,None,None,None,'ИТР:'],[None,None,None,None,'Мастер',None,4,5]]:
        sheet.append(row)
    output=BytesIO();book.save(output);return output.getvalue()


class GprParserTests(SimpleTestCase):
    def test_classic_year_mapping_and_precision(self):
        payload=parse_gpr(classic(first=1.2345),'plan.xlsx',year=2027)
        works=[row for row in payload['entries'] if row['kind']=='work']
        self.assertEqual(works[0]['quantity'],'1.235')
        self.assertEqual(works[0]['month'],'2027-01-01')
        self.assertEqual({row['kind'] for row in payload['entries']},{'work','labor','equipment'})
        self.assertTrue(payload['warnings'])

    def test_xls_reader(self):
        content=base64.b64decode(Path(__file__).with_name('test_data').joinpath('gpr_synthetic.xls.b64').read_text())
        payload=parse_gpr(content,'plan.xls',year=2027)
        self.assertEqual(len(payload['entries']),6)
        self.assertEqual(payload['entries'][0]['quantity'],'10.000')

    def test_resource_count_is_integer_and_negative_work_rejected(self):
        with self.assertRaises(ValidationError):parse_gpr(classic(first=-1),'plan.xlsx')
        book=Workbook();sheet=book.active
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(classic()))
        book.active['G6']=1.5
        out=BytesIO();book.save(out)
        with self.assertRaises(ValidationError):parse_gpr(out.getvalue(),'plan.xlsx')

    def test_chooses_gpr_and_uses_only_final_resource_block(self):
        book=Workbook();sheet=book.active;sheet.title='ГПР'
        sheet.append([None,'№','Вид ресурса','Подобъект','Наименование',None,'Кол-во',None,None,None,None,None,2026,2027])
        sheet.append([None]*12+['декабрь','январь'])
        sheet.append([None,None,None,'Object','Монтаж',None,'м',None,None,None,None,None,10,20])
        sheet.append([None,None,'техника','Object','Экскаватор',None,5,None,None,None,None,None,5,5])
        sheet.append([None,None,None,None,'Технические ресурсы'])
        sheet.append([None,None,'техника',None,'Специализированная техника:'])
        sheet.append([None,None,'техника',None,'Экскаватор',None,2,None,None,None,None,None,2,2])
        out=BytesIO();book.save(out)
        payload=parse_gpr(out.getvalue(),'mixed.xlsx',year=2027)
        self.assertEqual(payload['months'][0]['month'],'2026-12-01')
        equipment=[row for row in payload['entries'] if row['kind']=='equipment']
        self.assertEqual([row['quantity'] for row in equipment],['2.000','2.000'])

    def test_camp_service_subtotal_is_not_imported_as_person(self):
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(classic()))
        book.active.append([None,None,None,None,'Обслуживание городка (СБО)',None,12,15])
        book.active.append([None,None,None,None,'Повар',None,3,4])
        out=BytesIO();book.save(out)
        payload=parse_gpr(out.getvalue(),'camp.xlsx',year=2027)
        self.assertFalse(any(row['name'].startswith('Обслуживание городка') for row in payload['entries']))
        self.assertEqual(sum(Decimal(row['quantity']) for row in payload['entries'] if row['kind']=='labor' and row['column']==0),Decimal(7))

    def test_missing_months_or_unknown_sheet_rejected(self):
        with self.assertRaises(ValidationError):parse_gpr(classic(),'plan.xlsx',sheet_name='Unknown')
        with self.assertRaises(ValidationError):parse_gpr(b'invalid','plan.xls')


class SuzunHeaderTests(SimpleTestCase):
    def test_unit_header_and_aggregate_work_children(self):
        w=Workbook();s=w.active;s.title='ГПР_23км без МОП'
        s.append(['№','Вид ресурса','Наименование','Ед. изм.','Проектный объем',None,None,None,None,None,None,None,None,2026,None,2027])
        s.append([None]*13+['ноябрь','декабрь','январь'])
        s.append([1,None,'Расключение','точка',20]+[None]*8+[5,10,5])
        s.append([None,None,'УЗА 1','точка',10]+[None]*8+[5,5,None])
        s.append([None,None,'УЗА 2','точка',10]+[None]*8+[None,5,5])
        s.append([2,None,'Проезд','шт',2]+[None]*8+[None,'проезд',2])
        out=BytesIO();w.save(out)
        d=parse_gpr(out.getvalue(),'test.xlsx',sheet_name=s.title)
        self.assertEqual({r['name'] for r in d['entries']},{'УЗА 1','УЗА 2','Проезд'})
        self.assertEqual(next(r for r in d['entries'] if r['name']=='Проезд')['section'],s.title)
        self.assertTrue(any('проезд' in warning for warning in d['warnings']))
        self.assertTrue(all(r['unit'] in {'точка','шт'} for r in d['entries']))


class GprImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):fixtures.WorkspaceTests.setUpTestData.__func__(cls)
    def setUp(self):self.client.force_login(self.planner)

    def preview(self):
        response=self.client.post(reverse('planning:gpr_excel_import'),{'construction_object':self.obj.pk,'year':2027,'file':SimpleUploadedFile('plan.xlsx',classic())})
        self.assertEqual(response.status_code,200)
        return response

    def test_preview_is_read_only_confirm_dates_and_replay(self):
        response=self.preview()
        self.assertEqual(PlanningWorkspace.objects.count(),0)
        token=response.context['token']
        data={'preview_token':token,'month_0':'2027-03-01','month_1':'2027-04-01'}
        response=self.client.post(reverse('planning:gpr_excel_import'),data)
        self.assertEqual(response.status_code,302)
        workspace=PlanningWorkspace.objects.get()
        self.assertEqual(workspace.start_date,date(2027,3,1))
        self.assertEqual(workspace.end_date,date(2027,4,30))
        self.assertEqual(workspace.baseline_version.status,'DRAFT')
        self.assertEqual(workspace.baseline_version.work_allocations.count(),2)
        self.assertEqual(workspace.baseline_version.resource_allocations.count(),4)
        self.assertEqual(self.client.post(reverse('planning:gpr_excel_import'),data).status_code,302)
        self.assertEqual(PlanningWorkspace.objects.count(),1)

    def test_invalid_mapping_and_signature_do_not_save(self):
        response=self.preview();token=response.context['token']
        data={'preview_token':token,'month_0':'2027-01-01','month_1':'2027-01-01'}
        response=self.client.post(reverse('planning:gpr_excel_import'),data)
        self.assertEqual(response.status_code,200)
        self.assertTrue(response.context['month_form'].errors)
        self.assertEqual(PlanningWorkspace.objects.count(),0)
        self.assertEqual(self.client.post(reverse('planning:gpr_excel_import'),{'preview_token':token+'tampered'}).status_code,400)

    def test_duplicate_source_months_can_be_corrected_in_preview(self):
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(classic()))
        book.active['H1']=book.active['G1'].value
        out=BytesIO();book.save(out)
        response=self.client.post(reverse('planning:gpr_excel_import'),{'construction_object':self.obj.pk,'file':SimpleUploadedFile('plan.xlsx',out.getvalue())})
        self.assertEqual(response.status_code,200)
        token=response.context['token']
        response=self.client.post(reverse('planning:gpr_excel_import'),{'preview_token':token,'month_0':'2027-01-01','month_1':'2027-02-01'})
        self.assertEqual(response.status_code,302)

    def test_foreign_object_and_observer_denied(self):
        company=Company.objects.create(name='Foreign import')
        from apps.projects.models import Project
        project=Project.objects.create(company=company,name='Foreign project')
        obj=ConstructionObject.objects.create(company=company,project=project,name='Foreign')
        response=self.client.post(reverse('planning:gpr_excel_import'),{'construction_object':obj.pk,'file':SimpleUploadedFile('plan.xlsx',classic())})
        self.assertEqual(response.status_code,400)
        observer=User.objects.create_user(username='gpr-import-observer',company=self.company,role=Role.objects.get(code='OBSERVER'))
        self.client.force_login(observer)
        self.assertEqual(self.client.get(reverse('planning:gpr_excel_import')).status_code,403)

    def test_same_position_in_different_groups_is_not_collapsed(self):
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(classic()))
        book.active.append([None,None,None,None,'Основной рабочий состав:'])
        book.active.append([None,None,None,None,'Мастер',None,2,1])
        out=BytesIO();book.save(out)
        payload=parse_gpr(out.getvalue(),'groups.xlsx',year=2027)
        dates={f"month_{item['column']}":date.fromisoformat(item['month']) for item in payload['months']}
        workspace=apply_gpr(self.planner,self.obj,payload,dates,str(uuid4()))
        people=workspace.baseline_version.resource_allocations.filter(kind='labor',month=date(2027,1,1))
        self.assertEqual(people.count(),2)
        self.assertEqual(sum(people.values_list('count',flat=True)),6)
        self.assertEqual(len(set(people.values_list('brigade__group_id',flat=True))),2)

    def test_explicit_fuel_block_imports_liters(self):
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(classic()))
        book.active.append([None,None,None,None,'ГСМ'])
        book.active.append([None,None,None,None,'Дизельное топливо','л',100,200])
        out=BytesIO();book.save(out)
        payload=parse_gpr(out.getvalue(),'fuel.xlsx',year=2027)
        dates={f"month_{item['column']}":date.fromisoformat(item['month']) for item in payload['months']}
        workspace=apply_gpr(self.planner,self.obj,payload,dates,str(uuid4()))
        fuel=workspace.baseline_version.resource_allocations.get(kind='fuel',month=date(2027,1,1))
        self.assertEqual(fuel.fuel_type,'DIESEL')
        self.assertEqual(fuel.liters,Decimal(100))

    def test_existing_composite_accepts_parent_monthly_quantity(self):
        payload=parse_gpr(classic(name=self.composite.name,unit=self.composite.unit,section=self.composite.section.name),'plan.xlsx',year=2027)
        dates={f"month_{item['column']}":date.fromisoformat(item['month']) for item in payload['months']}
        workspace=apply_gpr(self.planner,self.obj,payload,dates,str(uuid4()))
        self.assertEqual(workspace.baseline_version.work_allocations.filter(work=self.composite).count(),2)
        self.assertEqual(workspace.baseline_version.work_allocations.filter(work=self.composite,month=date(2027,1,1)).get().quantity,Decimal(10))

    def test_import_into_existing_version_and_update_resources(self):
        from .workspace_services import WorkspaceService
        workspace=WorkspaceService.create(self.planner,self.obj,'Existing',date(2027,1,1),date(2027,2,28))
        version=workspace.baseline_version
        url=reverse('planning:workspace_gpr_import',args=[version.pk])
        self.assertEqual(self.client.get(url).status_code,200)
        for qty in [10,25]:
            preview=self.client.post(url,{'file':SimpleUploadedFile('gpr.xlsx',classic(first=qty)),'year':2027})
            self.assertEqual(preview.status_code,200)
            self.assertTrue(preview.context['token'])
            data={'preview_token':preview.context['token'],'month_0':'2027-01-01','month_1':'2027-02-01'}
            response=self.client.post(url,data)
            self.assertEqual(response.status_code,302)
            self.assertEqual(response.url,reverse('planning:workspace_edit',args=[version.pk]))
        self.assertEqual(PlanningWorkspace.objects.count(),1)
        self.assertEqual(version.work_allocations.count(),2)
        self.assertEqual(version.work_allocations.get(month=date(2027,1,1)).quantity,25)
        self.assertEqual(version.resource_allocations.count(),4)

    def test_existing_target_outside_period_and_token_binding(self):
        from .workspace_services import WorkspaceService
        workspace=WorkspaceService.create(self.planner,self.obj,'Existing',date(2027,1,1),date(2027,1,31))
        url=reverse('planning:workspace_gpr_import',args=[workspace.baseline_version_id])
        preview=self.client.post(url,{'file':SimpleUploadedFile('gpr.xlsx',classic()),'year':2027})
        data={'preview_token':preview.context['token'],'month_0':'2027-01-01','month_1':'2027-02-01'}
        self.assertEqual(self.client.post(reverse('planning:gpr_excel_import'),data).status_code,400)
        self.assertEqual(self.client.post(url,data).status_code,400)
        self.assertEqual(workspace.baseline_version.work_allocations.count(),0)
        self.assertEqual(workspace.baseline_version.resource_allocations.count(),0)
