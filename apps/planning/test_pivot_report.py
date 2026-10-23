import json
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from datetime import date
from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook
from .pivot_engine import DEFAULT, validate_config, formula_tree, evaluate, pivot_grid
from .pivot_report import records
from .models import PivotReportPreset, DailySchedule
from . import test_decomposition as fixtures

D=Decimal

class PivotArithmeticTests(SimpleTestCase):
    def test_formula_is_decimal_safe_and_missing_values_are_unknown(self):
        tree,dim=formula_tree('fact / plan')
        self.assertEqual(dim,(0,0));self.assertEqual(evaluate(tree,{'fact':D(3),'plan':D(4)}),D('.75'))
        self.assertIsNone(evaluate(tree,{'fact':D(3),'plan':D(0)}))
        self.assertIsNone(evaluate(tree,{'fact':None,'plan':D(4)}))
        for expression in ["__import__('os').system('id')",'plan.real','plan ** 2','plan + real_plan','[plan]','1e999']:
            with self.assertRaises(ValidationError):formula_tree(expression)

    def test_axes_are_unique_and_measures_valid(self):
        config=deepcopy(DEFAULT);config['columns']=['work']
        with self.assertRaises(ValidationError):validate_config(config)
        config=deepcopy(DEFAULT);config['measures']=['unknown']
        with self.assertRaises(ValidationError):validate_config(config)

    def test_ratio_after_aggregation_units_separate_and_axes_reversible(self):
        def record(label,unit,p,f):
            return {'building':'КТП','complex_object':'УЗА','package':'Комплекс','work':label,'kind':'works','row':{'unit':unit,'operation':'sum','cells':[{'plan':D(p),'fact':D(f)}],'total':{'plan':D(p),'fact':D(f)}}}
        rows=[record('А','шт',10,5),record('Б','шт',90,90),record('В','м',100,20)]
        config={'rows':['building'],'columns':['month'],'measures':['completion'],'formulas':[]}
        report={'columns':[{'date':date(2026,1,1)}]}
        grid=pivot_grid(rows,report,config)
        self.assertEqual(sorted(r['values'][0] for r in grid['rows']),[D(20),D(95)])
        config['rows'],config['columns']=config['columns'],config['rows']
        self.assertEqual(sorted(r['values'][0] for r in pivot_grid(rows,report,config)['rows']),[D(20),D(95)])


class PivotReportTests(TestCase):
    setUpTestData=classmethod(fixtures.DecompositionTests.setUpTestData.__func__)
    setUp=fixtures.DecompositionTests.setUp

    def url(self,kind='plan_fact',pk=None):
        return reverse({'plan_fact':'planning:pivot_report_list','facts':'planning:pivot_fact_list','gpr':'planning:pivot_gpr_detail','smg':'planning:pivot_smg_detail'}[kind],args=[pk] if pk else [])

    def test_each_version_report_stays_on_its_source_and_exports_excel(self):
        url=self.url('gpr',self.version.pk)
        config={'rows':['work'],'columns':['month'],'measures':['real_plan','estimated_plan'],'formulas':[]}
        response=self.client.get(url,{'start':'2026-01-01','end':'2026-01-31','config':json.dumps(config)})
        self.assertEqual(response.status_code,200)
        self.assertEqual(sum(row['values'][0] for row in response.context['grid']['rows'] if row['label']=='Реальная · план, ₽'),D('1050'))
        response=self.client.get(url,{'start':'2026-01-01','end':'2026-01-31','config':json.dumps(config),'export':'xlsx'})
        self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(response.content));self.assertEqual(book.active.title,'Сводная таблица')
        self.assertEqual(sum(cell.value for row in book.active.iter_rows(min_row=4) for cell in row[-1:] if isinstance(cell.value,(int,float))),2100)

    def test_personal_default_is_per_section_and_cannot_be_loaded_by_another_user(self):
        payload={'action':'save','name':'Мой вариант','config':json.dumps(DEFAULT),'make_default':'1'}
        self.assertEqual(self.client.post(self.url(),payload).status_code,302)
        self.client.post(self.url(),{**payload,'name':'Другой'})
        self.client.post(self.url('facts'),payload)
        self.assertEqual(PivotReportPreset.objects.filter(user=self.planner,is_default=True).count(),2)
        default=PivotReportPreset.objects.get(user=self.planner,section='plan_fact',is_default=True)
        response=self.client.get(self.url());self.assertEqual(response.context['selected_preset'],default)
        other=type(self.planner).objects.create_user(username='pivot-other',password='test-pass',company=self.company,role=self.planner.role)
        self.client.force_login(other)
        self.assertEqual(self.client.get(self.url(),{'preset':default.pk}).status_code,404)
        self.assertEqual(self.client.post(self.url(),{'action':'delete','preset':default.pk}).status_code,404)
        self.assertTrue(PivotReportPreset.objects.filter(pk=default.pk).exists())
        self.assertEqual(self.client.get(self.url('gpr',self.version.pk)).status_code,404)

    def test_fact_source_and_default_formula_validation(self):
        response=self.client.get(self.url('facts'),{'start':'2026-01-01','end':'2026-01-31'})
        self.assertEqual(response.status_code,200);self.assertContains(response,'Сводная таблица')
        config=deepcopy(DEFAULT);config['formulas']=[{'name':'Выполнение','expression':'fact / plan','percent':True}];config['measures']=['custom_0']
        self.client.post(self.url(),{'action':'save','name':'Выполнение','config':json.dumps(config)})
        self.assertTrue(PivotReportPreset.objects.filter(name='Выполнение').exists())
        config['formulas'][0]['expression']='__import__("os")'
        self.client.post(self.url(),{'action':'save','name':'Плохая','config':json.dumps(config)})
        self.assertFalse(PivotReportPreset.objects.filter(name='Плохая').exists())

    def test_smg_source_uses_schedule_even_without_gpr(self):
        from .independent_smg import create
        from .daily_schedule_table import add_entry
        from .daily_schedule import save_entry, identities
        schedule=create(self.planner,self.obj,date(2026,1,1))
        entry=add_entry(self.planner,schedule,{'kind':'work','work':self.simple})
        save_entry(self.planner,entry,[{'quantity':7 if day.day==5 else 0} for day,item in identities(entry)],entry.revision)
        config={'rows':['work'],'columns':['date'],'measures':['plan'],'formulas':[]}
        response=self.client.get(self.url('smg',schedule.pk),{'config':json.dumps(config)})
        self.assertEqual(response.status_code,200)
        self.assertEqual(sum(v or 0 for row in response.context['grid']['rows'] for v in row['values']),D(7))
        self.assertIsNone(response.context['owner'].source_id)
        self.assertEqual(len(response.context['grid']['columns']),31)
