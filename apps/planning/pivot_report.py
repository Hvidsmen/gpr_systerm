"""Scoped pivot report with personal, per-section saved configurations."""
import json
from calendar import monthrange
from datetime import date
from io import BytesIO
from django.contrib import messages
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views import View
from core.permissions import READ_ROLES, require_roles, scope_queryset
from .models import GlobalPlanVersion, DailySchedule, PivotReportPreset
from .report_forms import MatrixReportForm
from .report_services import build_matrix, Source
from .pivot_engine import FIELDS, BASE, DEFAULT, validate_config
from .pivot_layout import build_layout


def read_config(raw):
    try:return validate_config(json.loads(raw))
    except (ValueError,TypeError,RecursionError):raise ValidationError('Не удалось прочитать настройки сводной таблицы.')


class PivotReport(View):
    source_kind='plan_fact'
    def owner(self,request,pk):
        if self.source_kind not in {'gpr','smg'}:return None
        model=DailySchedule if self.source_kind=='smg' else GlobalPlanVersion
        return get_object_or_404(scope_queryset(model.objects.all(),request.user),pk=pk)

    def presets(self,user):
        return PivotReportPreset.objects.filter(company=user.company,user=user,section=self.source_kind)

    def post(self,request,pk=None):
        require_roles(request.user,READ_ROLES);self.owner(request,pk)
        query=self.presets(request.user);action=request.POST.get('action')
        try:
            if action not in {'save','default','delete'}:raise ValidationError('Неизвестное действие.')
            with transaction.atomic():
                # Serialize default changes, including concurrent first saves.
                get_user_model().objects.select_for_update().get(pk=request.user.pk)
                preset=get_object_or_404(query,pk=request.POST.get('preset')) if request.POST.get('preset','').isdigit() else None
                if action in {'default','delete'} and preset is None:raise ValidationError('Выберите сохранённую настройку.')
                if action=='delete':preset.delete()
                elif action=='default':
                    query.update(is_default=False);preset.is_default=True;preset.save(update_fields=['is_default','updated_at'])
                else:
                    config=read_config(request.POST.get('config',''))
                    name=request.POST.get('name','').strip()
                    if not name or len(name)>120:raise ValidationError('Укажите название настройки до 120 символов.')
                    conflict=query.filter(name=name)
                    if preset:conflict=conflict.exclude(pk=preset.pk)
                    if conflict.exists():raise ValidationError('Настройка с таким названием уже существует.')
                    if preset is None:preset=PivotReportPreset(company=request.user.company,user=request.user,section=self.source_kind)
                    preset.name=name;preset.config=config
                    if request.POST.get('make_default')=='1':query.update(is_default=False);preset.is_default=True
                    preset.save()
            messages.success(request,'Личная настройка отчёта сохранена.' if action!='delete' else 'Настройка удалена.')
        except ValidationError as exc:messages.error(request,' '.join(exc.messages))
        if action in {'save','default'} and 'preset' in locals() and preset and preset.pk:
            from urllib.parse import urlencode
            selection={k:request.POST[k] for k in ('plan_type','plan_version','start','end','project') if request.POST.get(k)}
            selection['preset']=preset.pk
            return redirect(request.path+'?'+urlencode(selection))
        from urllib.parse import urlencode
        selection={k:request.POST[k] for k in ('plan_type','plan_version','start','end','project') if request.POST.get(k)}
        return redirect(request.path+('?' + urlencode(selection) if selection else ''))

    def get(self,request,pk=None):
        require_roles(request.user,READ_ROLES);owner=self.owner(request,pk)
        params=request.GET.copy();today=timezone.localdate()
        start=today.replace(day=1);end=start.replace(day=monthrange(start.year,start.month)[1])
        if owner:
            if self.source_kind=='smg':start=owner.month;end=start.replace(day=monthrange(start.year,start.month)[1])
            else:start,end=owner.start_date,owner.end_date
            params.setlist('objects',[str(owner.construction_object_id)])
            params.pop('consolidated',None);params.pop('project',None)
            for key in list(params):
                if key.startswith('version_'):params.pop(key)
        for key,value in {'start':start.isoformat(),'end':end.isoformat(),'mode':'latest'}.items():params.setdefault(key,value)
        if self.source_kind=='plan_fact' and request.GET.get('plan_version','').isdigit():
            kind=request.GET.get('plan_type','gpr')
            model={'gpr':GlobalPlanVersion,'smg':DailySchedule}.get(kind)
            selected=scope_queryset(model.objects.all(),request.user).filter(pk=request.GET['plan_version']).first() if model else None
            if selected:
                a=selected.start_date if kind=='gpr' else selected.month
                b=selected.end_date if kind=='gpr' else a.replace(day=monthrange(a.year,a.month)[1])
                if 'start' not in request.GET:params['start']=str(a)
                if 'end' not in request.GET:params['end']=str(b)
        if 'sections_set' not in params:params.setlist('sections',['works'])
        if self.source_kind=='plan_fact':
            from .pivot_sources import PivotSourceForm
            params.setdefault('plan_type','gpr')
            form=PivotSourceForm(params,user=request.user)
        else:form=MatrixReportForm(params,user=request.user)
        presets=self.presets(request.user);selected=None;config={**DEFAULT,'measures':['fact','real_fact','estimated_fact']} if self.source_kind=='facts' else {**DEFAULT,'measures':['plan','real_plan','estimated_plan']} if owner else DEFAULT
        if request.GET.get('preset','').isdigit():selected=get_object_or_404(presets,pk=request.GET['preset'])
        elif 'config' not in request.GET:selected=presets.filter(is_default=True).first()
        try:
            config=read_config(request.GET['config']) if 'config' in request.GET else validate_config(selected.config if selected else config)
        except ValidationError as exc:form.add_error(None,exc)
        if self.source_kind=='plan_fact' and params.get('plan_type')=='object':
            config={**config,'rows':[x for x in config['rows'] if x not in {'month','date'}],'columns':[x for x in config['columns'] if x not in {'month','date'}]}
            if not config['rows']:config['rows']=['building','work']
            config['totals']={**config['totals'],'row_levels':[x for x in config['totals']['row_levels'] if x in config['rows'][:-1]],'column_levels':[x for x in config['totals']['column_levels'] if x in config['columns'][:-1]]}
        if self.source_kind=='plan_fact' and not selected and 'config' not in request.GET:
            config={**config,'rows':['building','section','complex_object','package','work'],'totals':{**config['totals'],'row_levels':['building','section','complex_object','package']}}
        context={'form':form,'owner':owner,'source_kind':self.source_kind,'presets':presets,'selected_preset':selected,
                 'pivot_config':config,'pivot_fields':FIELDS,'pivot_measures':BASE}
        context['pivot_source_periods']=getattr(form,'version_periods',{})
        # is_valid must run before constructing the scoped sources.
        if form.is_valid():
            filters={**form.cleaned_data,'daily':'date' in config['rows']+config['columns'],'flat_objects':True,'work_levels':'','month':None}
            if self.source_kind=='plan_fact' and filters.get('selected_plan'):
                filters['objects']=form.fields['objects'].queryset.filter(pk=filters['selected_plan'].construction_object_id)
            if owner and (filters['start']<start or filters['end']>end):form.add_error(None,'Период должен находиться внутри выбранной версии плана.')
            else:
                try:
                    overrides=None
                    if owner:
                        filters['objects']=form.fields['objects'].queryset.filter(pk=owner.construction_object_id)
                        if self.source_kind=='gpr':
                            from .global_services import build_snapshot
                            snapshot=build_snapshot(owner,validate_inputs=False) if owner.status in ['DRAFT','REJECTED'] else owner.snapshot
                        else:
                            from .daily_schedule import overlay_schedules
                            snapshot=overlay_schedules({'works':[],'resources':{}},[owner],start,end)
                        overrides={owner.construction_object_id:[Source(owner,snapshot,start,end)]}
                    source_note=''
                    if self.source_kind=='plan_fact':
                        from .pivot_sources import sources
                        overrides=sources(request.user,filters)
                        selected=filters.get('selected_plan')
                        source_note=(dict(form.fields['plan_type'].choices)[filters['plan_type']]+' · '+str(selected)+(' · '+selected.get_status_display() if hasattr(selected,'get_status_display') else '')) if selected else dict(form.fields['plan_type'].choices)[filters['plan_type']]+' · актуальные версии по объектам'
                    if self.source_kind=='facts':
                        from .journal_services import journal_matrix
                        report=journal_matrix(request.user,filters,'facts')
                    else:report=build_matrix(request.user,filters,source_overrides=overrides,plan_only=bool(owner))
                    data=records(report)
                    if self.source_kind=='plan_fact' and filters.get('plan_type')=='object':
                        from .pivot_sources import object_records
                        data=object_records(request.user,filters,report)
                        source_note+=' · весь объём плана; факт за выбранный период. План по месяцам не распределён.'
                    grid=build_layout(data,report,config)
                    if self.source_kind=='plan_fact' and filters.get('plan_type')=='object':
                        for obj in report['objects']:obj['warnings']=[w for w in obj.get('warnings',[]) if 'Согласованный план' not in w]
                    context['source_note']=source_note
                    context['source_count']=len(data)
                    from .pivot_sources import friendly_period
                    for header in grid['headers']:
                        for cell in header:cell['label']=friendly_period(cell['label'])
                    for row in grid['rows']:
                        if row.get('field') in {'Месяц','День'}:row['label']=friendly_period(row.get('label',''))
                        row['leaf']=row['kind']=='group' and row.get('depth')==len(config['rows'])-1
                        row['resource_label']=row.get('indicator','').split(' · ')[0]
                    context.update(grid=grid,warnings=[w for obj in report['objects'] for w in obj.get('warnings',[])])
                    if request.GET.get('export')=='xlsx':return pivot_workbook(grid,filters,owner)
                except ValidationError as exc:form.add_error(None,exc)
        return render(request,'planning/pivot_report.html',context)

def records(report):
    """Only leaves contribute; priced package totals are never added twice."""
    result=[]
    def walk(nodes, obj, kind, section, work_section=None):
        for node in nodes:
            for parent in node['rows']:
                dims=parent.get('decomposition', {})
                for row in parent.get('children') or [parent]:
                    composite=dims.get('kind')=='COMPOSITE'
                    result.append({'section':work_section or node['label'] if kind=='works' else section,'building':str(obj['object'].project)+' · '+obj['name'],
                        'complex_object':dims.get('complex_object_name') or ('Без комплексного объекта' if kind=='works' else section),
                        'package':dims.get('package_name') if composite else ('Самостоятельные работы' if kind=='works' else section),
                        'work':row['label'], 'kind':kind, 'row':row})
            walk(node.get('children',[]),obj,kind,section,work_section or node['label'])
    for obj in report['objects']:
        for section in obj['sections']:walk(section['groups'],obj,section['kind'],section['label'])
    return result



def pivot_workbook(grid,filters,owner=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    book=Workbook();sheet=book.active;sheet.title='Сводная таблица'
    sheet.append(['Сводная таблица',str(filters['start'])+' — '+str(filters['end'])])
    if owner:sheet.append(['Версия',str(owner)])
    hierarchy=grid['view']=='hierarchy'
    labels=['Объект / работа / ресурс','Показатель','Ед.','Вид ресурса'] if hierarchy else grid['row_labels']+['Показатель','Ед.','Тип объёма']
    left=len(labels);header=sheet.max_row+1;depth=grid['header_depth']
    for index,label in enumerate(labels,1):
        sheet.cell(header,index,label)
        if depth>1:sheet.merge_cells(start_row=header,end_row=header+depth-1,start_column=index,end_column=index)
    occupied=set()
    for level,cells in enumerate(grid['headers']):
        column=left+1
        for cell in cells:
            while (level,column) in occupied:column+=1
            sheet.cell(header+level,column,cell['label'])
            if cell['colspan']>1 or cell['rowspan']>1:sheet.merge_cells(start_row=header+level,end_row=header+level+cell['rowspan']-1,start_column=column,end_column=column+cell['colspan']-1)
            for r in range(level,level+cell['rowspan']):
                for c in range(column,column+cell['colspan']):occupied.add((r,c))
            column+=cell['colspan']
    for row in grid['rows']:
        kind=row['kind'];label=row['label']
        if hierarchy:
            if kind=='subtotal':label='Итого: '+row['levels'][-1]
            elif kind=='grand':label='Общий итог'
            elif kind=='detail':label=row['levels'][-1] if row.get('levels') else label
            values=[label,row.get('measure_label') or row['label'],row.get('unit',''),row.get('resource_label',row.get('indicator',''))]
        else:values=list(row.get('levels',[]))+[label,row.get('unit',''),row.get('indicator','')]
        sheet.append(values+row['values']);number=sheet.max_row
        sheet.row_dimensions[number].outlineLevel=min(row['depth'] if kind in {'detail','group'} else max(0,row['depth']-1),7) if hierarchy else 0
        summary=kind!='detail' and not row.get('leaf')
        for cell in sheet[number]:
            cell.font=Font(name='Calibri',size=11,color='273B53',bold=summary)
            if summary:cell.fill=PatternFill('solid',fgColor='E2EFF8' if kind=='grand' else 'F2F7FB')
            if cell.column>left:cell.number_format='#,##0.00' if row.get('unit')=='₽' else '#,##0.0'+chr(34)+'%'+chr(34) if row.get('unit')=='%' else '#,##0.###'
        if hierarchy:sheet.cell(number,1).alignment=Alignment(indent=min(row['depth'],7),wrap_text=True)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value,str):cell.data_type='s'
            if header<=cell.row<header+depth:cell.font=Font(name='Calibri',size=11,bold=True,color='355C7C');cell.fill=PatternFill('solid',fgColor='EDF5FC')
    for i in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(i)].width=50 if i==1 and hierarchy else 32 if i<=left-2 else 18
    sheet.freeze_panes=get_column_letter(left+1)+str(header+depth);sheet.print_title_rows=f'1:{header+depth-1}';sheet.sheet_view.showGridLines=False
    sheet.sheet_properties.outlinePr.summaryBelow=False
    stream=BytesIO();book.save(stream)
    response=HttpResponse(stream.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet');response['Content-Disposition']='attachment; filename="pivot-report.xlsx"';return response
