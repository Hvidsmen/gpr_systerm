from io import BytesIO
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from apps.accounts.models import Company, User, Role
from .models import Brigade, BrigadeGroup, BrigadeMacroGroup, EquipmentType, EquipmentCategory


class CatalogExportTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name='Excel export')
        self.user = User.objects.create_user(username='export-planner', company=self.company, role=Role.objects.get(code='PLANNER'))
        self.client.force_login(self.user)
        self.group = BrigadeGroup.objects.create(company=self.company,name='Group')
        self.macro = BrigadeMacroGroup.objects.create(company=self.company,name='Macro')
        self.category = EquipmentCategory.objects.create(company=self.company,name='Category')
        for i in range(25):
            Brigade.objects.create(staffing_current=True,company=self.company,name='Монтаж '+str(i),group=self.group,macro_group=self.macro,description='=1+1')
            EquipmentType.objects.create(company=self.company,name='Кран '+str(i),category=self.category)
        Brigade.objects.create(staffing_current=True,company=self.company,name='Other')
        EquipmentType.objects.create(company=self.company,name='Other')
        other = Company.objects.create(name='Foreign export')
        Brigade.objects.create(staffing_current=True,company=other,name='Монтаж Foreign')
        EquipmentType.objects.create(company=other,name='Кран Foreign')

    def test_all_filtered_rows_exported_not_just_current_page(self):
        for route, params in [('brigade_export',{'name':'монтаж','group':self.group.pk,'macro_group':self.macro.pk,'page':2}),('equipment_type_export',{'name':'кран','category':self.category.pk,'page':2})]:
            response = self.client.get(reverse('resources:'+route),params)
            self.assertEqual(response.status_code,200)
            workbook = load_workbook(BytesIO(response.content))
            sheet = workbook.active
            self.assertEqual(sheet.max_row,26)
            rows = list(sheet.values)
            header = rows[0]
            self.assertTrue(all('Foreign' not in str(row) and 'Other' not in str(row) for row in rows[1:]))
            self.assertEqual(rows[1][header.index('Единица измерения')], 'чел.' if route=='brigade_export' else 'ед.')
            if route=='brigade_export':
                cell = sheet.cell(2,header.index('Описание')+1)
                self.assertEqual(cell.value,'=1+1')
                self.assertEqual(cell.data_type,'s')
            workbook.close()

    def test_empty_invalid_filters_and_manager_access(self):
        response = self.client.get(reverse('resources:brigade_export'),{'name':'No matches'})
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.active.max_row,1)
        workbook.close()
        self.assertEqual(self.client.get(reverse('resources:brigade_export'),{'group':'bad'}).status_code,400)
        manager = User.objects.create_user(username='export-manager',company=self.company,role=Role.objects.get(code='MANAGER'))
        self.client.force_login(manager)
        self.assertEqual(self.client.get(reverse('resources:equipment_type_export')).status_code,200)
        foreman = User.objects.create_user(username='export-foreman',company=self.company,role=Role.objects.get(code='FOREMAN'))
        self.client.force_login(foreman)
        for route in ('brigade_export','equipment_type_export'):
            self.assertEqual(self.client.get(reverse('resources:'+route)).status_code,403)

    def test_planning_type_column_uses_root_name_for_root_and_parent_name_for_model(self):
        root=EquipmentType.objects.create(company=self.company,name='АТЗ',category=self.category)
        EquipmentType.objects.create(company=self.company,name='АТЗ / Камаз 4671М2-11',parent=root,brand_model='Камаз 4671М2-11')
        no_category=EquipmentType.objects.create(company=self.company,name='АТЗ без группы')
        response=self.client.get(reverse('resources:equipment_type_export'),{'name':'АТЗ'})
        self.assertEqual(response.status_code,200)
        workbook=load_workbook(BytesIO(response.content));sheet=workbook.active
        headers=[cell.value for cell in sheet[1]];column=headers.index('Техника для планирования')
        rows={row[0]:row for row in list(sheet.values)[1:]}
        self.assertEqual(rows['АТЗ'][column],'АТЗ')
        self.assertEqual(rows['АТЗ / Камаз 4671М2-11'][column],'АТЗ')
        self.assertEqual(rows['АТЗ без группы'][column],no_category.name)
        self.assertNotEqual(rows['АТЗ'][column],self.category.name)
        workbook.close();root.refresh_from_db();self.assertIsNone(root.parent_id)
