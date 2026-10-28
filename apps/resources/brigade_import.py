"""Excel changes only groups of existing company-scoped positions."""
from io import BytesIO
from zipfile import ZipFile
from django import forms
from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import render
from django.views import View
from openpyxl import Workbook,load_workbook
from core.permissions import require_roles,PLAN_ROLES
from .models import Brigade,BrigadeGroup

HEADERS=['ID','Название','Группа']
MAX_ROWS=5000


class BrigadeImportForm(forms.Form):
    file=forms.FileField(label='Excel с группами должностей (.xlsx)',widget=forms.FileInput(attrs={'class':'form-control','accept':'.xlsx'}))
    def clean_file(self):
        file=self.cleaned_data['file']
        if not file.name.lower().endswith('.xlsx'):raise forms.ValidationError('Выберите файл .xlsx.')
        if file.size>5*1024*1024:raise forms.ValidationError('Максимальный размер файла — 5 МБ.')
        return file


def text(value):return '' if value is None else str(value).strip()


def read_rows(file):
    try:
        with ZipFile(file) as archive:
            if sum(info.file_size for info in archive.infolist())>50*1024*1024:return [],['Файл слишком большой после распаковки.']
        file.seek(0);book=load_workbook(file,read_only=True,data_only=False)
    except Exception:return [],['Не удалось прочитать исправный файл .xlsx без пароля.']
    rows=[];errors=[]
    try:
        sheet=book.worksheets[0];sheet.reset_dimensions();iterator=sheet.iter_rows();header=next(iterator,())
        columns={text(c.value).casefold():i for i,c in enumerate(header)}
        if not {'название','группа'}<=columns.keys():return [],['Нужны заголовки «Название» и «Группа». Используйте экспорт справочника или шаблон.']
        recognized=[text(c.value).casefold() for c in header if text(c.value).casefold() in {'id','название','группа'}]
        if len(recognized)!=len(set(recognized)):return [],['Заголовки ID, «Название» и «Группа» не должны повторяться.']
        for number,cells in enumerate(iterator,2):
            if number>MAX_ROWS+1:errors.append(f'Допускается не более {MAX_ROWS} строк.');break
            values={}
            for title in HEADERS:
                index=columns.get(title.casefold());cell=cells[index] if index is not None and index<len(cells) else None
                if cell is not None and cell.data_type in {'f','e'}:errors.append(f'Строка {number}, «{title}»: замените формулу или ошибку значением.')
                values[title]=text(cell.value) if cell else ''
            if not any(values.values()):continue
            if not values['Название'] or len(values['Название'])>150:errors.append(f'Строка {number}: укажите название до 150 символов.')
            if len(values['Группа'])>150:errors.append(f'Строка {number}: группа — не более 150 символов.')
            if values['ID'] and (not values['ID'].isdigit() or not 0<int(values['ID'])<=9223372036854775807):errors.append(f'Строка {number}: некорректный ID.')
            rows.append({'line':number,'id':values['ID'],'name':values['Название'],'group':values['Группа']})
        if not rows and not errors:errors.append('В файле нет данных для загрузки.')
        return rows,errors
    finally:book.close()


@transaction.atomic
def import_rows(company,rows,user=None):
    from apps.accounts.models import Company
    from apps.audit.models import AuditLog
    Company.objects.select_for_update().get(pk=company.pk)
    locked_groups=list(BrigadeGroup.objects.select_for_update().filter(company=company))
    existing=list(Brigade.objects.select_for_update().filter(company=company,merge_source__isnull=True));by_id={b.pk:b for b in existing};by_name={}
    for b in existing:by_name.setdefault(b.name.strip().casefold(),[]).append(b)
    selected={};unknown=[]
    for row in rows:
        matches=[by_id[int(row['id'])]] if row['id'] and int(row['id']) in by_id else [] if row['id'] else by_name.get(row['name'].casefold(),[])
        if not matches:unknown.append(row['name']);continue
        if len(matches)>1:raise ValidationError(f"Строка {row['line']}: несколько должностей «{row['name']}». Используйте экспорт с ID.")
        item=matches[0]
        if item.name.strip().casefold()!=row['name'].casefold():raise ValidationError(f"Строка {row['line']}: название не совпадает с ID. Переименование через Excel недоступно.")
        if item.pk in selected and selected[item.pk][1].casefold()!=row['group'].casefold():raise ValidationError(f"Строка {row['line']}: для одной должности указаны разные группы.")
        selected[item.pk]=(item,row['group'])
    groups={g.name.strip().casefold():g for g in locked_groups}
    result={'updated':0,'unchanged':0,'unknown':len(unknown),'unknown_names':unknown[:100],'warnings':['Должность не найдена: '+name for name in unknown[:100]],'groups':0}
    for item,label in selected.values():
        group=groups.get(label.casefold()) if label else None
        if label and group is None:
            result['warnings'].append(f'Группа «{label}» не найдена: должность «{item.name}» не изменена.');continue
        target=group.pk if group else None
        if item.group_id==target:result['unchanged']+=1;continue
        before=item.group_id;item.group=group;item.save(update_fields=['group','updated_at']);result['updated']+=1
        AuditLog.objects.create(user=user,action='UPDATE',entity_type='resources.Brigade',entity_id=str(item.pk),old_value={'group_id':before},new_value={'group_id':target},comment='Изменение группы должности из Excel')
    return result


class BrigadeImport(View):
    def get(self,request):
        require_roles(request.user,PLAN_ROLES)
        return render(request,'resources/brigade_import.html',{'form':BrigadeImportForm()})
    def post(self,request):
        require_roles(request.user,PLAN_ROLES);form=BrigadeImportForm(request.POST,request.FILES);errors=[];result=None
        if form.is_valid():
            rows,errors=read_rows(form.cleaned_data['file'])
            if not errors:
                try:result=import_rows(request.user.company,rows,request.user)
                except ValidationError as exc:errors=exc.messages
        return render(request,'resources/brigade_import.html',{'form':form,'errors':errors[:100],'error_count':len(errors),'result':result},status=400 if errors or form.errors else 200)


def brigade_import_template(request):
    require_roles(request.user,PLAN_ROLES)
    book=Workbook();sheet=book.active;sheet.title='Группы должностей';sheet.append(HEADERS);sheet.freeze_panes='A2'
    for column,width in [('A',15),('B',65),('C',40)]:sheet.column_dimensions[column].width=width
    guide=book.create_sheet('Инструкция')
    for line in ['Импорт меняет только группу существующей должности. Новые должности не добавляются.',
        'Можно загрузить экспорт справочника, изменив только колонку «Группа». Остальные колонки игнорируются.',
        'ID сохраняйте из экспорта. Без ID используется единственное совпадение названия в вашей компании.',
        'Пустая группа снимает назначение. Группы должны существовать в справочнике; неизвестные группы пропускаются с предупреждением.',
        'Неизвестные должности пропускаются и перечисляются в результате. Ошибка в строке отменяет весь импорт.',
        'Не более 5000 строк; размер файла до 5 МБ.']:
        guide.append([line])
    guide.column_dimensions['A'].width=110;output=BytesIO();book.save(output);book.close()
    response=HttpResponse(output.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet');response['Content-Disposition']='attachment; filename="position_groups.xlsx"';return response
