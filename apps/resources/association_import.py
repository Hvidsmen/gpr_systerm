"""Preview and atomically append Excel associations to existing company resources."""
import hashlib,json
from io import BytesIO
from zipfile import BadZipFile,ZipFile
from django import forms
from django.contrib import messages
from django.core import signing
from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import Http404
from django.shortcuts import render,redirect
from django.views.decorators.http import require_http_methods
from openpyxl import load_workbook
from core.permissions import require_roles,PERSONNEL_CATALOG_EDIT_ROLES,EQUIPMENT_CATALOG_EDIT_ROLES
from core.excel_mapping import name_key,validate_excel_aliases
from core.catalogue_excel import xlsx_response
from apps.audit.models import AuditLog
from .models import Brigade,EquipmentType

SALT='resources.excel-associations.v1'
MAX_ROWS=10000


class UploadForm(forms.Form):
    file=forms.FileField(label='Файл Excel',widget=forms.ClearableFileInput(attrs={'accept':'.xlsx','class':'form-control'}))


def catalogue(user,kind,lock=False):
    model=Brigade if kind=='labor' else EquipmentType
    query=model.objects.filter(company=user.company,merge_source__isnull=True)
    if kind=='equipment':query=query.filter(parent__isnull=True)
    return query.select_for_update() if lock else query


def stamp(row):
    return hashlib.sha256(json.dumps([row.pk,row.name,row.excel_aliases,str(row.updated_at)],ensure_ascii=False).encode()).hexdigest()


def parse(upload,query):
    if upload.size>5*1024*1024:raise ValidationError('Размер файла не должен превышать 5 МБ.')
    if not upload.name.lower().endswith('.xlsx'):raise ValidationError('Загрузите файл .xlsx.')
    raw=upload.read()
    try:
        with ZipFile(BytesIO(raw)) as archive:
            if sum(info.file_size for info in archive.infolist())>50*1024*1024:raise ValidationError('Слишком большой объём данных внутри Excel.')
        book=load_workbook(BytesIO(raw),read_only=True,data_only=False)
    except (BadZipFile,ValueError,KeyError,OSError) as exc:raise ValidationError('Не удалось прочитать Excel. Проверьте файл .xlsx.') from exc
    try:
        sheet=book.active;iterator=sheet.iter_rows();header=next(iterator,None)
        if not header:raise ValidationError('Файл пуст.')
        headers={name_key(cell.value):index for index,cell in enumerate(header) if cell.value is not None}
        alias_col=next((headers[key] for key in ['ассоциация','ассоциация из excel','название в excel'] if key in headers),None)
        name_col=headers.get('название');id_col=headers.get('id')
        if alias_col is None or name_col is None and id_col is None:
            raise ValidationError('Нужны колонки «Название» (или «ID») и «Ассоциация». Скачайте шаблон.')
        records=list(query);by_id={r.pk:r for r in records};by_name={}
        for row in records:by_name.setdefault(name_key(row.name),[]).append(row)
        updates={};errors=[];skipped=0
        for line,cells in enumerate(iterator,2):
            if line>MAX_ROWS+1:raise ValidationError('Не более 10000 строк в одном файле.')
            def value(index):return cells[index].value if index is not None and index<len(cells) else None
            if not any(value(i) is not None for i in [alias_col,name_col,id_col]):continue
            if any(i is not None and i<len(cells) and cells[i].data_type=='f' for i in [alias_col,name_col,id_col]):
                errors.append(f'Строка {line}: замените формулы обычными значениями.');continue
            alias=str(value(alias_col) or '').strip();name=str(value(name_col) or '').strip();raw_id=value(id_col)
            if not alias:skipped+=1;continue
            if raw_id not in (None,''):
                try:pk=int(str(raw_id).strip())
                except (ValueError,TypeError):pk=None
                target=by_id.get(pk)
                if target is None:errors.append(f'Строка {line}: ID не найден в этом справочнике компании.');continue
                if name and name_key(name)!=name_key(target.name):errors.append(f'Строка {line}: название не соответствует ID.');continue
            else:
                candidates=by_name.get(name_key(name),[])
                if len(candidates)!=1:
                    errors.append(f'Строка {line}: «{name}» — '+('найдено несколько записей; укажите ID.' if candidates else 'запись не найдена.'));continue
                target=candidates[0]
            item=updates.setdefault(target.pk,{'id':target.pk,'name':target.name,'before':list(target.excel_aliases),'after':list(target.excel_aliases),'stamp':stamp(target),'added':[]})
            if name_key(alias) in {name_key(target.name),*(name_key(v) for v in item['after'])}:skipped+=1;continue
            try:validate_excel_aliases([*item['after'],alias])
            except ValidationError as exc:errors.append(f'Строка {line}: '+str(exc.messages[0]));continue
            item['after'].append(alias);item['added'].append(alias)
        if errors:raise ValidationError(errors[:50]+([f'Ещё ошибок: {len(errors)-50}'] if len(errors)>50 else []))
        return [row for row in updates.values() if row['added']],skipped
    finally:book.close()


@transaction.atomic
def apply(user,kind,token):
    require_roles(user,PERSONNEL_CATALOG_EDIT_ROLES if kind=='labor' else EQUIPMENT_CATALOG_EDIT_ROLES)
    try:payload=signing.loads(token,salt=SALT,max_age=1800)
    except signing.BadSignature:raise ValidationError('Предпросмотр устарел. Загрузите файл снова.')
    if (payload.get('user'),payload.get('company'),payload.get('kind'))!=(user.pk,user.company_id,kind):raise ValidationError('Предпросмотр относится к другому пользователю или справочнику.')
    entries=payload.get('rows',[])
    if not isinstance(entries,list) or not 0<len(entries)<=MAX_ROWS:raise ValidationError('Нет данных для загрузки.')
    ids=[r['id'] for r in entries];rows={r.pk:r for r in catalogue(user,kind,lock=True).filter(pk__in=ids)}
    if len(rows)!=len(ids) or any(stamp(rows[r['id']])!=r['stamp'] for r in entries):raise ValidationError('Справочник изменился после предпросмотра. Загрузите файл снова.')
    for item in entries:
        row=rows[item['id']];validate_excel_aliases(item['after']);row.excel_aliases=item['after'];row.save(update_fields=['excel_aliases','updated_at'])
    AuditLog.objects.create(user=user,action='UPDATE',entity_type='resources.ExcelAssociations',entity_id=kind,
        old_value={str(r['id']):r['before'] for r in entries},new_value={str(r['id']):r['after'] for r in entries},comment='Массовая загрузка ассоциаций из Excel')
    return len(entries),sum(len(r['added']) for r in entries)


@require_http_methods(['GET','POST'])
def import_associations(request,kind):
    if kind not in {'labor','equipment'}:raise Http404
    require_roles(request.user,PERSONNEL_CATALOG_EDIT_ROLES if kind=='labor' else EQUIPMENT_CATALOG_EDIT_ROLES)
    query=catalogue(request.user,kind)
    if request.method=='GET' and request.GET.get('template')=='1':
        rows=((row.pk,row.name,alias) for row in query.order_by('name','pk') for alias in (row.excel_aliases or ['']))
        return xlsx_response('excel-associations-'+kind+'.xlsx',[('Ассоциации',['ID','Название','Ассоциация'],rows)])
    form=UploadForm(request.POST or None,request.FILES or None);preview=None;token='';error=[]
    if request.method=='POST':
        if request.POST.get('action')=='apply':
            try:
                count,aliases=apply(request.user,kind,request.POST.get('token',''))
                messages.success(request,f'Ассоциации загружены: {aliases}. Обновлено записей: {count}.')
                return redirect('resources:brigade_list' if kind=='labor' else 'resources:equipment_type_list')
            except ValidationError as exc:error=exc.messages
        elif form.is_valid():
            try:
                preview,skipped=parse(form.cleaned_data['file'],query)
                if preview:token=signing.dumps({'user':request.user.pk,'company':request.user.company_id,'kind':kind,'rows':preview},salt=SALT,compress=True)
                else:error=['Новых ассоциаций нет. Пустые строки и уже существующие названия пропущены.']
            except ValidationError as exc:form.add_error('file',exc)
    return render(request,'resources/association_import.html',{'form':form,'preview':preview,'token':token,'errors':error,'kind':kind,
        'back_url':'resources:brigade_list' if kind=='labor' else 'resources:equipment_type_list'})
