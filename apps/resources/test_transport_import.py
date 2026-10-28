from datetime import timedelta
from io import BytesIO
from unittest.mock import patch
from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.exceptions import ValidationError,PermissionDenied
from django.urls import reverse
from openpyxl import Workbook
from apps.production.test_transport import TransportTests
from apps.production.transport_details_models import EquipmentDayDetails
from apps.production.transport_models import EquipmentCampStay,EquipmentDayRecord
from apps.production.history_models import FactRevision
from .models import Equipment
from . import transport_import as service
from .transport_import_models import TransportImportRow


def upload(rows,name='Сводка.xlsm'):
    book=Workbook();sheet=book.active;sheet.title='Сводка';sheet.append(service.HEADERS);sheet.append(list(range(1,22)))
    for row in rows:sheet.append(row)
    out=BytesIO();book.save(out)
    return SimpleUploadedFile(name,out.getvalue(),content_type='application/vnd.ms-excel.sheet.macroEnabled.12')


class TransportImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):TransportTests.setUpTestData.__func__(cls)

    def row(self,**kwargs):
        data=dict(zip(service.FIELDS,[1,2021,'CAT 320','INV-100','A100','VIN-100',self.camp.name,'Экскаватор',
            'Иванов Иван','Петров Пётр','Да','Испр.','В работе','Неисправность','27.10.2026','Заявка №1','Оплачено','В пути','01.11.2026','ТО выполнено','Примечание']))
        data.update(kwargs);return [data[k] for k in service.FIELDS]

    def preview(self,rows=None,user=None,second=None):return service.preview(user or self.admin,[upload(rows or [self.row()]),*([upload(second,'Вторая сводка.xlsm')] if second else [])],self.day)

    def test_first_vin_loaded_later_collision_queued_from_both_files(self):
        batch=self.preview([self.row(),self.row(inventory='INV-200',plate='A200')],second=[self.row()])
        self.assertEqual(batch.rows.filter(status='READY').count(),1)
        self.assertFalse(Equipment.objects.filter(inventory_number='INV-100').exists())
        service.apply(self.admin,batch)
        self.assertEqual(batch.rows.filter(status='APPLIED').count(),1);self.assertEqual(batch.rows.filter(status='PENDING').count(),2)
        machine=Equipment.objects.get(inventory_number='INV-100')
        self.assertEqual(machine.vin,'VIN-100')
        details=EquipmentDayDetails.objects.get(equipment=machine)
        self.assertEqual(details.parts_status,'В пути');self.assertEqual(details.shift_two,'Петров Пётр');self.assertTrue(details.in_gpr)
        self.assertTrue(EquipmentCampStay.objects.filter(equipment=machine,camp=self.camp).exists())
        self.assertEqual(EquipmentDayRecord.objects.get(equipment=machine).status,'WORKED')
        self.assertEqual(batch.rows.first().original['Состояние поставки запчастей'],'В пути')
        with self.assertRaises(ValidationError):service.apply(self.admin,batch)

    def test_dispatcher_corrects_vin_preserving_original_and_actor_history(self):
        batch=self.preview([self.row(),self.row(inventory='INV-200',plate='A200')]);service.apply(self.admin,batch)
        row=batch.rows.get(status='PENDING');original=row.original.copy();version=row.updated_at.isoformat()
        service.resolve(self.operator,row,row.data|{'vin':'VIN-200'},version)
        row.refresh_from_db();self.assertEqual(row.status,'APPLIED');self.assertEqual(row.original,original)
        self.assertEqual(Equipment.objects.get(inventory_number='INV-200').vin,'VIN-200')
        change=FactRevision.objects.filter(entity_type='resources.transportimportrow',entity_id=str(row.pk),actor=self.operator).last()
        self.assertIsNotNone(change);self.assertEqual(change.direction,'mechanization')
        with self.assertRaises(ValidationError):service.resolve(self.colleague,row,row.data,version)

    def test_repeated_same_machine_can_be_closed_without_creating_another(self):
        batch=self.preview([self.row()],second=[self.row()]);service.apply(self.admin,batch)
        row=batch.rows.get(status='PENDING');count=Equipment.objects.count()
        service.resolve(self.operator,row,row.data,row.updated_at.isoformat(),close_duplicate=True)
        self.assertEqual(Equipment.objects.count(),count);row.refresh_from_db();self.assertEqual(row.status,'DUPLICATE')
        again=self.preview();service.apply(self.admin,again);self.assertEqual(again.rows.get().status,'APPLIED')

    def test_invalid_row_does_not_block_good_rows_or_discard_any_column(self):
        bad=self.row(vin='VIN-BAD',inventory='INV-BAD',plate='A300',year='190',condition='Не испр.',repair_status='В ожидании ЗЧ')
        batch=self.preview([bad,self.row()]);service.apply(self.admin,batch)
        self.assertEqual(batch.rows.filter(status='APPLIED').count(),1)
        self.assertIn('год',batch.rows.get(status='PENDING').error)
        self.assertEqual(len(batch.rows.get(status='PENDING').original),21)

    def test_master_changes_stale_preview_and_foreign_camp_or_company_rejected(self):
        batch=self.preview();self.machine.name='Изменено';self.machine.save()
        with self.assertRaises(ValidationError):service.apply(self.admin,batch)
        batch=self.preview([self.row(site=self.unrelated.name)],user=self.operator)
        service.apply(self.operator,batch);row=batch.rows.get()
        self.assertEqual(row.status,'PENDING')
        self.assertFalse(Equipment.objects.filter(inventory_number='INV-100').exists())

    def test_ui_queue_export_correction_and_readonly_roles(self):
        batch=self.preview([self.row(),self.row(inventory='INV-200',plate='A200')]);service.apply(self.admin,batch)
        self.client.force_login(self.operator)
        url=reverse('resources:transport_import_errors');self.assertContains(self.client.get(url),'VIN-100')
        self.assertEqual(self.client.get(url,{'export':'xlsx'}).status_code,200)
        row=batch.rows.get(status='PENDING');url=reverse('resources:transport_import_correction',args=[row.pk])
        response=self.client.get(url);self.assertContains(response,'Исходная строка')
        data={k:v for k,v in row.data.items() if k!='number'}|{'vin':'VIN-200','version':row.updated_at.isoformat(),'action':'apply'}
        self.assertEqual(self.client.post(url,data).status_code,302)
        self.client.force_login(self.users['DISPATCHER']);self.assertEqual(self.client.get(url).status_code,403)

    def test_details_can_be_updated_daily_and_latest_appears_in_catalogue(self):
        batch=self.preview();service.apply(self.admin,batch);row=batch.rows.get()
        self.client.force_login(self.operator);url=reverse('production:transport_details',args=[row.equipment_id]);params={'camp':self.camp.pk,'date':str(self.day)}
        response=self.client.get(url,params);self.assertEqual(response.status_code,200)
        data=params.copy()
        from apps.production.transport_details_models import EquipmentServiceStatus
        delivered=EquipmentServiceStatus.objects.get(company=self.company,kind='PARTS',name='Доставлены')
        data.update({'defects':'Уточнение поломки','parts_stage':delivered.pk,'signed':response.context['form'].initial['signed']})
        self.assertEqual(self.client.post(url,data).status_code,302)
        self.assertEqual(EquipmentDayDetails.objects.get(equipment_id=row.equipment_id).parts_stage.name,'Доставлены')
        self.assertEqual(self.client.post(url,data).status_code,400)
        self.client.force_login(self.admin)
        response=self.client.get(reverse('resources:equipment_unit_detail',args=[row.equipment_id]));self.assertContains(response,'Доставлены')
        response=self.client.get(reverse('resources:equipment_unit_list'),{'export':'xlsx'})
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(response.content),data_only=True)
        self.assertIn('Ремонт и обслуживание',book.sheetnames)
        self.assertIn('Доставлены',[c.value for row in book['Ремонт и обслуживание'] for c in row])
        response=self.client.get(reverse('production:transport_board'),params|{'export':'xlsx'})
        book=load_workbook(BytesIO(response.content),data_only=True)
        self.assertIn('Доставлены',[c.value for row in book['Ремонт и обслуживание'] for c in row])
