from io import BytesIO
from django.test import TestCase
from django.urls import reverse
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook,load_workbook
from apps.accounts.models import Company,Role,User
from .models import Brigade,EquipmentType
from .association_import import apply


def upload(rows):
    book=Workbook();sheet=book.active
    for row in rows:sheet.append(row)
    stream=BytesIO();book.save(stream);book.close()
    return SimpleUploadedFile('associations.xlsx',stream.getvalue())


class AssociationImportTests(TestCase):
    def setUp(self):
        self.company=Company.objects.create(name='Import associations')
        self.user=User.objects.create_user(username='import-associations',company=self.company,role=Role.objects.get(code='PLANNER'))
        self.job=Brigade.objects.create(company=self.company,name='Мастер',excel_aliases=['Прораб'],staffing_current=True)
        self.gear=EquipmentType.objects.create(company=self.company,name='Экскаватор')
        self.client.force_login(self.user)

    def preview(self,kind,rows):
        return self.client.post(reverse('resources:catalogue_association_import',args=[kind]),{'file':upload(rows)})

    def test_preview_additive_apply_and_search_then_safe_duplicate_skip(self):
        url=reverse('resources:catalogue_association_import',args=['labor'])
        rows=[['ID','Название','Ассоциация'],[self.job.pk,self.job.name,'Начальник участка'],[self.job.pk,self.job.name,'прораб'],['',self.job.name,'Старший мастер']]
        page=self.preview('labor',rows);self.assertEqual(page.status_code,200)
        self.assertContains(page,'Предпросмотр');self.assertEqual(len(page.context['preview'][0]['added']),2)
        self.job.refresh_from_db();self.assertEqual(self.job.excel_aliases,['Прораб'])
        response=self.client.post(url,{'action':'apply','token':page.context['token']})
        self.assertEqual(response.status_code,302)
        self.job.refresh_from_db();self.assertEqual(self.job.excel_aliases,['Прораб','Начальник участка','Старший мастер'])
        self.assertContains(self.client.get(reverse('resources:brigade_list'),{'name':'начальник участка'}),self.job.name)
        page=self.preview('labor',rows);self.assertEqual(page.context['preview'],[])
        self.assertEqual(Brigade.objects.filter(company=self.company).count(),1)

    def test_equipment_template_has_current_ids_and_imports_root_only(self):
        url=reverse('resources:catalogue_association_import',args=['equipment'])
        child=EquipmentType.objects.create(company=self.company,name='Модель',parent=self.gear)
        response=self.client.get(url,{'template':'1'});book=load_workbook(BytesIO(response.content));rows=list(book.active.values);book.close()
        self.assertEqual(rows[0],('ID','Название','Ассоциация'));self.assertEqual(len(rows),2)
        page=self.preview('equipment',[rows[0],rows[1][:2]+('ЭО-4225',)])
        self.assertEqual(self.client.post(url,{'action':'apply','token':page.context['token']}).status_code,302)
        self.gear.refresh_from_db();self.assertEqual(self.gear.excel_aliases,['ЭО-4225'])
        bad=self.preview('equipment',[['ID','Ассоциация'],[child.pk,'Нельзя']]);self.assertFalse(bad.context['token']);self.assertTrue(bad.context['form'].errors)

    def test_unmatched_ambiguous_foreign_formulas_and_duplicate_validation_block_entire_file(self):
        foreign=Brigade.objects.create(company=Company.objects.create(name='Other associations'),name='Чужая')
        Brigade.objects.create(company=self.company,name=self.job.name)
        cases=[[['Название','Ассоциация'],['Нет такой','Новая']],
            [['Название','Ассоциация'],[self.job.name,'Новая']],
            [['ID','Ассоциация'],[foreign.pk,'Новая']],
            [['ID','Ассоциация'],[self.job.pk,'=1+1']],
            [['ID','Ассоциация'],[self.job.pk,'x'*501]]]
        for rows in cases:
            page=self.preview('labor',rows);self.assertFalse(page.context['token']);self.assertTrue(page.context['form'].errors)
        self.job.refresh_from_db();self.assertEqual(self.job.excel_aliases,['Прораб'])

    def test_stale_preview_token_tampering_and_direction_scoping(self):
        page=self.preview('labor',[['ID','Ассоциация'],[self.job.pk,'Новая']]);token=page.context['token']
        with self.assertRaises(ValidationError):apply(self.user,'equipment',token)
        with self.assertRaises(ValidationError):apply(self.user,'labor',token+'x')
        self.job.name='Уточнено';self.job.save()
        with self.assertRaises(ValidationError):apply(self.user,'labor',token)
        self.job.refresh_from_db();self.assertEqual(self.job.excel_aliases,['Прораб'])
        for role,allowed,denied in [('PERSONNEL_CATALOG_MANAGER','labor','equipment'),('EQUIPMENT_CATALOG_MANAGER','equipment','labor')]:
            self.user.role=Role.objects.get(code=role);self.user.save()
            self.assertEqual(self.client.get(reverse('resources:catalogue_association_import',args=[allowed])).status_code,200)
            self.assertEqual(self.client.post(reverse('resources:catalogue_association_import',args=[denied]),{}).status_code,403)
