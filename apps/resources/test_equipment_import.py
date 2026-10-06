from io import BytesIO
from django.test import TestCase
from django.urls import reverse
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from apps.accounts.models import Company, User, Role
from .models import EquipmentType, EquipmentCategory
from .equipment_import import HEADERS


class EquipmentImportTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name='Excel company')
        self.user = User.objects.create_user(username='excel-planner', company=self.company, role=Role.objects.get(code='PLANNER'))
        self.client.force_login(self.user)
        self.url = reverse('resources:equipment_type_import')

    def file(self, rows, headers=HEADERS):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return SimpleUploadedFile('types.xlsx', output.getvalue())

    def test_import_defaults_duplicates_categories_and_company_scope(self):
        foreign = Company.objects.create(name='Foreign Excel')
        category = EquipmentCategory.objects.create(company=foreign, name='Землеройная')
        EquipmentType.objects.create(company=foreign, name='Экскаватор', category=category)
        old = EquipmentType.objects.create(company=self.company, name='Кран', unit='старое')
        response = self.client.post(self.url, {'file':self.file([
            ['Экскаватор', 'Землеройная', '', ''],
            ['экскаватор', 'Землеройная', 'ч', 'Да'],
            ['Кран', 'Подъёмная', 'ч', 'Нет'],
            ['Бульдозер', 'землеройная', 'маш.-ч', 'Нет'],
        ]), 'company':foreign.pk})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['result'], {'created':2, 'skipped':2, 'categories':1})
        row = EquipmentType.objects.get(company=self.company, name='Экскаватор')
        self.assertEqual(row.unit, 'ед.')
        self.assertTrue(row.is_active)
        self.assertNotEqual(row.category_id, category.pk)
        bulldozer = EquipmentType.objects.get(company=self.company, name='Бульдозер')
        self.assertEqual(bulldozer.category_id, row.category_id)
        self.assertFalse(bulldozer.is_active)
        old.refresh_from_db()
        self.assertEqual(old.unit, 'старое')
        self.assertFalse(EquipmentCategory.objects.filter(company=self.company, name='Подъёмная').exists())
        response = self.client.post(self.url, {'file':self.file([['Экскаватор']])})
        self.assertEqual(response.context['result']['created'], 0)

    def test_errors_cancel_entire_import(self):
        for rows in ([['Valid', 'New'], ['', 'Other']], [['Valid'], ['Bad', '', '', 'sometimes']], [['=1+1']]):
            response = self.client.post(self.url, {'file':self.file(rows)})
            self.assertEqual(response.status_code, 400)
            self.assertTrue(response.context['errors'])
            self.assertFalse(EquipmentType.objects.exists())
            self.assertFalse(EquipmentCategory.objects.exists())

    def test_template_and_invalid_files(self):
        response = self.client.get(reverse('resources:equipment_import_template'))
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(list(next(workbook.worksheets[0].values)), HEADERS)
        workbook.close()
        for upload in (SimpleUploadedFile('bad.xlsx', b'broken'), SimpleUploadedFile('old.xls', b'broken'), self.file([]), self.file([['Name']], headers=['Wrong'])):
            self.assertEqual(self.client.post(self.url, {'file':upload}).status_code, 400)
        self.assertEqual(self.client.post(self.url, {'file':self.file([['Only name']], headers=['Название'])}).status_code, 200)
        self.assertEqual(EquipmentType.objects.get().unit, 'ед.')

    def test_roles_cannot_import(self):
        for code in ('MANAGER', 'FOREMAN'):
            user = User.objects.create_user(username='excel-'+code, company=self.company, role=Role.objects.get(code=code))
            self.client.force_login(user)
            self.assertEqual(self.client.get(self.url).status_code, 403)
            self.assertEqual(self.client.post(self.url, {'file':self.file([['Unauthorized']])}).status_code, 403)
            self.assertEqual(self.client.get(reverse('resources:equipment_import_template')).status_code, 403)
        self.assertFalse(EquipmentType.objects.exists())
