"""Reclassify existing numbered machines without creating catalogue entries."""
from collections import defaultdict
from django import forms
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import transaction
from django.shortcuts import render
from django.views.decorators.http import require_http_methods
from core.permissions import EQUIPMENT_CATALOG_EDIT_ROLES,require_roles
from .models import Equipment,EquipmentType,EquipmentCategory
from .brigade_import import BrigadeImportForm,text


class ImportForm(BrigadeImportForm):
    file=forms.FileField(label='Excel с группами техники (.xlsx)',widget=forms.FileInput(attrs={'class':'form-control','accept':'.xlsx'}))


def read_rows(file):
    from zipfile import ZipFile
    from openpyxl import load_workbook
    try:
        with ZipFile(file) as archive:
            if sum(i.file_size for i in archive.infolist())>50*1024*1024:return [],['Файл слишком большой после распаковки.']
        file.seek(0);book=load_workbook(file,read_only=True,data_only=False)
    except Exception:return [],['Не удалось прочитать файл .xlsx.']
    rows=[];errors=[]
    try:
        sheet=book.worksheets[0];sheet.reset_dimensions();iterator=sheet.iter_rows();header=next(iterator,())
        columns={text(c.value).casefold():i for i,c in enumerate(header)}
        recognized={'id','госномер','инвентарный номер','vin / заводской номер','vin','техника для планирования','группа техники','группа'}
        if not({'техника для планирования','группа техники','группа'}&columns.keys()):return [],['Нужна колонка «Техника для планирования» или «Группа техники». Используйте экспорт машин.']
        if not({'id','госномер','инвентарный номер','vin / заводской номер','vin'}&columns.keys()):return [],['Нужен ID, VIN, инвентарный номер или госномер для поиска существующей машины.']
        present=[text(c.value).casefold() for c in header if text(c.value).casefold() in recognized]
        if len(present)!=len(set(present)):return [],['Заголовки изменяемых колонок и идентификаторов не должны повторяться.']
        for number,cells in enumerate(iterator,2):
            if number>5001:errors.append('Допускается не более 5000 строк.');break
            values={}
            for name in recognized:
                index=columns.get(name);cell=cells[index] if index is not None and index<len(cells) else None
                if cell and cell.data_type in {'f','e'}:errors.append(f'Строка {number}: замените формулу или ошибку в «{name}» значением.')
                values[name]=text(cell.value) if cell else ''
            if not any(values.values()):continue
            values['line']=number;rows.append(values)
        if not rows and not errors:errors.append('В файле нет данных.')
        return rows,errors
    finally:book.close()


@transaction.atomic
def apply(user,rows):
    from apps.accounts.models import Company
    from apps.audit.models import AuditLog
    Company.objects.select_for_update().get(pk=user.company_id)
    locked_groups=list(EquipmentCategory.objects.select_for_update().filter(company=user.company))
    locked_types=list(EquipmentType.objects.select_for_update().filter(company=user.company,merge_source__isnull=True).select_related('category'))
    machines=list(Equipment.objects.select_for_update().filter(company=user.company).select_related('type__parent','type__category','type__parent__category'))
    maps={name:defaultdict(list) for name in ['id','госномер','инвентарный номер','vin']}
    for machine in machines:
        for name,value in [('id',str(machine.pk)),('госномер',machine.plate_number),('инвентарный номер',machine.inventory_number),('vin',machine.vin)]:
            if value:maps[name][text(value).casefold()].append(machine)
    roots=defaultdict(list);categories=defaultdict(list)
    for root in locked_types:
        if root.parent_id is None:roots[root.name.strip().casefold()].append(root)
    for group in locked_groups:categories[group.name.strip().casefold()].append(group)
    variants=[row for row in locked_types if row.parent_id]
    warnings=[];updated=unchanged=0;seen=set()
    for row in rows:
        label=f"Строка {row['line']}";candidates=None;invalid=False
        # Every supplied identifier must agree; an unknown ID never falls back to another machine.
        for name in ['id','vin','инвентарный номер','госномер']:
            value=row.get('vin / заводской номер') or row.get('vin') if name=='vin' else row.get(name)
            if not value:continue
            matches=maps[name].get(value.casefold(),[])
            ids={m.pk for m in matches}
            candidates=ids if candidates is None else candidates&ids
        if not candidates or len(candidates)!=1:warnings.append(label+': машина не найдена однозначно; изменений нет.');continue
        machine=next(m for m in machines if m.pk in candidates)
        if machine.pk in seen:warnings.append(label+': повтор машины в файле; изменений нет.');continue
        seen.add(machine.pk)
        root=machine.type.parent or machine.type
        planning=row.get('техника для планирования','');group_name=row.get('группа техники') or row.get('группа','')
        if planning:
            matches=roots.get(planning.casefold(),[])
            if len(matches)!=1:warnings.append(label+f': плановая позиция «{planning}» не найдена однозначно; изменений нет.');continue
            root=matches[0]
        target=machine.type if (machine.type.parent_id or machine.type_id)==root.pk else root
        if target==root and machine.type.parent_id and machine.type.parent_id!=root.pk:
            model_label=(machine.brand_model or machine.type.brand_model).strip().casefold()
            matches=[v for v in variants if v.parent_id==root.pk and model_label and (v.brand_model.strip().casefold()==model_label or v.name.strip().casefold()==model_label)]
            if len(matches)==1:target=matches[0]
            else:warnings.append(label+': для выбранной плановой позиции нет подходящей существующей модели; изменений нет.');continue
        group=machine.catalogue_group
        if group_name and group_name!='Без группы':
            matches=categories.get(group_name.casefold(),[])
            if len(matches)!=1:warnings.append(label+f': группа «{group_name}» не найдена однозначно; изменений нет.');continue
            group=matches[0]
        if target.pk==machine.type_id and (group.pk if group else None)==machine.catalogue_group_id:unchanged+=1;continue
        before={'type_id':machine.type_id,'catalogue_group_id':machine.catalogue_group_id};machine.type=target;machine.catalogue_group=group;machine.full_clean();machine.save(update_fields=['type','catalogue_group','updated_at']);updated+=1
        AuditLog.objects.create(user=user,action='UPDATE',entity_type='resources.Equipment',entity_id=str(machine.pk),old_value=before,new_value={'type_id':target.pk,'catalogue_group_id':machine.catalogue_group_id},comment='Изменение классификации существующей машины из Excel')
    return {'updated':updated,'unchanged':unchanged,'warnings':warnings,'skipped':len(warnings)}


@login_required
@require_http_methods(['GET','POST'])
def import_groups(request):
    require_roles(request.user,EQUIPMENT_CATALOG_EDIT_ROLES);form=ImportForm(request.POST or None,request.FILES or None);errors=[];result=None
    if request.method=='POST' and form.is_valid():
        rows,errors=read_rows(form.cleaned_data['file'])
        if not errors:
            try:result=apply(request.user,rows)
            except ValidationError as exc:errors=exc.messages
    return render(request,'resources/equipment_group_import.html',{'form':form,'errors':errors,'result':result},status=400 if errors or form.errors else 200)
