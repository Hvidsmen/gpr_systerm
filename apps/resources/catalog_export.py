"""Export the same company-scoped filtered queryset as the catalog list."""
from io import BytesIO
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from core.permissions import require_roles, READ_ROLES
from .views import BrigadeListView, EquipmentTypeListView


class CatalogExcelExport:
    headers = []
    filename = ''
    sheet_title = ''

    def get(self, request, *args, **kwargs):
        require_roles(request.user, READ_ROLES)
        queryset = self.get_queryset()
        if self.filter_form.errors:
            return HttpResponse('Некорректные фильтры. Исправьте значения в списке и повторите экспорт.', status=400, content_type='text/plain; charset=utf-8')
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self.sheet_title
        sheet.freeze_panes = 'A2'
        sheet.append(self.headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='0D6EFD')
            sheet.column_dimensions[cell.column_letter].width = 28 if cell.value != 'Описание' else 60
        for index, row in enumerate(queryset.iterator(chunk_size=1000), 2):
            # Treat every catalog value as text, including names starting '='.
            for column, value in enumerate(self.values(row), 1):
                cell = sheet.cell(index, column, ILLEGAL_CHARACTERS_RE.sub('', str(value)))
                cell.data_type = 's'
        sheet.auto_filter.ref = sheet.dimensions
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{self.filename}"'
        return response


class BrigadeExport(CatalogExcelExport, BrigadeListView):
    headers = ['ID', 'Название', 'Группа', 'Описание', 'Активна', 'Единица измерения', 'Штатное расписание ЗУП', 'Обновлено из ЗУП']
    filename = 'brigades.xlsx'
    sheet_title = 'Бригады'

    def values(self, row):
        return [row.pk, row.name, row.group.name if row.group_id else '', row.description, 'Да' if row.is_active else 'Нет', row.unit, row.staffing_status_label, row.staffing_synced_at.isoformat() if row.staffing_synced_at else '']


class EquipmentTypeExport(CatalogExcelExport, EquipmentTypeListView):
    headers = ['Название', 'Техника для планирования', 'Группа техники', 'Единица измерения', 'Активна', 'Категории водителя', 'Классы водителя']
    filename = 'equipment_types.xlsx'
    sheet_title = 'Техника для планирования'

    def values(self, row):
        return [row.name, row.parent.name if row.parent_id else row.name, row.category.name if row.category_id else '', row.unit, 'Да' if row.is_active else 'Нет', ', '.join(row.driver_category_labels), ', '.join(row.driver_class_labels)]
