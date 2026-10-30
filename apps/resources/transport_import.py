"""Import the first Summary sheet; collisions stay editable in a dispatcher queue."""
import hashlib,json,re,unicodedata
from collections import defaultdict
from datetime import date,datetime
from io import BytesIO
from openpyxl.utils import get_column_letter
from zipfile import ZipFile
from openpyxl import load_workbook
from django.core.exceptions import ValidationError,PermissionDenied
from django.db import transaction
from django.db.models import F
from django.utils import timezone
from core.fact_history import with_history_actor
from core.fact_edit_window import require_fact_day
from core.permissions import require_roles,role_code
from apps.accounts.models import Company
from apps.personnel.models import Camp
from apps.production import transport
from apps.production.transport_models import EquipmentCampStay,EquipmentTransportEvent,EquipmentDayRecord
from apps.production.transport_details_models import EquipmentDayDetails
from apps.production.transport_detail_values import SOURCE_DETAIL_FIELDS as DETAIL_FIELDS,from_source,lookup_context
from .models import Equipment,EquipmentType
from .transport_import_models import TransportImportBatch,TransportImportRow
from .fleet_import import index,key

HEADERS=['№','Год выпуска','Марка/модель','Инвентарный номер','Гос.рег.знак','VIN№(заводской номер)','Участок','Категория',
    '1 смена','2 смена','В ГПР','Тех. состояние','Состояние','Краткое описание неисправностей, поломок','Дата постановки на ремонт',
    'Номер и дата заявки','Статус заявок (оплачено, не оплачено), дата оплаты','Состояние поставки запчастей','Плановая дата поставки ЗЧ в Дудинку',
    'Проведение ТО весна-лето 2026г. (да,нет), готовность ТС к эксплуатации..','Примечание']
FIELDS=['number','year','model','inventory','plate','vin','site','planning_type','shift_one','shift_two','in_gpr','condition','repair_status',
    'defects','repair_date','request_details','request_status','parts_status','expected_parts_date','maintenance','comment']
PLACEHOLDERS={'','-','—','б/н','нет','н/д','отсутствует'}


def text(value):
    if value is None:return ''
    if isinstance(value,(date,datetime)):return value.isoformat()
    return str(value).strip()


def normalized(value):return re.sub(r'[\s\-]+','',unicodedata.normalize('NFKC',text(value))).upper().translate(str.maketrans('АВЕКМНОРСТУХ','ABEKMHOPCTYX'))
def usable(value):return key(value) not in PLACEHOLDERS

def fingerprint(company):
    # A change in master data or the transport journal invalidates the preview.
    from apps.production.transport_models import EquipmentTransportControl
    models=[Equipment,EquipmentType,Camp,EquipmentTransportControl]
    payload={m._meta.label:list(m.objects.filter(company=company).order_by('pk').values()) for m in models}
    return hashlib.sha256(json.dumps(payload,default=str,sort_keys=True).encode()).hexdigest()


def read_files(files):
    rows=[]
    for upload in files:
        if not upload.name.lower().endswith(('.xlsx','.xlsm')) or upload.size>5*1024*1024:raise ValidationError('Нужны .xlsx / .xlsm до 5 МБ каждый.')
        try:
            with ZipFile(upload) as archive:
                if sum(i.file_size for i in archive.infolist())>50*1024*1024:raise ValidationError('Слишком большой файл после распаковки.')
            upload.seek(0);raw=upload.read()
            book=load_workbook(BytesIO(raw),read_only=True,data_only=False,keep_links=False)
            cached=load_workbook(BytesIO(raw),read_only=True,data_only=True,keep_links=False)
        except ValidationError:raise
        except Exception:raise ValidationError(f'Не удалось прочитать {upload.name}.')
        try:
            sheet=book.worksheets[0]
            if sheet.title.strip().casefold()!='сводка':raise ValidationError(f'{upload.name}: первый лист должен называться «Сводка».')
            if sheet.max_row>5002 or sheet.max_column>100:raise ValidationError('Слишком много строк или столбцов.')
            headers=[text(c.value) for c in next(sheet.iter_rows(max_row=1))]
            if len(headers)<21:raise ValidationError(f'{upload.name}: нужны все 21 колонка сводки.')
            if headers[:8]!=HEADERS[:8]:raise ValidationError(f'{upload.name}: не найдены основные 8 колонок сводки.')
            if headers[8:19]!=HEADERS[8:19] or not headers[19].startswith('Проведение ТО') or headers[20]!='Примечание':raise ValidationError(f'{upload.name}: структура дополнительных колонок отличается от сводки. Данные не загружены.')
            for number,(cells,results) in enumerate(zip(sheet.iter_rows(min_row=2),cached.worksheets[0].iter_rows(min_row=2)),2):
                values=[c.value for c in cells]
                if number==2 and (values[:21]==list(range(1,22)) or [c.value for c in results[:21]]==list(range(1,22))):continue
                if not any(v is not None for v in values):continue
                if not values[2] and not any(values[3:6]):continue
                original={label or f'Столбец {n+1}':text(values[n]) for n,label in enumerate(headers)}
                data={};formula_errors={}
                for index,(name,cell,result) in enumerate(zip(FIELDS,cells,results)):
                    value=cell.value
                    if cell.data_type=='f':
                        value=result.value
                        if index and result.data_type=='e':formula_errors[name]=f'Ячейка {cell.coordinate} («{HEADERS[index]}»): результат формулы содержит ошибку Excel {result.value}.';value=None
                        elif index and value is None and result.data_type not in {'s','str','inlineStr'}:formula_errors[name]=missing_formula_result(name,number)
                    elif index and cell.data_type=='e':formula_errors[name]=f'Ячейка {cell.coordinate} («{HEADERS[index]}»): ошибка Excel {cell.value}.';value=None
                    data[name]=('Да' if value else 'Нет') if name=='in_gpr' and isinstance(value,bool) else text(value)
                if formula_errors:data['_formula_errors']=formula_errors
                if key(data['year']) in PLACEHOLDERS:data['year']=''
                rows.append({'source_file':upload.name[:255],'source_sheet':sheet.title,'source_row':number,'original':original,'data':data,'error':''})
        finally:book.close();cached.close()
    if not rows:raise ValidationError('На первых листах «Сводка» нет машин.')
    return rows


def missing_formula_result(name,row):
    index=FIELDS.index(name)
    return f'Ячейка {get_column_letter(index+1)}{row} («{HEADERS[index]}»): в книге не сохранён результат формулы. Пересчитайте и сохраните файл в Excel, затем загрузите его повторно.'


def validate(data):
    errors=[message for name,message in data.get('_formula_errors',{}).items() if not data.get(name)]
    for field,label,limit in [('model','Марка / модель',150),('planning_type','Техника для планирования',100),('site','ВЖГ',150)]:
        if not data.get(field):errors.append(f'Не заполнено: {label}.')
        elif len(data[field])>limit:errors.append(f'{label}: слишком длинное значение.')
    for field,limit in [('plate',50),('inventory',50),('vin',100)]:
        if len(data.get(field,''))>limit:errors.append(f'{field}: слишком длинное значение.')
    if data.get('year') and (not data['year'].isdigit() or not 1900<=int(data['year'])<=2100):errors.append('Некорректный год выпуска.')
    if not any(usable(data.get(k)) and (k!='inventory' or key(data[k])!='сск') for k in ['vin','inventory','plate']):errors.append('Нет надёжного номера машины.')
    if len(data.get('planning_type','')+' / '+data.get('model',''))>150:errors.append('Название техники и модели длиннее 150 символов.')
    condition=key(data.get('condition',''))
    if condition not in {'испр','не испр','исправна','в ремонте','под списание','на запчасти'}:errors.append('Укажите техническое состояние: Испр., Не испр. или Под списание.')
    status=key(data.get('repair_status',''))
    if status not in {'в работе','в простое','в ремонте','в ожидании р','в ожидании зч','под списание','на запчасти'}:errors.append('Неизвестная отметка работы / ремонта за день.')
    if condition in {'испр','исправна'} and status not in {'в работе','в простое'}:errors.append('Исправная машина имеет отметку ремонта. Уточните состояние.')
    if condition not in {'испр','исправна'} and status=='в работе':errors.append('Неисправная машина отмечена в работе. Уточните состояние.')
    if key(data.get('in_gpr','')) not in {'','да','нет'}:errors.append('В ГПР: укажите Да, Нет или оставьте пустым.')
    for name,label in DETAIL_FIELDS:
        limit=400 if name in {'shift_one','shift_two'} else 200 if name in {'repair_date','expected_parts_date','repair_status'} else 10000
        if name!='in_gpr' and len(data.get(name,''))>limit:errors.append(f'{label}: слишком длинное значение.')
    return errors


def collisions(company,rows):
    masters={name:defaultdict(list) for name in ['vin','inventory','plate']};seen={name:{} for name in masters};result=[]
    units={u.pk:u for u in Equipment.objects.filter(company=company).order_by('pk')}
    attrs={'vin':'vin','plate':'plate_number','inventory':'inventory_number'}
    for unit in units.values():
        for field,attr in attrs.items():
            value=getattr(unit,attr)
            if usable(value) and not (field=='inventory' and key(value)=='сск'):masters[field][normalized(value)].append(unit.pk)
    for row in rows:
        data=row['data'];errors=validate(data);candidates=set();duplicate=None
        if row.get('error'):errors.insert(0,row['error'])
        for field in masters:
            value=data.get(field,'')
            if not usable(value) or field=='inventory' and key(value)=='сск':continue
            ids=masters[field].get(normalized(value),[]);candidates.update(ids)
            if len(ids)>1:errors.append(f'{value}: номер уже указан у нескольких машин справочника.')
            if normalized(value) in seen[field]:errors.append(f'{value}: повтор строки {seen[field][normalized(value)]}.')
        if len(candidates)>1:errors.append('Номера указывают на разные машины справочника.')
        elif candidates:
            duplicate=next(iter(candidates));unit=units[duplicate]
            for field,attr in attrs.items():
                incoming=data.get(field,'');actual=getattr(unit,attr)
                if usable(incoming) and usable(actual) and not (field=='inventory' and (key(incoming)=='сск' or key(actual)=='сск')) and normalized(incoming)!=normalized(actual):
                    errors.append(f'{incoming}: не совпадает с номером существующей машины {actual}.')
        result.append({**row,'error':' '.join(errors),'duplicate_of_id':duplicate,'equipment_id':duplicate if not errors else None})
        if not errors:
            for field in seen:
                value=data.get(field,'')
                if usable(value) and not (field=='inventory' and key(value)=='сск'):seen[field][normalized(value)]=f'{row["source_file"]}, строка {row["source_row"]}'
    return result


@with_history_actor
@transaction.atomic
def preview(user,files,day):
    require_roles(user,transport.OPERATORS);require_fact_day(user,day)
    transport._company_mutex(user)
    rows=collisions(user.company,read_files(files))
    batch=TransportImportBatch.objects.create(company=user.company,date=day,files=[f.name for f in files],fingerprint=fingerprint(user.company),created_by=user)
    for row in rows:
        TransportImportRow.objects.create(company=user.company,batch=batch,date=day,status='PENDING' if row['error'] else 'READY',**row)
    return batch


def _camp(user,site):
    candidates=[c for c in Camp.objects.filter(company=user.company) if key(c.name)==key(site)]
    if len(candidates)>1:raise ValidationError('Неоднозначное название ВЖГ. Уточните справочник ВЖГ.')
    camp=candidates[0] if candidates else None
    if camp is None:
        if role_code(user)!='ADMIN':raise ValidationError('ВЖГ не найден. Администратор должен создать городок и связать его с объектами.')
        camp=Camp.objects.create(company=user.company,name=site)
    transport.authorize(user,camp)
    return camp


def _apply_row(user,row,lookups=None):
    data=row.data;errors=validate(data)
    data['_existing_card']=bool(row.equipment_id)
    if errors:raise ValidationError(errors)
    camp=_camp(user,data['site'])
    if row.equipment_id:
        machine=Equipment.objects.select_for_update().get(company=user.company,pk=row.equipment_id)
    else:
        roots=index(EquipmentType.objects.filter(company=user.company,merge_source__isnull=True,parent=None),True)
        root=roots.get(key(data['planning_type']))
        if root is None:root=EquipmentType.objects.create(company=user.company,name=data['planning_type'])
        if not root.is_active:raise ValidationError('Техника для планирования неактивна.')
        name=data['planning_type']+' / '+data['model']
        variants=index(EquipmentType.objects.filter(company=user.company,parent=root,merge_source__isnull=True),True)
        model=variants.get(key(name))
        if model is None:model=EquipmentType.objects.create(company=user.company,parent=root,name=name,brand_model=data['model'],category=root.category)
        if not model.is_active:raise ValidationError('Модель неактивна.')
        machine=Equipment(company=user.company,type=model,name=data['model'],brand_model=data['model'],manufacture_year=int(data['year']) if data['year'] else None,
            inventory_number=data['inventory'],plate_number=data['plate'] if usable(data['plate']) else '',vin=data['vin'] if usable(data['vin']) else '',site_name=camp.name)
        machine.full_clean();machine.save()
    condition='READY' if key(data['condition']) in {'испр','исправна'} else 'SCRAP' if key(data['condition']) in {'под списание','на запчасти'} else 'REPAIR'
    stay=transport.present(user.company,row.date).filter(equipment=machine).select_related('camp').first()
    if stay and stay.camp_id!=camp.pk:
        transport.move(user,stay.camp,row.date,[machine.pk],'TRANSFER',target=camp,signed=transport.snapshot(user,stay.camp,row.date),comment=f'Импорт: {row.source_file}, строка {row.source_row}')
    elif not stay:
        transport.move(user,camp,row.date,[machine.pk],'ARRIVAL',condition=condition,signed=transport.snapshot(user,camp,row.date),comment=f'Импорт: {row.source_file}, строка {row.source_row}')
    if transport.condition_on(machine,row.date)!=condition:
        if condition!='READY' and machine.transport_days.filter(company=user.company,date=row.date).exists():transport.save_day(user,camp,row.date,[machine.pk],'CLEAR',signed=transport.snapshot(user,camp,row.date))
        transport.set_condition(user,camp,row.date,[machine.pk],condition,signed=transport.snapshot(user,camp,row.date))
    if condition=='READY':
        previous=machine.transport_days.filter(company=user.company,date=row.date,camp=camp).first()
        status='WORKED' if key(data['repair_status'])=='в работе' else 'IDLE'
        target={'construction_object':previous.construction_object,'work':previous.work,'work_item':previous.work_item} if previous and status=='WORKED' else {}
        transport.save_day(user,camp,row.date,[machine.pk],status,signed=transport.snapshot(user,camp,row.date),comment=data.get('comment',''),**target)
    values=from_source(user.company,data,lookups)
    values.update({name:data.get(name,'') for name in ['shift_one','shift_two','repair_date','expected_parts_date','repair_status','request_status','parts_status','maintenance']})
    details=EquipmentDayDetails.objects.filter(company=user.company,equipment=machine,date=row.date).first() or EquipmentDayDetails(company=user.company,equipment=machine,camp=camp,date=row.date,recorded_by=user)
    details.camp=camp
    for field,value in values.items():setattr(details,field,value)
    details.full_clean();details.save()
    row.status='APPLIED';row.error='';row.equipment=machine;row.camp=camp;row.resolved_by=user;row.resolved_at=timezone.now();row.save()
    return machine


@with_history_actor
@transaction.atomic
def apply(user,batch):
    require_roles(user,transport.OPERATORS);require_fact_day(user,batch.date)
    transport._company_mutex(user);Company.objects.select_for_update().get(pk=user.company_id)
    batch=TransportImportBatch.objects.select_for_update().get(pk=batch.pk,company=user.company)
    if batch.applied_at:raise ValidationError('Этот импорт уже применён.')
    if fingerprint(user.company)!=batch.fingerprint:raise ValidationError('Справочники или ведение техники изменились. Создайте новый предпросмотр.')
    names={key(name) for name in batch.files}
    newer=TransportImportBatch.objects.filter(company=user.company,applied_at__isnull=False,pk__gt=batch.pk)
    if any(names & {key(name) for name in record.files} for record in newer):raise ValidationError('Уже применена более новая загрузка этого файла. Создайте новый предпросмотр.')
    lookups=lookup_context(user.company)
    for row in batch.rows.select_for_update().order_by('pk'):
        issues=collisions(user.company,[{'data':row.data,'source_file':row.source_file,'source_row':row.source_row,'error':''}])[0]
        if row.status!='READY' or issues['error']:
            row.status='PENDING';row.error=row.error or issues['error'];row.duplicate_of_id=issues['duplicate_of_id'];row.save();continue
        try:
            row.equipment_id=issues['equipment_id']
            with transaction.atomic():_apply_row(user,row,lookups)
        except (ValidationError,PermissionDenied) as exc:
            row.refresh_from_db();row.status='PENDING';row.error='; '.join(exc.messages) if isinstance(exc,ValidationError) else str(exc);row.save()
    # Replace only errors of these source filenames, after processing succeeds.
    # Keep original rows and their audit trail, never touch machines or facts.
    names={key(name) for name in batch.files}
    previous=TransportImportRow.objects.filter(company=user.company,status='PENDING',batch__applied_at__isnull=False,batch_id__lt=batch.pk)
    for row in previous.select_for_update():
        if key(row.source_file) in names:
            row.status='SUPERSEDED';row.resolved_by=user;row.resolved_at=timezone.now();row.save()
    batch.applied_at=timezone.now();batch.save();require_fact_day(user,batch.date)
    return batch


@with_history_actor
@transaction.atomic
def resolve(user,row,data,version,*,close_duplicate=False):
    require_roles(user,transport.OPERATORS)
    transport._company_mutex(user);Company.objects.select_for_update().get(pk=user.company_id)
    row=TransportImportRow.objects.select_for_update().select_related('batch').get(pk=row.pk,company=user.company)
    if row.updated_at.isoformat()!=version or row.status!='PENDING' or not row.batch.applied_at:raise ValidationError('Запись уже изменена или импорт ещё не применён. Обновите страницу.')
    if close_duplicate:
        if not row.duplicate_of_id:raise ValidationError('Сначала укажите существующую машину, совпадающую с этой записью.')
        row.status='DUPLICATE';row.resolved_by=user;row.resolved_at=timezone.now();row.save();return row
    require_fact_day(user,row.date)
    checked=collisions(user.company,[{'data':data,'source_file':row.source_file,'source_row':row.source_row,'error':''}])[0]
    if checked['error']:raise ValidationError(checked['error'])
    for earlier in row.batch.rows.filter(status='APPLIED').exclude(pk=row.pk):
        if any(usable(data.get(field)) and normalized(data.get(field))==normalized(earlier.data.get(field)) for field in ['vin','plate']):raise ValidationError('Эта машина уже загружена из другой строки этого импорта. Закройте повтор либо исправьте номера.')
    row.data=data;row.equipment_id=checked['equipment_id'];_apply_row(user,row);require_fact_day(user,row.date)
    return row


def _queued_values(row):
    data=dict(row.data)
    # Older versions rejected formula-based row numbering too. The order number
    # is informational; the actual workbook row and original expression are kept.
    if str(data.get('number','')).startswith('='):data['number']=''
    if 'Формулы не загружаются как реквизиты' in row.error:
        errors=dict(data.get('_formula_errors',{}))
        for name in FIELDS[1:]:
            if str(data.get(name,'')).startswith('='):
                data[name]='';errors[name]=missing_formula_result(name,row.source_row)
        if errors:data['_formula_errors']=errors
    return data


@with_history_actor
@transaction.atomic
def retry_pending(user,*,existing_only=False):
    require_roles(user,transport.OPERATORS);transport._company_mutex(user)
    Company.objects.select_for_update().get(pk=user.company_id)
    count=0
    for batch in TransportImportBatch.objects.filter(company=user.company,applied_at__isnull=False).order_by('pk'):
        successful=list(batch.rows.filter(status='APPLIED').order_by('pk'))
        for row in batch.rows.select_for_update().filter(status='PENDING').order_by('pk'):
            data=_queued_values(row)
            repeated=next((previous for previous in successful if any(usable(data.get(field)) and not (field=='inventory' and key(data[field])=='сск') and normalized(data[field])==normalized(previous.data.get(field,'')) for field in ['vin','plate','inventory'])),None)
            if repeated:
                row.data=data;row.error='; '.join(['Эта машина уже загружена из другой строки этого импорта. Закройте повтор либо исправьте номера.',*validate(data)]);row.duplicate_of_id=repeated.equipment_id;row.save();continue
            check=collisions(user.company,[{'data':data,'source_file':row.source_file,'source_row':row.source_row,'error':''}])[0]
            if check['error']:
                row.data=data;row.error=check['error'];row.duplicate_of_id=check['duplicate_of_id'];row.save();continue
            if existing_only and not check['equipment_id']:continue
            try:
                with transaction.atomic():
                    require_fact_day(user,row.date);row.data=data;row.equipment_id=check['equipment_id'];_apply_row(user,row)
            except (ValidationError,PermissionDenied) as exc:
                row.refresh_from_db();row.error='; '.join(exc.messages) if isinstance(exc,ValidationError) else str(exc);row.save()
            else:successful.append(row);count+=1
    return count


def retry_existing(user):
    return retry_pending(user,existing_only=True)


@with_history_actor
@transaction.atomic
def clear_errors(user):
    require_roles(user,transport.OPERATORS)
    transport._company_mutex(user)
    count=0
    for row in TransportImportRow.objects.select_for_update().filter(company=user.company,status='PENDING',batch__applied_at__isnull=False):
        row.status='CLEARED';row.resolved_by=user;row.resolved_at=timezone.now();row.save();count+=1
    return count
