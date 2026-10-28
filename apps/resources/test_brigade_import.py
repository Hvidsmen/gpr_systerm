from io import BytesIO
from django.test import TestCase
from django.urls import reverse
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from apps.accounts.models import Company, User, Role
from .models import Brigade, BrigadeGroup, BrigadeMacroGroup
from .brigade_import import HEADERS


class BrigadeImportTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name='Brigade Excel')
        self.user = User.objects.create_user(username='brigade-excel-planner', company=self.company, role=Role.objects.get(code='PLANNER'))
        self.client.force_login(self.user)
        self.url = reverse('resources:brigade_import')

    def file(self, rows, headers=HEADERS):
        workbook = Workbook()
        workbook.active.append(headers)
        for row in rows:
            workbook.active.append(row)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return SimpleUploadedFile('brigades.xlsx', output.getvalue())

    def test_only_groups_of_existing_positions_change(self):
        group=BrigadeGroup.objects.create(company=self.company,name='Монтаж')
        item=Brigade.objects.create(company=self.company,name='Бригада А',description='Сохранить',is_active=False)
        response=self.client.post(self.url,{'file':self.file([[item.pk,item.name,group.name],['','Новая должность','Новая группа']])})
        self.assertEqual(response.status_code,200);item.refresh_from_db()
        self.assertEqual(item.group_id,group.pk);self.assertEqual(item.description,'Сохранить');self.assertFalse(item.is_active)
        self.assertEqual(Brigade.objects.count(),1);self.assertEqual(BrigadeGroup.objects.count(),1)
        self.assertEqual(response.context['result']['updated'],1);self.assertEqual(response.context['result']['unknown'],1)

    def test_unknown_groups_are_warned_and_other_columns_ignored(self):
        item=Brigade.objects.create(company=self.company,name='Бригада А',description='Сохранить',is_active=False)
        file=self.file([[item.pk,item.name,'Неизвестная','Изменить','Да']],headers=['ID','Название','Группа','Описание','Активна'])
        response=self.client.post(self.url,{'file':file});self.assertEqual(response.status_code,200);item.refresh_from_db()
        self.assertIsNone(item.group_id);self.assertFalse(item.is_active);self.assertEqual(item.description,'Сохранить')
        self.assertTrue(response.context['result']['warnings']);self.assertFalse(BrigadeGroup.objects.exists())

    def test_errors_cancel_import_and_template(self):
        item=Brigade.objects.create(company=self.company,name='Бригада А')
        group=BrigadeGroup.objects.create(company=self.company,name='Монтаж')
        response=self.client.post(self.url,{'file':self.file([[item.pk,item.name,group.name],[item.pk,'Чужое название',group.name]])})
        self.assertEqual(response.status_code,400);item.refresh_from_db();self.assertIsNone(item.group_id)
        response=self.client.get(reverse('resources:brigade_import_template'));book=load_workbook(BytesIO(response.content))
        self.assertEqual(list(next(book.worksheets[0].values)),HEADERS);book.close()
        self.assertEqual(self.client.post(self.url,{'file':self.file([[item.name,'']],headers=['Название','Группа'])}).status_code,200)

    def test_bad_files_and_roles(self):
        for file in (SimpleUploadedFile('broken.xlsx',b'broken'),SimpleUploadedFile('old.xls',b'broken'),self.file([])):
            self.assertEqual(self.client.post(self.url,{'file':file}).status_code,400)
        for code in ('MANAGER','FOREMAN'):
            user = User.objects.create_user(username='brigade-excel-'+code,company=self.company,role=Role.objects.get(code=code))
            self.client.force_login(user)
            self.assertEqual(self.client.get(self.url).status_code,403)
            self.assertEqual(self.client.post(self.url,{'file':self.file([['Unauthorized']])}).status_code,403)
            self.assertEqual(self.client.get(reverse('resources:brigade_import_template')).status_code,403)
