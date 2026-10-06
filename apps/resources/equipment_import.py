"""Validate an entire workbook before importing company-scoped equipment types."""
from io import BytesIO
from zipfile import ZipFile
from django import forms
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import render
from django.views import View
from openpyxl import Workbook, load_workbook
from core.permissions import require_roles, PLAN_ROLES
from .models import EquipmentType, EquipmentCategory

HEADERS = ['Название', 'Категория', 'Единица измерения', 'Активна']
MAX_ROWS = 5000


class EquipmentImportForm(forms.Form):
    file = forms.FileField(label='Файл Excel (.xlsx)', widget=forms.FileInput(attrs={'class':'form-control', 'accept':'.xlsx'}))

    def clean_file(self):
        file = self.cleaned_data['file']
        if not file.name.lower().endswith('.xlsx'):
            raise forms.ValidationError('Выберите файл .xlsx. Старый формат .xls нужно пересохранить.')
        if file.size > 5 * 1024 * 1024:
            raise forms.ValidationError('Максимальный размер файла — 5 МБ.')
        return file


def text_value(value):
    return '' if value is None else str(value).strip()


def read_rows(file):
    errors, rows = [], []
    try:
        with ZipFile(file) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                return [], ['Файл слишком большой после распаковки.']
        file.seek(0)
        workbook = load_workbook(file, read_only=True, data_only=False)
    except Exception:
        return [], ['Не удалось прочитать Excel. Проверьте, что это исправный файл .xlsx без пароля.']
    try:
        sheet = workbook.worksheets[0]
        # Ignore inaccurate worksheet dimensions from third-party exports.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        header = next(iterator, ())
        columns = {text_value(cell.value).casefold(): index for index, cell in enumerate(header)}
        if 'название' not in columns:
            return [], ['В первой строке должен быть заголовок «Название». Скачайте шаблон.']
        recognized = [text_value(cell.value).casefold() for cell in header if text_value(cell.value).casefold() in {h.casefold() for h in HEADERS}]
        if len(recognized) != len(set(recognized)):
            return [], ['Заголовки столбцов не должны повторяться.']
        for number, cells in enumerate(iterator, 2):
            if number > MAX_ROWS + 1:
                errors.append(f'В файле допускается не более {MAX_ROWS} строк данных.')
                break
            if not any(text_value(cell.value) for cell in cells):
                continue
            values = []
            for title in HEADERS:
                index = columns.get(title.casefold())
                cell = cells[index] if index is not None and index < len(cells) else None
                if cell is not None and cell.data_type in ('f', 'e'):
                    errors.append(f'Строка {number}, «{title}»: замените формулу или ошибку Excel обычным значением.')
                values.append(text_value(cell.value) if cell is not None else '')
            name, category, unit, active = values
            unit = unit or 'ед.'
            for title, value, limit in [('Название',name,150), ('Категория',category,100), ('Единица измерения',unit,50)]:
                if len(value) > limit:
                    errors.append(f'Строка {number}: «{title}» — не более {limit} символов.')
            if not name:
                errors.append(f'Строка {number}: укажите название.')
            token = active.casefold()
            if token not in ('','да','нет','1','0','true','false','yes','no','активна','неактивна'):
                errors.append(f'Строка {number}: «Активна» должна быть «Да» или «Нет».')
            rows.append((name, category, unit, token not in ('нет','0','false','no','неактивна')))
        if not rows and not errors:
            errors.append('В файле нет данных для загрузки.')
        return rows, errors
    except Exception:
        return [], ['Не удалось прочитать строки Excel. Проверьте файл или перенесите данные в шаблон.']
    finally:
        workbook.close()


@transaction.atomic
def import_rows(company, rows):
    from apps.accounts.models import Company
    Company.objects.select_for_update().get(pk=company.pk)
    existing = {name.strip().casefold() for name in EquipmentType.objects.filter(company=company).values_list('name', flat=True)}
    categories = {category.name.strip().casefold(): category for category in EquipmentCategory.objects.filter(company=company)}
    created = skipped = added_categories = 0
    for name, category_name, unit, active in rows:
        if name.casefold() in existing:
            skipped += 1
            continue
        category = None
        if category_name:
            category = categories.get(category_name.casefold())
            if category is None:
                category = EquipmentCategory.objects.create(company=company, name=category_name)
                categories[category_name.casefold()] = category
                added_categories += 1
        EquipmentType.objects.create(company=company, name=name, category=category, unit=unit, is_active=active)
        existing.add(name.casefold())
        created += 1
    return {'created':created, 'skipped':skipped, 'categories':added_categories}


class EquipmentTypeImport(View):
    def get(self, request):
        require_roles(request.user, PLAN_ROLES)
        return render(request, 'resources/equipment_type_import.html', {'form':EquipmentImportForm()})

    def post(self, request):
        require_roles(request.user, PLAN_ROLES)
        form = EquipmentImportForm(request.POST, request.FILES)
        errors, result = [], None
        if form.is_valid():
            rows, errors = read_rows(form.cleaned_data['file'])
            if not errors:
                result = import_rows(request.user.company, rows)
        return render(request, 'resources/equipment_type_import.html', {'form':form, 'errors':errors[:100], 'error_count':len(errors), 'result':result}, status=400 if errors or form.errors else 200)


def equipment_import_template(request):
    require_roles(request.user, PLAN_ROLES)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Виды техники'
    sheet.append(HEADERS)
    sheet.freeze_panes = 'A2'
    for column, width in [('A',45),('B',30),('C',24),('D',16)]:
        sheet.column_dimensions[column].width = width
    guide = workbook.create_sheet('Инструкция')
    for line in [
        'Заполните первый лист. Заголовки первой строки не изменяйте.',
        'Название обязательно. Категория и единица измерения необязательны.',
        'Новые категории создаются автоматически в вашей компании.',
        'Единица измерения по умолчанию: ед. Активна: Да или Нет, по умолчанию Да.',
        'Существующие названия и повторы в файле пропускаются без обновления.',
        'При ошибке в любой строке весь импорт отменяется. Не более 5000 строк.',
    ]:
        guide.append([line])
    guide.column_dimensions['A'].width = 110
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="equipment_types_template.xlsx"'
    return response
