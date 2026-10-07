from copy import deepcopy
from datetime import date
from decimal import Decimal as D
from io import BytesIO
from unittest.mock import patch
from django.test import TestCase
from django.urls import reverse
from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from apps.planning import test_workspace as fixtures
from apps.planning.workspace_services import WorkspaceService
from apps.planning.models import ResourceMonthAllocation, WorkMonthAllocation
from apps.production.models import DailyFact, LaborFact, EquipmentFact, FuelFact
from apps.works.models import WorkPrice
from apps.works.prices import change_price, price_on
from .revenue import build_dashboard
from .revenue_detail import work_detail

JAN = fixtures.JAN
END = date(2026, 1, 31)


class RevenueTests(TestCase):
    setUpTestData = classmethod(fixtures.WorkspaceTests.setUpTestData.__func__)
    create = fixtures.WorkspaceTests.create
    approve = fixtures.WorkspaceTests.approve

    def data(self, **kwargs):
        return {
            "start": JAN,
            "end": END,
            "mode": "latest",
            "sections": ["works", "labor", "equipment", "fuel"],
            **kwargs,
        }

    def test_month_plan_and_cutoff_percentages_and_price_history(self):
        work = self.simple
        workspace = self.create([work])
        self.approve(workspace.baseline_version)
        workspace.baseline_version.refresh_from_db()
        frozen = deepcopy(workspace.baseline_version.snapshot)
        DailyFact.objects.create(
            company=self.company, project_work=work, date=JAN, actual_quantity=4
        )
        original = work.price_history.first()
        change_price(
            self.planner, work.pk, D(20), date.min, "Correct original price", original
        )
        report = build_dashboard(self.planner, self.data(), today=date(2026, 1, 15))
        row = report["works"][0]
        self.assertEqual(
            row["plan"], sum((r["plan_quantity"] or 0 for r in row["daily"]), D(0)) * 10
        )
        self.assertEqual(report["fact"], D(80))
        self.assertEqual(report["percent_date"], D(80) / report["plan_to_date"] * 100)
        self.assertEqual(report["percent_full"], D(80) / report["plan"] * 100)
        self.assertEqual(
            report["delta"], report["volume_delta"] + report["price_delta"]
        )
        workspace.baseline_version.refresh_from_db()
        self.assertEqual(workspace.baseline_version.snapshot, frozen)
        self.assertEqual(work.price_history.count(), 2)

    def test_composite_is_counted_once_and_uses_prior_history(self):
        workspace = self.create([self.composite])
        self.approve(workspace.baseline_version)
        for item, day, qty in [
            (self.a, date(2025, 12, 31), 10),
            (self.b, date(2025, 12, 31), 12),
            (self.b, JAN, 3),
        ]:
            DailyFact.objects.create(
                company=self.company,
                project_work=self.composite,
                work_item=item,
                date=day,
                actual_quantity=qty,
            )
        report = build_dashboard(self.planner, self.data(), today=JAN)
        self.assertEqual(
            report["fact"], D(100)
        )  # 4 -> 5 units, not sum of item revenue.
        self.assertEqual(len(report["works"]), 1)
        detail = work_detail(self.planner, report["works"][0], JAN, JAN)
        self.assertEqual([c["fact"] for c in detail["children"]], [D(10), D(15)])
        self.assertTrue(all(c["limiting"] for c in detail["children"]))
        self.assertEqual(detail["records_count"], 1)

    def test_resources_are_snapshot_at_period_end_and_missing_is_not_zero(self):
        workspace = self.create([self.simple])
        for kind, kwargs in [
            ("labor", {"brigade": self.brigade, "count": 10}),
            ("equipment", {"equipment_type": self.equipment, "count": 3}),
            ("fuel", {"fuel_type": "DIESEL", "balance": 500, "liters": 310}),
        ]:
            ResourceMonthAllocation.objects.create(
                company=self.company,
                version=workspace.baseline_version,
                kind=kind,
                month=JAN,
                **kwargs,
            )
        self.approve(workspace.baseline_version)
        LaborFact.objects.create(
            company=self.company,
            construction_object=self.obj,
            brigade=self.brigade,
            date=END,
            actual_workers=0,
        )
        FuelFact.objects.create(
            company=self.company,
            construction_object=self.obj,
            date=END,
            fuel_type="DIESEL",
            actual_balance=400,
            actual_liters=20,
        )
        report = build_dashboard(self.planner, self.data(), today=date(2026, 2, 5))
        self.assertEqual(report["cutoff"], END)
        self.assertEqual(report["resources"]["labor"]["fact"], D(0))
        self.assertIsNone(report["resources"]["equipment"]["fact"])
        self.assertEqual(report["resources"]["fuel"]["plan"], D(500))
        self.assertEqual(report["resources"]["fuel"]["fact"], D(400))
        self.assertTrue(
            all(r["label"] != "Расход" for r in report["resource_objects"][0]["rows"])
        )

    def test_future_period_has_no_actuals_or_resources(self):
        report = build_dashboard(self.planner, self.data(), today=date(2025, 12, 1))
        self.assertTrue(report["future"])
        self.assertEqual(report["fact"], 0)
        self.assertEqual(report["resource_objects"], [])
        self.assertIsNone(report["plan_to_date"])

    def test_prices_schedule_audit_correction_and_validation(self):
        work = self.simple
        with patch("apps.works.prices.timezone.localdate", return_value=JAN):
            change_price(self.planner, work.pk, D(12), JAN, "New price")
            change_price(self.planner, work.pk, D(20), JAN.replace(day=15), "Future")
            with self.assertRaises(ValidationError):
                change_price(self.planner, work.pk, D(1), date(2025, 12, 1))
        self.assertEqual(price_on(work, JAN), D(12))
        self.assertEqual(price_on(work, JAN.replace(day=15)), D(20))
        entry = work.price_history.last()
        entry.price = 2
        with self.assertRaises(ValidationError):
            entry.save()
        with self.assertRaises(ValidationError):
            entry.delete()
        self.assertEqual(work.price_history.count(), 3)
        with self.assertRaises(ValidationError):
            change_price(self.planner, work.pk, D(-1), date(2027, 1, 1))

    def test_ui_drill_export_and_company_filters(self):
        workspace = self.create([self.simple, self.composite])
        self.approve(workspace.baseline_version)
        self.client.force_login(self.planner)
        params = {"start": str(JAN), "end": str(END), "work": self.composite.pk}
        response = self.client.get(reverse("dashboard:revenue"), params)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Разбор отклонения")
        self.assertContains(response, "Выполнение всего периода")
        self.assertContains(response, "История цен")
        self.assertNotContains(response, "Расход ГСМ")
        params["work"] = 999999
        self.assertEqual(
            self.client.get(reverse("dashboard:revenue"), params).status_code, 403
        )
        response = self.client.get(
            reverse("dashboard:revenue_export"), {"start": str(JAN), "end": str(END)}
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["Выручка", "Ресурсы"])
        from apps.accounts.models import Company
        from apps.projects.models import Project

        other = Company.objects.create(name="other")
        project = Project.objects.create(company=other, name="Secret")
        response = self.client.get(
            reverse("dashboard:revenue"), {"project": project.pk}
        )
        self.assertContains(response, "Исправьте фильтры")
        self.assertNotContains(response, "Secret")

    def test_price_ui_rights_and_forged_edit_price(self):
        self.client.force_login(self.manager)
        self.assertEqual(
            self.client.get(
                reverse("works:price_change", args=[self.simple.pk])
            ).status_code,
            403,
        )
        self.client.force_login(self.planner)
        self.assertEqual(
            self.client.get(
                reverse("works:price_change", args=[self.simple.pk])
            ).status_code,
            200,
        )
        response = self.client.post(
            reverse("works:price_change", args=[self.simple.pk]),
            {"price": 12, "effective_from": "2027-01-01", "reason": "Contract"},
        )
        self.assertEqual(response.status_code, 302)
        entry = self.simple.price_history.last()
        self.assertEqual(entry.created_by, self.planner)
        self.assertEqual(entry.price, D(12))
        original = self.simple.price_history.first()
        response = self.client.post(
            reverse("works:price_change", args=[self.simple.pk]),
            {"price": 11, "reason": "Fix original", "corrects": original.pk},
        )
        self.assertEqual(response.status_code, 302)
        corrected = self.simple.price_history.get(corrects=original)
        self.assertEqual(corrected.effective_from, date.min)
        self.assertEqual(corrected.created_by, self.planner)

    def test_foreign_user_and_unapproved_plan(self):
        workspace = self.create([self.simple])
        report = build_dashboard(self.planner, self.data(), today=JAN)
        self.assertIsNone(report["plan"])
        from apps.accounts.models import User, Role

        master = User.objects.create_user(
            username="revenue-master",
            company=self.company,
            role=Role.objects.get(code="FOREMAN"),
        )
        self.client.force_login(master)
        response = self.client.get(reverse("dashboard:revenue"))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("production:fact_day_workspace"))

    def test_zero_plan_percent_is_undefined(self):
        workspace = self.create([self.simple])
        workspace.baseline_version.work_allocations.update(quantity=0)
        self.approve(workspace.baseline_version)
        report = build_dashboard(self.planner, self.data(), today=JAN)
        self.assertEqual(report["plan"], 0)
        self.assertIsNone(report["percent_full"])

    def test_frozen_scheduled_prices_and_daily_actual_rates(self):
        with patch("apps.works.prices.timezone.localdate", return_value=JAN):
            scheduled = change_price(
                self.planner,
                self.simple.pk,
                D(20),
                JAN.replace(day=15),
                "New agreement",
            )
        workspace = self.create([self.simple])
        self.approve(workspace.baseline_version)
        for day in [JAN.replace(day=14), JAN.replace(day=15)]:
            DailyFact.objects.create(
                company=self.company,
                project_work=self.simple,
                date=day,
                actual_quantity=1,
            )
        change_price(
            self.planner,
            self.simple.pk,
            D(30),
            scheduled.effective_from,
            "Correction",
            scheduled,
        )
        report = build_dashboard(self.planner, self.data(), today=JAN.replace(day=15))
        row = report["works"][0]
        self.assertEqual(row["daily"][13]["plan_rate"], D(10))
        self.assertEqual(row["daily"][14]["plan_rate"], D(20))
        self.assertEqual(row["daily"][14]["fact_rate"], D(30))
        self.assertEqual(report["fact"], D(40))
        self.assertEqual(report["price_delta"], D(10))

    def test_fact_after_cutoff_is_excluded(self):
        workspace = self.create([self.simple])
        self.approve(workspace.baseline_version)
        for day, qty in [(JAN, 2), (JAN.replace(day=20), 100)]:
            DailyFact.objects.create(
                company=self.company,
                project_work=self.simple,
                date=day,
                actual_quantity=qty,
            )
        report = build_dashboard(self.planner, self.data(), today=JAN.replace(day=15))
        self.assertEqual(report["fact"], D(20))
        self.assertIsNone(report["works"][0]["daily"][19]["fact"])

    def test_explicit_version_and_legacy_snapshot_price(self):
        workspace = self.create([self.simple])
        self.approve(workspace.baseline_version)
        version = workspace.baseline_version
        version.refresh_from_db()
        snapshot = deepcopy(version.snapshot)
        for spec in snapshot["works"]:
            spec.pop("revenue_prices", None)
        from apps.planning.models import GlobalPlanVersion

        GlobalPlanVersion.objects.filter(pk=version.pk).update(snapshot=snapshot)
        version.refresh_from_db()
        original = self.simple.price_history.first()
        change_price(
            self.planner,
            self.simple.pk,
            D(99),
            original.effective_from,
            "Fix",
            original,
        )
        report = build_dashboard(
            self.planner, self.data(**{f"version_{self.obj.pk}": version}), today=JAN
        )
        self.assertEqual(report["works"][0]["daily"][0]["plan_rate"], D(10))

    def test_filter_applies_to_revenue_and_export(self):
        workspace = self.create([self.simple, self.composite])
        self.approve(workspace.baseline_version)
        report = build_dashboard(self.planner, self.data(works_q="Simple"), today=JAN)
        self.assertEqual([w["id"] for w in report["works"]], [self.simple.pk])
        self.client.force_login(self.planner)
        response = self.client.get(
            reverse("dashboard:revenue_export"),
            {"start": str(JAN), "end": str(END), "works_q": "Simple"},
        )
        book = load_workbook(BytesIO(response.content))
        text = str(list(book["Выручка"].values))
        self.assertIn("Simple", text)
        self.assertNotIn("Composite", text)

    def test_unplanned_work_is_volume_deviation_not_price_deviation(self):
        workspace = self.create([self.simple])
        self.approve(workspace.baseline_version)
        DailyFact.objects.create(
            company=self.company,
            project_work=self.composite,
            work_item=self.a,
            date=JAN,
            actual_quantity=2,
        )
        DailyFact.objects.create(
            company=self.company,
            project_work=self.composite,
            work_item=self.b,
            date=JAN,
            actual_quantity=3,
        )
        report = build_dashboard(self.planner, self.data(), today=JAN)
        row = next(w for w in report["works"] if w["id"] == self.composite.pk)
        self.assertEqual(row["volume_delta"], D(100))
        self.assertEqual(row["price_delta"], D(0))
