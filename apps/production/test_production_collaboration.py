from datetime import timedelta
from uuid import uuid4
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from core.fact_history import history_actor
from .test_production_entry import ProductionEntryTests
from .dispatch_models import EmployeeActivity,EmployeeWorkLink,WorkFactClaim,EmployeeFactClaim,ProductionPresence,FactSectionDay
from .models import DailyFact
from .history_models import FactRevision


class ProductionCollaborationTests(TestCase):
    setUpTestData=classmethod(ProductionEntryTests.setUpTestData.__func__)
    payload=ProductionEntryTests.payload

    def setUp(self):
        from apps.planning.models import DailyScheduleEntry
        DailyScheduleEntry.objects.create(schedule=self.schedule,key='labor:group',kind='labor',label=self.brigade.name,rows=[{'date':str(self.day),'brigade_id':self.brigade.pk,'planned_count':'3'}])
        self.user=self.users['DISPATCHER'];self.client.force_login(self.user)
        self.session=str(uuid4());self.scope={'construction_object':self.obj.pk,'date':str(self.day),'session':self.session}
        self.person=EmployeeActivity.objects.create(company=self.company,construction_object=self.obj,date=self.day,employee=self.employee,position=self.employee.position,assignment_reviewed=False,entry_direction='works',reported_by=self.user)
        self.fact=DailyFact.objects.create(company=self.company,date=self.day,project_work=self.work,actual_quantity=5,reported_by=self.user)
        self.second_fact=DailyFact.objects.create(company=self.company,date=self.day,project_work=self.second_work,actual_quantity=3,reported_by=self.user)

    def edge(self,action='add',**kwargs):
        return self.client.post(reverse('production:production_links'),{**self.scope,'activity':self.person.pk,'fact':self.fact.pk,'action':action,**kwargs},content_type='application/json')

    def heartbeat(self,**kwargs):
        return self.client.post(reverse('production:production_live'),{**self.scope,**kwargs},content_type='application/json')

    def claim(self,key):
        response=self.client.post(reverse('production:work_claims'),{**self.scope,'key':key,'action':'take'},content_type='application/json')
        self.assertEqual(response.status_code,200,response.content)

    def test_relationships_ignore_work_and_employee_leases_and_preserve_data_version(self):
        self.claim(f'employee:{self.employee.pk}:1');self.claim(f'work:{self.work.pk}')
        original=self.person.updated_at
        self.client.force_login(self.second)
        response=self.edge();self.assertEqual(response.status_code,200,response.content)
        self.person.refresh_from_db();self.assertEqual(self.person.updated_at,original)
        self.assertEqual(self.person.links_revision,1);self.assertTrue(self.person.assignment_reviewed)
        self.assertEqual(self.person.reported_by_id,self.user.pk)
        self.assertTrue(FactRevision.objects.filter(entity_type='production.employeeworklink',actor=self.second,operation='CREATE').exists())

    def test_distinct_edges_merge_and_same_edge_is_idempotent(self):
        self.assertEqual(self.edge().status_code,200);self.client.force_login(self.second)
        self.assertEqual(self.edge(fact=self.second_fact.pk).status_code,200);self.assertEqual(self.edge().status_code,200)
        self.assertEqual(self.person.work_links.count(),2)
        self.assertEqual(self.edge('remove').status_code,200)
        self.person.refresh_from_db();self.assertEqual(self.person.work_id,self.second_work.pk);self.assertTrue(self.person.assignment_reviewed)

    def test_clear_all_rejects_stale_revision_without_losing_new_edge(self):
        self.edge();response=self.edge('mode',mode='general',revision=0)
        self.assertEqual(response.status_code,409);self.assertEqual(self.person.work_links.count(),1)
        self.person.refresh_from_db();response=self.edge('mode',mode='general',revision=self.person.links_revision)
        self.assertEqual(response.status_code,200);self.person.refresh_from_db();self.assertTrue(self.person.general);self.assertFalse(self.person.work_links.exists())

    def test_delta_is_atomic_if_one_work_was_deleted(self):
        response=self.edge('delta',revision=0,edges=[{'action':'add','fact':self.fact.pk},{'action':'add','fact':999999}])
        self.assertEqual(response.status_code,409);self.assertFalse(self.person.work_links.exists())

    def test_zero_does_not_remove_links_or_reason(self):
        self.edge();data=self.payload('DISPATCHER','works');data.update(collaborative='1',autosave='1')
        prefix=next(k[:-len('project_work')] for k,v in data.items() if k.endswith('-project_work') and str(v)==str(self.work.pk))
        data[prefix+'actual_quantity']='0';data[prefix+'deviation_reason']=self.reason.pk
        response=self.client.post(self.url,data);self.assertEqual(response.status_code,200,response.content)
        self.fact.refresh_from_db();self.assertEqual(self.fact.actual_quantity,0);self.assertEqual(self.fact.deviation_reason_id,self.reason.pk);self.assertTrue(self.person.work_links.exists())

    def test_person_delete_removes_links_but_keeps_work_and_dictionary_and_history(self):
        self.edge();data=self.payload('DISPATCHER','works');data.update({'autosave':'1','collaborative':'1','labor-0-DELETE':'on'})
        identity=self.person.pk;response=self.client.post(self.url,data);self.assertEqual(response.status_code,200,response.content)
        self.assertFalse(EmployeeActivity.objects.filter(pk=identity).exists());self.assertFalse(EmployeeWorkLink.objects.exists())
        self.assertTrue(DailyFact.objects.filter(pk=self.fact.pk).exists());self.assertTrue(type(self.employee).objects.filter(pk=self.employee.pk,is_active=True).exists())
        self.assertTrue(FactRevision.objects.filter(entity_type='production.employeeactivity',entity_id=str(identity),actor=self.user,operation='DELETE').exists())
        self.assertEqual(self.edge().status_code,409)

    def test_work_delete_clears_only_its_edges_and_never_marks_general(self):
        from apps.works.models import ProjectWork
        outside=ProjectWork.objects.create(company=self.company,section=self.work.section,name='Вне плана',unit='м')
        fact=DailyFact.objects.create(company=self.company,date=self.day,project_work=outside,actual_quantity=1,deviation_reason=self.reason,outside_smg=True)
        self.edge(fact=fact.pk)
        data=self.payload('DISPATCHER','works');prefix=next(k[:-len('project_work')] for k,v in data.items() if k.endswith('-project_work') and str(v)==str(outside.pk))
        data.update({'autosave':'1','collaborative':'1',prefix+'DELETE':'on'})
        response=self.client.post(self.url,data);self.assertEqual(response.status_code,200,response.content)
        self.person.refresh_from_db();self.assertFalse(self.person.general);self.assertFalse(self.person.assignment_reviewed);self.assertFalse(self.person.work_links.exists())
        self.assertEqual(self.edge(fact=fact.pk).status_code,409)

    def test_planned_work_cannot_be_deleted(self):
        data=self.payload('DISPATCHER','works');prefix=next(k[:-len('project_work')] for k,v in data.items() if k.endswith('-project_work') and str(v)==str(self.work.pk))
        data.update({'autosave':'1',prefix+'DELETE':'on'})
        response=self.client.post(self.url,data);self.assertEqual(response.status_code,400);self.assertTrue(DailyFact.objects.filter(pk=self.fact.pk).exists())

    def test_data_edit_does_not_restore_stale_relationships(self):
        data=self.payload('DISPATCHER','works');data.update({'autosave':'1','collaborative':'1','labor-0-activity':'Комментарий диспетчера'})
        self.client.force_login(self.second);self.edge();self.client.force_login(self.user)
        response=self.client.post(self.url,data);self.assertEqual(response.status_code,200,response.content)
        self.person.refresh_from_db();self.assertEqual(self.person.activity,'Комментарий диспетчера');self.assertEqual(self.person.work_id,self.work.pk);self.assertTrue(self.person.work_links.exists())

    def test_presence_deduplicates_same_user_tabs_and_does_not_claim(self):
        self.assertEqual(self.heartbeat(step=2).status_code,200)
        response=self.heartbeat(session=str(uuid4()),step=3);self.assertEqual(response.status_code,200)
        users=response.json()['users'];self.assertEqual(len(users),1);self.assertEqual(users[0]['steps'],[2,3]);self.assertTrue(users[0]['self']);self.assertFalse(users[0]['editing'])
        self.assertFalse(WorkFactClaim.objects.exists());self.assertFalse(EmployeeFactClaim.objects.exists())

    def test_idle_heartbeat_renews_claim_and_internal_tab_switch_keeps_it(self):
        self.heartbeat();self.claim(f'work:{self.work.pk}')
        WorkFactClaim.objects.update(expires_at=timezone.now()+timedelta(seconds=2))
        self.heartbeat(step=4);self.assertGreater(WorkFactClaim.objects.get().expires_at,timezone.now()+timedelta(minutes=4))

    def test_leaving_production_and_crash_release_session_only(self):
        self.heartbeat();self.claim(f'work:{self.work.pk}');self.claim(f'employee:{self.employee.pk}:1')
        self.heartbeat(action='close');ProductionPresence.objects.update(closing_at=timezone.now()-timedelta(seconds=6))
        self.client.force_login(self.second);response=self.heartbeat(session=str(uuid4()))
        self.assertEqual(response.status_code,200);self.assertFalse(WorkFactClaim.objects.exists());self.assertFalse(EmployeeFactClaim.objects.exists())
        self.assertEqual(len(response.json()['users']),1)

    def test_live_rebase_updates_only_unprotected_rows(self):
        page=self.client.get(self.url,{'construction_object':self.obj.pk,'date':str(self.day),'section':'works'})
        signed=page.context['production_snapshot']
        response=self.heartbeat(production_snapshot=signed,keep=[f'employee:{self.employee.pk}:1'])
        self.assertEqual(response.status_code,200,response.content);self.assertIn('rows',response.json());self.assertIn('production_snapshot',response.json())
        second=self.heartbeat(digest=response.json()['digest']);self.assertNotIn('rows',second.json())

    def test_other_object_and_date_and_roles_are_scoped(self):
        self.assertEqual(self.edge(construction_object=self.other.pk).status_code,400)
        self.assertEqual(self.edge(date=str(self.day-timedelta(days=10))).status_code,400)
        self.client.force_login(self.users['PERSONNEL_DISPATCHER']);self.assertEqual(self.edge().status_code,403)

    def test_no_manual_claim_or_confirmation_is_visible_and_plan_groups_exist(self):
        response=self.client.get(self.url,{'construction_object':self.obj.pk,'date':str(self.day),'section':'works'})
        self.assertNotContains(response,'data-confirm-production');self.assertContains(response,'data-production-filters="people"');self.assertContains(response,'data-production-presence')
        self.assertContains(response,'data-position-group=')

    def test_excel_is_formatted_real_workbook_and_preserves_text_as_text(self):
        from openpyxl import load_workbook
        from io import BytesIO
        self.edge()
        response=self.client.get(reverse('production:production_report_download'),{'construction_object':self.obj.pk,'date':str(self.day)})
        self.assertEqual(response.status_code,200)
        self.assertIn('attachment',response['Content-Disposition'])
        book=load_workbook(BytesIO(b''.join(response.streaming_content)),rich_text=True)
        sheet=book.active
        self.assertTrue(sheet['A1'].font.bold);self.assertGreater(len(sheet.merged_cells.ranges),10)
        text=' '.join(str(cell.value or '') for row in sheet for cell in row)
        self.assertIn(self.employee.display_name,text);self.assertIn(self.brigade.name,text)
        self.assertIn('Работы и сотрудники',text);self.assertIn('Сотрудники',text)
        self.assertEqual(sheet.page_setup.fitToWidth,1)

    def test_excel_does_not_interpret_employee_names_as_formulas(self):
        from .report_export import workbook
        from openpyxl import load_workbook
        data={'works':[],'people':[{'employee':1,'name':'=HYPERLINK("https://example.invalid")','position':'@Position','group':'-Group','shift':1,'general':True,'reviewed':True,'targets':[],'owner':'+Owner'}]}
        sheet=load_workbook(workbook(self.obj,self.day,data)).active
        for row in sheet:
            for cell in row:self.assertNotEqual(cell.data_type,'f')
