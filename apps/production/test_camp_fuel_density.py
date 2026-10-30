from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from django.test import TestCase
from django.urls import reverse
from django.core.exceptions import ValidationError
from django.db import IntegrityError,transaction
from openpyxl import load_workbook
from . import test_camp_fuel as fixtures,camp_fuel as service
from .camp_fuel_models import CampFuelDay,mass_tonnes
from .history_models import FactRevision

class FuelDensityTests(TestCase):
    setUpTestData=classmethod(fixtures.CampFuelTests.setUpTestData.__func__)
    save=fixtures.CampFuelTests.save
    def test_distinct_densities_calculate_masses_in_tonnes(self):
        row=self.save(opening=100,received=20,consumed=10,opening_density='0.840',closing_density='0.830')
        self.assertEqual(row.opening_tonnes,Decimal('84.000'))
        self.assertEqual(row.closing_tonnes,Decimal('91.300'))
        row.refresh_from_db();self.assertEqual(row.opening_density,Decimal('.840000'))
        self.assertEqual(row.actual_balance,110)
    def test_missing_density_is_unknown_even_for_zero_volume(self):
        row=self.save(opening=10)
        self.assertIsNone(row.opening_density);self.assertIsNone(row.opening_tonnes);self.assertIsNone(row.closing_tonnes)
        self.assertIsNone(mass_tonnes(0,None))
        self.assertEqual(mass_tonnes(0,Decimal('.840')),0)
    def test_mass_rounding_is_half_up_and_keeps_full_density_precision(self):
        self.assertEqual(mass_tonnes(Decimal('.001'),Decimal('.5')),Decimal('.001'))
        self.assertEqual(mass_tonnes(Decimal('1.125'),Decimal('.840123')),Decimal('.945'))
        self.assertEqual(mass_tonnes(Decimal('999999999999.999'),Decimal('999.999999')),Decimal('999999998999999.000'))
    def test_next_day_prefills_previous_density_but_keeps_saved_measurements(self):
        self.save(opening=100,closing_density='.830',opening_density='.840')
        next_day=self.day+timedelta(days=1)
        row=self.save(day=next_day,consumed=10,closing_density='.820')
        self.assertEqual(row.opening_density,Decimal('.830'));self.assertEqual(row.opening_tonnes,83)
        self.assertEqual(row.closing_tonnes,Decimal('73.800'))
        row=self.save(day=next_day,opening_density='.835',consumed=10,closing_density='.825')
        self.save(received=20,opening_density='.841',closing_density='.831')
        row.refresh_from_db();self.assertEqual(row.opening_density,Decimal('.835'));self.assertEqual(row.closing_density,Decimal('.825'))
        self.assertEqual(row.opening_balance,120);self.assertEqual(row.opening_tonnes,Decimal('100.200'))
        self.assertEqual(row.closing_tonnes,Decimal('90.750'))
    def test_omitted_density_in_service_edit_preserves_it_and_explicit_blank_clears(self):
        self.save(opening=10,opening_density='.840',closing_density='.830')
        row=self.save(received=1);self.assertEqual(row.opening_density,Decimal('.840'));self.assertEqual(row.closing_density,Decimal('.830'))
        row=self.save(received=1,opening_density=None,closing_density=None)
        self.assertIsNone(row.opening_tonnes);self.assertIsNone(row.closing_tonnes)
    def test_invalid_density_rejected_atomically(self):
        self.save(opening=10,opening_density='.840')
        for value in ['0','-0.8','NaN','Infinity','0.8400001','1000']:
            with self.subTest(value=value),self.assertRaises(ValidationError):self.save(received=20,closing_density=value)
        row=CampFuelDay.objects.get();self.assertEqual(row.received_liters,0);self.assertIsNone(row.closing_density)
    def test_positive_density_database_constraint(self):
        with self.assertRaises(IntegrityError),transaction.atomic():
            CampFuelDay.objects.create(company=self.company,camp=self.camp,date=self.day,opening_density=0)
    def test_density_changes_are_audited_and_invalidate_stale_forms(self):
        self.save(opening=10,opening_density='.840',closing_density='.830')
        old=service.token(self.operator,self.camp,self.day,'TOTAL')
        self.save(user=self.colleague,opening_density='.845',closing_density='.835')
        change=FactRevision.objects.filter(entity_type='production.campfuelday',operation='UPDATE').latest('pk')
        self.assertEqual(change.actor_id,self.colleague.pk);self.assertEqual(Decimal(change.after['fields']['opening_density']),Decimal('.845'))
        with self.assertRaisesMessage(ValidationError,'другой диспетчер'):self.save(signed=old,received=1)
    def test_page_and_excel_have_density_and_computed_tonnes_and_ignore_forged_masses(self):
        self.client.force_login(self.operator);url=reverse('production:camp_fuel');params={'camp':self.camp.pk,'date':str(self.day)}
        response=self.client.get(url,params);data=params|{'opening':'100','received':'20','consumed':'10','opening_density':'0.840','closing_density':'0.830','opening_tonnes':'999','closing_tonnes':'999','signed':response.context['form'].initial['signed']}
        self.assertEqual(self.client.post(url,data).status_code,302)
        response=self.client.get(url,params);self.assertEqual(response.context['opening_tonnes'],84);self.assertEqual(response.context['closing_tonnes'],Decimal('91.300'))
        self.assertContains(response,'data-fuel-opening-mass');self.assertContains(response,'т/м³')
        sheet=load_workbook(BytesIO(self.client.get(url,params|{'export':'xlsx'}).content),data_only=True).active
        self.assertEqual(sheet.cell(2,6).value,110);self.assertEqual(sheet.cell(2,9).value,.84)
        self.assertEqual(sheet.cell(2,10).value,.83);self.assertEqual(sheet.cell(2,11).value,84);self.assertEqual(sheet.cell(2,12).value,91.3)
        self.assertEqual(sheet.cell(2,12).number_format,'0.000')
    def test_new_day_form_prefills_only_start_density(self):
        self.save(opening=10,closing_density='.830')
        self.client.force_login(self.operator)
        page=self.client.get(reverse('production:camp_fuel'),{'camp':self.camp.pk,'date':str(self.day+timedelta(days=1))})
        self.assertEqual(page.context['form'].initial['opening_density'],Decimal('.830'))
        self.assertIsNone(page.context['form'].initial['closing_density'])
        self.assertFalse(page.context['form'].fields['opening_density'].disabled)
        self.assertEqual(page.context['opening_tonnes'],Decimal('8.300'))
        self.assertIsNone(page.context['closing_tonnes'])
    def test_read_only_day_still_displays_saved_density_and_mass(self):
        day=self.day-timedelta(days=3);self.save(user=self.admin,day=day,opening=10,opening_density='.84',closing_density='.83')
        self.client.force_login(self.operator);page=self.client.get(reverse('production:camp_fuel'),{'camp':self.camp.pk,'date':str(day)})
        self.assertNotContains(page,'data-camp-fuel-form');self.assertEqual(page.context['opening_tonnes'],Decimal('8.400'))
        self.assertEqual(page.context['closing_tonnes'],Decimal('8.300'))

    def test_omitted_post_fields_preserve_measured_density_and_blank_fields_clear(self):
        self.save(opening=10,opening_density='.840',closing_density='.830')
        self.client.force_login(self.operator);url=reverse('production:camp_fuel');params={'camp':self.camp.pk,'date':str(self.day)}
        form=self.client.get(url,params).context['form']
        data=params|{'opening':'10','received':'1','consumed':'0','signed':form.initial['signed']}
        self.assertEqual(self.client.post(url,data).status_code,302)
        row=CampFuelDay.objects.get();self.assertEqual(row.opening_tonnes,Decimal('8.400'));self.assertEqual(row.closing_tonnes,Decimal('9.130'))
        data.update(signed=self.client.get(url,params).context['form'].initial['signed'],opening_density='',closing_density='')
        self.assertEqual(self.client.post(url,data).status_code,302)
        row.refresh_from_db();self.assertIsNone(row.opening_tonnes);self.assertIsNone(row.closing_tonnes)
