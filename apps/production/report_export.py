"""A real, formatted XLSX export of the saved production day report."""
from io import BytesIO
import math
from django.contrib.auth.decorators import login_required
from django.http import FileResponse, JsonResponse
from django.views.decorators.http import require_GET
from core.permissions import require_roles,READ_ROLES
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from .forms import ObjectDateForm
from .entry_concurrency import sqlite_serialized
from .production_entry import report_data,work_catalogue,token

BLUE='173F57';TEAL='0087A4';PALE='EDF5FA';LINE='D7E5EE'


def workbook(obj,day,report):
    book=Workbook();sheet=book.active;sheet.title='Сводный отчёт'
    widths=[15,15,15,12,16,9,20,20,28]
    for n,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(n)].width=width
    sheet.sheet_view.showGridLines=False;sheet.freeze_panes='A9'
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    sheet.oddFooter.center.text='Страница &P из &N'
    sheet.print_options.horizontalCentered=True
    thin=Side(style='thin',color=LINE)
    row=1

    def write(line,start,end,value,*,fill=None,bold=False,size=11):
        if end>start:sheet.merge_cells(start_row=line,start_column=start,end_row=line,end_column=end)
        cell=sheet.cell(line,start,value)
        # Catalogue names and comments are text, never Excel formulas.
        if isinstance(value,str):cell.data_type='s'
        cell.font=Font(name='Calibri',size=size,color=BLUE,bold=bold)
        cell.alignment=Alignment(vertical='top',wrap_text=True)
        for col in range(start,end+1):
            target=sheet.cell(line,col);target.border=Border(bottom=thin)
            if fill:target.fill=PatternFill('solid',fgColor=fill)
        return cell

    def heading(label,level=1):
        nonlocal row
        write(row,1,9,label,fill=PALE if level==1 else 'F6F9FC',bold=True,size=12 if level==1 else 11)
        sheet.row_dimensions[row].height=25;row+=1

    write(row,1,9,f'Сводный отчёт за {day:%d.%m.%Y}',bold=True,size=20);sheet.row_dimensions[row].height=34;row+=1
    write(row,1,9,obj.name,bold=True,size=13);sheet.row_dimensions[row].height=25;row+=1
    write(row,1,9,'Фактические объёмы, сотрудники и связи. Объёмы разных единиц измерения не суммируются.',size=10);row+=2
    people=report['people'];works=report['works'];choices={w['token']:w for w in work_catalogue(obj)}
    cards=[('Работ',len(works),1,2),('Сотрудников',len({p['employee'] for p in people}),3,4),
        ('Без привязки',len({p['employee'] for p in people if p['general']}),5,6),('Связи не проверены',sum(not p['reviewed'] for p in people),7,9)]
    for label,value,start,end in cards:
        write(row,start,end,label,fill=PALE,size=10);write(row+1,start,end,value,fill=PALE,bold=True,size=18)
    sheet.row_dimensions[row+1].height=30;row+=3
    pending=[p['name'] for p in people if not p['reviewed']]
    write(row,1,9,'Проверьте занятость: '+', '.join(pending) if pending else 'Связи проверены.',fill='FFF4DA' if pending else 'EAF6F0',size=10)
    sheet.row_dimensions[row].height=max(22,15*math.ceil(len(', '.join(pending))/145));row+=2
    heading('Работы и сотрудники')
    for start,end,text in [(1,4,'Работа'),(5,5,'Объём'),(6,6,'Ед.'),(7,8,'Сотрудники и смены'),(9,9,'Ввод')]:write(row,start,end,text,fill=PALE,bold=True)
    row+=1;last=None
    for work in sorted(works,key=lambda w:(choices.get(token(w['project_work_id'],w['work_item_id']),{}).get('object',''),choices.get(token(w['project_work_id'],w['work_item_id']),{}).get('package',''),w['project_work_id'],w['work_item_id'] or 0)):
        key=token(work['project_work_id'],work['work_item_id']);meta=choices.get(key,{})
        group=(meta.get('object') or 'Без объекта')+' → '+(meta.get('package') or 'Самостоятельные работы')
        if group!=last:heading(group,2);last=group
        title=meta.get('title',meta.get('label',key));path=meta.get('label',key)
        prefix=path[:-len(title)].rstrip(' →') if path.endswith(title) else ''
        rich=CellRichText(prefix+'\n' if prefix else '',TextBlock(InlineFont(b=True,rFont='Calibri',color=BLUE),title))
        write(row,1,4,rich);write(row,5,5,work['actual_quantity']);sheet.cell(row,5).number_format='#,##0.###'
        write(row,6,6,meta.get('unit',''))
        linked='\n'.join(p['name']+f" · смена {p['shift']}" for p in people if key in p['targets']) or 'Сотрудники не привязаны'
        write(row,7,8,linked);write(row,9,9,work.get('owner') or 'Общий ввод')
        sheet.row_dimensions[row].height=max(38,15*max(math.ceil(len(path)/55)+1,math.ceil(len(linked)/39)))
        row+=1
    if not works:write(row,1,9,'Объёмы работ пока не внесены.');row+=1
    row+=1;heading('Сотрудники')
    for start,end,text in [(1,3,'Сотрудник'),(4,4,'Смена'),(5,8,'Выполненные работы / занятость'),(9,9,'Ввод')]:write(row,start,end,text,fill=PALE,bold=True)
    row+=1;last_group=last_role=None
    from apps.resources.group_order import order_key
    labor_group_order=order_key(obj.company)
    for person in sorted(people,key=lambda p:(labor_group_order(p.get('group','')),p.get('position',''),p['name'])):
        group=person.get('group') or 'Без группы плана';role=person.get('position') or 'Должность не указана'
        if group!=last_group:heading(group);last_group=group;last_role=None
        if role!=last_role:heading(role,2);last_role=role
        assignment='Связи не проверены' if not person['reviewed'] else 'Работал без привязки к работе' if person['general'] else '\n'.join(choices.get(t,{}).get('label',t) for t in person['targets'])
        name=person['name'];staff=person.get('staff_position')
        if staff and staff!=role:name+='\nШтатная должность: '+staff
        write(row,1,3,name);write(row,4,4,person['shift']);write(row,5,8,assignment);write(row,9,9,person['owner'])
        sheet.row_dimensions[row].height=max(38,15*max(math.ceil(len(name)/42),math.ceil(len(assignment)/78)))
        row+=1
    if not people:write(row,1,9,'Сотрудники пока не добавлены.');row+=1
    sheet.print_area=f'A1:I{row}';sheet.print_title_rows='1:2'
    output=BytesIO();book.save(output);output.seek(0)
    return output


@login_required
@require_GET
@sqlite_serialized
def download(request):
    require_roles(request.user,READ_ROLES|{'DISPATCHER','ALL_FACT_DISPATCHER'})
    selector=ObjectDateForm(request.GET,user=request.user)
    if not selector.is_valid():return JsonResponse({'error':'Выберите доступный объект и дату.'},status=400)
    obj=selector.cleaned_data['construction_object'];day=selector.cleaned_data['date']
    return FileResponse(workbook(obj,day,report_data(obj,day)),as_attachment=True,filename=f'Производство_{day:%Y-%m-%d}.xlsx',content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
