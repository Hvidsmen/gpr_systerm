from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from django.core.exceptions import ValidationError,PermissionDenied
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from .test_transport import TransportTests
from . import camp_fuel as service
from .camp_fuel_models import CampFuelDay
from .history_models import FactRevision


class CampFuelTests(TestCase):
    @classmethod
    def setUpTestData(cls):TransportTests.setUpTestData.__func__(cls)

    def save(self,day=None,user=None,camp=None,fuel='DIESEL',**kwargs):
        user=user or self.operator;camp=camp or self.camp;day=day or self.day
        return service.save_day(user,camp,day,fuel,signed=kwargs.pop('signed',service.token(user,camp,day,fuel)),received=kwargs.pop('received',0),consumed=kwargs.pop('consumed',0),**kwargs)

    def test_carry_forward_separates_camps_and_fuels_and_skipped_days(self):
        row=self.save(opening=100,received=50,consumed=20)
        self.assertEqual(row.actual_balance,130)
        later=self.save(day=self.day+timedelta(days=1),consumed=30,opening=999)
        self.assertEqual(later.opening_balance,130);self.assertEqual(later.actual_balance,100)
        self.save(fuel='PETROL_92',opening=20,consumed=1)
        self.save(camp=self.target,opening=50)
        self.assertEqual(CampFuelDay.objects.count(),4)
        self.assertEqual(service.balance_before(self.camp,self.day+timedelta(days=5),'DIESEL')[0].actual_balance,100)

    def test_fractional_liters_remain_balanced_in_database(self):
        row=self.save(opening='0.10',received='0.20',consumed='0.05')
        self.assertEqual(row.actual_balance,Decimal('0.25'))
        row=self.save(day=self.day+timedelta(days=1),consumed='0.02')
        self.assertEqual(row.actual_balance,Decimal('0.23'))

    def test_colleagues_edit_same_entry_with_audit_and_stale_form_rejected(self):
        self.save(opening=100)
        stale=service.token(self.operator,self.camp,self.day,'DIESEL')
        row=self.save(user=self.colleague,received=20,consumed=10,comment='Уточнение')
        self.assertEqual(row.recorded_by_id,self.operator.pk)
        change=FactRevision.objects.filter(entity_type='production.campfuelday',operation='UPDATE').get()
        self.assertEqual(change.actor_id,self.colleague.pk);self.assertEqual(change.direction,'fuel')
        self.assertEqual(change.before['fields']['received_liters'],'0.00')
        self.assertEqual(change.after['fields']['received_liters'],'20.00')
        with self.assertRaises(ValidationError):self.save(signed=stale,received=200)
        self.assertEqual(CampFuelDay.objects.get().received_liters,20)

    def test_backdated_correction_updates_future_balances_atomically(self):
        self.save(opening=100,consumed=10)
        self.save(day=self.day+timedelta(days=1),consumed=50)
        self.save(received=30,consumed=10)
        later=CampFuelDay.objects.order_by('date').last()
        self.assertEqual(later.opening_balance,120);self.assertEqual(later.actual_balance,70)
        with self.assertRaises(ValidationError):self.save(consumed=90)
        self.assertEqual(CampFuelDay.objects.order_by('date').first().received_liters,30)
        self.assertTrue(FactRevision.objects.filter(reason__startswith='Перенос остатка').exists())

    def test_negative_precision_invalid_fuel_and_closed_days_rejected(self):
        for data in [{'consumed':1},{'received':-1},{'received':'NaN'},{'received':'0.001'}]:
            with self.assertRaises(ValidationError):self.save(**data)
        with self.assertRaises(ValidationError):self.save(day=self.day-timedelta(days=2))
        with self.assertRaises(ValidationError):self.save(day=self.day+timedelta(days=2))
        self.save(user=self.admin,day=self.day+timedelta(days=2),opening=20)
        with self.assertRaises(ValidationError):self.save(user=self.admin,day=self.day)
        with self.assertRaises(PermissionDenied):self.save(camp=self.unrelated)

    def test_ui_export_and_view_only_day(self):
        self.client.force_login(self.operator);url=reverse('production:camp_fuel');params={'camp':self.camp.pk,'date':str(self.day),'fuel_type':'DIESEL'}
        response=self.client.get(url,params);self.assertEqual(response.status_code,200)
        data=params|{'opening':'100','received':'50','consumed':'20','signed':response.context['form'].initial['signed']}
        self.assertEqual(self.client.post(url,data).status_code,302)
        response=self.client.get(url,params|{'export':'xlsx'});self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(response.content),data_only=True);self.assertEqual(book.active.cell(2,7).value,130)
        self.assertContains(self.client.get(url,params),'130')
        closed=params|{'date':str(self.day-timedelta(days=2))}
        self.assertNotContains(self.client.get(url,closed),'data-camp-fuel-form')
        self.assertEqual(self.client.post(url,data|closed).status_code,403)
