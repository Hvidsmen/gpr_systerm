from datetime import date
from decimal import Decimal
from io import BytesIO
from copy import deepcopy

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import User, Role
from apps.planning import test_workspace
from apps.planning.test_meeting_import import upload, JAN
from apps.planning.meeting_import import parse_meeting_workbook
from apps.works.models import ProjectWork
from .meeting_import import resolved_preview, apply_facts
from .models import DailyFact, LaborFact, EquipmentFact


class FactMeetingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        test_workspace.WorkspaceTests.setUpTestData.__func__(cls)
        cls.admin = User.objects.create_user(
            username="fact-import-admin",
            company=cls.company,
            role=Role.objects.get(code="ADMIN"),
        )
        cls.foreman = User.objects.create_user(
            username="fact-import-foreman",
            company=cls.company,
            role=Role.objects.get(code="FOREMAN"),
        )

    def sheet(self, **kwargs):
        return parse_meeting_workbook(
            upload(**kwargs),
            JAN,
            date(2026, 1, 2),
            object_name=self.obj.name,
            facts=True,
        )[0]

    def payload(self, sheet=None, existing="keep"):
        sheet = sheet or self.sheet()
        return {
            "object": self.obj.pk,
            "sheet": sheet,
            "rows": resolved_preview(self.admin, self.obj, sheet),
            "start": JAN.isoformat(),
            "end": "2026-01-02",
            "existing": existing,
        }

    def test_facts_only_exact_dates_daily_resources_and_units(self):
        sheet = self.sheet(work_names=[("Simple", "m.", 2)])
        self.assertFalse(sheet["errors"])
        self.assertEqual(len(sheet["entries"]), 6)
        self.assertEqual(
            [r["quantity"] for r in sheet["entries"] if r["kind"] == "work"],
            ["999.000000", "999.000000"],
        )
        self.assertEqual(
            {r["month"] for r in sheet["entries"]}, {"2026-01-01", "2026-01-02"}
        )
        self.assertEqual(sheet["entries"][0]["unit"], "m")

    def test_preview_has_no_writes_confirm_creates_missing_once(self):
        payload = self.payload()
        self.assertFalse(DailyFact.objects.exists())
        self.assertFalse(ProjectWork.objects.filter(name="Imported work").exists())
        self.assertEqual(apply_facts(self.admin, payload), (6, 0))
        self.assertEqual(ProjectWork.objects.filter(name="Imported work").count(), 1)
        self.assertEqual(DailyFact.objects.count(), 2)
        self.assertEqual(LaborFact.objects.count(), 2)
        self.assertEqual(EquipmentFact.objects.count(), 2)
        self.assertEqual(DailyFact.objects.first().reported_by, self.admin)

    def test_keep_replace_and_stale_preview(self):
        sheet = self.sheet(work_names=[("Simple", "m", 2)])
        fact = DailyFact.objects.create(
            company=self.company,
            project_work=self.simple,
            date=JAN,
            actual_quantity=3,
            reported_by=self.admin,
            comment="keep note",
        )
        payload = self.payload(sheet)
        self.assertEqual(apply_facts(self.admin, payload), (5, 1))
        fact.refresh_from_db()
        self.assertEqual(fact.actual_quantity, 3)
        payload = self.payload(sheet, existing="replace")
        apply_facts(self.admin, payload)
        fact.refresh_from_db()
        self.assertEqual(fact.actual_quantity, 999)
        self.assertEqual(fact.comment, "keep note")
        with self.assertRaises(ValidationError):
            apply_facts(self.admin, payload)

    def test_blank_skipped_zero_written_negative_blocks(self):
        book = load_workbook(BytesIO(upload().read()))
        sheet = book.active
        sheet.cell(5, 15).value = 0
        sheet.cell(5, 16).value = None
        data = BytesIO()
        book.save(data)
        parsed = parse_meeting_workbook(
            SimpleUploadedFile("f.xlsx", data.getvalue()),
            JAN,
            date(2026, 1, 2),
            facts=True,
        )[0]
        rows = [r for r in parsed["entries"] if r["kind"] == "work"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(Decimal(rows[0]["quantity"]), 0)
        sheet.cell(5, 15).value = -1
        data = BytesIO()
        book.save(data)
        parsed = parse_meeting_workbook(
            SimpleUploadedFile("f.xlsx", data.getvalue()),
            JAN,
            date(2026, 1, 2),
            facts=True,
        )[0]
        self.assertTrue(parsed["errors"])

    def test_partial_composite_facts_allowed(self):
        sheet = self.sheet(work_names=[("A", "m", 2)])
        payload = self.payload(sheet)
        self.assertEqual(payload["rows"][0]["target_type"], "item")
        apply_facts(self.admin, payload)
        self.assertEqual(
            DailyFact.objects.filter(
                project_work=self.composite, work_item=self.a
            ).count(),
            2,
        )

    def test_multiple_fact_files_with_individual_periods(self):
        self.client.force_login(self.admin)
        url = reverse("production:fact_meeting_import")
        response = self.client.post(url, {
            "construction_object": self.obj.pk, "start": "2026-01-01", "end": "2026-01-02",
            "existing": "keep", "file": [upload(name="Один"), upload(name="Два")],
            "file_start": ["2026-01-01", "2026-01-02"], "file_end": ["2026-01-01", "2026-01-02"],
            "file_sheet": ["Один", "Два"],
        })
        self.assertContains(response, "Подтвердить импорт факта")
        result = self.client.post(url, {"action": "confirm", "preview": response.context["preview"]})
        self.assertEqual(result.status_code, 302)
        self.assertTrue(DailyFact.objects.filter(date=JAN).exists())
        self.assertTrue(DailyFact.objects.filter(date=date(2026, 1, 2)).exists())

    def test_multiple_fact_files_conflict_blocks_preview(self):
        self.client.force_login(self.admin)
        response = self.client.post(reverse("production:fact_meeting_import"), {
            "construction_object": self.obj.pk, "start": "2026-01-01", "end": "2026-01-02",
            "existing": "keep", "file": [upload(), upload()],
        })
        self.assertContains(response, "Конфликт файлов")
        self.assertFalse(DailyFact.objects.exists())

    def test_repeated_equipment_fact_keeps_first_and_warns(self):
        sheet = self.sheet()
        row = next(row for row in sheet['entries'] if row['kind'] == 'equipment')
        duplicate = dict(row, row=169, quantity='88')
        sheet['entries'].append(duplicate)
        resolved = resolved_preview(self.admin, self.obj, sheet)
        matching = [r for r in resolved if r['kind'] == 'equipment' and r['month'] == row['month']]
        self.assertEqual(len(matching), 1)
        self.assertNotEqual(matching[0]['quantity'], '88')
        self.assertTrue(any('169' in warning for warning in sheet['warnings']))

    def test_repeated_people_fact_keeps_first_across_groups(self):
        sheet = self.sheet()
        row = next(row for row in sheet['entries'] if row['kind'] == 'labor')
        sheet['entries'].append(dict(row, row=314, section='Другая группа', quantity='88'))
        rows = resolved_preview(self.admin, self.obj, sheet)
        matching = [r for r in rows if r['kind'] == 'labor' and r['month'] == row['month']]
        self.assertEqual(len(matching), 1)
        self.assertNotEqual(matching[0]['quantity'], '88')
        self.assertTrue(any('314' in warning for warning in sheet['warnings']))

    def test_new_resources_in_different_groups_create_one_catalogue_position(self):
        sheet = self.sheet()
        for row in sheet['entries']:
            if row['kind'] in {'labor', 'equipment'}:
                row['section'] = 'Первая' if row['month'] == '2026-01-01' else 'Вторая'
        payload = self.payload(sheet)
        apply_facts(self.admin, payload)
        from apps.resources.models import Brigade, EquipmentType
        self.assertEqual(Brigade.objects.filter(company=self.company, name='Imported brigade').count(), 1)
        self.assertEqual(EquipmentType.objects.filter(company=self.company, name='Imported equipment').count(), 1)

    def test_new_equipment_repeat_across_categories_uses_first(self):
        sheet = self.sheet()
        row = next(row for row in sheet['entries'] if row['kind'] == 'equipment')
        sheet['entries'].append(dict(row, row=169, section='Другая категория', quantity='88'))
        rows = resolved_preview(self.admin, self.obj, sheet)
        matching = [r for r in rows if r['kind'] == 'equipment' and r['month'] == row['month']]
        self.assertEqual(len(matching), 1)
        self.assertNotEqual(matching[0]['quantity'], '88')

    def test_admin_imports_renamed_sheet_in_past(self):
        from unittest.mock import patch
        self.client.force_login(self.admin)
        url = reverse("production:fact_meeting_import")
        with patch("django.utils.timezone.localdate", return_value=date(2026, 10, 10)):
            response = self.client.post(url, {
                "construction_object": self.obj.pk, "start": "2026-01-01",
                "end": "2026-01-02", "existing": "keep",
                "sheet_name": "Переименованный лист", "file": upload(name="Переименованный лист"),
            })
            self.assertContains(response, "Подтвердить импорт факта")
            result = self.client.post(url, {"action": "confirm", "preview": response.context["preview"]})
            self.assertEqual(result.status_code, 302)
        self.assertTrue(DailyFact.objects.filter(project_work__section__construction_object=self.obj, date=JAN).exists())
        self.assertFalse(type(self.obj).objects.filter(name="Переименованный лист").exists())

    def test_permissions_assignment_and_replay(self):
        url = reverse("production:fact_meeting_import")
        self.client.force_login(self.planner)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.client.force_login(self.foreman)
        data = {
            "construction_object": self.obj.pk,
            "start": "2026-01-01",
            "end": "2026-01-02",
            "existing": "keep",
            "file": upload(),
        }
        response = self.client.post(url, data)
        self.assertContains(response, "Выберите корректный вариант")
        self.foreman.assigned_objects.add(self.obj)
        data["file"] = upload()
        response = self.client.post(url, data)
        self.assertContains(response, "Подтвердить импорт факта")
        token = response.context["preview"]
        self.foreman.assigned_objects.clear()
        self.assertEqual(
            self.client.post(url, {"action": "confirm", "preview": token}).status_code,
            404,
        )
        self.assertFalse(DailyFact.objects.exists())
        self.foreman.assigned_objects.add(self.obj)
        self.assertEqual(
            self.client.post(url, {"action": "confirm", "preview": token}).status_code,
            302,
        )
        self.assertEqual(DailyFact.objects.count(), 2)
        self.client.post(url, {"action": "confirm", "preview": token})
        self.assertEqual(DailyFact.objects.count(), 2)

    def test_invalid_date_rolls_back_catalog_creation(self):
        payload = self.payload()
        payload["start"] = "2026-01-02"
        with self.assertRaises(ValidationError):
            apply_facts(self.admin, payload)
        self.assertFalse(ProjectWork.objects.filter(name="Imported work").exists())
        self.assertFalse(DailyFact.objects.exists())

    def test_same_named_work_in_another_section_creates_separate_facts(self):
        sheet = self.sheet(work_names=[('Simple', 'm', 2), ('Simple', 'm', 7)],
                           sections=[self.simple.section.name, 'Another section'])
        payload = self.payload(sheet)
        work_rows = [r for r in payload['rows'] if r['kind'] == 'work']
        self.assertEqual({r['target_id'] for r in work_rows if r['section'] == self.simple.section.name}, {self.simple.pk})
        self.assertTrue(all(r['target_id'] is None for r in work_rows if r['section'] == 'Another section'))
        apply_facts(self.admin, payload)
        other = ProjectWork.objects.get(section__construction_object=self.obj,
                                       section__name='Another section', name='Simple')
        self.assertNotEqual(other.pk, self.simple.pk)
        self.assertEqual(DailyFact.objects.filter(project_work=self.simple).count(), 2)
        self.assertEqual(DailyFact.objects.filter(project_work=other).count(), 2)
