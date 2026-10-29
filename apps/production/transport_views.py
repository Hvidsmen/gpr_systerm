from calendar import monthrange
from datetime import date, timedelta
from urllib.parse import urlencode
from collections import defaultdict
from uuid import uuid4
from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError,PermissionDenied
from core.fact_edit_window import can_edit_fact_day,require_fact_day,window_context
from django.db import transaction
from django.db.models import Q, F
from django.http import JsonResponse, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_POST, require_http_methods
from core.permissions import require_roles, role_code, scope_queryset
from apps.resources.models import Equipment
from apps.resources.equipment_units import UnitForm
from apps.personnel.models import Camp
from apps.projects.models import ConstructionObject
from apps.works.models import ProjectWork, ProjectWorkItem
from .transport_models import EquipmentCampStay, EquipmentTransportEvent, EquipmentDayRecord, EquipmentTransportControl, CONDITIONS, USES
from . import transport as service

DATE_INPUT = forms.DateInput(format='%Y-%m-%d', attrs={'type': 'date'})


class Selector(forms.Form):
    camp = forms.ModelChoiceField(label='ВЖГ', queryset=Camp.objects.none())
    date = forms.DateField(label='Дата', widget=DATE_INPUT)
    q = forms.CharField(label='Поиск техники', required=False, max_length=150, widget=forms.TextInput(attrs={'type':'search','placeholder':'Номер, VIN, название или модель'}))
    condition = forms.ChoiceField(label='Состояние', required=False, choices=[('', 'Все состояния'), *CONDITIONS])
    use = forms.ChoiceField(label='За день', required=False, choices=[('', 'Все отметки'), ('EMPTY', 'Не заполнено'), *USES])
    def __init__(self, *args, user, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields['camp'].queryset = service.camps_for(user)
        for field in self.fields.values(): field.widget.attrs['class'] = 'form-select' if isinstance(field.widget, forms.Select) else 'form-control'


def target_choices(user, camp, day, current=()):
    """Plans and previously recorded actual work, with a retained saved selection."""
    from .dispatch import plans
    from .models import DailyFact
    objects = scope_queryset(ConstructionObject.objects.filter(pk__in=camp.object_links.filter(company=user.company).values('construction_object_id')), user)
    objects = list(objects.order_by('name', 'pk'))
    retained = [r for r in current if r.construction_object_id]
    known = {o.pk for o in objects}
    for row in retained:
        if row.construction_object_id not in known and scope_queryset(ConstructionObject.objects.all(),user).filter(pk=row.construction_object_id).exists():
            objects.append(row.construction_object); known.add(row.construction_object_id)
    result = [('', 'Без привязки к объекту и работе')]; mapping = {'': (None, None, None)}
    for obj in objects:
        planned, _ = plans(obj, day)
        keys = set(planned)
        keys.update(DailyFact.objects.filter(company=user.company, project_work__section__construction_object=obj, date__lte=day).values_list('project_work_id','work_item_id'))
        keys.update(EquipmentDayRecord.objects.filter(company=user.company, construction_object=obj, date__lte=day).exclude(work=None).values_list('work_id','work_item_id'))
        keys.update((r.work_id,r.work_item_id) for r in retained if r.construction_object_id==obj.pk and r.work_id)
        works = {w.pk:w for w in ProjectWork.objects.filter(company=user.company, section__construction_object=obj, pk__in=[k[0] for k in keys]).select_related('section','complex_object')}
        items = {i.pk:i for i in ProjectWorkItem.objects.filter(company=user.company, pk__in=[k[1] for k in keys if k[1]])}
        group = [(f'object:{obj.pk}', 'Объект в целом · без привязки к работе')]
        mapping[f'object:{obj.pk}']=(obj,None,None)
        for work_id,item_id in sorted(keys,key=lambda k:((works[k[0]].name if k[0] in works else ''),k[1] or 0)):
            work=works.get(work_id);item=items.get(item_id)
            if not work or (work.kind=='COMPOSITE' and not item) or (item and item.project_work_id!=work.pk): continue
            token=f'work:{work.pk}:{item.pk if item else 0}'
            label=' → '.join(filter(None,[work.section.name,str(work.complex_object) if work.complex_object_id else '',work.name,item.name if item else '']))
            group.append((token,label));mapping[token]=(obj,work,item)
        result.append((obj.name,group))
    return result,mapping


class ActionForm(forms.Form):
    action = forms.ChoiceField(widget=forms.HiddenInput, choices=[('ARRIVAL','Приезд'),('DEPARTURE','Выезд'),('TRANSFER','Перевод'),('WRITEOFF','Списание'),('CONDITION','Состояние'),('DAY','Отметка за день')])
    equipment = forms.ModelMultipleChoiceField(label='Техника', queryset=Equipment.objects.none())
    target_camp = forms.ModelChoiceField(label='ВЖГ назначения', queryset=Camp.objects.none(), required=False)
    condition = forms.ChoiceField(label='Техническое состояние',choices=CONDITIONS,required=False,initial='READY')
    status = forms.ChoiceField(label='Работа за день',choices=[*USES,('CLEAR','Очистить отметку за день')],required=False,initial='WORKED')
    target = forms.ChoiceField(label='Где работала — объект / работа (необязательно)',required=False)
    repeat = forms.BooleanField(label='Повторить последнюю привязку к работе',required=False,help_text='Отметка «Работала» подтверждается за выбранный день. Прежний комментарий не копируется.')
    comment = forms.CharField(label='Примечание / причина (необязательно)',required=False,max_length=4000,widget=forms.Textarea(attrs={'rows':2}))
    signed = forms.CharField(widget=forms.HiddenInput)
    session = forms.UUIDField(widget=forms.HiddenInput,required=False)
    def __init__(self,*args,user,camp,day,current=(),load_targets=True,**kwargs):
        super().__init__(*args,**kwargs)
        self.fields['equipment'].queryset=Equipment.objects.filter(company=user.company).select_related('type__parent').order_by('type__name','plate_number','inventory_number','pk')
        self.fields['equipment'].label_from_instance=lambda e:f'{e.plate_number or e.inventory_number or e.vin} · {e.type.parent.name if e.type.parent_id else e.type.name} · {e.name}'
        self.fields['target_camp'].queryset=service.camps_for(user).filter(is_active=True).exclude(pk=camp.pk)
        options,self.targets=target_choices(user,camp,day,current) if load_targets else ([('','Без привязки к объекту и работе')],{'':(None,None,None)})
        self.fields['target'].choices=options
        for field in self.fields.values():
            field.widget.attrs['class']='form-check-input' if isinstance(field.widget,forms.CheckboxInput) else 'form-select' if isinstance(field.widget,forms.Select) else 'form-control'
            if isinstance(field.widget,forms.Select):field.widget.attrs['data-native-select']='1'
        self.fields['equipment'].widget.attrs['size']='8'
    def clean(self):
        data=super().clean()
        if len(data.get('equipment',()))>100:self.add_error('equipment','За один раз можно выбрать не более 100 машин.')
        if data.get('action')=='DAY' and not data.get('status'):self.add_error('status','Выберите отметку за день.')
        if data.get('action')=='CONDITION' and not data.get('condition'):self.add_error('condition','Выберите состояние.')
        return data


def _rows(user, camp, day):
    stays=list(service.present(user.company,day).filter(camp=camp).select_related('equipment__type__parent','equipment__type__category','equipment__type__parent__category','equipment__catalogue_group'))
    ids=[s.equipment_id for s in stays]
    records={r.equipment_id:r for r in EquipmentDayRecord.objects.filter(company=user.company,equipment_id__in=ids,date=day).select_related('construction_object','work','work_item','recorded_by')}
    controls={c.equipment_id:c for c in EquipmentTransportControl.objects.filter(company=user.company,equipment_id__in=ids).select_related('locked_by')}
    rows=[]
    conditions=service.conditions_for(user.company,ids,[day])
    for stay in stays:
        machine=stay.equipment;record=records.get(machine.pk);control=controls.get(machine.pk)
        condition=conditions[machine.pk,day];status=record.status if record else ''
        owner=control.locked_by if control and control.expires_at and control.expires_at>timezone.now() else None
        rows.append({'machine':machine,'stay':stay,'record':record,'condition':condition,'condition_label':dict(CONDITIONS).get(condition,'Не указано'),
            'status':status,'status_label':dict(USES).get(status,'Не заполнено'),'owner':owner,
            'planning_type':machine.type.parent if machine.type.parent_id else machine.type,
            'target':f'work:{record.work_id}:{record.work_item_id or 0}' if record and record.work_id else f'object:{record.construction_object_id}' if record and record.construction_object_id else '',
            'details_url':reverse('production:transport_details',args=[machine.pk])+'?'+urlencode({'camp':camp.pk,'date':str(day)}),
            'history_url':reverse('production:transport_history',args=[machine.pk])})
    from apps.resources.group_order import machine_order
    return sorted(rows,key=lambda row:machine_order(row['machine']))


def _summary(rows):
    groups={}
    for row in rows:
        group=groups.setdefault(row['planning_type'].pk,{'type':row['planning_type'].name,'total':0,'WORKED':0,'IDLE':0,'EMPTY':0,'REPAIR':0,'SCRAP':0,'UNKNOWN':0})
        group['total']+=1
        group[row['status'] or 'EMPTY']+=1
        if row['condition']!='READY':group[row['condition'] or 'UNKNOWN']+=1
    return list(groups.values())


def _excel(camp,day,rows):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    workbook=Workbook();ws=workbook.active;ws.title='Техника за день'
    ws.append(['Дата','ВЖГ','Техника для планирования','Название','Госномер','Инвентарный номер','VIN','Состояние','За день','Объект','Работа','Работа в комплексе','Комментарий','Кто внёс'])
    for row in rows:
        m=row['machine'];r=row['record']
        values=[day.isoformat(),camp.name,row['planning_type'].name,m.name,m.plate_number,m.inventory_number,m.vin,row['condition_label'],row['status_label'],r.construction_object.name if r and r.construction_object_id else '',r.work.name if r and r.work_id else '',r.work_item.name if r and r.work_item_id else '',r.comment if r else '',(r.recorded_by.get_full_name() or r.recorded_by.username) if r and r.recorded_by_id else '']
        ws.append(values)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='087F99')
    ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(60,max(14,max(len(str(c.value or '')) for c in col)+2))
    from .transport_details_models import EquipmentDayDetails,DETAIL_FIELDS
    from django.db.models import OuterRef,Subquery
    ids=[r['machine'].pk for r in rows]
    details=EquipmentDayDetails.objects.filter(company=camp.company,equipment_id__in=ids,date__lte=day)
    latest=EquipmentDayDetails.objects.filter(company=camp.company,equipment_id=OuterRef('equipment_id'),date__lte=day).order_by('-date','-pk').values('pk')[:1]
    notes=workbook.create_sheet('Ремонт и обслуживание');notes.append(['Дата сведений','Машина','ВЖГ',*[label for name,label in DETAIL_FIELDS]])
    for detail in details.filter(pk=Subquery(latest)).select_related('equipment','camp'):
        notes.append([detail.date.isoformat(),str(detail.equipment),detail.camp.name,*[str(getattr(detail,name)) if hasattr(getattr(detail,name),'_meta') else getattr(detail,'get_repair_stage_display')() if name=='repair_stage' else getattr(detail,name) for name,label in DETAIL_FIELDS]])
        for cell in notes[notes.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    notes.freeze_panes='D2';notes.auto_filter.ref=notes.dimensions
    for col in notes.columns:notes.column_dimensions[col[0].column_letter].width=30
    out=BytesIO();workbook.save(out)
    response=HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']=f'attachment; filename="transport-{day.isoformat()}.xlsx"'
    return response


@login_required
@require_http_methods(['GET','POST'])
def board(request):
    require_roles(request.user,service.OPERATORS if request.method=='POST' else service.READERS)
    raw=request.POST if request.method=='POST' else request.GET
    default_camp=service.camps_for(request.user).filter(is_active=True).first() or service.camps_for(request.user).first()
    selector_data={**{'camp':str(default_camp.pk) if default_camp else '', 'date':timezone.localdate().isoformat()},**raw.dict()}
    selector=Selector(selector_data,user=request.user)
    tab=request.GET.get('tab','day')
    if tab not in {'day','state','calendar','events','report'}:tab='day'
    context={'selector':selector,'editable':role_code(request.user) in service.OPERATORS,'tab':tab,'can_create':role_code(request.user) in service.OPERATORS}
    if not selector.is_valid():return render(request,'production/transport_board.html',context,status=400 if raw else 200)
    camp,day=selector.cleaned_data['camp'],selector.cleaned_data['date'];service.authorize(request.user,camp,False)
    context.update(window_context(request.user,day))
    context['editable']=context['editable'] and context['fact_day_editable']
    if request.method=='POST':
        try:require_fact_day(request.user,day)
        except ValidationError as exc:raise PermissionDenied('; '.join(exc.messages))
    rows=_rows(request.user,camp,day);current=[r['record'] for r in rows if r['record']]
    form=ActionForm(request.POST if request.method=='POST' else None,user=request.user,camp=camp,day=day,current=current,load_targets=tab=='day' or request.method=='POST' and request.POST.get('action')=='DAY',initial={'signed':service.snapshot(request.user,camp,day),'session':uuid4()})
    if request.method=='POST' and form.is_valid():
        values=form.cleaned_data;common={'signed':values['signed'],'session':values['session'],'comment':values['comment']};ids=list(values['equipment'].values_list('pk',flat=True))
        try:
            action=values['action']
            if action=='DAY':
                obj,work,item=form.targets[values['target']]
                count=service.save_day(request.user,camp,day,ids,values['status'],construction_object=obj,work=work,work_item=item,repeat=values['repeat'],**common)
            elif action=='CONDITION':count=service.set_condition(request.user,camp,day,ids,values['condition'],**common)
            else:count=service.move(request.user,camp,day,ids,action,target=values['target_camp'],condition=values['condition'] or 'READY',**common)
            messages.success(request,f'Сохранено машин: {count}.')
            return redirect(reverse('production:transport_board')+'?'+urlencode({'camp':camp.pk,'date':str(day),'tab':tab}))
        except ValidationError as exc:form.add_error(None,ValidationError('; '.join(exc.messages)))
    data=selector.cleaned_data
    filtered=[r for r in rows if (not data['condition'] or r['condition']==data['condition']) and (tab=='state' or not data['use'] or ((r['status'] or 'EMPTY')==data['use'])) and all(term.casefold() in ' '.join(str(value or '') for value in [r['machine'],r['machine'].vin,r['machine'].inventory_number,r['machine'].brand_model,r['planning_type'].name]).casefold() for term in data['q'].split())]
    if request.GET.get('export')=='xlsx':return _excel(camp,day,filtered)
    from .transport_inline import prepare_rows
    from .transport_details_models import DETAIL_FIELDS
    from .transport_detail_values import ROSTER_DETAIL_FIELDS,REPAIR_FIELD_GROUPS
    from django.core.paginator import Paginator
    page=Paginator(filtered,50).get_page(request.GET.get('page'))
    display_rows=list(page.object_list)
    if context['tab'] in {'day','state'}:
        prepare_rows(request.user,camp,day,display_rows,context['editable'])
        for row in display_rows:
            if tab=='day':row['target_widget']=forms.Select(choices=form.fields['target'].choices,attrs={'class':'form-select form-select-sm','data-detail-field':'target','aria-label':'Объект / работа: '+str(row['machine']),'disabled':not context['editable']}).render('target-'+str(row['machine'].pk),row['target'])
            details=row['details_form']
            details['comment'].label='Примечание по ремонту'
            row['visible_details']=[details[name] for name in (('in_gpr',) if tab=='state' else ROSTER_DETAIL_FIELDS)]
            row['repair_groups']=[{'title':title,'fields':[details[name] for name in names]} for title,names in REPAIR_FIELD_GROUPS]
    context.update(detail_columns=[label for name,label in DETAIL_FIELDS if name in (('in_gpr',) if tab=='state' else ROSTER_DETAIL_FIELDS)],page_obj=page,filtered_count=len(filtered))
    new_form=NewMachineForm(company=request.user.company,prefix='new-machine')
    new_form.fields.pop('site_name');new_form.fields.pop('is_active')
    context['new_machine_form']=new_form
    context.update(camp=camp,day=day,rows=display_rows,summary=_summary(rows),action_form=form,conditions=CONDITIONS,uses=USES,
        total=len(rows),unknown=sum(r['condition'] not in {'READY','REPAIR','SCRAP'} for r in rows),ready=sum(r['condition']=='READY' for r in rows),worked=sum(r['status']=='WORKED' for r in rows),idle=sum(r['status']=='IDLE' for r in rows),
        missing=sum(not r['status'] for r in rows),repair=sum(r['condition']=='REPAIR' for r in rows),scrap=sum(r['condition']=='SCRAP' for r in rows),
        query=urlencode({'camp':camp.pk,'date':str(day)}),export_query=urlencode({'camp':camp.pk,'date':str(day),'q':data['q'],'condition':data['condition'],'use':data['use']}),error_open=request.method=='POST',can_create=role_code(request.user) in service.OPERATORS)
    occupied=service.present(request.user.company,day).values('equipment_id')
    disposed=EquipmentTransportEvent.objects.filter(company=request.user.company,kind='WRITEOFF',cancelled_at=None).values('equipment_id')
    context['arrival_ids']=list(Equipment.objects.filter(company=request.user.company,is_active=True).exclude(pk__in=occupied).exclude(pk__in=disposed).values_list('pk',flat=True))
    from apps.resources.transport_import_models import TransportImportRow
    context['import_errors']=TransportImportRow.objects.filter(company=request.user.company,status='PENDING',batch__applied_at__isnull=False).count()
    context['events']=EquipmentTransportEvent.objects.filter(company=request.user.company).filter(Q(from_camp=camp)|Q(to_camp=camp)).filter(date__lte=day).select_related('equipment','from_camp','to_camp','recorded_by','cancelled_by')[:100]
    if context['tab']=='calendar':
        start=day.replace(day=1);end=day.replace(day=monthrange(day.year,day.month)[1]);days=[start+timedelta(days=i) for i in range(end.day)]
        stays=list(EquipmentCampStay.objects.filter(company=request.user.company,camp=camp,cancelled_at=None,arrived_on__lte=end).filter(Q(departed_on=None)|Q(departed_on__gt=start)).exclude(departed_on=F('arrived_on')).select_related('equipment__type__parent'))
        machines={s.equipment_id:s.equipment for s in stays};records={(r.equipment_id,r.date):r for r in EquipmentDayRecord.objects.filter(company=request.user.company,camp=camp,date__range=(start,end))}
        calendar=[]
        conditions=service.conditions_for(request.user.company,machines,days)
        for machine in sorted(machines.values(),key=lambda m:str(m)):
            cells=[]
            for date_value in days:
                inside=any(s.equipment_id==machine.pk and s.arrived_on<=date_value and (s.departed_on is None or s.departed_on>date_value) for s in stays)
                state=conditions[machine.pk,date_value] if inside else '';record=records.get((machine.pk,date_value))
                code= ('✓' if record.status=='WORKED' else 'П') if inside and record else {'READY':'—','REPAIR':'Р','SCRAP':'З'}.get(state,'?') if inside else '·'
                if date_value>timezone.localdate() and not record and state=='READY':code='·'
                cells.append({'date':date_value,'code':code,'state':state,'status':record.status if record else ''})
            calendar.append({'machine':machine,'cells':cells})
        context.update(calendar=calendar,dates=days)
    return render(request,'production/transport_board.html',context,status=400 if request.method=='POST' else 200)


@login_required
@require_POST
def claims(request):
    require_roles(request.user,service.OPERATORS)
    camp=get_object_or_404(service.camps_for(request.user),pk=request.POST.get('camp'))
    try:
        day=date.fromisoformat(request.POST.get('date',''));ids=[int(value) for value in request.POST.getlist('equipment')]
        owners=service.claim(request.user,camp,day,ids,request.POST.get('action'),request.POST.get('session'))
        return JsonResponse({'owners':owners})
    except (ValueError,TypeError,ValidationError) as exc:return JsonResponse({'error':'; '.join(exc.messages) if isinstance(exc,ValidationError) else 'Неверная дата или техника.'},status=409)


@login_required
@require_http_methods(['GET','POST'])
def history(request,pk):
    require_roles(request.user,service.OPERATORS if request.method=='POST' else service.READERS)
    machine=get_object_or_404(Equipment.objects.filter(company=request.user.company),pk=pk)
    visible=service.camps_for(request.user)
    events=machine.transport_events.filter(Q(from_camp__in=visible)|Q(to_camp__in=visible)).select_related('from_camp','to_camp','recorded_by','cancelled_by')
    if not events.exists() and role_code(request.user)!='ADMIN':
        from django.http import Http404
        raise Http404('История техники недоступна.')
    error=''
    if request.method=='POST':
        event=get_object_or_404(events,pk=request.POST.get('event'))
        try:
            service.cancel_event(request.user,event,request.POST.get('reason',''),signed=request.POST.get('signed'),session=request.POST.get('session'))
            messages.success(request,'Ошибочная запись отменена. История сохранена.');return redirect('production:transport_history',pk=machine.pk)
        except ValidationError as exc:error='; '.join(exc.messages)
    event_rows=[]
    for event in events[:150]:
        camp=event.to_camp or event.from_camp
        event_rows.append({'event':event,'signed':service.snapshot(request.user,camp,event.date) if camp else '', 'can_cancel': can_edit_fact_day(request.user,event.date) and all(visible.filter(pk=value).exists() for value in [event.from_camp_id,event.to_camp_id] if value)})
    days=machine.transport_days.filter(camp__in=visible).select_related('camp','construction_object','work','work_item','recorded_by').order_by('-date')[:150]
    return render(request,'production/transport_history.html',{'machine':machine,'event_rows':event_rows,'days':days,'editable':role_code(request.user) in service.OPERATORS,'error':error})


class NewMachineForm(UnitForm):
    receive=forms.BooleanField(label='Сразу оформить приезд в выбранный ВЖГ',required=False,initial=True)
    condition=forms.ChoiceField(label='Техническое состояние при приезде',choices=CONDITIONS,initial='READY')


@login_required
@require_http_methods(['GET','POST'])
def create_machine(request):
    require_roles(request.user,service.OPERATORS)
    camp=get_object_or_404(service.camps_for(request.user),pk=request.POST.get('camp',request.GET.get('camp')))
    try:day=date.fromisoformat(request.POST.get('date',request.GET.get('date','')))
    except ValueError:day=timezone.localdate()
    ajax=request.headers.get('Accept')=='application/json'
    form=NewMachineForm(request.POST if request.method=='POST' else None,company=request.user.company,prefix='new-machine' if ajax else None)
    # Transport creates a physical machine; planning catalogues remain managed separately.
    form.fields.pop('site_name');form.fields.pop('is_active')
    if request.method=='POST' and form.is_valid():
        try:
            with transaction.atomic():
                service.authorize(request.user,camp)
                machine=form.save()
                if form.cleaned_data['receive']:
                    service.move(request.user,camp,day,[machine.pk],'ARRIVAL',condition=form.cleaned_data['condition'],signed=service.snapshot(request.user,camp,day))
                service.audit(request.user,machine,'CREATE',{'equipment':machine.pk,'received':form.cleaned_data['receive']})
            if ajax:return JsonResponse({'saved':True,'id':machine.pk,'received':form.cleaned_data['receive']})
            messages.success(request,'Машина добавлена в справочник'+(' и принята в ВЖГ.' if form.cleaned_data['receive'] else '. Приезд можно оформить позднее.'))
            return redirect(reverse('production:transport_board')+'?'+urlencode({'camp':camp.pk,'date':str(day)}))
        except ValidationError as exc:form.add_error(None,ValidationError('; '.join(exc.messages)))
    if ajax:return JsonResponse({'errors':form.errors},status=400)
    return render(request,'production/transport_create.html',{'form':form,'camp':camp,'day':day},status=400 if request.method=='POST' else 200)
