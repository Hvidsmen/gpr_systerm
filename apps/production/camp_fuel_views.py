from decimal import Decimal
from io import BytesIO
from urllib.parse import urlencode
from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied,ValidationError
from django.http import HttpResponse
from django.shortcuts import render,redirect
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_http_methods
from core.permissions import require_roles,role_code
from core.fact_edit_window import window_context,require_fact_day
from . import transport,camp_fuel,fuel_access
from .camp_fuel_models import CampFuelDay,FUEL_CHOICES


class Selector(forms.Form):
    camp=forms.ModelChoiceField(label='ВЖГ',queryset=None)
    date=forms.DateField(label='Дата',widget=forms.DateInput(format='%Y-%m-%d',attrs={'type':'date'}))
    def __init__(self,*args,user,**kwargs):
        super().__init__(*args,**kwargs);self.fields['camp'].queryset=fuel_access.camps_for(user)


class DayForm(forms.Form):
    opening=forms.DecimalField(label='Остаток на начало, м³',min_value=0,max_digits=15,decimal_places=3)
    received=forms.DecimalField(label='Приход за день, м³',min_value=0,max_digits=15,decimal_places=3)
    consumed=forms.DecimalField(label='Расход за день, м³',min_value=0,max_digits=15,decimal_places=3)
    comment=forms.CharField(label='Комментарий',max_length=4000,required=False,widget=forms.Textarea(attrs={'rows':2}))
    signed=forms.CharField(widget=forms.HiddenInput)
    def __init__(self,*args,previous=None,**kwargs):
        super().__init__(*args,**kwargs)
        if previous:
            self.fields['opening'].disabled=True
            self.initial['opening']=previous.actual_balance
            self.fields['opening'].help_text=f'Остаток за {previous.date:%d.%m.%Y}. Переносится автоматически.'
        else:self.fields['opening'].help_text='Начальный запас задаётся при начале учёта. Дальше остаток переносится автоматически.'


def excel(camp,fuel_type,rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book=Workbook();sheet=book.active;sheet.title='ГСМ ВЖГ'
    sheet.append(['Дата','ВЖГ','Начальный остаток, м³','Приход, м³','Расход, м³','Конечный остаток, м³','Комментарий','Первоначальный автор'])
    for row in rows:
        sheet.append([row.date,camp.name,row.opening_balance,row.received_liters,row.actual_liters,row.actual_balance,row.comment,str(row.recorded_by or '')])
        sheet.cell(sheet.max_row,1).number_format='DD.MM.YYYY'
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    for cell in sheet[1]:cell.font=Font(bold=True)
    sheet.freeze_panes='C2';sheet.auto_filter.ref=sheet.dimensions
    for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(55,max(18,max(len(str(c.value or '')) for c in col)+2))
    stream=BytesIO();book.save(stream)
    response=HttpResponse(stream.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']='attachment; filename="camp-fuel.xlsx"';return response


@login_required
@require_http_methods(['GET','POST'])
def board(request):
    require_roles(request.user,fuel_access.OPERATORS if request.method=='POST' else fuel_access.READERS)
    source=request.POST if request.method=='POST' else request.GET
    first=fuel_access.camps_for(request.user).first()
    defaults={'camp':str(first.pk) if first else '', 'date':str(timezone.localdate())}
    selector=Selector({**defaults,**source.dict()},user=request.user)
    context={'selector':selector}
    if not first:return render(request,'production/camp_fuel.html',context,status=400 if request.method=='POST' else 200)
    if not selector.is_valid():return render(request,'production/camp_fuel.html',context,status=400)
    camp,day=(selector.cleaned_data[k] for k in ('camp','date'))
    fuel_type='TOTAL'
    fuel_access.authorize(request.user,camp,False)
    context.update(window_context(request.user,day))
    editable=context['fact_day_editable'] and role_code(request.user) in fuel_access.OPERATORS
    if request.method=='POST':
        try:require_fact_day(request.user,day)
        except ValidationError as exc:raise PermissionDenied('; '.join(exc.messages))
    previous,current,signed=camp_fuel.entry_snapshot(request.user,camp,day,fuel_type)
    opening=previous.actual_balance if previous else current.opening_balance if current else Decimal(0)
    initial={'opening':opening,'received':current.received_liters if current else 0,'consumed':current.actual_liters if current else 0,
        'comment':current.comment if current else '', 'signed':signed}
    form=DayForm(request.POST if request.method=='POST' else None,initial=initial,previous=previous)
    if request.method=='POST' and form.is_valid():
        try:camp_fuel.save_day(request.user,camp,day,fuel_type,**form.cleaned_data)
        except ValidationError as exc:form.add_error(None,'; '.join(exc.messages))
        else:
            messages.success(request,'Баланс ГСМ сохранён. Остатки следующих дней пересчитаны; история изменений сохранена.')
            return redirect(reverse('production:camp_fuel')+'?'+urlencode(selector.cleaned_data|{'camp':camp.pk}))
    rows=camp_fuel.ledger(camp,fuel_type).filter(date__lte=day).select_related('recorded_by')
    if request.GET.get('export')=='xlsx':return excel(camp,fuel_type,rows)
    closing=current.actual_balance if current else opening
    from .history_views import visible_revisions,changes
    revisions=visible_revisions(request.user).filter(entity_type='production.campfuelday',camp_ids__0=camp.pk).order_by('-created_at')[:30]
    for revision in revisions:revision.changes=changes(revision)
    context.update(camp=camp,day=day,fuel_type=fuel_type,fuel_label=dict(FUEL_CHOICES)[fuel_type],form=form,editable=editable,current=current,
        opening=opening,closing=closing,rows=rows.order_by('-date')[:60],revisions=revisions,
        query=urlencode({'camp':camp.pk,'date':str(day)}),camp_query=urlencode({'camp':camp.pk,'date':str(day)}))
    return render(request,'production/camp_fuel.html',context,status=400 if request.method=='POST' else 200)
