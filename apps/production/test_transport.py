from datetime import timedelta
from io import BytesIO
from uuid import uuid4
from unittest.mock import patch
from django.core.exceptions import PermissionDenied,ValidationError
from django.test import TestCase,Client
from django.urls import reverse
from django.utils import timezone
from apps.accounts.models import User,Role,Company
from apps.personnel.models import Camp,CampObjectLink
from apps.resources.models import Equipment
from . import test_dispatch as fixtures
from . import transport as service
from .transport_models import EquipmentCampStay,EquipmentTransportEvent,EquipmentDayRecord,EquipmentTransportControl
from .models import EquipmentActivity,EquipmentFact,FactSectionDay


class TransportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        fixtures.DispatcherTests.setUpTestData.__func__(cls)
        cls.operator=cls.users['MECHANIZATION_DISPATCHER'];cls.admin=cls.users['ADMIN'];cls.camp=cls.housing_camp
        cls.target=Camp.objects.create(company=cls.company,name='ВЖГ 2')
        CampObjectLink.objects.create(company=cls.company,camp=cls.target,construction_object=cls.obj)
        cls.unrelated=Camp.objects.create(company=cls.company,name='Недоступный ВЖГ')
        CampObjectLink.objects.create(company=cls.company,camp=cls.unrelated,construction_object=cls.other)
        cls.second=Equipment.objects.create(company=cls.company,type=cls.type,name='Second crane',inventory_number='INV-2',vin='TEST-002')
        cls.colleague=User.objects.create_user(username='transport-colleague',company=cls.company,role=Role.objects.get(code='MECHANIZATION_DISPATCHER'))
        cls.colleague.assigned_objects.add(cls.obj)
        cls.board=reverse('production:transport_board')

    def sign(self,user=None,camp=None,day=None):return service.snapshot(user or self.operator,camp or self.camp,day or self.day)
    def receive(self,machine=None,day=None,camp=None,condition='READY'):
        machine=machine or self.machine;day=day or self.day;camp=camp or self.camp
        return service.move(self.operator,camp,day,[machine.pk],'ARRIVAL',condition=condition,signed=self.sign(camp=camp,day=day))
    def save(self,status='WORKED',machine=None,day=None,**kwargs):
        machine=machine or self.machine;day=day or self.day
        return service.save_day(self.operator,self.camp,day,[machine.pk],status,signed=self.sign(day=day),**kwargs)

    def test_arrival_persists_presence_and_condition_but_never_copies_daily_work(self):
        self.receive();self.save()
        tomorrow=self.day+timedelta(days=1)
        self.assertTrue(service.present(self.company,tomorrow).filter(equipment=self.machine,camp=self.camp).exists())
        self.assertEqual(service.condition_on(self.machine,tomorrow),'READY')
        self.assertFalse(EquipmentDayRecord.objects.filter(equipment=self.machine,date=tomorrow).exists())
        self.assertFalse(EquipmentActivity.objects.exists())
        self.assertEqual(EquipmentDayRecord.objects.get().status,'WORKED')

    def test_movement_has_half_open_dates_and_preserves_history(self):
        self.receive();day=self.day+timedelta(days=1)
        service.move(self.operator,self.camp,day,[self.machine.pk],'TRANSFER',target=self.target,signed=self.sign(day=day))
        self.assertFalse(service.present(self.company,day).filter(equipment=self.machine,camp=self.camp).exists())
        self.assertTrue(service.present(self.company,day).filter(equipment=self.machine,camp=self.target).exists())
        day+=timedelta(days=1)
        with patch('django.utils.timezone.localdate',return_value=day):
            service.move(self.operator,self.target,day,[self.machine.pk],'DEPARTURE',signed=self.sign(camp=self.target,day=day))
        self.assertFalse(service.present(self.company,day).filter(equipment=self.machine).exists())
        self.assertEqual(self.machine.camp_stays.count(),2)
        self.assertEqual(self.machine.transport_events.exclude(kind='CONDITION').count(),3)

    def test_technical_state_and_actual_writeoff_are_separate_and_reversible(self):
        self.receive(condition='SCRAP');self.machine.refresh_from_db();self.assertTrue(self.machine.is_active)
        self.save();self.save('CLEAR')
        service.move(self.operator,self.camp,self.day,[self.machine.pk],'WRITEOFF',signed=self.sign())
        self.machine.refresh_from_db();self.assertFalse(self.machine.is_active)
        event=self.machine.transport_events.get(kind='WRITEOFF')
        service.cancel_event(self.operator,event,'Ошибочно списана',signed=self.sign())
        self.machine.refresh_from_db();self.assertTrue(self.machine.is_active)
        self.assertTrue(service.present(self.company,self.day).filter(equipment=self.machine).exists())
        self.assertEqual(service.condition_on(self.machine,self.day),'SCRAP')
        event.refresh_from_db();self.assertTrue(event.cancelled_at)

    def test_repair_carries_forward_and_day_use_is_independent(self):
        self.receive(condition='REPAIR')
        self.save('IDLE')
        day=self.day+timedelta(days=1)
        self.assertEqual(service.condition_on(self.machine,day),'REPAIR')
        self.save('WORKED',day=day,construction_object=self.obj,work=self.work)
        service.set_condition(self.operator,self.camp,day,[self.machine.pk],'READY',signed=self.sign(day=day))
        self.assertEqual(list(self.machine.transport_days.values_list('status',flat=True)),['IDLE','WORKED'])

    def test_project_attribution_uses_existing_fact_totals_and_clearing_removes_projection(self):
        self.receive();self.save(construction_object=self.obj,work=self.work)
        row=EquipmentDayRecord.objects.get();self.assertTrue(row.activity_id)
        self.assertEqual(EquipmentActivity.objects.get().work_id,self.work.pk)
        self.assertEqual(EquipmentFact.objects.get().actual_count,1)
        section=FactSectionDay.objects.get(section='mechanization');section.confirmed_at=timezone.now();section.save()
        self.save('IDLE',construction_object=self.obj)
        self.assertEqual(EquipmentFact.objects.get().actual_count,0)
        section.refresh_from_db();self.assertIsNone(section.confirmed_at)
        self.save('CLEAR');self.assertFalse(EquipmentDayRecord.objects.exists());self.assertFalse(EquipmentActivity.objects.exists())
        self.assertTrue(EquipmentFact.objects.filter(actual_count=0).exists() or not EquipmentFact.objects.exists())

    def test_repeat_copies_only_previous_work_target_with_explicit_new_day_confirmation(self):
        self.receive();self.save(construction_object=self.obj,work=self.work,comment='Вчерашнее примечание')
        day=self.day+timedelta(days=1);self.save(day=day,repeat=True)
        record=EquipmentDayRecord.objects.get(date=day)
        self.assertEqual(record.work_id,self.work.pk);self.assertEqual(record.status,'WORKED');self.assertEqual(record.comment,'')

    def test_repair_keeps_daily_fact_and_projection_but_inconsistent_movement_is_blocked(self):
        self.receive();self.save(construction_object=self.obj,work=self.work,comment='Работала утром')
        record=EquipmentDayRecord.objects.get(equipment=self.machine,date=self.day);activity_id=record.activity_id
        with self.assertRaises(ValidationError):service.move(self.operator,self.camp,self.day,[self.machine.pk],'DEPARTURE',signed=self.sign())
        service.set_condition(self.operator,self.camp,self.day,[self.machine.pk],'REPAIR',signed=self.sign())
        record.refresh_from_db();self.assertEqual(record.status,'WORKED');self.assertEqual(record.work_id,self.work.pk);self.assertEqual(record.activity_id,activity_id)
        self.assertEqual(service.condition_on(self.machine,self.day),'REPAIR')
        self.save(construction_object=self.obj,work=self.work,comment='Уточнена работа до поломки')
        record.refresh_from_db();self.assertEqual(record.activity_id,activity_id);self.assertEqual(record.activity.state,'WORKING')
        self.assertTrue(service.present(self.company,self.day).filter(equipment=self.machine).exists())

    def test_concurrent_disjoint_saves_merge_and_same_machine_stale_save_rejects(self):
        self.receive();self.receive(self.second)
        first=self.sign();second=self.sign(self.colleague)
        service.save_day(self.operator,self.camp,self.day,[self.machine.pk],'WORKED',signed=first)
        service.save_day(self.colleague,self.camp,self.day,[self.second.pk],'IDLE',signed=second)
        self.assertEqual(EquipmentDayRecord.objects.count(),2)
        with self.assertRaises(ValidationError):service.save_day(self.colleague,self.camp,self.day,[self.machine.pk],'IDLE',signed=second)
        self.assertEqual(EquipmentDayRecord.objects.get(equipment=self.machine).status,'WORKED')

    def test_active_lease_shows_owner_blocks_other_session_and_expires(self):
        self.receive();session=uuid4();other_session=uuid4()
        owners=service.claim(self.operator,self.camp,self.day,[self.machine.pk],'acquire',session)
        self.assertTrue(owners[str(self.machine.pk)]['mine'])
        with self.assertRaisesMessage(ValidationError,self.operator.username):service.claim(self.colleague,self.camp,self.day,[self.machine.pk],'acquire',other_session)
        with self.assertRaises(ValidationError):service.save_day(self.colleague,self.camp,self.day,[self.machine.pk],'WORKED',signed=self.sign(self.colleague),session=other_session)
        self.assertEqual(service.claim(self.colleague,self.camp,self.day,[self.machine.pk],'status',other_session)[str(self.machine.pk)]['owner'],self.operator.username)
        EquipmentTransportControl.objects.filter(equipment=self.machine).update(expires_at=timezone.now()-timedelta(seconds=1))
        self.save(session=other_session)

    def test_scope_future_dates_foreign_ids_and_signed_form_tampering(self):
        with self.assertRaises(PermissionDenied):service.move(self.operator,self.unrelated,self.day,[self.machine.pk],'ARRIVAL',signed=self.sign(camp=self.unrelated))
        foreign_company=Company.objects.create(name='Other transport company')
        foreign=Equipment.objects.create(company=foreign_company,type=self.type,name='Foreign',plate_number='FOREIGN')
        for ids,day,signed in [([foreign.pk],self.day,self.sign()),([self.machine.pk],timezone.localdate()+timedelta(days=1),self.sign()),([self.machine.pk],self.day,self.sign(self.colleague))]:
            with self.assertRaises(ValidationError):service.move(self.operator,self.camp,day,ids,'ARRIVAL',signed=signed)
        self.assertFalse(EquipmentCampStay.objects.exists())

    def test_bulk_failure_is_atomic_and_machines_cannot_overlap_camps(self):
        self.receive(self.second)
        with self.assertRaises(ValidationError):service.move(self.operator,self.camp,self.day,[self.machine.pk,self.second.pk],'ARRIVAL',signed=self.sign())
        self.assertFalse(self.machine.camp_stays.exists())
        with self.assertRaises(ValidationError):service.move(self.operator,self.target,self.day,[self.second.pk],'ARRIVAL',signed=self.sign(camp=self.target))

    def test_day_page_tabs_excel_filter_and_admin_only_legacy_marker_remain_safe(self):
        self.receive();self.receive(self.second,condition='REPAIR');self.save()
        self.client.force_login(self.operator)
        for tab in ['day','calendar','events','report']:
            response=self.client.get(self.board,{'camp':self.camp.pk,'date':self.day,'tab':tab});self.assertEqual(response.status_code,200)
            self.assertNotContains(response,'data-legacy-structure')
        response=self.client.get(self.board,{'camp':self.camp.pk,'date':self.day,'use':'EMPTY','view':'list'});self.assertEqual(len(response.context['rows']),1)
        response=self.client.get(self.board,{'camp':self.camp.pk,'date':self.day,'q':'INV-2','view':'list'});self.assertEqual(len(response.context['rows']),1)
        response=self.client.get(self.board,{'camp':self.camp.pk,'date':self.day,'export':'xlsx'})
        from openpyxl import load_workbook
        ws=load_workbook(BytesIO(response.content)).active
        self.assertEqual(ws.max_row,3);self.assertIn('VIN',[cell.value for cell in ws[1]])
        self.assertEqual(self.client.get(reverse('production:fact_day_workspace')).status_code,302)

    def test_board_post_validates_optional_target_and_roles_csrf(self):
        self.receive();self.client.force_login(self.operator)
        payload={'camp':self.camp.pk,'date':str(self.day),'action':'DAY','equipment':[self.machine.pk],'status':'WORKED','target':f'work:{self.work.pk}:0','signed':self.sign(),'session':str(uuid4())}
        self.assertEqual(self.client.post(self.board,payload).status_code,302)
        self.assertEqual(EquipmentDayRecord.objects.get().work_id,self.work.pk)
        payload['signed']=self.sign();payload['target']='work:999999:0'
        self.assertEqual(self.client.post(self.board,payload).status_code,400)
        csrf=Client(enforce_csrf_checks=True);csrf.force_login(self.operator)
        self.assertEqual(csrf.post(self.board,payload).status_code,403)
        self.client.force_login(self.users['DISPATCHER']);self.assertEqual(self.client.get(self.board).status_code,403)
        self.client.force_login(self.users['PERSONNEL_DISPATCHER']);self.assertEqual(self.client.post(self.board,payload).status_code,403)

    def test_transport_can_add_physical_machine_with_or_without_arrival(self):
        self.client.force_login(self.operator);url=reverse('production:transport_create')
        payload={'camp':self.camp.pk,'date':str(self.day),'type':self.type.pk,'name':'Новая машина','plate_number':'NEW-123','condition':'READY','receive':'on'}
        response=self.client.post(url,payload);self.assertEqual(response.status_code,302)
        machine=Equipment.objects.get(plate_number='NEW-123');self.assertTrue(machine.camp_stays.exists())
        payload.update(name='Только справочник',plate_number='NEW-456');payload.pop('receive')
        self.assertEqual(self.client.post(url,payload).status_code,302)
        self.assertFalse(Equipment.objects.get(plate_number='NEW-456').camp_stays.exists())

    def test_cancel_requires_reason_and_protects_later_events(self):
        self.receive();arrival=self.machine.transport_events.get(kind='ARRIVAL')
        with self.assertRaises(ValidationError):service.cancel_event(self.operator,arrival,'',signed=self.sign())
        day=self.day+timedelta(days=1)
        service.move(self.operator,self.camp,day,[self.machine.pk],'DEPARTURE',signed=self.sign(day=day))
        with self.assertRaises(ValidationError):service.cancel_event(self.operator,arrival,'Ошибка',signed=self.sign())
        departure=self.machine.transport_events.get(kind='DEPARTURE')
        service.cancel_event(self.operator,departure,'Ошибочный выезд',signed=self.sign(day=day))
        service.cancel_event(self.operator,arrival,'Ошибочный приезд',signed=self.sign())
        self.assertFalse(service.present(self.company,self.day).filter(equipment=self.machine).exists())

    def test_legacy_actual_cannot_be_overwritten_and_new_projection_is_protected_from_old_editor(self):
        self.receive()
        EquipmentActivity.objects.create(company=self.company,equipment=self.machine,construction_object=self.obj,date=self.day,state='WORKING',general=True,activity='Прежняя запись')
        with self.assertRaises(ValidationError):self.save(construction_object=self.obj)
        self.assertFalse(EquipmentDayRecord.objects.exists());self.assertEqual(EquipmentActivity.objects.get().activity,'Прежняя запись')
        EquipmentActivity.objects.all().delete();self.save(construction_object=self.obj,work=self.work)
        data=fixtures.DispatcherTests.payload(self,'MECHANIZATION_DISPATCHER','mechanization')
        data['equipment-0-activity']='Перезапись';response=self.client.post(self.url,data)
        self.assertContains(response,'Измените её через',status_code=400)
        self.assertEqual(EquipmentActivity.objects.get().work_id,self.work.pk)

    def test_physical_catalogue_reflects_new_condition_camp_and_historical_writeoff(self):
        from apps.resources.equipment_unit_catalogue import unit_queryset
        from apps.resources.equipment_units import UnitForm
        self.receive(condition='REPAIR')
        row=unit_queryset(self.company,self.day).get(pk=self.machine.pk)
        self.assertEqual(row.unit_state,'REPAIR');self.assertEqual(row.transport_camp_name,self.camp.name)
        day=self.day+timedelta(days=1)
        service.move(self.operator,self.camp,day,[self.machine.pk],'WRITEOFF',signed=self.sign(day=day))
        self.assertEqual(unit_queryset(self.company,day).get(pk=self.machine.pk).unit_state,'WRITTEN_OFF')
        self.assertEqual(unit_queryset(self.company,self.day).get(pk=self.machine.pk).unit_state,'REPAIR')
        form=UnitForm({'type':self.type.pk,'plate_number':'CHANGED','name':self.machine.name,'is_active':'on'},instance=self.machine,company=self.company)
        self.assertFalse(form.is_valid())

    def test_model_link_projects_to_planning_type_and_unlinked_object_is_rejected(self):
        from apps.resources.models import EquipmentType
        model=EquipmentType.objects.create(company=self.company,parent=self.type,name='Model-1')
        Equipment.objects.filter(pk=self.machine.pk).update(type=model);self.machine.refresh_from_db()
        self.receive();self.save(construction_object=self.obj)
        self.assertEqual(EquipmentFact.objects.get().equipment_type_id,self.type.pk)
        with self.assertRaises(ValidationError):self.save(construction_object=self.other)

    def test_old_actual_form_cannot_bypass_camp_presence_or_technical_condition(self):
        self.receive(condition='REPAIR')
        row=EquipmentActivity(company=self.company,equipment=self.machine,construction_object=self.obj,date=self.day,state='WORKING',general=True,activity='Попытка обхода')
        with self.assertRaises(ValidationError):row.full_clean()
        day=self.day+timedelta(days=1)
        service.move(self.operator,self.camp,day,[self.machine.pk],'DEPARTURE',signed=self.sign(day=day))
        row.date=day;row.state='REPAIR'
        with self.assertRaises(ValidationError):row.full_clean()
