"""Read explicit project volumes and historical stock from paired schedule rows.

Formula support is deliberately limited to this format; no Excel expressions are
executed and cached totals are never mistaken for opening stock.
"""
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from django.core.exceptions import ValidationError
from openpyxl.utils import get_column_letter


def amount(value, *, opening=False):
    try:
        if value is None or isinstance(value, bool):
            raise ValueError
        result = Decimal(str(value).replace('\xa0', '').replace(' ', '').replace(',', '.'))
        places = Decimal('.001') if opening else Decimal('.000001')
        if not result.is_finite() or result < 0 or result >= Decimal('1000000000000' if opening else '10000000000') or result != result.quantize(places):
            raise ValueError
        return str(result.quantize(places))
    except (ValueError, InvalidOperation, TypeError):
        raise ValidationError('Нужен неотрицательный объём: до 3 дробных знаков для ВНО, до 6 для плана. Пустая ячейка не считается нулём.')


def period(sheet, header, offset=0):
    for r in range(header - 1, max(0, header - 5), -1):
        value = sheet.cell(r, 11+offset).value
        if isinstance(value, (date, datetime)):
            return value.replace(day=1).date().isoformat() if isinstance(value, datetime) else value.replace(day=1).isoformat()
    return None


def read_measure(sheet, n, header, metric, norm, *, offset=0):
    month = period(sheet, header, offset)
    fact_row = n + 1
    if str(sheet.cell(fact_row, 10+offset).value or '').strip().casefold() != 'факт':
        raise ValidationError(f'{sheet.title}, строка {n}: после «План» ожидается связанная строка «Факт».')
    for col in (1+offset, 2+offset):
        top = ' '.join(str(sheet.cell(n, col).value or '').split()).casefold()
        bottom = ' '.join(str(sheet.cell(fact_row, col).value or '').split()).casefold()
        if bottom and bottom != top:
            raise ValidationError(f'{sheet.title}, строка {fact_row}: участок или комплекс не совпадает со строкой «План».')
    target = n if metric == 'plan' else fact_row
    cell = sheet.cell(target, 8+offset)
    value = cell.value
    daily_columns = []
    for col in range(11+offset, sheet.max_column + 1):
        if sheet.cell(header, col).value != len(daily_columns) + 1:
            break
        daily_columns.append(col)
    filled = sum(sheet.cell(fact_row, col).value is not None for col in daily_columns)
    norm_column=get_column_letter(7+offset);value_column=get_column_letter(8+offset);day_column=get_column_letter(11+offset)
    if cell.data_type == 'f':
        expression = re.sub(r'\s+', '', value).upper().replace('$', '')
        if metric == 'plan' and expression == f'={norm_column}{n}':
            value = norm
        elif metric == 'opening':
            prefix = r'([+]?(?:\d+(?:\.\d*)?|\.\d+))'
            match = re.fullmatch(r'=' + prefix + r'\+SUM\(' + day_column + str(target) + r':([A-Z]+)' + str(target) + r'\)', expression)
            if not match or not daily_columns or match.group(2) != get_column_letter(daily_columns[-1]):
                raise ValidationError(f'{sheet.title}, {value_column}{target}: ожидается начальный факт числом либо формула «число + SUM(дни этой строки)».')
            value = match.group(1)
        else:
            raise ValidationError(f'{sheet.title}, {value_column}{target}: для плана разрешена ссылка «={norm_column}{n}» или число. Сохраните иной расчёт значением.')
    elif metric == 'opening' and filled:
        raise ValidationError(f'{sheet.title}, {value_column}{target}: есть дневной факт и числовой итог. Нельзя определить входящий остаток без задвоения; укажите его формулой «начальный факт + SUM(дни)».')
    try:
        volume = amount(value, opening=metric == 'opening')
    except ValidationError as exc:
        raise ValidationError(f'{sheet.title}, {value_column}{target}: ' + '; '.join(exc.messages))
    return {'volume': volume, 'value_source': f'{sheet.title}:{value_column}{target}', 'file_month': month, 'daily_values_ignored': filled if metric == 'opening' else 0}
