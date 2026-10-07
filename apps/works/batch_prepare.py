"""Editable work/price worksheet and atomic Excel round trip."""

from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
from uuid import uuid4
from urllib.parse import urlencode
from zipfile import ZipFile, BadZipFile

from django import forms
from django.contrib import messages
from django.core import signing
from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.views import View
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from core.permissions import PLAN_ROLES, require_roles
from apps.projects.models import ConstructionObject
from apps.planning.models import GlobalPlanVersion, WorkMonthAllocation, MonthlyPlan
from .models import ProjectWork, ProjectWorkItem, WorkPrice
from .merge_service import (
    source_works,
    preview as merge_preview,
    apply_merge,
    fingerprint,
)
from .merge_views import MergeForm
from .prices import price_on, change_price

SALT = "works.batch-worksheet.v1"
PREVIEW_SALT = "works.batch-preview.v1"
HEADERS = [
    "Проект",
    "Объект",
    "Раздел",
    "Работа",
    "Подработа",
    "Тип работы",
    "Ед. изм.",
    "Цена",
    "Сделать составной работой",
    "Цена составной работы",
    "Норматив",
    "Единица составной работы",
    "ID строки",
]


class SelectionForm(forms.Form):
    construction_object = forms.ModelChoiceField(
        queryset=ConstructionObject.objects.none(), label="Строительный объект"
    )
    version = forms.ModelChoiceField(
        queryset=GlobalPlanVersion.objects.none(),
        label="Версия плана",
        required=False,
        empty_label="Все работы объекта",
    )
    start = forms.DateField(
        label="Начало периода",
        required=False,
        widget=forms.DateInput(attrs={"type": "date"}, format="%Y-%m-%d"),
    )
    end = forms.DateField(
        label="Конец периода",
        required=False,
        widget=forms.DateInput(attrs={"type": "date"}, format="%Y-%m-%d"),
    )

    def __init__(self, *args, user, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["construction_object"].queryset = ConstructionObject.objects.filter(
            company=user.company
        ).select_related("project")
        object_id = self.data.get("construction_object") if self.is_bound else None
        if str(object_id or "").isdigit():
            self.fields["version"].queryset = GlobalPlanVersion.objects.filter(
                company=user.company, construction_object_id=object_id
            ).order_by("-version_number")
        for field in self.fields.values():
            field.widget.attrs["class"] = (
                "form-select form-select-sm"
                if isinstance(field, forms.ModelChoiceField)
                else "form-control form-control-sm"
            )

    def clean(self):
        data = super().clean()
        if data.get("start") and data.get("end") and data["start"] > data["end"]:
            raise ValidationError("Конец периода раньше начала.")
        return data


def selected_works(user, values):
    obj = values["construction_object"]
    qs = (
        ProjectWork.objects.filter(
            company=user.company,
            section__construction_object=obj,
            merged_source__isnull=True,
        )
        .select_related("section__construction_object__project", "load_profile")
        .prefetch_related("items", "price_history")
        .order_by("section__name", "name", "pk")
    )
    version, start, end = values.get("version"), values.get("start"), values.get("end")
    if version or start or end:
        ids = set()
        allocations = WorkMonthAllocation.objects.filter(
            company=user.company, version__construction_object=obj
        )
        plans = MonthlyPlan.objects.filter(
            company=user.company, project_work__section__construction_object=obj
        )
        if version:
            allocations = allocations.filter(version=version)
            plans = plans.none()
        if start:
            allocations = allocations.filter(month__gte=start.replace(day=1))
            plans = plans.filter(end_date__gte=start)
        if end:
            allocations = allocations.filter(month__lte=end.replace(day=1))
            plans = plans.filter(start_date__lte=end)
        ids.update(allocations.values_list("work_id", flat=True))
        ids.update(plans.values_list("project_work_id", flat=True))
        versions = (
            [version]
            if version
            else GlobalPlanVersion.objects.filter(
                company=user.company, construction_object=obj
            )
        )
        for v in versions:
            if start and v.end_date < start or end and v.start_date > end:
                continue
            for spec in v.snapshot.get("works", []):
                daily = spec.get("plans", [])
                if not (start or end) or any(
                    (not start or r["date"] >= start.isoformat())
                    and (not end or r["date"] <= end.isoformat())
                    for r in daily
                ):
                    ids.add(spec["id"])
            lower = v.source_versions.select_related("monthly_plan")
            if start:
                lower = lower.filter(monthly_plan__end_date__gte=start)
            if end:
                lower = lower.filter(monthly_plan__start_date__lte=end)
            ids.update(lower.values_list("monthly_plan__project_work_id", flat=True))
        qs = qs.filter(pk__in=ids)
    return list(qs)


def table_rows(works):
    rows = []
    for work in works:
        obj = work.section.construction_object
        price = str(price_on(work))
        for item in list(work.items.all()) or [None]:
            row_id = f"{work.pk}:{item.pk if item else 0}"
            cells = [
                obj.project.name,
                obj.name,
                work.section.name,
                work.name,
                item.name if item else "",
                work.get_kind_display(),
                item.unit if item else work.unit,
                price,
                "",
                "",
                str(item.quantity_per_unit) if item else "",
                "",
                row_id,
            ]
            rows.append(
                {
                    "id": row_id,
                    "work": work.pk,
                    "item": item.pk if item else None,
                    "cells": cells,
                }
            )
    if len(rows) > 5000:
        raise ValidationError("Больше 5000 строк. Уточните период или версию.")
    return rows


def metadata(user, obj, rows):
    return {"company": user.company_id, "object": obj.pk, "rows": rows}


def decode_metadata(user, token):
    try:
        data = signing.loads(token, salt=SALT, max_age=30 * 86400)
        if data["company"] != user.company_id:
            raise signing.BadSignature()
        get_object_or_404(ConstructionObject, pk=data["object"], company=user.company)
        return data
    except (signing.BadSignature, KeyError, TypeError):
        raise ValidationError(
            "Шаблон устарел или изменён. Выгрузите новый файл из системы."
        )


def export_book(meta):
    book = Workbook()
    sheet = book.active
    sheet.title = "Работы"
    sheet.append(HEADERS)
    for row in meta["rows"]:
        sheet.append(row["cells"])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for c in sheet[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1769AA")
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    for col, width in zip(
        "ABCDEFGHIJKLM", [22, 25, 35, 55, 40, 18, 12, 16, 40, 22, 15, 24, 15]
    ):
        sheet.column_dimensions[col].width = width
    sheet.column_dimensions["M"].hidden = True
    info = book.create_sheet("_Служебное")
    info.sheet_state = "veryHidden"
    token = signing.dumps(meta, salt=SALT, compress=True)
    for offset in range(0, len(token), 30000):
        info.append([token[offset : offset + 30000]])
    stream = BytesIO()
    book.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="works-prices.xlsx"'
    return response


def import_book(user, upload, object_id):
    if upload.size > 8 * 1024 * 1024 or not upload.name.lower().endswith(".xlsx"):
        raise ValidationError("Нужен .xlsx размером до 8 МБ.")
    try:
        with ZipFile(upload) as archive:
            if (
                len(archive.infolist()) > 3000
                or sum(e.file_size for e in archive.infolist()) > 80 * 1024 * 1024
            ):
                raise ValidationError("Excel слишком большой после распаковки.")
        upload.seek(0)
        book = load_workbook(upload, read_only=True, data_only=False, keep_links=False)
        try:
            info = book["_Служебное"]
            if info.max_row > 200:
                raise ValidationError("Слишком большой служебный лист.")
            token = "".join(
                str(r[0] or "") for r in info.iter_rows(max_col=1, values_only=True)
            )
            meta = decode_metadata(user, token)
            if meta["object"] != object_id:
                raise ValidationError("Файл принадлежит другому объекту.")
            sheet = book["Работы"]
            if sheet.max_row > 5001 or sheet.max_column > 13:
                raise ValidationError("Неверный размер таблицы.")
            if (
                list(next(sheet.iter_rows(max_row=1, max_col=13, values_only=True)))
                != HEADERS
            ):
                raise ValidationError("Не изменяйте заголовки столбцов шаблона.")
            entries = []
            for num, row in enumerate(sheet.iter_rows(min_row=2, max_col=13), 2):
                if any(cell.data_type == "f" for cell in row):
                    raise ValidationError(f"Строка {num}: замените формулы значениями.")
                cells = [cell.value if cell.value is not None else "" for cell in row]
                if any(v != "" for v in cells):
                    entries.append(cells)
            return meta, entries
        finally:
            book.close()
    except (BadZipFile, KeyError, ValueError, OSError, TypeError) as error:
        raise ValidationError(
            "Не удалось прочитать шаблон. Используйте файл, выгруженный из этой формы."
        ) from error


def number(value, label, digits=2, positive=False):
    try:
        result = Decimal(
            str(value).replace(" ", "").replace("\xa0", "").replace(",", ".")
        )
        if (
            not result.is_finite()
            or result < 0
            or positive
            and result <= 0
            or result >= (Decimal("10000000") if positive else Decimal("10000000000"))
            or result.as_tuple().exponent < -digits
        ):
            raise ValueError()
        return result
    except (ValueError, TypeError, InvalidOperation):
        raise ValidationError(
            f"{label}: неверное число (до {digits} знаков после запятой)."
        )


def current_state(user, works):
    histories = list(
        WorkPrice.objects.filter(company=user.company, work__in=works)
        .order_by("pk")
        .values()
    )
    return sha256(
        (
            fingerprint(works)
            + json.dumps(
                {
                    "prices": histories,
                    "items": list(
                        ProjectWorkItem.objects.filter(
                            company=user.company, project_work__in=works
                        )
                        .order_by("pk")
                        .values()
                    ),
                },
                default=str,
                sort_keys=True,
            )
        ).encode()
    ).hexdigest()


def build_preview(user, meta, entries, effective):
    require_roles(user, PLAN_ROLES)
    obj = get_object_or_404(ConstructionObject, company=user.company, pk=meta["object"])
    if effective < timezone.localdate():
        raise ValidationError("Для новой цены выберите сегодняшнюю или будущую дату.")
    originals = {r["id"]: r for r in meta["rows"]}
    seen = set()
    grouped = defaultdict(list)
    prices = {}
    selected = set()
    active = {
        w.pk: w
        for w in ProjectWork.objects.filter(
            company=user.company,
            section__construction_object=obj,
            merged_source__isnull=True,
        )
        .select_related("section", "load_profile")
        .prefetch_related("items")
    }
    row_cache = {}
    for index, cells in enumerate(entries, 2):
        rid = str(cells[12])
        original = originals.get(rid)
        if not original or rid in seen:
            raise ValidationError(f"Строка {index}: неверный или повторный ID строки.")
        seen.add(rid)
        work = active.get(original["work"])
        if not work:
            raise ValidationError(f"Строка {index}: работа удалена или уже объединена.")
        if work.pk not in row_cache:
            row_cache[work.pk] = table_rows([work])
        fresh = row_cache[work.pk]
        current = next((r for r in fresh if r["id"] == rid), None)
        if not current or [str(c or "") for c in cells[:7]] != current["cells"][:7]:
            raise ValidationError(
                f"Строка {index}: данные работы изменились. Выгрузите новый шаблон; изменяйте только цены и настройки объединения."
            )
        if work.kind == "COMPOSITE" and (
            any(cells[j] != "" for j in [8, 9, 11])
            or str(cells[10]) != str(current["cells"][10])
        ):
            raise ValidationError(
                f"Строка {index}: у существующей составной работы изменяется только цена; настройки её подработ здесь не изменяются."
            )
        selected.add(work.pk)
        if cells[7] != "":
            price = number(cells[7], f"Строка {index}, цена")
            if work.pk in prices and prices[work.pk] != price:
                raise ValidationError(
                    f"«{work.name}»: укажите одну цену основной работы во всех её строках."
                )
            prices[work.pk] = price
        name = str(cells[8]).strip()
        if name:
            if work.kind != "SIMPLE" or original["item"]:
                raise ValidationError(
                    "Объединять можно только простые работы, не подработы существующих составных."
                )
            if len(name) > 255:
                raise ValidationError("Название составной работы — до 255 символов.")
            grouped[name.casefold()].append(
                {
                    "id": work.pk,
                    "name": name,
                    "norm": str(
                        number(
                            cells[10],
                            f"Строка {index}, норматив",
                            digits=3,
                            positive=True,
                        )
                    ),
                    "unit": str(cells[11]).strip().rstrip(" ."),
                    "price": (
                        str(
                            number(
                                cells[9], f"Строка {index}, цена составной работы"
                            ).quantize(Decimal(".01"))
                        )
                        if cells[9] != ""
                        else None
                    ),
                }
            )
        elif work.kind == "SIMPLE" and any(cells[j] != "" for j in [9, 10, 11]):
            raise ValidationError(
                f"Строка {index}: заполните название составной работы или очистите её настройки."
            )
    if not selected:
        raise ValidationError("Нет строк для обработки.")
    groups = []
    merged_ids = set()
    for members in grouped.values():
        ids = [r["id"] for r in members]
        works = source_works(user, ids)
        name = members[0]["name"]
        units = {r["unit"] for r in members if r["unit"]}
        group_prices = {r["price"] for r in members if r["price"] is not None}
        if len(units) != 1 or len(group_prices) != 1:
            raise ValidationError(
                f"«{name}»: укажите единую единицу и цену составной работы; заполненные значения должны совпадать."
            )
        if any(w.name.casefold() == name.casefold() for w in active.values()):
            raise ValidationError(
                f"Работа «{name}» уже существует. Выберите новое название составной работы."
            )
        unit = next(iter(units))
        data = {
            "name": name,
            "section": works[0].section_id,
            "unit": unit,
            "allow_fractional": "on",
            "work_group": works[0].work_group_id or "",
        }
        for member in members:
            w = active[member["id"]]
            data[f"norm_{w.pk}"] = member["norm"]
            data[f"profile_{w.pk}"] = w.load_profile_id or ""
        form = MergeForm(data, company=user.company, works=works)
        if not form.is_valid():
            raise ValidationError(
                f"«{name}»: "
                + "; ".join(
                    f'{field}: {", ".join(errors)}'
                    for field, errors in form.errors.items()
                )
            )
        report = merge_preview(user, works, form.cleaned_data)
        groups.append(
            {
                "ids": ids,
                "data": data,
                "price": next(iter(group_prices)),
                "name": name,
                "unit": unit,
                "sources": [w.name for w in works],
                "section": works[0].section.name,
                "members": [
                    {
                        "name": w.name,
                        "norm": str(form.cleaned_data[f"norm_{w.pk}"]),
                        "unit": w.unit,
                    }
                    for w in works
                ],
                "fact_quantity": str(report["fact_quantity"]),
                "fact_records": report["fact_records"],
                "plan_count": len(report["plans"]),
            }
        )
        merged_ids.update(ids)
    changes = []
    for wid, value in prices.items():
        if wid in merged_ids:
            continue
        work = active[wid]
        if value != price_on(work, effective):
            changes.append(
                {
                    "id": wid,
                    "name": work.name,
                    "old": str(price_on(work, effective)),
                    "price": str(value),
                }
            )
    if not groups and not changes:
        raise ValidationError(
            "Нет изменений: заполните цены или настройки объединения."
        )
    works = [active[wid] for wid in sorted(selected)]
    return {
        "object": obj.pk,
        "ids": [w.pk for w in works],
        "groups": groups,
        "prices": changes,
        "effective": effective.isoformat(),
        "state": current_state(user, works),
    }


@transaction.atomic
def apply_batch(user, payload):
    require_roles(user, PLAN_ROLES)
    get_object_or_404(
        ConstructionObject.objects.select_for_update(),
        company=user.company,
        pk=payload["object"],
    )
    works = list(
        ProjectWork.objects.select_for_update()
        .filter(
            company=user.company,
            section__construction_object_id=payload["object"],
            pk__in=payload["ids"],
        )
        .order_by("pk")
    )
    if (
        len(works) != len(payload["ids"])
        or current_state(user, works) != payload["state"]
    ):
        raise ValidationError(
            "Работы, цены, план или факт изменились после просмотра. Подготовьте импорт заново."
        )
    effective = date.fromisoformat(payload["effective"])
    revision_count = 0
    if effective < timezone.localdate():
        raise ValidationError(
            "Дата действия цены уже прошла. Подготовьте импорт заново."
        )
    for group in payload["groups"]:
        sources = source_works(user, group["ids"])
        form = MergeForm(group["data"], company=user.company, works=sources)
        if not form.is_valid():
            raise ValidationError(
                "Настройки объединения изменились. Подготовьте импорт заново."
            )
        parent, revisions = apply_merge(
            user, group["ids"], form.cleaned_data, fingerprint(sources)
        )
        change_price(
            user,
            parent.pk,
            Decimal(group["price"]),
            effective,
            reason="Excel: объединение и назначение цены",
        )
        revision_count += len(revisions)
    for entry in payload["prices"]:
        change_price(
            user,
            entry["id"],
            Decimal(entry["price"]),
            effective,
            reason="Массовое назначение цены из таблицы работ",
        )
    return revision_count


class WorkBatchPrepareView(View):
    template_name = "works/batch_prepare.html"

    def get(self, request):
        require_roles(request.user, PLAN_ROLES)
        form = SelectionForm(request.GET or None, user=request.user)
        context = {"selection": form, "effective": timezone.localdate()}
        if form.is_valid():
            try:
                works = selected_works(request.user, form.cleaned_data)
                rows = table_rows(works)
                meta = metadata(
                    request.user, form.cleaned_data["construction_object"], rows
                )
                if request.GET.get("action") == "export":
                    return export_book(meta)
                context.update(
                    rows=rows,
                    meta=signing.dumps(meta, salt=SALT, compress=True),
                    object=form.cleaned_data["construction_object"],
                    export_query=request.GET.urlencode(),
                    loaded=True,
                )
            except ValidationError as error:
                form.add_error(None, error)
        return render(request, self.template_name, context)

    def post(self, request):
        require_roles(request.user, PLAN_ROLES)
        form = meta = entries = None
        try:
            if request.POST.get("action") == "confirm":
                payload = signing.loads(
                    request.POST.get("preview", ""), salt=PREVIEW_SALT, max_age=3600
                )
                if (
                    payload["user"] != request.user.pk
                    or payload["company"] != request.user.company_id
                    or payload["nonce"] != request.session.get("work_batch_nonce")
                ):
                    raise signing.BadSignature()
                count = apply_batch(request.user, payload)
                request.session.pop("work_batch_nonce", None)
                messages.success(
                    request,
                    f'Объединено групп: {len(payload["groups"])}. Изменено цен: {len(payload["prices"])+len(payload["groups"])}. Новых черновиков планов: {count}.',
                )
                return redirect("works:work_list")
            form = SelectionForm(request.POST, user=request.user)
            if not form.is_valid():
                return render(request, self.template_name, {"selection": form})
            if request.POST.get("action") == "import":
                upload = request.FILES.get("file")
                if not upload:
                    raise ValidationError("Выберите заполненный Excel.")
                meta, entries = import_book(
                    request.user, upload, form.cleaned_data["construction_object"].pk
                )
            else:
                meta = decode_metadata(request.user, request.POST.get("meta", ""))
                if meta["object"] != form.cleaned_data["construction_object"].pk:
                    raise ValidationError("Таблица другого объекта.")
                edits = json.loads(request.POST.get("table_data", "{}"))
                if not isinstance(edits, dict):
                    raise ValidationError("Неверный формат изменений таблицы.")
                entries = []
                for row in meta["rows"]:
                    cells = list(row["cells"])
                    for j in range(7, 12):
                        cells[j] = edits.get(
                            f'cell_{row["id"]}_{j}',
                            request.POST.get(f'cell_{row["id"]}_{j}', ""),
                        )
                    entries.append(cells)
            effective = forms.DateField().clean(request.POST.get("effective"))
            payload = build_preview(request.user, meta, entries, effective)
            nonce = uuid4().hex
            request.session["work_batch_nonce"] = nonce
            payload.update(
                user=request.user.pk, company=request.user.company_id, nonce=nonce
            )
            return render(
                request,
                self.template_name,
                {
                    "report": payload,
                    "preview": signing.dumps(payload, salt=PREVIEW_SALT, compress=True),
                },
            )
        except (signing.BadSignature, KeyError, TypeError, ValueError):
            messages.error(
                request, "Просмотр устарел или изменён. Подготовьте таблицу заново."
            )
        except ValidationError as error:
            messages.error(request, "; ".join(error.messages))
            if (
                form is not None
                and form.is_valid()
                and meta is not None
                and entries is not None
            ):
                edited = {str(c[12]): c for c in entries if len(c) == 13}
                display_rows = []
                for row in meta["rows"]:
                    cells = list(row["cells"])
                    if row["id"] in edited:
                        cells[7:12] = edited[row["id"]][7:12]
                    display_rows.append({**row, "cells": cells})
                query = {name: request.POST.get(name, "") for name in form.fields}
                try:
                    display_effective = date.fromisoformat(
                        request.POST.get("effective", "")
                    )
                except ValueError:
                    display_effective = timezone.localdate()
                return render(
                    request,
                    self.template_name,
                    {
                        "selection": form,
                        "loaded": True,
                        "rows": display_rows,
                        "meta": signing.dumps(meta, salt=SALT, compress=True),
                        "object": form.cleaned_data["construction_object"],
                        "effective": display_effective,
                        "export_query": urlencode(query),
                    },
                )
        return redirect("works:batch_prepare")
