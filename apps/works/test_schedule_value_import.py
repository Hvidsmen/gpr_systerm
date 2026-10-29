from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from openpyxl import Workbook
from django.test import TestCase, SimpleTestCase
from django.core import signing
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from apps.accounts.models import Company, User, Role
from apps.projects.models import Project, ConstructionObject, Section
from apps.planning.models import LoadProfile, ObjectWorkPlan, ObjectWorkPlanRow, WorkMonthAllocation
from apps.production.models import WorkOpeningBalance, DailyFact
from apps.production import work_opening
from .models import ProjectWork, ProjectWorkItem
from .schedule_import import HEADERS, parse_schedule, apply_rows
from .schedule_value_views import SALT
from . import schedule_loading

DAY = date(2026, 11, 1)


def source_book(plan=145, stock=3, daily=2, norm=290):
    book = Workbook(); sheet = book.active
    sheet['K2'] = DAY
    for col, value in enumerate(HEADERS + ['Факт с начала строительства', 'Остаток работ', 'План /факт'] + list(range(1, 31)), 1):
        sheet.cell(3, col, value)
    sheet.append(['УЗА 26', 'Видеонаблюдение', '210-СВН', '3.6', 'Кабель', 'м.п.', norm, plan, None, 'План'])
    sheet.append(['УЗА 26', 'Видеонаблюдение', None, None, None, None, None, f'={stock}+SUM(K5:AN5)', None, 'Факт', daily])
    sheet.append(['УЗА 26', 'Видеонаблюдение', '210-СВН', '3.6', 'Оборудование', 'шт', 5, '=G6', None, 'План'])
    sheet.append(['УЗА 26', 'Видеонаблюдение', None, None, None, None, None, '=0+SUM(K7:AN7)', None, 'Факт'])
    for row in (4, 6):
        for col in range(3, 8): sheet.merge_cells(start_row=row, start_column=col, end_row=row + 1, end_column=col)
    stream = BytesIO(); book.save(stream); stream.seek(0)
    return stream


class ValueParserTests(SimpleTestCase):
    def test_norm_project_volume_and_opening_are_independent_and_daily_fact_excluded(self):
        norm, _ = parse_schedule(source_book())
        plan, _ = parse_schedule(source_book(), metric='plan')
        opening, _ = parse_schedule(source_book(), metric='opening')
        self.assertEqual(norm[0]['quantity'], '290.000')
        self.assertEqual(plan[0]['volume'], '145.000000')
        self.assertEqual(opening[0]['volume'], '3.000')
        self.assertEqual(opening[0]['daily_values_ignored'], 1)
        self.assertEqual(opening[0]['file_month'], str(DAY))
        self.assertEqual(plan[1]['volume'], '5.000000')

    def test_blank_negative_unsupported_and_nonfinite_project_values_are_rejected(self):
        for value in (None, -1, 'NaN', '=SUM(K4:AN4)', '=HYPERLINK("https://example.com")'):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                parse_schedule(source_book(plan=value), metric='plan')

    def test_wrong_formula_row_or_unmatched_fact_cannot_become_stock(self):
        from openpyxl import load_workbook
        for cell, value in [('H5', '=3+SUM(K7:AN7)'), ('J5', 'План'), ('A5', 'Другой участок'), ('H5', 5)]:
            book = load_workbook(source_book()); book.active[cell] = value
            stream = BytesIO(); book.save(stream); stream.seek(0)
            with self.subTest(cell=cell, value=value), self.assertRaises(ValidationError):
                parse_schedule(stream, metric='opening')


class ValueImportTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name='График: три назначения')
        self.user = User.objects.create_user(username='schedule-admin', company=self.company, role=Role.objects.get(code='ADMIN'))
        self.planner = User.objects.create_user(username='schedule-planner', company=self.company, role=Role.objects.get(code='PLANNER'))
        project = Project.objects.create(company=self.company, name='Ямал')
        self.obj = ConstructionObject.objects.create(company=self.company, project=project, name='КТП')
        self.section = Section.objects.create(company=self.company, construction_object=self.obj, name='УЗА')
        self.planner.assigned_objects.add(self.obj)
        self.profile = LoadProfile.objects.create(company=self.company, name='Равномерный')
        with source_book() as stream: rows, _ = parse_schedule(stream)
        apply_rows(self.user, self.section, self.profile, rows)
        self.work = ProjectWork.objects.get(section=self.section)
        self.client.force_login(self.user)
        self.plan_url = reverse('planning:object_plan_import')
        self.opening_url = reverse('production:work_opening_import')

    def preview(self, mode='plan', url=None, **fields):
        return self.client.post(url or (self.plan_url if mode == 'plan' else self.opening_url), {
            'construction_object': self.obj.pk, 'section': self.section.pk,
            'name': 'План проекта из графика', 'reason': 'Входящий факт из исходного графика',
            'file': SimpleUploadedFile('График.xlsx', source_book().read()), **fields})

    def confirm(self, response, url=None, mode='plan'):
        return self.client.post(url or (self.plan_url if mode == 'plan' else self.opening_url), {'action': 'confirm', 'token': response.context['token']})

    def test_plan_preview_then_confirm_uses_h_and_replay_does_not_duplicate(self):
        response = self.preview()
        self.assertEqual(response.status_code, 200); self.assertTrue(response.context['token'])
        self.assertFalse(ObjectWorkPlan.objects.exists())
        self.assertEqual(self.confirm(response).status_code, 302)
        plan = ObjectWorkPlan.objects.get()
        self.assertEqual(plan.rows.get(work_item__name='Кабель').quantity, Decimal(145))
        self.assertEqual(plan.rows.get(work_item__isnull=True).quantity, Decimal('.5'))
        self.assertFalse(WorkOpeningBalance.objects.exists()); self.assertFalse(DailyFact.objects.exists())
        self.assertFalse(WorkMonthAllocation.objects.exists())
        self.assertEqual(self.work.items.get(name='Кабель').quantity_per_unit, 290)
        self.assertEqual(self.confirm(response).status_code, 302); self.assertEqual(ObjectWorkPlan.objects.count(), 1)

    def test_existing_plan_updates_only_file_rows_and_preserves_links_and_other_work(self):
        plan = ObjectWorkPlan.objects.create(company=self.company, construction_object=self.obj, name='Существующий', created_by=self.user)
        extra = ProjectWorkItem.objects.create(company=self.company, project_work=self.work, name='Другая работа', unit='шт', quantity_per_unit=2, load_profile=self.profile)
        ObjectWorkPlanRow.objects.create(company=self.company, plan=plan, work=self.work, work_item=extra, quantity=7)
        url = reverse('planning:object_plan_import_into', args=[plan.pk])
        response = self.preview(url=url)
        self.assertEqual(self.confirm(response, url=url).status_code, 302)
        plan.refresh_from_db(); self.assertEqual(plan.name, 'Существующий')
        self.assertEqual(plan.rows.get(work_item=extra).quantity, 7)
        self.assertEqual(plan.rows.get(work_item__name='Кабель').quantity, 145)

    def test_unknown_work_or_norm_conflict_blocks_whole_preview(self):
        self.work.items.filter(name='Кабель').update(quantity_per_unit=300)
        response = self.preview()
        self.assertEqual(response.context['issues'], 1); self.assertNotIn('token', response.context)
        self.assertFalse(ObjectWorkPlan.objects.exists())

    def test_catalogue_changed_after_preview_blocks_confirmation(self):
        response = self.preview()
        self.work.items.filter(name='Оборудование').update(name='Новое имя')
        self.assertContains(self.confirm(response), 'Справочник или доступ')
        self.assertFalse(ObjectWorkPlan.objects.exists())

    def test_stale_existing_plan_does_not_overwrite_colleague(self):
        plan = ObjectWorkPlan.objects.create(company=self.company, construction_object=self.obj, created_by=self.user)
        row = ObjectWorkPlanRow.objects.create(company=self.company, plan=plan, work=self.work, work_item=self.work.items.first(), quantity=20)
        url = reverse('planning:object_plan_import_into', args=[plan.pk]); response = self.preview(url=url)
        row.quantity = 99; row.save()
        self.assertContains(self.confirm(response, url=url), 'План изменился')
        row.refresh_from_db(); self.assertEqual(row.quantity, 99)

    def test_opening_date_from_file_and_base_only_without_daily_facts(self):
        response = self.preview('opening'); self.assertEqual(response.context['opening_date'], DAY)
        self.assertEqual(self.confirm(response, mode='opening').status_code, 302)
        balance = WorkOpeningBalance.objects.get()
        self.assertEqual((balance.quantity, balance.date), (Decimal(3), DAY))
        self.assertFalse(DailyFact.objects.exists()); self.assertFalse(ObjectWorkPlan.objects.exists())
        self.assertEqual(self.confirm(response, mode='opening').status_code, 302)
        self.assertEqual(WorkOpeningBalance.objects.count(), 1)

    def test_earlier_fact_blocks_opening_preview(self):
        DailyFact.objects.create(company=self.company, project_work=self.work, work_item=self.work.items.get(name='Кабель'), date=DAY-timedelta(days=1), actual_quantity=1)
        response = self.preview('opening')
        self.assertNotIn('token', response.context); self.assertGreater(response.context['issues'], 0)
        self.assertFalse(WorkOpeningBalance.objects.exists())

    def test_stock_changed_after_preview_and_tampered_token_cannot_overwrite(self):
        response = self.preview('opening')
        item = self.work.items.get(name='Кабель')
        WorkOpeningBalance.objects.create(company=self.company, project_work=self.work, work_item=item, date=DAY, quantity=8)
        self.assertContains(self.confirm(response, mode='opening'), 'изменились')
        self.assertEqual(WorkOpeningBalance.objects.get().quantity, 8)
        self.assertContains(self.client.post(self.opening_url, {'action': 'confirm', 'token': 'bad'}), 'недействителен')

    def test_zero_explicitly_clears_and_partial_file_preserves_other_openings(self):
        other = self.work.items.get(name='Оборудование')
        WorkOpeningBalance.objects.create(company=self.company, project_work=self.work, work_item=other, date=DAY, quantity=4)
        response = self.preview('opening'); self.assertEqual(self.confirm(response, mode='opening').status_code, 302)
        self.assertEqual(WorkOpeningBalance.objects.get(work_item=other).quantity, 0)

    def test_confirmation_rolls_back_when_any_row_fails(self):
        response = self.preview()
        with patch.object(ObjectWorkPlanRow, 'full_clean', side_effect=ValidationError('Ошибка второй стадии')):
            self.assertContains(self.confirm(response), 'Ошибка второй стадии')
        self.assertFalse(ObjectWorkPlan.objects.exists())

    def test_wrong_user_or_import_kind_and_revoked_scope_are_rejected(self):
        response = self.preview()
        self.client.force_login(self.planner)
        self.assertContains(self.confirm(response), 'другому пользователю')
        response = self.preview(); self.planner.assigned_objects.clear()
        self.assertEqual(self.confirm(response).status_code, 404)
        self.assertFalse(ObjectWorkPlan.objects.exists())

    def test_admin_only_opening_and_discoverable_navigation(self):
        self.assertContains(self.client.get(reverse('production:work_opening_list')), 'Входящие остатки (ВНО)')
        self.assertContains(self.client.get(reverse('production:work_opening_list')), reverse('production:work_opening_import'))
        for code in ['PLANNER', 'MANAGER', 'FOREMAN', 'OBSERVER']:
            user = User.objects.create_user(username='import-' + code, company=self.company, role=Role.objects.get(code=code))
            user.assigned_objects.add(self.obj); self.client.force_login(user)
            self.assertEqual(self.client.get(self.opening_url).status_code, 403)
            self.assertEqual(self.client.post(self.opening_url, {'action': 'confirm'}).status_code, 403)
        self.assertFalse(WorkOpeningBalance.objects.exists())

    def test_other_company_objects_and_plans_are_unavailable(self):
        company = Company.objects.create(name='Другая компания')
        project = Project.objects.create(company=company, name='Чужой')
        obj = ConstructionObject.objects.create(company=company, project=project, name='Чужой')
        plan = ObjectWorkPlan.objects.create(company=company, construction_object=obj, created_by=self.user)
        self.assertEqual(self.client.get(reverse('planning:object_plan_import_into', args=[plan.pk])).status_code, 404)
        response = self.preview(construction_object=obj.pk)
        self.assertTrue(response.context['form'].errors)
        self.assertFalse(ObjectWorkPlanRow.objects.exists())
