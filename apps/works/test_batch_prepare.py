from copy import deepcopy
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.planning import test_workspace
from apps.planning.models import WorkMonthAllocation
from apps.planning.workspace_services import WorkspaceService
from apps.production.models import DailyFact
from . import test_merge
from .models import ProjectWork, WorkMergeSource
from .prices import change_price, price_on
from .batch_prepare import (
    table_rows,
    metadata,
    selected_works,
    export_book,
    import_book,
    build_preview,
    apply_batch,
    SelectionForm,
)


class BatchPrepareTests(TestCase):
    setUpTestData = classmethod(test_workspace.WorkspaceTests.setUpTestData.__func__)

    def setUp(self):
        test_merge.WorkMergeTests.setUp(self)
        self.today = timezone.localdate()
        self.rows = table_rows(self.works)
        self.meta = metadata(self.planner, self.obj, self.rows)
        self.entries = [list(r["cells"]) for r in self.rows]
        self.url = reverse("works:batch_prepare")

    def group_entries(self, name="New assembly", price="120"):
        entries = deepcopy(self.entries)
        for index, cells in enumerate(entries):
            cells[8] = name
            cells[9] = price if index == 0 else ""
            cells[10] = 2 if index == 0 else 3
            cells[11] = "шт" if index == 0 else ""
        return entries

    def test_selection_all_and_optional_version_period(self):
        allworks = selected_works(self.planner, {"construction_object": self.obj})
        self.assertEqual(
            {w.pk for w in allworks},
            {self.simple.pk, self.second.pk, self.composite.pk},
        )
        workspace = WorkspaceService.create(
            self.planner, self.obj, "Filter", date(2026, 1, 1), date(2026, 2, 28)
        )
        for month, work in [
            (date(2026, 1, 1), self.simple),
            (date(2026, 2, 1), self.second),
        ]:
            WorkMonthAllocation.objects.create(
                company=self.company,
                version=workspace.baseline_version,
                work=work,
                month=month,
                quantity=10,
            )
        selected = selected_works(
            self.planner,
            {
                "construction_object": self.obj,
                "version": workspace.baseline_version,
                "start": date(2026, 2, 15),
                "end": date(2026, 2, 28),
            },
        )
        self.assertEqual([w.pk for w in selected], [self.second.pk])
        form = SelectionForm(
            {
                "construction_object": self.obj.pk,
                "version": workspace.baseline_version.pk,
            },
            user=self.planner,
        )
        self.assertTrue(form.is_valid())

    def test_excel_round_trip_and_formula_rejected(self):
        response = export_book(self.meta)
        book = load_workbook(BytesIO(response.content))
        sheet = book["Работы"]
        self.assertTrue(sheet.column_dimensions["M"].hidden)
        sheet.cell(2, 8).value = 45
        stream = BytesIO()
        book.save(stream)
        meta, entries = import_book(
            self.planner,
            SimpleUploadedFile("prices.xlsx", stream.getvalue()),
            self.obj.pk,
        )
        payload = build_preview(self.planner, meta, entries, self.today)
        self.assertEqual(payload["prices"][0]["price"], "45")
        self.assertFalse(WorkMergeSource.objects.exists())
        sheet.cell(2, 8).value = "=1+2"
        stream = BytesIO()
        book.save(stream)
        with self.assertRaises(ValidationError):
            import_book(
                self.planner,
                SimpleUploadedFile("p.xlsx", stream.getvalue()),
                self.obj.pk,
            )

    def test_merge_fact_plan_and_explicit_parent_price_history(self):
        version = test_merge.WorkMergeTests.plan(self)
        fact = DailyFact.objects.create(
            company=self.company,
            project_work=self.simple,
            date=date(2026, 1, 1),
            actual_quantity=18,
            reported_by=self.planner,
        )
        payload = build_preview(
            self.planner, self.meta, self.group_entries(), self.today
        )
        apply_batch(self.planner, payload)
        parent = ProjectWork.objects.get(name="New assembly")
        self.assertEqual(price_on(parent), Decimal(120))
        self.assertEqual(parent.items.count(), 2)
        fact.refresh_from_db()
        self.assertEqual(fact.project_work, parent)
        self.assertEqual(fact.actual_quantity, 18)
        self.assertEqual(fact.reported_by, self.planner)
        self.assertEqual(version.work_allocations.filter(work=parent).count(), 2)
        self.assertTrue(
            parent.price_history.filter(created_by=self.planner, price=120).exists()
        )
        self.simple.refresh_from_db()
        self.assertEqual(self.simple.unit_price, 10)

    def test_existing_composite_repeated_price_and_history(self):
        rows = table_rows([self.composite])
        meta = metadata(self.planner, self.obj, rows)
        entries = [list(r["cells"]) for r in rows]
        entries[0][7] = "75"
        entries[1][7] = "75.00"
        payload = build_preview(self.planner, meta, entries, self.today)
        self.assertEqual(len(payload["prices"]), 1)
        apply_batch(self.planner, payload)
        self.assertEqual(price_on(self.composite), Decimal(75))
        entries[1][7] = "80"
        with self.assertRaises(ValidationError):
            build_preview(self.planner, meta, entries, self.today)

    def test_conflicts_duplicates_missing_norm_and_single_member(self):
        variants = []
        cells = self.group_entries()
        cells[1][9] = "121"
        variants.append(cells)
        cells = self.group_entries()
        cells[0][10] = ""
        variants.append(cells)
        cells = self.group_entries()
        cells[1][8] = ""
        cells[1][10] = ""
        variants.append(cells)
        cells = self.group_entries()
        cells[1][11] = "m"
        variants.append(cells)
        cells = self.group_entries()
        cells[1][12] = cells[0][12]
        variants.append(cells)
        for entries in variants:
            with self.subTest(entries=entries), self.assertRaises(ValidationError):
                build_preview(self.planner, self.meta, entries, self.today)
        self.assertFalse(WorkMergeSource.objects.exists())

    def test_stale_preview_and_atomic_rollback(self):
        payload = build_preview(
            self.planner, self.meta, self.group_entries(), self.today
        )
        change_price(
            self.planner, self.simple.pk, Decimal(55), self.today + timedelta(days=2)
        )
        with self.assertRaises(ValidationError):
            apply_batch(self.planner, payload)
        self.assertFalse(WorkMergeSource.objects.exists())
        payload = build_preview(
            self.planner, self.meta, self.group_entries(), self.today
        )
        payload["prices"].append({"id": self.simple.pk, "price": "-1"})
        with self.assertRaises(ValidationError):
            apply_batch(self.planner, payload)
        self.assertFalse(WorkMergeSource.objects.exists())
        self.assertFalse(ProjectWork.objects.filter(name="New assembly").exists())

    def test_multiple_groups_share_versions(self):
        version = test_merge.WorkMergeTests.plan(self)
        extras = []
        for name in ["third", "fourth"]:
            work = ProjectWork.objects.create(
                company=self.company,
                section=self.simple.section,
                name=name,
                unit="m",
                load_profile=self.profile,
            )
            extras.append(work)
            WorkMonthAllocation.objects.create(
                company=self.company,
                version=version,
                work=work,
                month=date(2026, 1, 1),
                quantity=8,
            )
        rows = table_rows(self.works + extras)
        meta = metadata(self.planner, self.obj, rows)
        entries = self.group_entries()
        for row in rows[2:]:
            cells = list(row["cells"])
            cells[8] = "Other assembly"
            cells[9] = "60"
            cells[10] = "1"
            cells[11] = "шт"
            entries.append(cells)
        payload = build_preview(self.planner, meta, entries, self.today)
        apply_batch(self.planner, payload)
        self.assertEqual(WorkMergeSource.objects.count(), 4)
        self.assertEqual(
            ProjectWork.objects.filter(
                name__in=["Other assembly", "New assembly"]
            ).count(),
            2,
        )

    def test_view_web_edit_confirmation_and_replay(self):
        import json

        response = self.client.get(self.url, {"construction_object": self.obj.pk})
        self.assertContains(response, "Выгрузить в Excel")
        self.assertContains(response, "Сделать составной работой")
        edits = {
            f'cell_{r["id"]}_{j}': str(r["cells"][j])
            for r in response.context["rows"]
            for j in range(7, 12)
        }
        edits[f"cell_{self.simple.pk}:0_7"] = "32"
        data = {
            "construction_object": self.obj.pk,
            "meta": response.context["meta"],
            "table_data": json.dumps(edits),
            "effective": self.today.isoformat(),
            "action": "table",
        }
        response = self.client.post(self.url, data)
        self.assertContains(response, "Подтвердить объединение и цены")
        token = response.context["preview"]
        self.client.post(self.url, {"action": "confirm", "preview": token})
        self.assertEqual(price_on(self.simple), 32)
        count = self.simple.price_history.count()
        self.client.post(self.url, {"action": "confirm", "preview": token})
        self.assertEqual(self.simple.price_history.count(), count)
        self.client.force_login(self.manager)
        self.assertEqual(self.client.get(self.url).status_code, 403)

    def test_import_view_and_invalid_edit_keeps_values(self):
        response = export_book(self.meta)
        book = load_workbook(BytesIO(response.content))
        for index, cells in enumerate(self.group_entries(), 2):
            for col, value in enumerate(cells, 1):
                book["Работы"].cell(index, col).value = value
        stream = BytesIO()
        book.save(stream)
        response = self.client.post(
            self.url,
            {
                "action": "import",
                "construction_object": self.obj.pk,
                "effective": self.today.isoformat(),
                "file": SimpleUploadedFile("filled.xlsx", stream.getvalue()),
            },
        )
        self.assertContains(response, "Норматив на единицу работы")
        self.assertEqual(len(response.context["report"]["groups"]), 1)
        self.assertFalse(WorkMergeSource.objects.exists())
        # Invalid values remain editable instead of discarding the entered group.
        for index, cells in enumerate(self.group_entries(), 2):
            for col, value in enumerate(cells, 1):
                book["Работы"].cell(index, col).value = value
        book["Работы"].cell(2, 11).value = None
        stream = BytesIO()
        book.save(stream)
        response = self.client.post(
            self.url,
            {
                "action": "import",
                "construction_object": self.obj.pk,
                "effective": self.today.isoformat(),
                "file": SimpleUploadedFile("filled.xlsx", stream.getvalue()),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "New assembly")
        self.assertContains(response, "неверное число")

    def test_wrong_object_company_and_tampered_identity_blocked(self):
        from apps.projects.models import ConstructionObject

        other = ConstructionObject.objects.create(
            company=self.company, project=self.obj.project, name="Other"
        )
        data = export_book(self.meta).content
        with self.assertRaises(ValidationError):
            import_book(self.planner, SimpleUploadedFile("filled.xlsx", data), other.pk)
        foreign = deepcopy(self.meta)
        foreign["company"] = self.company.pk + 999
        with self.assertRaises(ValidationError):
            import_book(
                self.planner,
                SimpleUploadedFile("filled.xlsx", export_book(foreign).content),
                self.obj.pk,
            )
        entries = deepcopy(self.entries)
        entries[0][3] = "Altered work name"
        with self.assertRaises(ValidationError):
            build_preview(self.planner, self.meta, entries, self.today)
        entries = deepcopy(self.entries)
        entries[0][7] = -1
        with self.assertRaises(ValidationError):
            build_preview(self.planner, self.meta, entries, self.today)

    def test_import_stale_object_returns_message_without_changes(self):
        stale = deepcopy(self.meta)
        stale["object"] = self.obj.pk + 99999
        before = self.simple.price_history.count()
        response = self.client.post(
            self.url,
            {
                "action": "import",
                "construction_object": self.obj.pk,
                "effective": self.today.isoformat(),
                "file": SimpleUploadedFile("old.xlsx", export_book(stale).content),
            },
            follow=True,
        )
        self.assertContains(response, "выгрузите новый шаблон Excel")
        self.assertEqual(self.simple.price_history.count(), before)
        self.assertFalse(WorkMergeSource.objects.exists())

    def test_missing_object_rejected_at_preview_and_confirmation(self):
        stale = deepcopy(self.meta)
        stale["object"] = self.obj.pk + 99999
        with self.assertRaisesMessage(ValidationError, "больше недоступен"):
            build_preview(self.planner, stale, self.entries, self.today)
        entries = deepcopy(self.entries)
        entries[0][7] = "77"
        payload = build_preview(self.planner, self.meta, entries, self.today)
        payload["object"] = stale["object"]
        with self.assertRaisesMessage(ValidationError, "больше недоступен"):
            apply_batch(self.planner, payload)
