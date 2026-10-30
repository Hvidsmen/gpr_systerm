"""Parse legacy and section-aware schedules; preview an atomic catalogue sync."""
from collections import OrderedDict
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException
from django import forms
from django.core import signing
from django.core.exceptions import ValidationError
from django.contrib import messages
from django.db import transaction
from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect
from openpyxl import load_workbook
from core.permissions import PLAN_ROLES, require_roles, scope_queryset
from apps.projects.models import Section, ConstructionObject
from apps.planning.models import LoadProfile
from apps.audit.models import AuditLog
from .models import ComplexObject, ProjectWork, ProjectWorkItem, MeasurementUnit

HEADERS=['Участок производства работ','Группа работ','Раздел РД','Пункт в РДЦ','Вид работ','Ед. изм.','Проектное количество']
SALT='works.schedule-catalogue.v2'

def text(value):
    return ' '.join(str(value).split()) if value is not None else ''

def parse_schedule(upload, *, metric="catalogue"):
    if metric not in {"catalogue", "plan", "opening"}:raise ValueError("Unknown schedule metric")
    raw=upload.read(5*1024*1024+1)
    if len(raw)>5*1024*1024:raise ValidationError('Размер файла не должен превышать 5 МБ.')
    try:
        with ZipFile(BytesIO(raw)) as archive:
            if len(archive.infolist())>5000 or sum(entry.file_size for entry in archive.infolist())>32*1024*1024:
                raise ValidationError('Слишком большой распакованный файл Excel.')
        book=load_workbook(BytesIO(raw),data_only=False)
    except (BadZipFile,ValueError,KeyError,OSError,ParseError,InvalidFileException) as exc:raise ValidationError('Не удалось прочитать XLSX.') from exc
    rows=OrderedDict();ignored=0;matched=0
    try:
        for sheet in book:
            if sheet.max_row>5000 or sheet.max_column>100:raise ValidationError('Лист превышает предел: 5000 строк и 100 столбцов.')
            layout=next(((n,offset) for n in range(1,min(sheet.max_row,30)+1) for offset in (0,1) if [text(sheet.cell(n,c+offset).value).casefold() for c in range(1,8)]==[h.casefold() for h in HEADERS] and (not offset or text(sheet.cell(n,1).value).casefold()=='раздел')),None)
            if not layout:continue
            header,offset=layout
            matched+=1
            merges=[area for area in sheet.merged_cells.ranges if area.min_col<=8]
            def cell(n,c):
                result=sheet.cell(n,c)
                if result.value is None:
                    for area in merges:
                        if area.min_row<=n<=area.max_row and area.min_col<=c<=area.max_col:
                            result=sheet.cell(area.min_row,area.min_col);break
                return result
            for n in range(header+1,sheet.max_row+1):
                label=text(sheet.cell(n,10+offset).value).casefold().replace(' ','')
                if label=='факт':ignored+=1;continue
                if label!='план':
                    if any(sheet.cell(n,c+offset).value is not None for c in range(1,8)):
                        raise ValidationError(f'{sheet.title}, строка {n}: ожидается «План» или «Факт» в столбце «План /факт».')
                    continue
                cells=[cell(n,c+offset) for c in range(1,8)]
                if any(c.data_type=='f' for c in cells):raise ValidationError(f'{sheet.title}, строка {n}: в реквизитах работы нужны значения, а не формулы.')
                values=[text(c.value) for c in cells]
                for required_index in (0,1,4,5):
                    if not values[required_index]:raise ValidationError(f'{sheet.title}, строка {n}: не заполнен столбец «{HEADERS[required_index]}».')
                if any(len(value)>255 for value in values[:5]) or len(values[5])>50:raise ValidationError(f'{sheet.title}, строка {n}: слишком длинное название или единица измерения.')
                try:quantity=Decimal(values[6].replace(' ','').replace(',','.'))
                except InvalidOperation:raise ValidationError(f'{sheet.title}, строка {n}: неверный норматив.')
                if not quantity.is_finite() or quantity<=0 or quantity>=Decimal('10000000') or quantity!=quantity.quantize(Decimal('.001')):
                    raise ValidationError(f'{sheet.title}, строка {n}: норматив должен быть положительным, до 7 целых и 3 дробных знаков.')
                row=dict(zip(['object','package','rd_section','rdc_point','work','unit'],values[:6]),quantity=str(quantity.quantize(Decimal('.001'))),source=f'{sheet.title}:{n}')
                if offset:
                    section_cell=cell(n,1)
                    section_name=text(section_cell.value)
                    if section_cell.data_type=='f' or not section_name or len(section_name)>255:raise ValidationError(f'{sheet.title}, строка {n}: укажите название раздела текстом.')
                    row['section']=section_name
                if metric != 'catalogue':
                    from .schedule_values import read_measure
                    row.update(read_measure(sheet,n,header,metric,row['quantity'],offset=offset))
                key=(row.get('section','').casefold(),row['object'].casefold(),row['package'].casefold(),row['work'].casefold())
                if key in rows and any(rows[key][field]!=row[field] for field in (('rd_section','rdc_point','unit','quantity') if metric == 'catalogue' else ('rd_section','rdc_point','unit','quantity','volume','file_month'))):
                    raise ValidationError(f'{row["source"]}: повтор работы с различающимися реквизитами или нормативом.')
                rows.setdefault(key,row)
                if len(rows)>2000:raise ValidationError('Можно импортировать до 2000 работ за раз.')
        if not matched or not rows:raise ValidationError('Не найдены строки «План» с заголовками графика (с разделом или без него).')
        return list(rows.values()),ignored
    finally:book.close()

class ImportForm(forms.Form):
    construction_object=forms.ModelChoiceField(queryset=ConstructionObject.objects.none(),label='Строительный объект',empty_label='Выберите объект')
    section=forms.ModelChoiceField(queryset=Section.objects.none(),required=False,label='Раздел для старого формата без столбца «Раздел»',help_text='В новом формате разделы берутся из файла и создаются автоматически.')
    load_profile=forms.ModelChoiceField(queryset=LoadProfile.objects.none(),label='Профиль нагрузки вложенных работ',help_text='В файле профилей нет. Этот профиль будет назначен только новым работам.')
    file=forms.FileField(label='График XLSX')
    def __init__(self,*args,user,**kwargs):
        super().__init__(*args,**kwargs)
        objects=scope_queryset(ConstructionObject.objects.select_related('project'),user)
        sections=scope_queryset(Section.objects.select_related('construction_object'),user)
        self.fields['construction_object'].queryset=objects
        self.fields['construction_object'].label_from_instance=lambda row:f'{row.project.name} / {row.name}'
        object_id=self.data.get(self.add_prefix('construction_object')) if self.is_bound else self.initial.get('construction_object')
        try:object_id=int(object_id)
        except (TypeError,ValueError):object_id=None
        self.fields['section'].queryset=sections.filter(construction_object_id=object_id,construction_object__in=objects) if object_id else sections.none()
        self.fields['section'].empty_label='Выберите раздел' if object_id else 'Сначала выберите объект'
        self.fields['section'].label_from_instance=lambda row:row.name
        self.section_options=list(sections.filter(construction_object__in=objects).order_by('name','pk').values('id','construction_object_id','name'))
        self.fields['load_profile'].queryset=LoadProfile.objects.filter(company=user.company)


def inspect_rows(company,section,rows):
    from .schedule_sync import inspect_target
    return inspect_target(company,section.construction_object,rows,section)['rows']


def apply_rows(user,section,profile,rows):
    from .schedule_sync import apply_catalogue
    return apply_catalogue(user,section.construction_object,profile,rows,section)['created']


@login_required
def schedule_import(request):
    from .schedule_sync import preview_catalogue,apply_catalogue
    require_roles(request.user,PLAN_ROLES)
    form=ImportForm(request.POST or None,request.FILES or None,user=request.user,initial={'construction_object':request.GET.get('construction_object'),'section':request.GET.get('section')})
    context={'form':form}
    if request.method=='POST':
        try:
            if request.POST.get('action')=='confirm':
                data=signing.loads(request.POST.get('token',''),salt=SALT,max_age=900)
                if data['user']!=request.user.pk or data['company']!=request.user.company_id:raise ValidationError('Предпросмотр принадлежит другому пользователю.')
                obj=scope_queryset(ConstructionObject.objects.all(),request.user).filter(pk=data['object']).first()
                section=scope_queryset(Section.objects.all(),request.user).filter(pk=data.get('section'),construction_object=obj).first() if data.get('section') else None
                profile=LoadProfile.objects.filter(company=request.user.company,pk=data['profile']).first()
                if not obj or not profile or (data.get('section') and not section):raise ValidationError('Раздел или профиль больше недоступен. Повторите предпросмотр.')
                counts=apply_catalogue(request.user,obj,profile,data['rows'],section,expected=data['revision'])
                messages.success(request,f"Объектов: {counts['objects']}. Добавлено работ: {counts['created']}, обновлено: {counts['updated']}, помечено старой структурой: {counts['archived']}.")
                return redirect('works:work_list')
            elif form.is_valid():
                rows,ignored=parse_schedule(form.cleaned_data['file'])
                obj=form.cleaned_data['construction_object'];section=form.cleaned_data['section']
                preview=preview_catalogue(request.user,obj,rows,section)
                data={'user':request.user.pk,'company':request.user.company_id,'object':obj.pk,'section':section.pk if section else None,'profile':form.cleaned_data['load_profile'].pk,'rows':rows,'revision':preview['revision']}
                context.update(preview,ignored=ignored,object=obj,token=signing.dumps(data,salt=SALT,compress=True))
        except signing.BadSignature:context['error']='Предпросмотр недействителен или истёк. Загрузите файл заново.'
        except ValidationError as exc:context['error']='; '.join(exc.messages)
    return render(request,'works/schedule_import.html',context)
