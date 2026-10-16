from io import BytesIO
from django.test import SimpleTestCase, TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from apps.accounts.models import Company, Role, User
from .excel import parse, export
from .models import EstimateDocument, EstimateCatalog


def sample():
    w=Workbook();s=w.active
    s['A1']='ЛОКАЛЬНЫЙ СМЕТНЫЙ РАСЧЕТ';s['B3']='Обоснование';s['C3']='Наименование работ и затрат'
    s['A7']='Раздел 1. Монтаж'
    for r,a,b,c,u,q,base,cur in [(8,'1','ФЕР','Работа','м',2,None,None),(9,None,'91.1','Кран','маш.час',3,10,None),(10,None,None,'Всего по позиции',None,None,20,100),(11,None,None,'Итого по разделу 1',None,None,20,100),(12,None,'91.1','Повторная ведомость',None,None,10,50),(13,None,None,'ВСЕГО по смете',None,None,20,100)]:
        for col,val in [(1,a),(2,b),(3,c),(6,u),(9,q),(12,base),(14,cur)]:s.cell(r,col,val)
    out=BytesIO();w.save(out);return out.getvalue()


class ParserTests(SimpleTestCase):
    def test_totals_and_resources_are_not_double_counted(self):
        d=parse(sample(),'sample.xlsx')
        self.assertEqual(d['current_total'],'100')
        self.assertEqual([r['kind'] for r in d['rows']],['section','item','resource','detail','total','appendix','total'])
        self.assertEqual(d['rows'][2]['parent_row'],8)
        w=load_workbook(BytesIO(export([d])),data_only=False)
        self.assertEqual(w['Свод смет']['D2'].value,100)
        self.assertEqual(w['Ресурсы'].max_row,2)
        self.assertEqual(w['Позиции'].max_row,2)
    def test_unknown_format(self):
        with self.assertRaises(ValueError):parse(b'not excel')


class AccessTests(TestCase):
    def setUp(self):
        self.company=Company.objects.create(name='Estimates')
        self.role=Role.objects.get_or_create(code='PLANNER',defaults={'name':'Planner'})[0]
        self.user=User.objects.create_user(username='estimate-planner',company=self.company,role=self.role)
        self.client.force_login(self.user)
    def test_import_duplicate_and_scope(self):
        url=reverse('estimates:estimate_import')
        for _ in range(2):
            response=self.client.post(url,{'object_name':'Independent object','file':SimpleUploadedFile('sample.xlsx',sample())})
            self.assertEqual(response.status_code,302)
        self.assertEqual(EstimateDocument.objects.count(),1)
        self.assertEqual(EstimateCatalog.objects.count(),4)
        self.assertEqual(self.client.get(response.url).status_code,200)
        self.assertEqual(self.client.get(reverse('estimates:estimate_export')).status_code,200)
        other=Company.objects.create(name='Other');self.user.company=other;self.user.save()
        self.assertEqual(self.client.get(response.url).status_code,404)
    def test_observer_read_only(self):
        self.user.role=Role.objects.get_or_create(code='OBSERVER',defaults={'name':'Observer'})[0];self.user.save()
        self.assertEqual(self.client.get(reverse('estimates:estimate_list')).status_code,200)
        self.assertEqual(self.client.get(reverse('estimates:catalog_list')).status_code,200)
        self.assertEqual(self.client.post(reverse('estimates:estimate_import')).status_code,403)

    def test_upload_button_and_multiple_selection(self):
        response=self.client.get(reverse('estimates:estimate_list'))
        self.assertContains(response,'Загрузить сметы (несколько файлов)')
        self.assertContains(self.client.get(reverse('estimates:estimate_import')),' multiple')
        from openpyxl import load_workbook
        second=load_workbook(BytesIO(sample()));second.active['C8']='Другая работа';out=BytesIO();second.save(out)
        response=self.client.post(reverse('estimates:estimate_import'),{'object_name':'Batch','file':[SimpleUploadedFile('one.xlsx',sample()),SimpleUploadedFile('two.xlsx',out.getvalue()),SimpleUploadedFile('repeat.xlsx',sample())]})
        self.assertRedirects(response,reverse('estimates:estimate_list'))
        self.assertEqual(EstimateDocument.objects.count(),2)

    def test_invalid_batch_saves_nothing_and_names_bad_file(self):
        response=self.client.post(reverse('estimates:estimate_import'),{'object_name':'Batch','file':[SimpleUploadedFile('one.xlsx',sample()),SimpleUploadedFile('broken.xlsx',b'broken')]})
        self.assertEqual(response.status_code,200)
        self.assertContains(response,'broken.xlsx')
        self.assertEqual(EstimateDocument.objects.count(),0)

    def test_empty_batch_and_read_only_explanation(self):
        response=self.client.post(reverse('estimates:estimate_import'),{'object_name':'Batch'})
        self.assertEqual(response.status_code,200)
        self.user.role=Role.objects.get(code='OBSERVER');self.user.save()
        response=self.client.get(reverse('estimates:estimate_list'))
        self.assertContains(response,'Ваша роль разрешает просмотр')
        self.assertNotContains(response,'Загрузить сметы (несколько файлов)')
