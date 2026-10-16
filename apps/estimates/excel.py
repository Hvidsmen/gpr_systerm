"""Read cached source values only. Never evaluate workbook formulas."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile, BadZipFile
import hashlib
import re
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

LABELS = {'section': 'Работа', 'item': 'Подработа', 'resource': 'Ресурс', 'detail': 'Расчет / пояснение', 'total': 'Итог', 'appendix': 'Ведомость / итоговая расшифровка'}
COLUMNS = ['№ / отметки', 'Обоснование', 'Наименование', 'Ед. изм.', 'Количество на единицу', 'Коэффициент количества', 'Количество всего', 'Цена базисная', 'Коэффициент цены', 'Стоимость базисная, ₽', 'Индекс', 'Стоимость текущая, ₽']

def number(value):
    try:
        return Decimal(str(value).replace(' ', '').replace('\xa0', '').replace(',', '.'))
    except (InvalidOperation, ValueError):
        return None


def parse(data, filename=''):
    if len(data) > 12 * 1024 * 1024:
        raise ValueError('Размер файла превышает 12 МБ.')
    try:
        with ZipFile(BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > 80 * 1024 * 1024:
                raise ValueError('Слишком большой распакованный Excel.')
        book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError('Не удалось открыть XLSX.') from exc
    try:
        candidates = []
        for sheet in book:
            if sheet.max_row > 15000 or sheet.max_column > 300:
                raise ValueError('Превышен допустимый размер листа.')
            rows = list(sheet.iter_rows(max_col=14, values_only=True))
            header = next((i for i, r in enumerate(rows[:60]) if r[1] == 'Обоснование' and 'Наименование' in str(r[2])), None)
            if header is not None:
                candidates.append((sheet.title, rows, header))
        if len(candidates) != 1:
            raise ValueError('Ожидается один лист ЛСР с колонками «Обоснование» и «Наименование».')
        sheet, rows, header = candidates[0]
        title = next((str(r[0]) for r in rows[:header] if 'ЛОКАЛЬНЫЙ СМЕТНЫЙ' in str(r[0])), filename)
        subject = next((str(r[0]) for r in rows[4:header] if r[0] and not str(r[0]).startswith((' ', 'Составлен', 'Основание'))), '')
        result = dict(filename=filename, digest=hashlib.sha256(data).hexdigest(), title=title + ' · ' + subject, sheet=sheet, rows=[], base_total=None, current_total=None, warnings=[])
        section = ''; parent = None; appendix = False
        for index, r in enumerate(rows[header+4:], start=header+5):
            vals = [r[i] for i in (0,1,2,5,6,7,8,9,10,11,12,13)]
            vals = [v if isinstance(v, (str,int,float)) else str(v) if v is not None else None for v in vals]
            if not any(v is not None and str(v).strip() for v in vals):
                continue
            a, code, name, unit = [str(v or '').strip() for v in vals[:4]]
            if a.startswith('Раздел '):
                section = a; parent = None; appendix = False; kind = 'section'
            elif re.match(r'^\d+(?:\s|$)', a) and code and name and unit:
                kind = 'item'; parent = index; appendix = False
            elif name.lower().startswith(('итого по разделу', 'итоги по смете', 'всего по смете')):
                kind = 'total'; appendix = True; parent = None
            elif appendix:
                kind = 'appendix'
            elif parent and code and name and unit and number(vals[6]) is not None and unit != '%' and not name.startswith(('НР ', 'СП ')):
                kind = 'resource'
            else:
                kind = 'detail'
            if name.lower().startswith('всего по смете'):
                result['base_total'] = str(number(vals[9])) if number(vals[9]) is not None else None
                result['current_total'] = str(number(vals[11])) if number(vals[11]) is not None else None
            result['rows'].append(dict(source_row=index, section=section, kind=kind, parent_row=parent if kind not in {'item','section','total','appendix'} else None, values=vals))
        if result['current_total'] is None:
            result['warnings'].append('Контрольный итог «ВСЕГО по смете» не найден; расчет отчета выполняется по подработам.')
        if not any(r['kind']=='item' for r in result['rows']):
            raise ValueError('В файле не найдены позиции сметы.')
        return result
    finally:
        book.close()


def safe(value):
    return "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value


def export(documents, mode="full"):
    from .report import build_report
    book = Workbook(); summary = book.active; summary.title = 'Свод смет'
    summary.append(['Смета / файл', 'Наименование', 'Базисная стоимость, ₽', 'Текущая стоимость, ₽', 'Позиций', 'Ресурсов', 'Примечания'])
    positions = book.create_sheet('Позиции'); resources = book.create_sheet('Ресурсы'); detail = book.create_sheet('Расчеты и итоги')
    for s in (positions, resources, detail):
        s.append(['Файл', 'Лист', 'Строка', 'Раздел', 'Тип', 'Родительская строка'] + COLUMNS)
    catalogs = {'work': {}, 'resource': {}, 'unit': {}}
    hierarchy = book.create_sheet('Работы и подработы')
    hierarchy.append(['Объект', 'Смета', 'Файл', 'Работа / подработа', 'Обоснование', 'Ед.', 'Объем', 'Базисная стоимость, ₽', 'Текущая стоимость, ₽', 'Примечания'])
    for n, doc in enumerate(documents, 1):
        report = build_report(doc['rows'])
        doc = dict(doc, base_total=report['base'], current_total=report['current'], warnings=doc['warnings'] + report['diagnostics'])
        for work in report['works']:
            hierarchy.append([safe(doc.get('object_name', '')), safe(doc['title']), safe(doc['filename']), safe(work['name']), None, None, None, work['base'], work['current']])
            for sub in work['subworks']:
                hierarchy.append([safe(doc.get('object_name', '')), safe(doc['title']), safe(doc['filename']), safe(sub['name']), safe(sub['code']), safe(sub['unit']), sub['quantity'], sub['base'], sub['current'], '; '.join(sub['warnings'])])
                hierarchy.row_dimensions[hierarchy.max_row].outlineLevel = 1
        hierarchy.sheet_properties.outlinePr.summaryBelow = False
        summary.append([safe(doc['filename']), safe(doc['title']), number(doc['base_total']), number(doc['current_total']), sum(r['kind']=='item' for r in doc['rows']), sum(r['kind']=='resource' for r in doc['rows']), '; '.join(doc['warnings'])])
        sheet = book.create_sheet('ЛСР %02d' % n)
        sheet.append(['Строка', 'Раздел', 'Тип', 'Родительская строка'] + COLUMNS)
        for r in doc['rows']:
            vals = [safe(v) for v in r['values']]
            sheet.append([r['source_row'], safe(r['section']), LABELS[r['kind']], r['parent_row']] + vals)
            sheet.row_dimensions[sheet.max_row].outlineLevel = 2 if r['kind'] in {'resource','detail'} else 1 if r['kind']=='item' else 0
            target = positions if r['kind']=='item' else resources if r['kind']=='resource' else detail
            target.append([safe(doc['filename']), safe(doc['sheet']), r['source_row'], safe(r['section']), LABELS[r['kind']], r['parent_row']] + vals)
            if r['kind'] in {'item', 'resource'}:
                k = 'work' if r['kind']=='item' else 'resource'
                key = tuple(str(v or '') for v in r['values'][1:4]);catalogs[k][key] = key
                u = str(r['values'][3] or '')
                if u: catalogs['unit'][(u,)] = (u,)
        sheet.sheet_properties.outlinePr.summaryBelow = False
    for kind, title in [('work','Справочник позиций'),('resource','Справочник ресурсов'),('unit','Единицы измерения')]:
        s=book.create_sheet(title);s.append(['Обоснование','Наименование','Ед. изм.'] if kind!='unit' else ['Ед. изм.'])
        for key in sorted(catalogs[kind]): s.append([safe(v) for v in key])
    info=book.create_sheet('Описание');info.append(['Правила']);info.append(['Справочники смет независимы от ГПР. Итоги вычислены по подработам; строки Итого и Всего не участвуют в расчете.']);info.append(['Ресурсная расшифровка входит в стоимость позиции: повторно не суммировать.']);info.append(['Материалы заказчика, справочные итоги и начисления сохранены в расчетах.']);info.append(['Исходные коэффициенты и значения сохранены; формулы не вычисляются, используются сохраненные значения Excel.'])
    if summary.max_row > 1:
        totals = []
        for col in [3,4]:
            vals = [summary.cell(row,col).value for row in range(2,summary.max_row+1)]
            totals.append(sum(vals, Decimal('0')) if all(v is not None for v in vals) else None)
        summary.append(['Итого', None] + totals)
    if mode == 'works':
        for sheet in list(book):
            if sheet.title not in {'Свод смет', 'Работы и подработы'}:
                book.remove(sheet)
        summary.delete_cols(6)  # Resource counts are excluded from the concise report.
    for s in book:
        s.freeze_panes='D2';s.auto_filter.ref=s.dimensions
        s.sheet_view.showGridLines=False
        for c in s[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='1765A0');c.alignment=Alignment(wrap_text=True)
        s.row_dimensions[1].height=34
        for col in range(1,s.max_column+1):
            name=str(s.cell(1,col).value or '');s.column_dimensions[get_column_letter(col)].width=65 if 'Наименование' in name else 38 if name in {'Файл','Смета / файл','Раздел'} else 19
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(c.value,(int,float,Decimal)):
                    c.number_format = '#,##0' if '₽' in str(s.cell(1,c.column).value) else '#,##0.########'
                if c.row%2==0:c.fill=PatternFill('solid',fgColor='F0F5FB')
        s.sheet_properties.pageSetUpPr.fitToPage=True
        s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0;s.print_title_rows='1:1'
    output=BytesIO();book.save(output);return output.getvalue()
