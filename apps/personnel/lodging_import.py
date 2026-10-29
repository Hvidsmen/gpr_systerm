"""Administrator imports a commandant's chessboard; unmatched names stay queued."""
from datetime import date,datetime
from hashlib import sha256
from io import BytesIO
import re
from zipfile import ZipFile
from collections import defaultdict
from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import transaction
from django.shortcuts import get_object_or_404,redirect,render
from django.views.decorators.http import require_http_methods
from django.utils import timezone
from core.permissions import require_roles
from core.fact_history import history_actor
from .models import Camp,LodgingBuilding,LodgingBed,LodgingStay,LodgingEvent
from .lodging_import_models import LodgingImportBatch,LodgingImportRow
from . import lodging

SHEET='Шахматка КТП 1'

def text(value):return ' '.join(str(value or '').split())
def key(value):return text(value).casefold().replace('ё','е')
def numbered_key(value):
    value=text(value)
    return value if re.match(r'^ключ(?:и|ик)?\s*(?:№\s*)?\S*\d',value,re.I) else ''
def day(value):
    if isinstance(value,datetime):return value.date()
    if isinstance(value,date):return value
    if not value:return None
    for fmt in ('%d.%m.%Y','%Y-%m-%d'):
        try:return datetime.strptime(text(value),fmt).date()
        except ValueError:pass
    raise ValidationError('Дата в шахматке не распознана.')


def parse(raw):
    try:
        with ZipFile(BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist())>50*1024*1024:raise ValidationError('Файл слишком большой после распаковки.')
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(raw),read_only=True,data_only=True,keep_links=False)
    except ValidationError:raise
    except Exception:raise ValidationError('Не удалось прочитать файл .xlsx.')
    try:
        if SHEET not in book.sheetnames:raise ValidationError(f'В книге отсутствует лист «{SHEET}».')
        sheet=book[SHEET];rows=[];building='';description='';snapshot=None;seen=set();expected=None;header=False
        for number,cells in enumerate(sheet.iter_rows(values_only=True),1):
            if number>5000:raise ValidationError('В листе более 5000 строк.')
            values=list(cells)+[None]*12
            if not header:
                for value in values:
                    if isinstance(value,(datetime,date)):snapshot=day(value)
                if key(values[0])=='итого' and isinstance(values[2],(int,float)):expected=int(values[2])
                if key(values[1]) in {'№ места','номер места'} and key(values[4]).replace('.','')=='фио':header=True
                continue
            if values[0]:building=text(values[0]);description=text(values[2])
            has_bed=isinstance(values[1],(int,float)) and values[1]>0
            if not has_bed and not text(values[4]):continue
            bed=str(int(values[1])) if has_bed else '';identity=(key(building),bed or f'row:{number}')
            if not building or identity in seen:raise ValidationError(f'Строка {number}: повтор или отсутствие номера вагон-дома/места.')
            seen.add(identity)
            arrived=day(values[8]);departed=day(values[9]);name=text(values[4])
            warning='' if has_bed else 'Не указан номер места. Уточните исходный файл.'
            if name and not arrived:warning='Не указана дата прибытия.'
            if arrived and snapshot and arrived>snapshot:warning='Прибытие позже даты среза.'
            if departed and (not arrived or departed<arrived):warning='Дата выбытия раньше прибытия.'
            if departed and snapshot and departed>snapshot:warning='Выбытие позже даты среза.'
            inventory=re.search(r'(?:инв\.?\s*№?\s*)([А-ЯA-Z]{1,4}[- ]?\d+)',description,re.I)
            note=' · '.join(v for v in [text(values[3]),text(values[11])] if v)
            key_source=text(values[7]);named_key=numbered_key(key_source)
            state='UNAVAILABLE' if 'опечатан' in key(building) else 'RESERVED' if not name and 'заказчик' in key(description+' · '+note) else 'AVAILABLE'
            rows.append({'source_row':number,'building':building,'description':description,'bed':bed,'inventory':inventory.group(1) if inventory else '',
                'kind':text(values[3]),'note':note,'bed_state':state,'name':name,'organization':text(values[5]),'position':text(values[6]),
                'keys':key(values[7]) in {'да','выданы','+','1'} or bool(named_key),'key_number':named_key,'keys_source':key_source,'arrived':str(arrived) if arrived else '',
                'departed':str(departed) if departed else '','department':text(values[10]),'pass_number':text(values[11]),'warning':warning})
        if not header or not snapshot or not rows:raise ValidationError('Не найдены заголовок шахматки, дата среза или места.')
        if snapshot>timezone.localdate():raise ValidationError('Дата среза находится в будущем.')
        named=sum(bool(row['name']) for row in rows)
        if expected is not None and named!=expected:raise ValidationError(f'Количество проживающих ({named}) не совпало с итогом листа ({expected}). Проверьте файл.')
        return {'snapshot':str(snapshot),'rows':rows,'people':named,'buildings':len({r['building'] for r in rows}),'beds':sum(bool(r['bed']) for r in rows)}
    finally:book.close()


class UploadForm(forms.Form):
    camp=forms.ModelChoiceField(label='ВЖГ',queryset=Camp.objects.none())
    file=forms.FileField(label=f'Книга Excel с листом «{SHEET}»',widget=forms.FileInput(attrs={'accept':'.xlsx'}))
    def __init__(self,*args,company,**kwargs):
        super().__init__(*args,**kwargs);self.fields['camp'].queryset=Camp.objects.filter(company=company)
    def clean_file(self):
        file=self.cleaned_data['file']
        if len(file.name)>250:raise forms.ValidationError('Имя файла должно быть не длиннее 250 символов.')
        if file.size>5*1024*1024 or not file.name.lower().endswith('.xlsx'):raise forms.ValidationError('Нужен .xlsx до 5 МБ.')
        return file


def bed_note(data,description=''):
    # Earlier import batches put the wagon description in every bed note.
    note=data.get('note','');description=data.get('description') or description
    if description and note==description:return ''
    if description and note.startswith(description+' · '):return note[len(description)+3:]
    return note


def refresh_metadata(user,batch,parsed):
    """Re-upload fills missing metadata without re-entering or replacing stays."""
    from apps.production.history_models import FactRevision
    source={data['source_row']:data for data in parsed['rows']}
    with history_actor(user,'Дополнение сведений из шахматки коменданта'):
        for row in batch.rows.select_for_update().select_related('stay__bed__building'):
            data=source.get(row.source_row)
            if not data:continue
            row.data.update({name:data[name] for name in ('description','key_number','keys')});row.save(update_fields=['data','updated_at'])
            stay=row.stay
            if not stay:continue
            building=stay.bed.building
            if not building.description and data['description']:building.description=data['description'];building.save(update_fields=['description','updated_at'])
            if not stay.key_number and data['key_number']:
                stay.key_number=data['key_number']
                # Preserve a commandant's subsequent correction of the issued flag.
                if not FactRevision.objects.filter(company=user.company,entity_type='personnel.lodgingstay',entity_id=str(stay.pk),operation='UPDATE').exists():stay.keys_issued=data['keys']
                stay.save(update_fields=['key_number','keys_issued','updated_at'])


@transaction.atomic
def apply(user,batch):
    require_roles(user,{'ADMIN'});lodging.mutex(user)
    batch=LodgingImportBatch.objects.select_for_update().get(pk=batch.pk,company=user.company)
    from apps.resources.models import Employee
    employees=defaultdict(list)
    for employee in Employee.objects.filter(company=user.company):employees[key(employee.display_name)].append(employee)
    buildings=defaultdict(list)
    for building in LodgingBuilding.objects.select_for_update().filter(company=user.company,camp=batch.camp):buildings[key(building.name)].append(building)
    for row in batch.rows.select_for_update().exclude(state='APPLIED'):
        data=row.data;original_stay_id=row.stay_id
        with transaction.atomic(),history_actor(user,'Загрузка шахматки коменданта'):
            try:
                matches=buildings.get(key(data['building']),[])
                if len(matches)>1:raise ValidationError('Несколько вагон-домов с одинаковым названием.')
                building=matches[0] if matches else LodgingBuilding.objects.create(company=user.company,camp=batch.camp,name=data['building'],description=data.get('description',''),inventory_number=data['inventory'],kind=data['kind'])
                if matches and not building.description and data.get('description'):
                    building.description=data['description'];building.save(update_fields=['description','updated_at'])
                buildings[key(building.name)]=[building]
                if not data['bed']:raise ValidationError('Не указан номер места. Уточните исходный файл.')
                bed,_=LodgingBed.objects.get_or_create(company=user.company,building=building,room='',number=data['bed'],defaults={'state':data['bed_state'],'note':bed_note(data,building.description)[:500]})
                if data['warning']:raise ValidationError(data['warning'])
                if data['name']:
                    matches=employees.get(key(data['name']),[])
                    if len(matches)!=1:raise ValidationError('Сотрудник не найден.' if not matches else 'Несколько сотрудников с одинаковым ФИО. Требуется ручное сопоставление.')
                    employee=matches[0]
                    if not employee.is_active:raise ValidationError('Сотрудник неактивен. Проверьте историческое заселение.')
                    arrived=date.fromisoformat(data['arrived']);departed=date.fromisoformat(data['departed']) if data['departed'] else None
                    from core.fact_edit_window import require_fact_day
                    require_fact_day(user,arrived)
                    if departed:require_fact_day(user,departed)
                    with transaction.atomic():
                        identical=LodgingStay.objects.filter(company=user.company,employee=employee,bed=bed,camp=batch.camp,arrived_on=arrived,departed_on=departed,cancelled_at=None).first()
                        if identical:row.stay=identical
                        else:
                            if not building.is_active or bed.state!='AVAILABLE':raise ValidationError('Место недоступно для заселения. Проверьте его состояние.')
                            controls=lodging.controls(user,[('bed',bed.pk),('employee',employee.pk)])
                            for control in controls:lodging.locked_elsewhere(control,user,None)
                            stay=LodgingStay(company=user.company,employee=employee,camp=batch.camp,bed=bed,arrived_on=arrived,departed_on=departed,keys_issued=data['keys'],key_number=data.get('key_number',''),pass_number=data['pass_number'],organization_snapshot=data['organization'],position_snapshot=data['position'],department_snapshot=data['department'],recorded_by=user,comment=f"{batch.filename} · {SHEET} · строка {row.source_row} · срез {batch.snapshot_date}. Ключи: {data['keys_source']}")
                            stay.full_clean();stay.save();row.stay=stay
                            LodgingEvent.objects.create(company=user.company,employee=employee,camp=batch.camp,date=arrived,kind='CHECKIN',to_bed=bed,stay=stay,comment=stay.comment,recorded_by=user)
                            if departed:LodgingEvent.objects.create(company=user.company,employee=employee,camp=batch.camp,date=departed,kind='CHECKOUT',from_bed=bed,stay=stay,comment=stay.comment,recorded_by=user)
                            lodging.bump(controls)
                row.state='APPLIED';row.warning='';row.save()
            except ValidationError as exc:
                row.stay_id=original_stay_id
                row.warning='; '.join(exc.messages);row.save(update_fields=['warning','updated_at'])
    return batch


@login_required
@require_http_methods(['GET','POST'])
def upload(request):
    require_roles(request.user,{'ADMIN'})
    form=UploadForm(request.POST or None,request.FILES or None,company=request.user.company,initial={'camp':request.GET.get('camp')})
    if request.method=='POST' and form.is_valid():
        try:
            file=form.cleaned_data['file'];raw=file.read();parsed=parse(raw);camp=form.cleaned_data['camp']
            camp_number=re.search(r'ктп[\s№-]*(\d+)',key(camp.name))
            if camp_number and int(camp_number.group(1))!=1:raise ValidationError('Лист «Шахматка КТП 1» относится к КТП-1. Выберите соответствующий ВЖГ.')
            with transaction.atomic():
                lodging.mutex(request.user)
                batch,created=LodgingImportBatch.objects.get_or_create(company=request.user.company,camp=camp,digest=sha256(raw).hexdigest(),defaults={'filename':file.name,'snapshot_date':date.fromisoformat(parsed['snapshot']),'created_by':request.user})
                if created:
                    for data in parsed['rows']:LodgingImportRow.objects.create(company=request.user.company,batch=batch,source_row=data['source_row'],data=data)
                else:refresh_metadata(request.user,batch,parsed)
                apply(request.user,batch)
            messages.success(request,'Шахматка обработана. Непроверенные строки сохранены отдельно; история перевахты не изменялась.')
            return redirect('personnel:lodging_import_review',pk=batch.pk)
        except ValidationError as exc:form.add_error(None,'; '.join(exc.messages))
    batches=LodgingImportBatch.objects.filter(company=request.user.company).select_related('camp','created_by').order_by('-pk')[:20]
    return render(request,'personnel/lodging_import.html',{'form':form,'batches':batches})


@login_required
@require_http_methods(['GET','POST'])
def review(request,pk):
    require_roles(request.user,{'ADMIN'});batch=get_object_or_404(LodgingImportBatch,company=request.user.company,pk=pk)
    if request.method=='POST':apply(request.user,batch);return redirect('personnel:lodging_import_review',pk=batch.pk)
    pending=batch.rows.filter(state='PENDING');rows=pending if request.GET.get('pending')!='0' else batch.rows.all()
    return render(request,'personnel/lodging_import_review.html',{'batch':batch,'pending':pending.count(),'applied':batch.rows.filter(state='APPLIED').count(),'people':batch.rows.filter(state='APPLIED',stay__isnull=False).count(),'beds':batch.rows.exclude(data__bed='').count(),'page_obj':Paginator(rows,50).get_page(request.GET.get('page'))})
