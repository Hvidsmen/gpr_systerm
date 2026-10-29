from datetime import timedelta
from io import BytesIO
from uuid import uuid4
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from apps.resources.models import Brigade,BrigadeGroup
from apps.resources.position_catalog import position_for
from . import test_lodging as fixtures,lodging,lodging_import
from .lodging_forms import LodgingMovementForm
from .models import LodgingBed,LodgingControl,LodgingStay


class LodgingMetadataTests(TestCase):
    setUpTestData=classmethod(fixtures.LodgingTests.setUpTestData.__func__)
    edit=fixtures.LodgingTests.edit
    def setUp(self):
        self.client.force_login(self.commandant)
        self.scope={'camp':self.a.pk,'date':str(self.day)}

    def test_only_free_available_places_on_selected_date_are_offered(self):
        self.edit()
        self.bed2.state='RESERVED';self.bed2.save()
        form=LodgingMovementForm(company=self.company,camp=self.a,day=self.day)
        self.assertEqual(list(form.fields['bed'].queryset),[self.bed3])
        # Planned checkout does not free a place; an actual departure does.
        stay=LodgingStay.objects.get();stay.planned_departure=self.day;stay.departed_on=self.day+timedelta(days=1);stay.save()
        endpoint=reverse('personnel:lodging_available_beds')
        self.assertEqual([r['id'] for r in self.client.get(endpoint,self.scope).json()['beds']],[self.bed3.pk])
        tomorrow={**self.scope,'date':str(self.day+timedelta(days=1))}
        self.assertIn(self.bed.pk,[r['id'] for r in self.client.get(endpoint,tomorrow).json()['beds']])
        self.bed3.building.is_active=False;self.bed3.building.save()
        self.assertEqual(self.client.get(endpoint,self.scope).json()['beds'],[])

    def test_temporary_reservations_excluded_but_own_post_can_save(self):
        session=str(uuid4());signed=lodging.claim(self.commandant,self.a,self.day,self.second_employee,self.bed,session)
        scope={**self.scope,'session':session}
        endpoint=reverse('personnel:lodging_available_beds')
        self.assertIn(self.bed.pk,[r['id'] for r in self.client.get(endpoint,scope).json()['beds']])
        self.client.force_login(self.colleague)
        self.assertNotIn(self.bed.pk,[r['id'] for r in self.client.get(endpoint,scope).json()['beds']])
        self.client.force_login(self.commandant)
        payload={**self.scope,'kind':'CHECKIN','employee':self.second_employee.pk,'bed':self.bed.pk,'session':session,'signed':signed,'key_number':'КЛЮЧ 2/19','keys_issued':'on','pass_number':'P-19'}
        self.assertEqual(self.client.post(reverse('personnel:lodging'),payload).status_code,302)
        stay=LodgingStay.objects.get();self.assertEqual(stay.key_number,'КЛЮЧ 2/19');self.assertEqual(stay.pass_number,'P-19')

    def test_gender_group_keys_and_pass_are_visible_and_exported_without_organization(self):
        group=BrigadeGroup.objects.create(company=self.company,name='ИТР Состав')
        brigade=Brigade.objects.create(company=self.company,name='Инженер',group=group)
        self.second_employee.position=position_for(brigade);self.second_employee.gender='F';self.second_employee.save()
        self.edit(keys_issued=True,key_number='КЛЮЧ 2/19',pass_number='P-19')
        endpoint=reverse('personnel:lodging');page=self.client.get(endpoint,self.scope)
        for text in ['Группа / должность','Женский','ИТР Состав','КЛЮЧ 2/19','Пропуск: P-19','Плановая дата выселения']:self.assertContains(page,text)
        self.assertNotContains(page,'Организация / должность')
        export=self.client.get(endpoint,{**self.scope,'export':'xlsx'})
        book=load_workbook(BytesIO(export.content));values=list(book.active.values);header=values[0];row=values[1]
        self.assertNotIn('Организация',header)
        for field,value in [('Пол','Женский'),('Группа должностей','ИТР Состав'),('Номер ключа','КЛЮЧ 2/19'),('Пропуск','P-19')]:self.assertEqual(row[header.index(field)],value)

    def test_gender_can_be_clarified_without_modifying_zup_fields(self):
        self.second_employee.zup_ref='a'*32;self.second_employee.save()
        self.edit();endpoint=reverse('personnel:lodging_employee_gender',args=[self.second_employee.pk])
        self.assertEqual(self.client.post(endpoint,{**self.scope,'gender':'M'}).status_code,302)
        self.second_employee.refresh_from_db();self.assertEqual(self.second_employee.gender,'M');self.assertEqual(self.second_employee.zup_ref,'a'*32)
        self.assertEqual(self.client.post(endpoint,{**self.scope,'gender':'guess'}).status_code,400)
        self.client.force_login(self.hr);self.assertEqual(self.client.post(endpoint,{**self.scope,'gender':'F'}).status_code,403)
        self.client.force_login(self.commandant);self.assertEqual(self.client.post(endpoint,{**self.scope,'camp':self.b.pk,'gender':'F'}).status_code,404)

    def test_named_keys_are_imported_separately_from_pass(self):
        from .test_lodging_import import workbook
        raw=workbook(self.day,[self.employee.display_name]);book=load_workbook(BytesIO(raw));sheet=book.active
        sheet.cell(4,8,'КЛЮЧ 2/19');sheet.cell(4,12,'P-19');output=BytesIO();book.save(output)
        row=lodging_import.parse(output.getvalue())['rows'][0]
        self.assertTrue(row['keys']);self.assertEqual(row['key_number'],'КЛЮЧ 2/19');self.assertEqual(row['pass_number'],'P-19')
        sheet.cell(4,8,'Нет');output=BytesIO();book.save(output);row=lodging_import.parse(output.getvalue())['rows'][0]
        self.assertFalse(row['keys']);self.assertEqual(row['key_number'],'')
