from datetime import timedelta
from django.test import TestCase
from django.urls import reverse
from apps.resources.models import Employee
from apps.projects.models import ConstructionObject
from .test_housing import HousingTests
from .models import StaffingSlot,Person,CampStay,CampObjectLink
from .rotation import RotationFilter,rotation_rows


class RotationTests(TestCase):
    @classmethod
    def setUpTestData(cls):HousingTests.setUpTestData.__func__(cls)
    def setUp(self):self.client.force_login(self.hr)
    def test_rest_and_vacancies_visible_even_without_smg_and_do_not_grant_production_access(self):
        StaffingSlot.objects.create(company=self.company,group=self.group,person=self.second_person,housing_camp=self.a,status='МО')
        vacancy=StaffingSlot.objects.create(company=self.company,group=self.group,housing_camp=self.a)
        rows,days,stats=rotation_rows(self.company,[self.a],self.day)
        self.assertEqual(stats['people'],2);self.assertEqual(stats['present'],1);self.assertEqual(stats['rest'],1);self.assertEqual(stats['vacancies'],1)
        self.assertEqual(len(rows),3)
        self.assertNotIn(self.second_employee.pk,set(__import__('apps.personnel.housing',fromlist=['available_employee_ids']).available_employee_ids(self.obj,self.day).values_list('person__employee_id',flat=True)))
        response=self.client.get(reverse('personnel:housing'),{'camp':self.a.pk,'date':str(self.day)})
        self.assertContains(response,'Таблица перевахты');self.assertContains(response,'Межвахтовый отдых');self.assertContains(response,'Открытая вакансия')
    def test_macro_object_selects_all_linked_camps_of_its_children(self):
        parent=ConstructionObject.objects.create(company=self.company,project=self.project,name='Общий КТП')
        self.obj.parent=parent;self.obj.save();self.obj3.parent=parent;self.obj3.save()
        form=RotationFilter({'construction_object':parent.pk,'date':str(self.day)},company=self.company)
        self.assertTrue(form.is_valid());self.assertEqual(set(form.selected_camps(self.company).values_list('pk',flat=True)),{self.a.pk,self.b.pk})
        form=RotationFilter({'construction_object':parent.pk,'date':str(self.day),'camps':[self.b.pk]},company=self.company)
        self.assertTrue(form.is_valid());self.assertEqual(list(form.selected_camps(self.company)),[self.b])
    def test_closed_stay_is_rest_and_transferred_employee_is_not_duplicated_in_previous_camp(self):
        self.stay.departed_on=self.day;self.stay.save()
        rows,_,stats=rotation_rows(self.company,[self.a],self.day)
        self.assertEqual(rows[0]['state'],'MO');self.assertEqual(stats['present'],0)
        CampStay.objects.create(company=self.company,person=self.person,camp=self.b,arrived_on=self.day)
        rows,_,stats=rotation_rows(self.company,[self.a],self.day);self.assertEqual(len(rows),0)
        rows,_,stats=rotation_rows(self.company,[self.a,self.b],self.day);self.assertEqual(len(rows),1)
    def test_multiple_camps_and_rotation_export_are_scoped(self):
        StaffingSlot.objects.create(company=self.company,group=self.group,person=self.second_person,housing_camp=self.b,status='МО')
        data={'camps':[self.a.pk,self.b.pk],'date':str(self.day)}
        page=self.client.get(reverse('personnel:housing'),data);self.assertEqual(page.status_code,200);self.assertEqual(page.context['rotation_stats']['people'],2)
        response=self.client.get(reverse('personnel:housing_export'),data);self.assertEqual(response.status_code,200)
        self.assertEqual(self.client.get(reverse('personnel:housing'),{'camp':self.foreign_camp.pk}).status_code,404)

    def test_work_and_travel_counts_use_selected_day_not_hours_or_vacancies(self):
        StaffingSlot.objects.create(company=self.company,group=self.group,person=self.second_person,housing_camp=self.a,status='МО',calendar={str(self.day):'ДП'})
        StaffingSlot.objects.create(company=self.company,group=self.group,housing_camp=self.a,pattern=self.pattern)
        rows,_,stats=rotation_rows(self.company,[self.a],self.day)
        self.assertEqual({k:stats[k] for k in ['people','working','travelling','on_shift','rest','vacancies']},
            {'people':2,'working':1,'travelling':1,'on_shift':1,'rest':1,'vacancies':1})
        _,_,tomorrow=rotation_rows(self.company,[self.a],self.day+timedelta(days=1))
        self.assertEqual(tomorrow['working'],0);self.assertEqual(tomorrow['travelling'],1)
        self.assertEqual(tomorrow['no_schedule'],1)

    def test_counts_follow_search_state_and_position_filters(self):
        from apps.resources.models import Brigade
        from apps.resources.position_catalog import position_for
        other_position=position_for(Brigade.objects.create(company=self.company,name='Инженер'))
        self.second_employee.position=other_position;self.second_employee.save()
        StaffingSlot.objects.create(company=self.company,group=self.group,person=self.second_person,housing_camp=self.a,status='МО',calendar={str(self.day):'ДП'})
        for filters,expected in [({'query':'001'},(1,1,0)),({'state':'TRAVEL'},(1,0,1)),({'state':'WORKING'},(1,1,0)),({'position':other_position},(1,0,1)),({'query':'нет совпадений'},(0,0,0))]:
            rows,_,stats=rotation_rows(self.company,[self.a],self.day,**filters)
            self.assertEqual((stats['people'],stats['working'],stats['travelling']),expected)
            self.assertEqual(stats['people'],len(rows))

    def test_missing_calendar_is_not_inferred_from_residence_and_people_are_not_duplicated(self):
        self.slot.pattern=None;self.slot.save()
        StaffingSlot.objects.create(company=self.company,group=self.group,person=self.person,housing_camp=self.b,pattern=self.pattern)
        rows,_,stats=rotation_rows(self.company,[self.a,self.b],self.day)
        self.assertEqual(len(rows),1);self.assertEqual(stats['on_shift'],1);self.assertEqual(stats['present'],1)
        self.assertEqual(stats['working'],0);self.assertEqual(stats['no_schedule'],1)

    def test_header_fact_and_total_cards_use_same_filtered_rows_without_plan(self):
        StaffingSlot.objects.create(company=self.company,group=self.group,person=self.second_person,housing_camp=self.a,status='МО',calendar={str(self.day):'ДП'})
        StaffingSlot.objects.create(company=self.company,group=self.group,housing_camp=self.a,pattern=self.pattern)
        page=self.client.get(reverse('personnel:housing'),{'camp':self.a.pk,'date':str(self.day)})
        self.assertContains(page,'План и факт');self.assertContains(page,'В работе по графику');self.assertContains(page,'В пути по графику')
        for key in ['working','travelling','people','on_shift','rest','vacancies','no_schedule']:
            self.assertEqual(sum(g[key] for g in page.context['position_groups']),page.context['rotation_stats'][key])
        self.assertContains(page,'data-position-fact-value="working">1</dd>',html=False)
        filtered=self.client.get(reverse('personnel:housing'),{'camp':self.a.pk,'date':str(self.day),'state':'TRAVEL'})
        self.assertEqual(filtered.context['rotation_stats']['people'],1);self.assertEqual(filtered.context['position_groups'][0]['travelling'],1)

    def test_calendar_work_codes_count_people_and_ignore_rest_zero_and_blank(self):
        for code,working,travel,missing in [('МВ',1,0,0),('10',1,0,0),(8.5,1,0,0),('0',0,0,0),(0,0,0,0),('В',0,0,0),('МО',0,0,0),(' дп ',0,1,0),('',0,0,1)]:
            self.slot.calendar={str(self.day):code};self.slot.save()
            _,_,stats=rotation_rows(self.company,[self.a],self.day)
            self.assertEqual((stats['working'],stats['travelling'],stats['no_schedule']),(working,travel,missing),code)

    def test_trip_filters_distinguish_multiday_incoming_and_outgoing_without_guessing(self):
        self.slot.pattern=None
        for following,expected in [('10','ARRIVING'),('МО','DEPARTING'),('',None)]:
            self.slot.calendar={str(self.day):'ДП',str(self.day+timedelta(days=1)):'ДП',str(self.day+timedelta(days=2)):following};self.slot.save()
            for state in ['ARRIVING','DEPARTING']:
                rows,_,_=rotation_rows(self.company,[self.a],self.day,state=state)
                self.assertEqual(len(rows),int(state==expected))

    def test_actual_arrivals_departures_and_transfers_are_scoped_to_selected_day(self):
        CampStay.objects.create(company=self.company,person=self.second_person,camp=self.a,arrived_on=self.day)
        rows,_,_=rotation_rows(self.company,[self.a],self.day,state='ARRIVING')
        self.assertEqual([r['person'].pk for r in rows],[self.second_person.pk])
        self.stay.departed_on=self.day;self.stay.save()
        CampStay.objects.create(company=self.company,person=self.person,camp=self.b,arrived_on=self.day)
        rows,_,_=rotation_rows(self.company,[self.a],self.day,state='DEPARTING')
        self.assertEqual([r['person'].pk for r in rows],[self.person.pk]);self.assertEqual(rows[0]['departure_camp'],self.a)
        self.assertEqual(rotation_rows(self.company,[self.a],self.day+timedelta(days=1),state='DEPARTING')[0],[])

    def test_departure_critical_filters_have_explicit_fourteen_day_boundary(self):
        self.stay.arrived_on=self.day-timedelta(days=5);self.stay.save()
        self.slot.pattern=None
        for offset,expected in [(-1,'DEPARTURE_OVERDUE'),(0,'DEPARTURE_PENDING'),(14,'DEPARTURE_PENDING'),(15,None)]:
            departure=self.day+timedelta(days=offset)
            self.slot.calendar={str(departure):'ДП',str(departure+timedelta(days=1)):'МО'};self.slot.save()
            for state in ['DEPARTURE_OVERDUE','DEPARTURE_PENDING']:
                rows,_,_=rotation_rows(self.company,[self.a],self.day,state=state)
                self.assertEqual(len(rows),int(expected==state),(offset,state))
        self.stay.departed_on=self.day;self.stay.save()
        self.assertEqual(rotation_rows(self.company,[self.a],self.day,state='DEPARTURE_PENDING')[0],[])

    def test_overdue_return_keeps_missed_date_until_actual_arrival(self):
        due=self.day-timedelta(days=2)
        StaffingSlot.objects.create(company=self.company,group=self.group,person=self.second_person,housing_camp=self.a,status='МО',departure_date=due-timedelta(days=2),calendar={str(due):'10',str(self.day):'10',str(self.day+timedelta(days=1)):'10'})
        rows,_,stats=rotation_rows(self.company,[self.a],self.day,state='ARRIVAL_OVERDUE')
        self.assertEqual(stats['people'],1);self.assertEqual(rows[0]['planned_return'],due)
        response=self.client.get(reverse('personnel:housing'),{'camp':self.a.pk,'date':str(self.day),'state':'ARRIVAL_OVERDUE'})
        self.assertContains(response,'housing-arrival-overdue');self.assertContains(response,due.strftime('%d.%m.%Y'))
        CampStay.objects.create(company=self.company,camp=self.a,person=self.second_person,arrived_on=self.day)
        self.assertEqual(rotation_rows(self.company,[self.a],self.day,state='ARRIVAL_OVERDUE')[0],[])

    def test_critical_filters_do_not_fabricate_missing_schedule_dates_and_export_matches(self):
        StaffingSlot.objects.create(company=self.company,group=self.group,person=self.second_person,housing_camp=self.a,status='МО')
        data={'camp':self.a.pk,'date':str(self.day),'state':'ARRIVAL_OVERDUE'}
        page=self.client.get(reverse('personnel:housing'),data)
        self.assertEqual(page.context['rotation_stats']['people'],0);self.assertContains(page,'data-rotation-filter')
        from openpyxl import load_workbook
        from io import BytesIO
        exported=self.client.get(reverse('personnel:housing_export'),data)
        self.assertEqual(load_workbook(BytesIO(exported.content)).active.max_row,1)
