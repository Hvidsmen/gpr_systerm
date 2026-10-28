from unittest.mock import patch
from datetime import timedelta
from io import BytesIO
from django.core.exceptions import ValidationError
from django.test import TestCase,Client
from django.urls import reverse
from django.utils import timezone
from apps.accounts.models import Company,Role,User
from apps.projects.models import Project,ConstructionObject
from apps.resources.models import Employee,Brigade
from apps.resources.position_catalog import position_for
from apps.production import test_dispatch as dispatch_fixtures
from apps.production.dispatch_models import EmployeeActivity
from .models import Camp,CampObjectLink,CampStay,Person,Board,StaffingGroup,StaffingSlot,ShiftPattern,HousingEvent
from . import housing


class HousingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.company=Company.objects.create(name='ВЖГ test')
        cls.project=Project.objects.create(company=cls.company,name='Проект')
        cls.obj=ConstructionObject.objects.create(company=cls.company,project=cls.project,name='КТП-1')
        cls.obj2=ConstructionObject.objects.create(company=cls.company,project=cls.project,name='КТП-2')
        cls.obj3=ConstructionObject.objects.create(company=cls.company,project=cls.project,name='КТП-3')
        cls.a=Camp.objects.create(company=cls.company,name='ВЖГ КТП-1')
        cls.b=Camp.objects.create(company=cls.company,name='Базовый ВЖГ')
        for camp,obj in [(cls.a,cls.obj),(cls.a,cls.obj2),(cls.b,cls.obj2),(cls.b,cls.obj3)]:CampObjectLink.objects.create(company=cls.company,camp=camp,construction_object=obj)
        cls.position=position_for(Brigade.objects.create(company=cls.company,name='Монтажник'))
        cls.employee=Employee.objects.create(company=cls.company,last_name='Иванов',first_name='Иван',position=cls.position,employee_number='001')
        cls.second_employee=Employee.objects.create(company=cls.company,last_name='Петров',first_name='Петр',position=cls.position,employee_number='002')
        cls.person=Person.objects.create(company=cls.company,employee=cls.employee,full_name=cls.employee.display_name,employee_number='001')
        cls.second_person=Person.objects.create(company=cls.company,employee=cls.second_employee,full_name=cls.second_employee.display_name,employee_number='002')
        cls.day=timezone.localdate()-timedelta(days=1)
        cls.stay=CampStay.objects.create(company=cls.company,camp=cls.a,person=cls.person,arrived_on=cls.day-timedelta(days=1))
        cls.hr=User.objects.create_user(username='housing-hr',company=cls.company,role=Role.objects.get(code='PERSONNEL_DISPATCHER'))
        cls.hr2=User.objects.create_user(username='housing-hr2',company=cls.company,role=Role.objects.get(code='PERSONNEL_DISPATCHER'))
        cls.dispatcher=User.objects.create_user(username='housing-production',company=cls.company,role=Role.objects.get(code='DISPATCHER'))
        cls.dispatcher.assigned_objects.add(cls.obj,cls.obj2,cls.obj3)
        cls.users={'DISPATCHER':cls.dispatcher};cls.url=reverse('production:fact_day_workspace')
        cls.board=Board.objects.create(company=cls.company,name='Существующая таблица')
        cls.group=StaffingGroup.objects.create(company=cls.company,board=cls.board,position='Монтажник')
        cls.pattern=ShiftPattern.objects.create(company=cls.company,board=cls.board,name='2/2',on_days=2,off_days=2,cycle_codes=['10','10','ДП','МО'],calendar={cls.day.isoformat():'10',(cls.day+timedelta(days=1)).isoformat():'ДП',(cls.day+timedelta(days=2)).isoformat():'МО'})
        cls.slot=StaffingSlot.objects.create(company=cls.company,group=cls.group,person=cls.person,housing_camp=cls.a,pattern=cls.pattern)
        cls.other=Company.objects.create(name='Secret company')
        cls.foreign_camp=Camp.objects.create(company=cls.other,name='Secret camp')
        cls.foreign_person=Person.objects.create(company=cls.other,full_name='Secret person')

    payload=dispatch_fixtures.DispatcherTests.payload

    def setUp(self):self.client.force_login(self.hr)

    def visible(self,obj,day=None):return set(housing.available_employee_ids(obj,day or self.day).values_list('person__employee_id',flat=True))

    def test_many_to_many_camps_restrict_and_deduplicate_employee_pool(self):
        self.assertEqual(self.visible(self.obj),{self.employee.pk})
        self.assertEqual(self.visible(self.obj2),{self.employee.pk})
        self.assertEqual(self.visible(self.obj3),set())
        housing.move(self.hr,'ARRIVAL',self.b,self.day,[self.second_person])
        self.assertEqual(self.visible(self.obj2),{self.employee.pk,self.second_employee.pk})
        self.assertEqual(self.visible(self.obj3),{self.second_employee.pk})
        self.assertEqual(len(self.visible(self.obj2)),2)

    def test_transfer_closes_old_camp_and_opens_new_from_same_date(self):
        token=housing.movement_snapshot(self.hr)
        housing.move(self.hr,'TRANSFER',self.a,self.day,[self.person],target=self.b,snapshot=token)
        self.assertEqual(self.visible(self.obj),set())
        self.assertEqual(self.visible(self.obj3),{self.employee.pk})
        self.assertEqual(self.visible(self.obj,self.day-timedelta(days=1)),{self.employee.pk})
        self.assertEqual(housing.present_stays(self.company,self.day).filter(person=self.person).count(),1)
        self.assertEqual(HousingEvent.objects.get().kind,'TRANSFER')
        self.assertEqual(self.slot.pattern_id,self.pattern.pk)

    def test_arrival_departure_repeat_cycles_and_future_actual_is_rejected(self):
        housing.move(self.hr,'DEPARTURE',self.a,self.day,[self.person])
        self.assertEqual(self.visible(self.obj),set())
        housing.move(self.hr,'ARRIVAL',self.a,self.day+timedelta(days=1),[self.person])
        self.assertEqual(self.visible(self.obj,self.day+timedelta(days=1)),{self.employee.pk})
        self.assertEqual(self.person.stays.count(),2)
        with self.assertRaises(ValidationError):housing.move(self.hr,'ARRIVAL',self.a,timezone.localdate()+timedelta(days=1),[self.second_person])

    def test_same_day_arrival_and_transfer_are_audited_without_double_presence(self):
        housing.move(self.hr,'ARRIVAL',self.a,self.day,[self.second_person])
        housing.move(self.hr,'TRANSFER',self.a,self.day,[self.second_person],target=self.b)
        self.assertEqual(self.visible(self.obj3),{self.second_employee.pk})
        self.assertEqual(housing.present_stays(self.company,self.day).filter(person=self.second_person).count(),1)
        self.assertEqual(self.second_person.housing_events.count(),2)

    def test_bulk_conflict_rolls_back_every_employee(self):
        with self.assertRaises(ValidationError):housing.move(self.hr,'ARRIVAL',self.a,self.day,[self.second_person,self.person])
        self.assertFalse(self.second_person.stays.exists());self.assertFalse(HousingEvent.objects.exists())

    def test_distinct_dispatchers_merge_distinct_people_but_same_person_conflicts(self):
        left=housing.movement_snapshot(self.hr);right=housing.movement_snapshot(self.hr2)
        housing.move(self.hr,'DEPARTURE',self.a,self.day,[self.person],snapshot=left)
        housing.move(self.hr2,'ARRIVAL',self.b,self.day,[self.second_person],snapshot=right)
        with self.assertRaises(ValidationError):housing.move(self.hr2,'ARRIVAL',self.b,self.day+timedelta(days=1),[self.person],snapshot=right)
        self.assertEqual(self.person.stays.count(),1)

    def test_employee_source_creates_one_linked_person_and_bulk_duplicates_are_deduplicated(self):
        employee=Employee.objects.create(company=self.company,last_name='Новый',first_name='Работник',position=self.position,employee_number='003')
        snapshot=housing.movement_snapshot(self.hr)
        housing.move(self.hr,'ARRIVAL',self.a,self.day,[],employees=[employee],snapshot=snapshot)
        self.assertEqual(Person.objects.filter(employee=employee).count(),1)
        self.assertIn(employee.pk,self.visible(self.obj))
        with self.assertRaises(ValidationError):housing.move(self.hr,'ARRIVAL',self.a,self.day,[self.person],employees=[self.employee])
        self.assertEqual(Person.objects.filter(employee=self.employee).count(),1)

    def test_cancel_requires_reason_and_version_and_preserves_history(self):
        with self.assertRaises(ValidationError):housing.cancel_stay(self.hr,self.stay,'',self.stay.updated_at.isoformat())
        with self.assertRaises(ValidationError):housing.cancel_stay(self.hr,self.stay,'Ошибка','old')
        with patch('django.utils.timezone.localdate',return_value=self.day):
            housing.cancel_stay(self.hr,self.stay,'Ошибка исходной даты',self.stay.updated_at.isoformat())
        self.assertEqual(self.visible(self.obj),set());self.stay.refresh_from_db();self.assertIsNotNone(self.stay.cancelled_at)
        self.assertEqual(self.person.stays.count(),1)
        self.assertEqual(HousingEvent.objects.get().kind,'CANCEL')

    def test_company_boundaries_permissions_and_csrf(self):
        for name,args in [('housing_history',[self.foreign_person.pk]),('camp_edit',[self.foreign_camp.pk])]:self.assertEqual(self.client.get(reverse('personnel:'+name,args=args)).status_code,404)
        self.assertEqual(self.client.get(reverse('personnel:housing'),{'camp':self.foreign_camp.pk}).status_code,404)
        response=self.client.post(reverse('personnel:camp_edit',args=[self.a.pk]),{'version':self.a.updated_at.isoformat(),'name':self.a.name,'objects_served':[999999],'is_active':'on'})
        self.assertEqual(response.status_code,200);self.assertTrue(response.context['form'].errors)
        self.client.force_login(self.dispatcher);self.assertEqual(self.client.get(reverse('personnel:housing')).status_code,403)
        client=Client(enforce_csrf_checks=True);client.force_login(self.hr)
        self.assertEqual(client.post(reverse('personnel:link_people')).status_code,403)

    def test_camp_form_saves_many_to_many_links_and_keeps_existing_fact(self):
        response=self.client.post(reverse('personnel:camp_edit',args=[self.a.pk]),{'version':self.a.updated_at.isoformat(),'name':self.a.name,'objects_served':[self.obj3.pk,self.obj2.pk],'is_active':'on'})
        self.assertEqual(response.status_code,302)
        self.assertEqual(set(self.a.objects_served.values_list('pk',flat=True)),{self.obj2.pk,self.obj3.pk})
        self.assertEqual(self.visible(self.obj),set());self.assertEqual(self.visible(self.obj3),{self.employee.pk})

    def test_production_accepts_active_employees_without_or_after_camp_residence(self):
        data=self.payload('DISPATCHER','works')
        data.update({'labor-0-employee':self.second_employee.pk,'labor-0-work_mode':'general'})
        response=self.client.post(self.url,data);self.assertEqual(response.status_code,302)
        self.assertTrue(EmployeeActivity.objects.filter(employee=self.second_employee).exists())
        data=self.payload('DISPATCHER','works');data.update({'labor-0-employee':self.employee.pk,'labor-0-work_mode':'general'})
        housing.move(self.hr,'DEPARTURE',self.a,self.day,[self.person])
        self.client.force_login(self.dispatcher);response=self.client.post(self.url,data);self.assertEqual(response.status_code,302)
        self.assertTrue(EmployeeActivity.objects.filter(employee=self.employee).exists())

    def test_employee_claim_without_camp_and_saved_historical_row_remain_editable(self):
        from uuid import uuid4
        self.client.force_login(self.dispatcher)
        response=self.client.post(reverse('production:work_claims'),{'construction_object':self.obj.pk,'date':str(self.day),'session':str(uuid4()),'key':f'employee:{self.second_employee.pk}:1','action':'take'},content_type='application/json')
        self.assertEqual(response.status_code,200)
        data=self.payload('DISPATCHER','works');data.update({'labor-0-employee':self.employee.pk,'labor-0-work_mode':'general'})
        self.assertEqual(self.client.post(self.url,data).status_code,302)
        housing.move(self.hr,'DEPARTURE',self.a,self.day,[self.person])
        data=self.payload('DISPATCHER','works');data['labor-0-activity']='Уточнение ранее сохранённого факта'
        self.assertEqual(self.client.post(self.url,data).status_code,302)
        self.assertEqual(EmployeeActivity.objects.get().activity,'Уточнение ранее сохранённого факта')

    def test_existing_calendar_is_used_and_overrides_do_not_move_employee(self):
        response=self.client.get(reverse('personnel:housing'),{'camp':self.a.pk,'date':str(self.day)})
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.context['rows'][0]['today_code'],'10')
        self.assertEqual(response.context['rows'][0]['planned_departure'],str(self.day+timedelta(days=1)))
        response=self.client.post(reverse('personnel:person_schedule',args=[self.person.pk]),{'version':housing.schedule_snapshot(Person.objects.get(pk=self.person.pk)),'slot':self.slot.pk,'pattern':self.pattern.pk,'start':str(self.day),'end':str(self.day),'code':'В'})
        self.assertEqual(response.status_code,302)
        self.slot.refresh_from_db();self.assertEqual(self.slot.calendar[self.day.isoformat()],'В')
        self.assertEqual(self.visible(self.obj),{self.employee.pk})
        self.assertEqual(self.person.stays.count(),1)
        response=self.client.post(reverse('personnel:person_schedule',args=[self.person.pk]),{'version':housing.schedule_snapshot(Person.objects.get(pk=self.person.pk)),'slot':self.slot.pk,'pattern':self.pattern.pk,'start':str(self.day),'end':str(self.day),'code':''})
        self.assertEqual(response.status_code,302);self.slot.refresh_from_db();self.assertEqual(self.slot.calendar,{})

    def test_new_person_calendar_uses_existing_table_structure(self):
        response=self.client.post(reverse('personnel:person_schedule',args=[self.second_person.pk]),{'version':housing.schedule_snapshot(self.second_person),'start':str(self.day),'end':str(self.day+timedelta(days=1)),'code':'10'})
        self.assertEqual(response.status_code,302);self.second_person.refresh_from_db()
        self.assertIsNotNone(self.second_person.schedule_slot_id)
        self.assertEqual(self.second_person.schedule_slot.calendar[self.day.isoformat()],'10')
        self.assertFalse(self.second_person.stays.exists())

    def test_linking_is_unambiguous_and_cannot_rebind_a_person_with_history(self):
        self.second_person.employee=None;self.second_person.save()
        self.assertEqual(housing.link_people(self.company),1)
        self.second_person.refresh_from_db();self.assertEqual(self.second_person.employee_id,self.second_employee.pk)
        response=self.client.post(reverse('personnel:person_edit',args=[self.person.pk]),{'full_name':self.person.full_name,'employee_number':'001','employee':''})
        self.assertEqual(response.status_code,200);self.person.refresh_from_db();self.assertEqual(self.person.employee_id,self.employee.pk)
        Employee.objects.create(company=self.company,last_name='Однофамилец',first_name='Иван',employee_number='009')
        Employee.objects.create(company=self.company,last_name='Однофамилец',first_name='Иван',employee_number='009')
        unresolved=Person.objects.create(company=self.company,full_name='Однофамилец Иван',employee_number='009')
        self.assertIsNone(housing.matching_employee(unresolved))

    def test_new_workspace_history_and_filtered_exports(self):
        import openpyxl
        for name,args in [('housing',[]),('camps',[]),('housing_history',[self.person.pk]),('person_schedule',[self.person.pk])]:
            self.assertEqual(self.client.get(reverse('personnel:'+name,args=args)).status_code,200,name)
        response=self.client.get(reverse('personnel:housing_export'),{'camp':self.a.pk,'date':str(self.day),'q':'Иванов'})
        book=openpyxl.load_workbook(BytesIO(response.content));self.assertEqual(book.active.max_row,2)
        self.assertEqual(book.active['C2'].value,self.person.full_name)
        self.employee.zup_full_name='=HYPERLINK("bad")';self.employee.save()
        response=self.client.get(reverse('personnel:housing_export'),{'camp':self.a.pk,'date':str(self.day)})
        self.assertTrue(openpyxl.load_workbook(BytesIO(response.content)).active['C2'].value.startswith("'="))

    def test_legacy_values_are_preserved_and_conflicting_camps_are_not_guessed(self):
        board=Board.objects.create(company=self.company,name='Legacy')
        group=StaffingGroup.objects.create(company=self.company,board=board,position='Сварщик')
        p=Person.objects.create(company=self.company,full_name='Неоднозначный Работник')
        for name in ['Старый ВЖГ','Другой ВЖГ']:
            StaffingSlot.objects.create(company=self.company,group=group,person=p,camp=name,status='МВ',arrival_date=self.day)
        housing.bootstrap_legacy(self.company,StaffingSlot.objects.filter(group=group))
        self.assertEqual(StaffingSlot.objects.filter(group=group,housing_camp__isnull=False).count(),2)
        self.assertFalse(p.stays.exists())
        self.assertEqual(list(StaffingSlot.objects.filter(group=group).values_list('camp',flat=True)),['Старый ВЖГ','Другой ВЖГ'])

    def test_schedule_and_camp_edits_reject_stale_versions(self):
        token=housing.schedule_snapshot(self.person)
        data={'version':token,'slot':self.slot.pk,'pattern':self.pattern.pk,'start':str(self.day),'end':str(self.day),'code':'В'}
        url=reverse('personnel:person_schedule',args=[self.person.pk])
        self.assertEqual(self.client.post(url,data).status_code,302)
        data['code']='МО';response=self.client.post(url,data)
        self.assertEqual(response.status_code,200);self.assertTrue(response.context['form'].errors)
        self.slot.refresh_from_db();self.assertEqual(self.slot.calendar[self.day.isoformat()],'В')
        camp_data={'name':self.a.name,'is_active':'on','version':self.a.updated_at.isoformat(),'objects_served':[self.obj.pk]}
        url=reverse('personnel:camp_edit',args=[self.a.pk]);self.assertEqual(self.client.post(url,camp_data).status_code,302)
        camp_data['objects_served']=[self.obj3.pk];response=self.client.post(url,camp_data)
        self.assertEqual(response.status_code,200);self.assertTrue(response.context['form'].errors)
        self.assertEqual(set(self.a.objects_served.values_list('pk',flat=True)),{self.obj.pk})

    def test_dispatcher_lands_in_housing_and_view_posts_are_atomic(self):
        self.assertRedirects(self.client.get(reverse('personnel:board_list')),reverse('personnel:housing'),fetch_redirect_response=False)
        self.assertEqual(self.client.get(reverse('personnel:board_list'),{'tables':'1'}).status_code,200)
        response=self.client.get(reverse('personnel:housing'),{'camp':self.a.pk,'date':str(self.day)})
        snapshot=response.context['form'].initial['snapshot']
        data={'camp':self.a.pk,'kind':'ARRIVAL','date':str(self.day),'employees':[self.second_employee.pk],'snapshot':snapshot}
        response=self.client.post(reverse('personnel:housing'),data);self.assertEqual(response.status_code,302)
        self.assertEqual(self.visible(self.obj),{self.employee.pk,self.second_employee.pk})
        response=self.client.post(reverse('personnel:housing'),data)
        self.assertEqual(response.status_code,200);self.assertTrue(response.context['form'].errors)
        self.assertEqual(self.second_person.stays.count(),1)

    def test_source_status_does_not_change_residence_without_movement(self):
        self.stay.departed_on=self.day;self.stay.save()
        self.slot.status='МВ';self.slot.arrival_date=self.day-timedelta(days=5);self.slot.save()
        self.assertEqual(self.visible(self.obj),set())
        response=self.client.get(reverse('personnel:housing'),{'camp':self.a.pk,'date':str(self.day)})
        self.assertEqual(response.context['rotation_stats']['present'],0)
        self.assertEqual(response.context['rotation_stats']['rest'],1)

    def test_homonymous_personnel_cards_are_not_automatically_bound_or_duplicated(self):
        e=Employee.objects.create(company=self.company,last_name='Тезка',first_name='Сотрудник',position=self.position)
        p=Person.objects.create(company=self.company,full_name=e.display_name)
        Person.objects.create(company=self.company,full_name=e.display_name)
        self.assertIsNone(housing.matching_employee(p))
        before=Person.objects.count()
        with self.assertRaises(ValidationError):housing.person_for(e)
        self.assertEqual(Person.objects.count(),before)
