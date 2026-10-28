from datetime import timedelta
from copy import deepcopy
from django.test import TestCase
from django.urls import reverse
from apps.accounts.models import User,Role
from apps.production.history_models import FactRevision
from .test_housing import HousingTests
from .models import CampStay,ShiftPattern,Person
from . import housing
from .rotation import rotation_rows
from .services import calendar_for,departure_days,report
from .rotation_calendar import cycle_definition,description


class ActualCycleTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        HousingTests.setUpTestData.__func__(cls)
        cls.admin=User.objects.create_user(username='cycle-admin',company=cls.company,role=Role.objects.get(code='ADMIN'))

    def setUp(self):
        self.client.force_login(self.hr)
        self.pattern.on_days=60;self.pattern.off_days=30;self.pattern.name='60/30'
        self.pattern.cycle_codes=['10']*60+['МО']*30;self.pattern.save()
        self.stay.arrived_on=self.day;self.stay.save()
        self.slot.pattern=self.pattern

    def calendar(self,start=None,end=None):
        return calendar_for(self.slot,start=start or self.day,end=end or self.day+timedelta(days=180))

    def test_existing_actual_arrival_restarts_full_cycle_and_extends_beyond_source(self):
        cal=self.calendar()
        self.assertEqual(cal[str(self.day)],'10');self.assertEqual(cal[str(self.day+timedelta(days=59))],'10')
        self.assertEqual(cal[str(self.day+timedelta(days=60))],'МО');self.assertEqual(cal[str(self.day+timedelta(days=89))],'МО')
        self.assertEqual(cal[str(self.day+timedelta(days=90))],'10')
        self.assertEqual(departure_days(cal,self.day)[0],str(self.day+timedelta(days=60)))

    def test_actual_rest_departure_reanchors_rest_instead_of_original_planned_day(self):
        actual=self.day+timedelta(days=1)
        housing.move(self.hr,'DEPARTURE',self.a,actual,[self.person])
        cal=self.calendar()
        self.assertEqual(cal[str(self.day)],'10');self.assertEqual(cal[str(actual)],'МО')
        self.assertEqual(cal[str(actual+timedelta(days=29))],'МО');self.assertEqual(cal[str(actual+timedelta(days=30))],'10')
        row=rotation_rows(self.company,[self.a],actual)[0][0]
        self.assertEqual(row['planned_return'],actual+timedelta(days=30))

    def test_transfer_preserves_original_anchor_and_overdue_departure(self):
        self.stay.arrived_on=self.day-timedelta(days=61);self.stay.save()
        housing.move(self.hr,'TRANSFER',self.a,self.day,[self.person],target=self.b)
        rows,_,_=rotation_rows(self.company,[self.b],self.day)
        self.assertEqual(rows[0]['planned_departure_display'],self.day-timedelta(days=1))
        self.assertEqual(rows[0]['departure_warning'],'overdue')
        self.assertEqual(rows[0]['state_days'],62)
        self.assertIn((self.day-timedelta(days=61)).strftime('%d.%m.%Y'),rows[0]['schedule_info'])

    def test_later_actual_arrival_restarts_work_and_preserves_previous_phase(self):
        self.stay.arrived_on=self.day-timedelta(days=120);self.stay.departed_on=self.day-timedelta(days=60);self.stay.save()
        actual=self.day-timedelta(days=2)
        housing.move(self.admin,'ARRIVAL',self.a,actual,[self.person])
        cal=self.calendar(start=self.day-timedelta(days=120))
        self.assertEqual(cal[str(actual)],'10');self.assertEqual(cal[str(actual+timedelta(days=59))],'10')
        self.assertEqual(cal[str(self.stay.departed_on)],'МО')
        self.assertEqual(cal[str(actual+timedelta(days=60))],'МО')

    def test_projection_does_not_write_sources_actuals_or_history_and_keeps_manual_exceptions(self):
        self.slot.calendar={str(self.day+timedelta(days=4)):'В'};self.slot.save()
        before=(deepcopy(self.pattern.calendar),deepcopy(self.slot.calendar),self.stay.updated_at,FactRevision.objects.count())
        cal=self.calendar();self.pattern.refresh_from_db();self.slot.refresh_from_db();self.stay.refresh_from_db()
        self.assertEqual(cal[str(self.day+timedelta(days=4))],'В')
        self.assertEqual((self.pattern.calendar,self.slot.calendar,self.stay.updated_at,FactRevision.objects.count()),before)
        original=self.pattern.calendar[str(self.day-timedelta(days=1))] if str(self.day-timedelta(days=1)) in self.pattern.calendar else None
        self.assertEqual(cal.get(str(self.day-timedelta(days=1))),original)

    def test_cancelling_wrong_actual_removes_its_anchor_without_deleting_history(self):
        old=self.calendar();self.assertIn(str(self.day+timedelta(days=60)),old)
        housing.cancel_stay(self.hr,self.stay,'Ошибочный заезд',self.stay.updated_at.isoformat())
        self.assertNotIn(str(self.day+timedelta(days=60)),self.calendar())
        self.assertTrue(FactRevision.objects.filter(entity_type='personnel.campstay',entity_id=str(self.stay.pk),operation='UPDATE').exists())

    def test_named_imported_cycle_uses_complete_template_including_weekends_and_travel(self):
        first=self.day-timedelta(days=100)
        codes=['10','В','10','ДП','МО','ДП']
        pattern=ShiftPattern(company=self.company,board=self.board,name='Вахта 1 (3/3)',calendar={str(first-timedelta(days=1)):'МО',**{str(first+timedelta(days=i)):code for i,code in enumerate(codes)}})
        self.assertEqual(cycle_definition(pattern)['codes'],codes)
        pattern.save();self.slot.pattern=pattern
        cal=self.calendar(end=self.day+timedelta(days=6))
        self.assertEqual([cal[str(self.day+timedelta(days=i))] for i in range(6)],codes)

    def test_cycle_without_any_known_hours_does_not_invent_hours(self):
        pattern=ShiftPattern(company=self.company,board=self.board,name='45/45')
        self.assertEqual(cycle_definition(pattern)['codes'][0],'МВ')
        pattern.name='Индивидуальный';self.assertIsNone(cycle_definition(pattern))
        pattern.save();self.slot.pattern=pattern
        self.assertEqual(self.calendar(),self.slot.calendar)
        self.assertIn('Цикл не задан',description(self.slot,self.day))

    def test_popup_and_rotation_use_forecast_while_opening_table_preserves_source(self):
        row=rotation_rows(self.company,[self.a],self.day)[0][0]
        self.assertEqual(row['today_code'],'10');self.assertEqual(row['planned_departure_display'],self.day+timedelta(days=60))
        popup=self.client.get(reverse('personnel:individual_calendar',args=[self.person.pk]),{'month':self.day.strftime('%Y-%m')}).json()
        self.assertEqual(popup['days'][self.day.day-1]['code'],row['today_code']);self.assertIn('фактического заезда',popup['schedule_info'])
        table=report(self.board,self.day.replace(day=1))
        self.assertEqual(table['groups'][0]['rows'][0]['codes'],[self.pattern.calendar.get(str(self.day.replace(day=n)),'') for n in range(1,len(table['dates'])+1)])

    def test_personal_cycle_can_be_set_without_changing_shared_pattern(self):
        before=deepcopy(self.pattern.cycle_codes)
        response=self.client.post(reverse('personnel:person_schedule',args=[self.person.pk]),{'version':housing.schedule_snapshot(Person.objects.get(pk=self.person.pk)),'slot':self.slot.pk,'pattern':self.pattern.pk,'on_days':14,'off_days':7,'hours':'8'})
        self.assertEqual(response.status_code,302);self.slot.refresh_from_db();self.pattern.refresh_from_db()
        self.assertNotEqual(self.slot.pattern_id,self.pattern.pk);self.assertEqual(self.pattern.cycle_codes,before)
        self.assertEqual(self.slot.pattern.on_days,14)
        cal=self.calendar();self.assertEqual(cal[str(self.day)],'8');self.assertEqual(cal[str(self.day+timedelta(days=14))],'МО')

    def test_missing_half_cycle_is_rejected_and_company_scope_remains_enforced(self):
        response=self.client.post(reverse('personnel:person_schedule',args=[self.person.pk]),{'version':housing.schedule_snapshot(self.person),'slot':self.slot.pk,'on_days':14})
        self.assertEqual(response.status_code,200);self.assertTrue(response.context['form'].errors)
        self.assertEqual(self.client.get(reverse('personnel:individual_calendar',args=[self.foreign_person.pk])).status_code,404)

    def test_long_unclosed_stay_keeps_original_deadline_without_expanding_every_historical_day(self):
        self.stay.arrived_on=self.day-timedelta(days=365*20);self.stay.save()
        cal=self.calendar()
        self.assertLess(len(cal),1200)
        self.assertEqual(departure_days(cal,self.stay.arrived_on)[0],str(self.stay.arrived_on+timedelta(days=60)))

    def test_imported_cycle_can_start_with_weekend_after_incoming_travel(self):
        first=self.day-timedelta(days=100);codes=['В','10','10','ДП','МО','ДП']
        pattern=ShiftPattern(company=self.company,board=self.board,name='3/3',calendar={str(first-timedelta(days=1)):'ДП',**{str(first+timedelta(days=i)):code for i,code in enumerate(codes)}})
        self.assertEqual(cycle_definition(pattern)['codes'],codes)

    def test_explicit_primary_schedule_is_used_in_other_camp_and_popup(self):
        from .models import StaffingSlot
        alternate=ShiftPattern.objects.create(company=self.company,board=self.board,name='3/2',on_days=3,off_days=2,cycle_codes=['8']*3+['МО']*2)
        other_slot=StaffingSlot.objects.create(company=self.company,group=self.group,person=self.person,housing_camp=self.b,pattern=alternate)
        self.person.schedule_slot=other_slot;self.person.save()
        row=rotation_rows(self.company,[self.a],self.day)[0][0]
        self.assertEqual(row['today_code'],'8');self.assertEqual(row['planned_departure_display'],self.day+timedelta(days=3))
        popup=self.client.get(reverse('personnel:individual_calendar',args=[self.person.pk]),{'month':self.day.strftime('%Y-%m')}).json()
        self.assertEqual(popup['days'][self.day.day-1]['code'],'8')

    def test_cycle_boundaries_do_not_require_hours_on_last_workday_or_first_return_day(self):
        self.pattern.on_days=3;self.pattern.off_days=2;self.pattern.cycle_codes=['В','10','В','МО','МО'];self.pattern.save()
        row=rotation_rows(self.company,[self.a],self.day)[0][0]
        self.assertEqual(row['planned_departure_display'],self.day+timedelta(days=3))
        housing.move(self.hr,'DEPARTURE',self.a,self.day,[self.person])
        row=rotation_rows(self.company,[self.a],self.day)[0][0]
        self.assertEqual(row['planned_return'],self.day+timedelta(days=2))

    def test_manual_hours_on_actual_departure_day_do_not_turn_into_immediate_return(self):
        self.slot.calendar={str(self.day):'7'};self.slot.save()
        housing.move(self.hr,'DEPARTURE',self.a,self.day,[self.person])
        row=rotation_rows(self.company,[self.a],self.day)[0][0]
        self.assertEqual(row['today_code'],'7');self.assertEqual(row['planned_return'],self.day+timedelta(days=30))

    def test_editing_imported_initial_schedule_creates_personal_copy(self):
        self.board.import_digest='initial-balance';self.board.save()
        original_pattern=self.slot.pattern_id;original_calendar=deepcopy(self.slot.calendar)
        response=self.client.post(reverse('personnel:person_schedule',args=[self.person.pk]),{'version':housing.schedule_snapshot(Person.objects.get(pk=self.person.pk)),'slot':self.slot.pk,'pattern':self.pattern.pk,'on_days':14,'off_days':7,'hours':'8','start':str(self.day),'end':str(self.day),'code':'В'})
        self.assertEqual(response.status_code,302)
        self.slot.refresh_from_db();self.person.refresh_from_db()
        self.assertEqual(self.slot.pattern_id,original_pattern);self.assertEqual(self.slot.calendar,original_calendar)
        self.assertNotEqual(self.person.schedule_slot_id,self.slot.pk)
        self.assertEqual(self.person.schedule_slot.pattern.on_days,14)
        self.assertEqual(self.person.schedule_slot.calendar[str(self.day)],'В')
        self.assertEqual(report(self.board,self.day.replace(day=1))['total']['employed'],1)

    def test_duration_counts_arrival_and_rest_days_inclusively_and_exports_full_schedule(self):
        self.pattern.name='Вахта 1.3 (60/30)';self.pattern.kind='М';self.pattern.save()
        row=rotation_rows(self.company,[self.a],self.day)[0][0]
        self.assertEqual(row['state_days'],1);self.assertEqual(row['schedule_label'],'М · Вахта 1.3 (60/30)')
        from io import BytesIO
        import openpyxl
        response=self.client.get(reverse('personnel:housing_export'),{'camp':self.a.pk,'date':str(self.day)})
        sheet=openpyxl.load_workbook(BytesIO(response.content)).active;headers=[c.value for c in sheet[1]]
        self.assertEqual(sheet.cell(2,headers.index('Дней на вахте / МО')+1).value,1)
        self.assertEqual(sheet.cell(2,headers.index('График')+1).value,'М · Вахта 1.3 (60/30)')
        housing.move(self.hr,'DEPARTURE',self.a,self.day,[self.person])
        row=rotation_rows(self.company,[self.a],self.day+timedelta(days=2))[0][0]
        self.assertEqual(row['state_days'],3);self.assertEqual(row['state_since'],self.day)

    def test_undated_legacy_status_does_not_invent_elapsed_days(self):
        from .models import StaffingSlot
        StaffingSlot.objects.create(company=self.company,group=self.group,person=self.second_person,housing_camp=self.a,status='МО')
        row=next(r for r in rotation_rows(self.company,[self.a],self.day)[0] if r['person'].pk==self.second_person.pk)
        self.assertIsNone(row['state_days'])

    def test_imported_holiday_extension_does_not_put_work_into_actual_rest_phase(self):
        first=self.day-timedelta(days=100)
        values=['10']*47+['ДП']*2+['МО']*39+['ДП']*2
        pattern=ShiftPattern(company=self.company,board=self.board,name='45/45',calendar={str(first-timedelta(days=1)):'МО',**{str(first+timedelta(days=i)):code for i,code in enumerate(values)}})
        cycle=cycle_definition(pattern)
        self.assertEqual(len(cycle['codes']),90);self.assertEqual(cycle['codes'][45],'ДП')
        self.assertEqual(cycle['codes'][-2:],['ДП','ДП'])
        self.assertFalse(any(code=='10' for code in cycle['codes'][45:]))
