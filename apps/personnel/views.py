from datetime import date, timedelta
from io import BytesIO
import secrets
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_http_methods
from core.permissions import require_roles,role_code
from .models import Board, Person, StaffingGroup, StaffingSlot
from .forms import BoardForm, PersonForm, GroupForm, SlotForm, CalendarForm, ImportForm
from .services import report
from . import importer


def access(request): require_roles(request.user, {'ADMIN','HR_HEAD','PERSONNEL_DISPATCHER'})
def scoped(request, model):
    access(request)
    return model.objects.filter(company_id=request.user.company_id)


@login_required
def board_list(request):
    if role_code(request.user)=='PERSONNEL_DISPATCHER' and not request.GET.get('tables'):return redirect('personnel:housing')
    return render(request,'personnel/boards.html',{'boards':scoped(request,Board)})


@login_required
@require_http_methods(['GET','POST'])
def board_create(request):
    access(request)
    obj=Board(company=request.user.company)
    form=BoardForm(request.POST if request.method=='POST' else None,instance=obj)
    if request.method=='POST' and form.is_valid():
        form.save(); return redirect('personnel:board_detail',pk=obj.pk)
    return render(request,'personnel/form.html',{'form':form,'title':'Новая таблица кадрового учета'})


def context(request, board):
    try: month=date.fromisoformat(request.GET.get('month',timezone.localdate().strftime('%Y-%m'))+'-01')
    except ValueError: month=timezone.localdate().replace(day=1)
    data=report(board,month,request.GET.get('q',''),request.GET.get('status',''),request.GET.get('camp',''))
    return dict(data,board=board,month=month,query=request.GET.get('q',''),selected_status=request.GET.get('status',''),selected_camp=request.GET.get('camp',''),camps=StaffingSlot.objects.filter(group__board=board).exclude(camp='').values_list('camp',flat=True).distinct().order_by('camp'))


@login_required
def board_detail(request,pk):
    board=get_object_or_404(scoped(request,Board),pk=pk)
    return render(request,'personnel/board.html',context(request,board))


@login_required
def people(request):
    queryset=scoped(request,Person).select_related('employee').prefetch_related('assignments__group__board')
    if request.GET.get('unlinked'):queryset=queryset.filter(employee=None)
    if request.GET.get('unplaced'):queryset=queryset.filter(stays__isnull=True)
    query=request.GET.get('q','').strip()
    if query: queryset=queryset.filter(Q(full_name__icontains=query)|Q(employee_number__icontains=query)|Q(phone__icontains=query))
    return render(request,'personnel/people.html',{'people':queryset,'query':query})


@login_required
@require_http_methods(['GET','POST'])
def person_edit(request,pk=None):
    queryset=scoped(request,Person)
    obj=get_object_or_404(queryset,pk=pk) if pk else Person(company=request.user.company)
    form=PersonForm(request.POST if request.method=='POST' else None,instance=obj)
    saved=False
    if request.method=='POST':
        with transaction.atomic():
            # Match movement/import lock order before checking identity against history.
            from apps.accounts.models import Company
            from django.db.models import F
            Company.objects.filter(pk=request.user.company_id).update(name=F('name'))
            if pk:obj=get_object_or_404(queryset.select_for_update(),pk=pk)
            form=PersonForm(request.POST,instance=obj)
            if form.is_valid():
                person=form.save()
                if not person.employee_id:
                    from .housing import matching_employee
                    employee=matching_employee(person)
                    if employee:person.employee=employee;person.save(update_fields=['employee','updated_at'])
                saved=True
        if saved:
            messages.success(request,'Карточка сотрудника сохранена.');return redirect('personnel:people')
    return render(request,'personnel/form.html',{'form':form,'title':'Карточка сотрудника'})


@login_required
@require_http_methods(['GET','POST'])
def group_edit(request,board_pk,pk=None):
    board=get_object_or_404(scoped(request,Board),pk=board_pk)
    obj=get_object_or_404(board.groups,pk=pk) if pk else StaffingGroup(company=request.user.company,board=board,order=board.groups.count()+1)
    form=GroupForm(request.POST if request.method=='POST' else None,instance=obj)
    if request.method=='POST' and form.is_valid():
        form.save();return redirect('personnel:board_detail',pk=board.pk)
    return render(request,'personnel/form.html',{'form':form,'board':board,'title':'Штатная должность'})


@login_required
@require_http_methods(['GET','POST'])
def slot_edit(request,group_pk,pk=None):
    group=get_object_or_404(scoped(request,StaffingGroup).select_related('board'),pk=group_pk)
    obj=get_object_or_404(group.slots,pk=pk) if pk else StaffingSlot(company=request.user.company,group=group)
    form=SlotForm(request.POST if request.method=='POST' else None,instance=obj,board=group.board)
    if request.method=='POST' and form.is_valid():
        slot=form.save()
        Person.objects.filter(schedule_slot=slot).exclude(pk=slot.person_id).update(schedule_slot=None)
        return redirect('personnel:board_detail',pk=group.board_id)
    return render(request,'personnel/form.html',{'form':form,'board':group.board,'title':group.position+' — сотрудник / вакансия','slot':obj if obj.pk else None,'person_link':True})


@login_required
@require_http_methods(['GET','POST'])
def calendar_edit(request,pk):
    slot=get_object_or_404(scoped(request,StaffingSlot).select_related('group__board'),pk=pk)
    if request.method=='GET' and slot.person_id:
        from django.urls import reverse
        return redirect(reverse('personnel:person_schedule',args=[slot.person_id])+f'?slot={slot.pk}')
    form=CalendarForm(request.POST if request.method=='POST' else None)
    if request.method=='POST' and form.is_valid():
        with transaction.atomic():
            slot=StaffingSlot.objects.select_for_update().get(pk=slot.pk)
            day=form.cleaned_data['start']; code=form.cleaned_data['code']; calendar=dict(slot.calendar)
            while day<=form.cleaned_data['end']:
                if code: calendar[day.isoformat()]=code
                else: calendar.pop(day.isoformat(),None)
                day+=timedelta(days=1)
            slot.calendar=calendar;slot.save(update_fields=['calendar','updated_at'])
        return redirect('personnel:board_detail',pk=slot.group.board_id)
    return render(request,'personnel/form.html',{'form':form,'board':slot.group.board,'title':'Календарь: '+str(slot.person or 'Вакансия')})


@login_required
@require_http_methods(['GET','POST'])
def import_excel(request):
    access(request)
    form=ImportForm(request.POST if request.method=='POST' else None,request.FILES or None)
    preview=None
    if request.method=='POST':
        if request.POST.get('action')=='confirm':
            pending=request.session.get('personnel_import')
            if not pending or not secrets.compare_digest(str(pending.get('token','')),request.POST.get('token','')) or pending.get('company')!=request.user.company_id or not request.POST.get('confirm'):
                messages.error(request,'Предпросмотр устарел. Выберите файл заново и подтвердите загрузку.')
            else:
                try:
                    board,created=importer.apply(request.user.company,pending['name'],pending['data'])
                    request.session.pop('personnel_import',None)
                    messages.success(request,'Данные загружены.' if created else 'Этот файл уже загружен. Открыта существующая таблица.')
                    return redirect('personnel:board_detail',pk=board.pk)
                except ValidationError as exc: messages.error(request,'; '.join(exc.messages))
            form=ImportForm()
        elif form.is_valid():
            try:
                preview=importer.read(form.cleaned_data['file'])
                token=secrets.token_urlsafe(24)
                request.session['personnel_import']=dict(company=request.user.company_id,name=form.cleaned_data['name'],data=preview,token=token)
            except ValidationError as exc: form.add_error('file','; '.join(exc.messages))
    pending=request.session.get('personnel_import',{})
    return render(request,'personnel/import.html',{'form':form,'preview':preview,'token':pending.get('token','')})


@login_required
def export(request,pk):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    board=get_object_or_404(scoped(request,Board),pk=pk)
    data=context(request,board)
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Нефтепровод'
    ws.append(['Проект','Объект','Должность','План явки','План штата','ФИО','ВЖГ','Телефон','Табельный №','Вакансии','Трудоустроен','Вид графика','График','Текущий статус','Выезд','Заезд','Дней МО','Дней МВ','Следующий рабочий день','Выезд по графику','Комментарии']+data['dates'])
    def safe(value):
        return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value
    for group in data['groups']:
        # Export the same filtered rows; summary is explicitly the full board.
        for index,row in enumerate(group['rows']):
            s=row['slot'];p=s.person
            values=[board.project,board.site,group['group'].position,group['group'].attendance_plan if index==0 else None,group['group'].staffing_plan if index==0 else None,p.full_name if p else 'Вакансия',s.camp,p.phone if p else '',p.employee_number if p else '',int(not s.person_id),int(bool(s.person_id)),s.pattern.kind if s.pattern_id else '',s.pattern.name if s.pattern_id else '',s.status,s.departure_date,s.arrival_date,row['elapsed'] if s.status=='МО' else None,row['elapsed'] if s.status=='МВ' else None,row['next_work'],row['next_leave'],s.comment]+row['codes']
            ws.append([safe(v) for v in values])
    ws.freeze_panes='F2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='174A7E')
    for row in ws:
        for c in row:
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(c.value,date):c.number_format='DD.MM.YYYY'
    for col in ['A','B','C','F','U']:ws.column_dimensions[col].width=32
    summary=wb.create_sheet('Сводка всего объекта');summary.append(['Должность','План явки','План штата','Трудоустроено','Явных вакансий','Дефицит штата','МВ','МО','Не уточнено'])
    for g in data['groups']:
        summary.append([safe(g['group'].position)]+[g['stats'][k] for k in ['attendance','staffing','employed','vacancies','shortage','on','off','unknown']])
    summary.append(['Итого']+[data['total'][k] for k in ['attendance','staffing','employed','vacancies','shortage','on','off','unknown']])
    summary.append(['Фактический статус — текущий. Пустая дата календаря означает отсутствие графика. В — выходной, ДП — день пути.'])
    out=BytesIO();wb.save(out)
    response=HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']='attachment; filename="personnel.xlsx"';return response


@login_required
@require_http_methods(['GET','POST'])
def pattern_create(request,board_pk):
    from .forms import PatternForm
    from .models import ShiftPattern
    board=get_object_or_404(scoped(request,Board),pk=board_pk)
    form=PatternForm(request.POST if request.method=='POST' else None)
    if request.method=='POST' and form.is_valid():
        d=form.cleaned_data;day=d['start'];calendar={};cycle=d['on_days']+d['off_days']
        while day<=d['end']:
            calendar[day.isoformat()]=str(d['hours']) if (day-d['anchor']).days%cycle<d['on_days'] else 'МО'
            day+=timedelta(days=1)
        ShiftPattern.objects.create(company=request.user.company,board=board,name=d['name'],kind=d['kind'],calendar=calendar)
        messages.success(request,'График создан. Выберите его в строке сотрудника; выходные и дни пути можно уточнить через «Изменить дни».')
        return redirect('personnel:board_detail',pk=board.pk)
    return render(request,'personnel/form.html',{'form':form,'board':board,'title':'Новый график перевахты: цикл вахта / отдых'})
