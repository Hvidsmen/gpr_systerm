from datetime import date
from io import BytesIO
from unittest.mock import patch
import openpyxl
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from apps.accounts.models import Company, Role, User
from .models import Board, Person, StaffingGroup, StaffingSlot, ShiftPattern
from .importer import read, apply
from .services import report


def workbook():
    wb=openpyxl.Workbook();s=wb.active;s.title='Нефтепровод'
    s.cell(5,7,'Наименование должности');s.cell(5,8,'Ф.И.О')
    s.merge_cells('D12:D14');s.merge_cells('E12:E14');s.cell(12,4,2);s.cell(12,5,3)
    for row,name in [(12,'Иванов Иван'),(13,'Петров Петр'),(14,'1вакансия')]:
        for col,value in [(2,'Нефтепровод'),(6,1),(7,'Механик'),(8,name),(9,'ВЖГ-1')]:s.cell(row,col,value)
    s.cell(12,15,'М');s.cell(12,16,'Вахта 1 (60/30)');s.cell(12,17,'МВ');s.cell(12,19,date(2026,10,1))
    s.cell(13,17,'МО');s.cell(13,18,date(2026,10,3));s.cell(13,6,None)
    s.cell(15,7,'Инженер');s.cell(15,6,1);s.cell(15,4,1);s.cell(15,5,1)
    s.cell(30,6,'________________');s.cell(30,7,'/________________/');s.cell(30,8,'Подпись')
    g=wb.create_sheet('Графики');g.cell(2,2,'М');g.cell(2,4,'Вахта 1 (60/30)')
    for c,day,code in [(5,1,10),(6,2,'В'),(7,3,'ДП'),(8,4,'МО')]:g.cell(1,c,date(2026,10,day));g.cell(2,c,code)
    out=BytesIO();wb.save(out);return out.getvalue()


class PersonnelTests(TestCase):
    def setUp(self):
        self.company=Company.objects.create(name='Personnel')
        self.other=Company.objects.create(name='Other Personnel')
        self.user=User.objects.create_user(username='hr',company=self.company,role=Role.objects.get(code='HR_HEAD'))
        self.client.force_login(self.user)
        self.data=read(BytesIO(workbook()))
        self.board,_=apply(self.company,'Нефтепровод',self.data)

    def test_import_and_totals_match_source_without_footer_or_duplicated_merged_plans(self):
        self.assertEqual((self.data['employed'],self.data['vacancies']),(2,2))
        self.assertEqual(Person.objects.filter(company=self.company).count(),2)
        self.assertEqual(self.board.groups.count(),2)
        self.assertEqual(StaffingSlot.objects.filter(group__board=self.board).count(),4)
        with patch('apps.personnel.services.timezone.localdate',return_value=date(2026,10,10)):
            d=report(self.board,date(2026,10,1))
        self.assertEqual(d['total'],dict(attendance=3,staffing=4,employed=2,vacancies=2,on=1,off=1,unknown=0,shortage=2))
        self.assertEqual(d['people_count'],2)
        self.assertEqual(d['daily'][:4],[1,0,0,0])
        self.assertEqual(d['groups'][0]['rows'][0]['elapsed'],9)
        self.assertEqual(d['groups'][0]['rows'][1]['elapsed'],7)
        d=report(self.board,date(2026,10,1),query='Иванов')
        self.assertEqual(d['total']['employed'],2)
        self.assertEqual(sum(len(g['rows']) for g in d['groups']),1)
        board,created=apply(self.company,'Another name',self.data)
        self.assertFalse(created);self.assertEqual(board.pk,self.board.pk)
        self.assertEqual(Board.objects.count(),1)

    def test_import_preview_requires_confirmation_and_is_idempotent(self):
        url=reverse('personnel:import')
        response=self.client.post(url,{'name':'Upload','file':SimpleUploadedFile('people.xlsx',workbook())})
        self.assertEqual(response.status_code,200);self.assertContains(response,'Подтверждаю загрузку')
        token=self.client.session['personnel_import']['token']
        self.client.post(url,{'action':'confirm','token':'wrong','confirm':'on'})
        self.assertEqual(Board.objects.count(),1)
        response=self.client.post(url,{'action':'confirm','token':token,'confirm':'on'})
        self.assertRedirects(response,reverse('personnel:board_detail',args=[self.board.pk]))
        self.assertNotIn('personnel_import',self.client.session)

    def test_pages_export_and_calendar_changes(self):
        for name,args in [('board_list',[]),('people',[]),('board_detail',[self.board.pk]),('board_create',[]),('person_create',[]),('group_create',[self.board.pk]),('pattern_create',[self.board.pk])]:
            response=self.client.get(reverse('personnel:'+name,args=args));self.assertEqual(response.status_code,200,name)
        slot=StaffingSlot.objects.filter(group__board=self.board,person__isnull=False).first()
        response=self.client.post(reverse('personnel:calendar_edit',args=[slot.pk]),{'start':'2026-10-01','end':'2026-10-02','code':'МО'})
        self.assertEqual(response.status_code,302);slot.refresh_from_db();self.assertEqual(slot.calendar,{'2026-10-01':'МО','2026-10-02':'МО'})
        self.client.post(reverse('personnel:calendar_edit',args=[slot.pk]),{'start':'2026-10-01','end':'2026-10-02','code':''})
        slot.refresh_from_db();self.assertEqual(slot.calendar,{})
        person=slot.person;person.full_name='=HYPERLINK("x")';person.save()
        response=self.client.get(reverse('personnel:export',args=[self.board.pk]),{'month':'2026-10'})
        self.assertEqual(response.status_code,200)
        w=openpyxl.load_workbook(BytesIO(response.content))
        self.assertTrue(w['Нефтепровод']['F2'].value.startswith("'="))
        self.assertEqual(w['Сводка всего объекта']['D4'].value,2)

    def test_access_is_limited_to_hr_admin_and_same_company(self):
        other_board=Board.objects.create(company=self.other,name='Secret')
        self.assertEqual(self.client.get(reverse('personnel:board_detail',args=[other_board.pk])).status_code,404)
        for code in ['PLANNER','DISPATCHER','ESTIMATOR','OBSERVER','CEO']:
            self.user.role=Role.objects.get(code=code);self.user.save()
            for name,args in [('board_list',[]),('people',[]),('board_detail',[self.board.pk]),('export',[self.board.pk]),('import',[])]:
                self.assertEqual(self.client.get(reverse('personnel:'+name,args=args)).status_code,403,code+name)
        self.user.role=Role.objects.get(code='ADMIN');self.user.save()
        self.assertEqual(self.client.get(reverse('personnel:board_list')).status_code,200)

    def test_forms_reject_other_company_person_and_vacancy_fact_status(self):
        group=self.board.groups.first()
        foreign=Person.objects.create(company=self.other,full_name='Foreign')
        url=reverse('personnel:slot_create',args=[group.pk])
        response=self.client.post(url,{'person':foreign.pk,'status':'МВ'})
        self.assertEqual(response.status_code,200);self.assertIn('person',response.context['form'].errors)
        response=self.client.post(url,{'status':'МВ'})
        self.assertEqual(response.status_code,200);self.assertTrue(response.context['form'].errors)
        self.assertEqual(StaffingSlot.objects.filter(group__board=self.board).count(),4)

    def test_manual_employee_group_slot_and_pattern_workflow(self):
        response=self.client.post(reverse('personnel:person_create'),{'full_name':'Новый Сотрудник','employee_number':'123','phone':'555'})
        self.assertEqual(response.status_code,302)
        person=Person.objects.get(employee_number='123')
        self.client.post(reverse('personnel:group_create',args=[self.board.pk]),{'position':'Сварщик','attendance_plan':1,'staffing_plan':2})
        group=self.board.groups.get(position='Сварщик')
        self.client.post(reverse('personnel:pattern_create',args=[self.board.pk]),dict(name='Цикл',kind='М',anchor='2026-10-01',start='2026-10-01',end='2026-10-05',on_days=2,off_days=2,hours=10))
        pattern=self.board.patterns.get(name='Цикл')
        self.assertEqual([pattern.calendar[f'2026-10-0{d}'] for d in range(1,6)],['10','10','МО','МО','10'])
        response=self.client.post(reverse('personnel:slot_create',args=[group.pk]),dict(person=person.pk,status='МВ',arrival_date='2026-10-01',pattern=pattern.pk))
        self.assertEqual(response.status_code,302)
        self.assertEqual(group.slots.get().person_id,person.pk)

    def test_bad_workbook_and_duplicate_employee_numbers_rollback(self):
        with self.assertRaises(ValidationError):read(BytesIO(b'not a workbook'))
        modified={**self.data,'digest':'another','rows':[dict(r) for r in self.data['rows']]}
        modified['rows'][0]['employee_number']='same'
        modified['rows'][1]['employee_number']='same'
        before=Person.objects.count()
        with self.assertRaises(ValidationError):apply(self.company,'Broken',modified)
        self.assertEqual(Board.objects.count(),1);self.assertEqual(Person.objects.count(),before)

    def test_csrf_and_company_validation(self):
        from django.test import Client
        client=Client(enforce_csrf_checks=True);client.force_login(self.user)
        response=client.post(reverse('personnel:person_create'),{'full_name':'CSRF'})
        self.assertEqual(response.status_code,403)
        person=Person.objects.create(company=self.other,full_name='Other')
        with self.assertRaises(ValidationError):StaffingSlot.objects.create(company=self.company,group=self.board.groups.first(),person=person)

    def test_departure_uses_outgoing_travel_day_and_missing_calendar_stays_empty(self):
        from .services import departure_days
        calendar={'2026-10-01':'МО','2026-10-02':'ДП','2026-10-03':'10','2026-10-04':'В','2026-10-05':'ДП','2026-10-06':'ДП','2026-10-07':'МО'}
        self.assertEqual(departure_days(calendar,date(2026,10,1)),['2026-10-05'])
        slot=StaffingSlot.objects.filter(group__board=self.board,person__isnull=False).first()
        slot.calendar={'2026-10-01':'МО'};slot.save()
        d=report(self.board,date(2027,1,1))
        self.assertEqual(d['groups'][0]['rows'][0]['codes'],['']*31)
        self.assertEqual(d['daily'],[0]*31)
