"""Read only personnel fields and schedule values, never execute workbook formulas."""
from collections import Counter
import hashlib
import io
import re
from datetime import date, datetime
from zipfile import ZipFile, BadZipFile
from django.core.exceptions import ValidationError
from django.db import transaction
import openpyxl
from .models import Board, Person, StaffingGroup, StaffingSlot, ShiftPattern


def text(value):
    if value is None: return ''
    return str(value).strip()


def read(upload):
    raw = upload.read(12 * 1024 * 1024 + 1)
    if len(raw) > 12 * 1024 * 1024: raise ValidationError('Файл больше 12 МБ.')
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024: raise ValidationError('Слишком большой распакованный файл.')
        workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    except (BadZipFile, ValueError, OSError, KeyError) as exc: raise ValidationError('Не удалось прочитать файл .xlsx.') from exc
    if 'Нефтепровод' not in workbook.sheetnames: raise ValidationError('В файле нет листа «Нефтепровод».')
    sheet = workbook['Нефтепровод']
    if sheet.max_row > 5000 or sheet.max_column > 1000: raise ValidationError('Слишком большая таблица.')
    if 'должност' not in text(sheet.cell(5,7).value).lower() or 'ф.' not in text(sheet.cell(5,8).value).lower(): raise ValidationError('Не найдены ожидаемые заголовки должности и ФИО в строке 5.')
    merged = {}
    for area in sheet.merged_cells.ranges:
        if area.min_col <= 5 and area.min_row >= 12:
            for col in range(max(2,area.min_col), min(5,area.max_col)+1):
                for row in range(area.min_row, area.max_row+1): merged[row,col] = (area.min_row, sheet.cell(area.min_row,area.min_col).value)
    def cell(row,col): return merged.get((row,col), (row,sheet.cell(row,col).value))[1]
    def number(value, row):
        if value is None or value == '': return 0
        try:
            result = int(value)
            if float(value) != result or result < 0: raise ValueError
            return result
        except (ValueError, TypeError, OverflowError): raise ValidationError(f'Строка {row}: некорректная численность {value}.')
    patterns = []
    if 'Графики' in workbook.sheetnames:
        schedule = workbook['Графики']
        if schedule.max_row > 1000 or schedule.max_column > 1000: raise ValidationError('Слишком большой лист графиков.')
        date_columns = [(c, schedule.cell(1,c).value.date()) for c in range(5,schedule.max_column+1) if isinstance(schedule.cell(1,c).value,datetime)]
        for row in range(2,schedule.max_row+1):
            name, kind = text(schedule.cell(row,4).value), text(schedule.cell(row,2).value)
            if not name: continue
            calendar = {day.isoformat(): text(schedule.cell(row,col).value) for col,day in date_columns if valid_code(schedule.cell(row,col).value)}
            patterns.append(dict(name=name,kind=kind,calendar=calendar))
    groups = {}; warnings = []
    rows = []
    for row in range(12,sheet.max_row+1):
        position = text(sheet.cell(row,7).value); name = text(sheet.cell(row,8).value)
        # Empty formatted rows, signatures and section headings are not employees.
        quantity = sheet.cell(row,6).value
        if not position or (not name and not isinstance(quantity, (int,float))): continue
        if quantity is not None and not isinstance(quantity, (int,float)): continue
        if isinstance(quantity, (int,float)) and quantity <= 0: continue
        anchor = merged.get((row,5),(row,None))[0]
        key = (anchor,position)
        if key not in groups:
            groups[key] = dict(key=str(anchor)+':'+position, position=position,attendance=number(cell(row,4),row),staffing=number(cell(row,5),row),order=row)
        vacancy = not name or bool(re.search(r'ваканси',name,re.I))
        if not vacancy and name.startswith(('#','=')): raise ValidationError(f'Строка {row}: вместо ФИО ошибка или формула без результата.')
        status = text(sheet.cell(row,17).value)
        original_status = status
        if status not in {'','МВ','МО'}:
            warnings.append(f'Строка {row}: статус «{status}» сохранен в комментарии.'); status = 'OTHER'
        def dt(col):
            value = sheet.cell(row,col).value
            if value in (None,''): return None
            if isinstance(value,datetime):return value.date().isoformat()
            if isinstance(value,date):return value.isoformat()
            raise ValidationError(f'Строка {row}: дата в колонке {col} не распознана.')
        rows.append(dict(row=row,position=position,group=groups[key]['key'],name='' if vacancy else name,phone=text(sheet.cell(row,10).value),employee_number=text(sheet.cell(row,11).value),camp=text(sheet.cell(row,9).value),status='' if vacancy else status,arrival=dt(19),departure=dt(18),kind=text(sheet.cell(row,15).value),pattern=text(sheet.cell(row,16).value),comment=(text(sheet.cell(row,24).value) + ('; Исходный статус: '+original_status if status=='OTHER' else '')).strip('; ')))
    if not rows: raise ValidationError('Сотрудники и вакансии на листе не найдены.')
    empty_rows = [r['row'] for r in rows if not r['name'] and not text(sheet.cell(r['row'],8).value)]
    if empty_rows: warnings.append('Пустые штатные строки считаются вакансиями: '+', '.join(map(str,empty_rows))+'.')
    duplicates = [name for name,count in Counter(r['name'] for r in rows if r['name']).items() if count>1]
    if duplicates: warnings.append('Повторные ФИО без различающих табельных номеров используют одну карточку сотрудника и отдельные штатные позиции: '+', '.join(duplicates)+'.')
    unknown = {(r['kind'],r['pattern']) for r in rows if r['pattern']} - {(p['kind'],p['name']) for p in patterns}
    warnings += [f'График {k} {n} не найден: календарь останется пустым.' for k,n in sorted(unknown)]
    return dict(digest=hashlib.sha256(raw).hexdigest(),groups=list(groups.values()),rows=rows,patterns=patterns,warnings=warnings,employed=sum(bool(r['name']) for r in rows),vacancies=sum(not r['name'] for r in rows))


def valid_code(value):
    if value in {'МВ','МО','В','ДП'}: return True
    try: return value is not None and 0 <= float(value) <= 24
    except (TypeError,ValueError): return False


@transaction.atomic
def apply(company, name, data):
    # Serialize same-company imports; the same workbook is never duplicated.
    from apps.accounts.models import Company
    Company.objects.select_for_update().get(pk=company.pk)
    previous = Board.objects.filter(company=company,import_digest=data['digest']).first()
    if previous: return previous,False
    board = Board.objects.create(company=company,name=name,project='Нефтепровод',import_digest=data['digest'])
    patterns = {}
    for item in data['patterns']:
        patterns[item['kind'],item['name']] = ShiftPattern.objects.create(company=company,board=board,**item)
    groups = {item['key']:StaffingGroup.objects.create(company=company,board=board,position=item['position'],attendance_plan=item['attendance'],staffing_plan=item['staffing'],order=item['order']) for item in data['groups']}
    for row in data['rows']:
        person = None
        if row['name']:
            people = Person.objects.filter(company=company)
            if row['employee_number']: people=people.filter(employee_number=row['employee_number'])
            else: people=people.filter(full_name=row['name'])
            matches = list(people[:2])
            if len(matches)>1: raise ValidationError(f"Строка {row['row']}: несколько сотрудников с одинаковым ФИО. Укажите табельный номер.")
            person = matches[0] if matches else Person.objects.create(company=company,full_name=row['name'],phone=row['phone'],employee_number=row['employee_number'])
            if matches and row['employee_number'] and person.full_name != row['name']: raise ValidationError(f"Строка {row['row']}: табельный номер уже принадлежит другому сотруднику.")
        StaffingSlot.objects.create(company=company,group=groups[row['group']],person=person,camp=row['camp'],status=row['status'],arrival_date=row['arrival'],departure_date=row['departure'],pattern=patterns.get((row['kind'],row['pattern'])),comment=row['comment'],source_row=row['row'])
    from .housing import bootstrap_legacy
    bootstrap_legacy(company,StaffingSlot.objects.filter(group__board=board))
    return board,True
