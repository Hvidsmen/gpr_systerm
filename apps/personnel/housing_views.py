from calendar import monthrange
from datetime import date,timedelta
from io import BytesIO
from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError,PermissionDenied
from core.fact_edit_window import can_edit_fact_day,require_fact_day,window_context
from django.db import transaction,IntegrityError
from django.db.models import Q
from django.http import HttpResponse,Http404
from django.shortcuts import get_object_or_404,redirect,render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_http_methods
from .views import access,scoped
from .models import Camp,CampObjectLink,CampStay,HousingEvent,Person,Board,StaffingGroup,StaffingSlot
from .forms import CampForm,MovementForm,PersonScheduleForm
from . import housing
from .services import calendar_for,departure_days


def selected_day(request):
    try:return forms.DateField().clean((request.POST if request.method=='POST' else request.GET).get('date') or timezone.localdate().isoformat())
    except ValidationError:return timezone.localdate()


def camp_url(camp,day):return reverse('personnel:housing')+f'?camp={camp.pk}&date={day.isoformat()}'


@login_required
def camps(request):
    from core.permissions import require_roles,READ_ROLES
    require_roles(request.user,READ_ROLES|{'HR_HEAD','PERSONNEL_DISPATCHER','MECHANIZATION_DISPATCHER','ALL_FACT_DISPATCHER'})
    qs=Camp.objects.filter(company=request.user.company).prefetch_related('objects_served')
    query=request.GET.get('q','').strip()
    if query:qs=qs.filter(Q(name__icontains=query)|Q(address__icontains=query)|Q(objects_served__name__icontains=query)).distinct()
    if request.GET.get('export')=='xlsx':
        return workbook_response('camps.xlsx','ВЖГ',[['ВЖГ','Расположение','Действующий','Связанные объекты'],*[[c.name,c.address,'Да' if c.is_active else 'Нет','; '.join(o.name for o in c.objects_served.all())] for c in qs]])
    return render(request,'personnel/camps.html',{'camps':qs,'query':query})


@login_required
@require_http_methods(['GET','POST'])
def camp_edit(request,pk=None):
    access(request)
    obj=get_object_or_404(scoped(request,Camp),pk=pk) if pk else Camp(company=request.user.company)
    form=CampForm(request.POST if request.method=='POST' else None,company=request.user.company,instance=obj)
    if request.method=='POST' and form.is_valid():
        try:
            with transaction.atomic():
                from apps.accounts.models import Company
                from django.db.models import F
                Company.objects.filter(pk=request.user.company_id).update(name=F('name'))
                # Match production's object-before-camp lock order, including FK checks.
                from apps.projects.models import ConstructionObject
                object_ids={row.pk for row in form.cleaned_data['objects_served']}
                if obj.pk:object_ids.update(obj.objects_served.values_list('pk',flat=True))
                list(ConstructionObject.objects.select_for_update().filter(company=request.user.company,pk__in=object_ids).order_by('pk'))
                if obj.pk:
                    current=Camp.objects.select_for_update().get(pk=obj.pk)
                    if current.updated_at.isoformat()!=form.cleaned_data['version']:raise ValidationError('ВЖГ или связи уже изменены другим пользователем. Обновите форму.')
                saved=form.save();selected={o.pk for o in form.cleaned_data['objects_served']}
                saved.object_links.exclude(construction_object_id__in=selected).delete()
                for object_id in selected:
                    CampObjectLink.objects.get_or_create(company=request.user.company,camp=saved,construction_object_id=object_id)
        except IntegrityError:form.add_error('name','Такой ВЖГ уже существует. Обновите справочник.')
        except ValidationError as exc:form.add_error(None,'; '.join(exc.messages))
        else:
            messages.success(request,'ВЖГ и связи с объектами сохранены.');return redirect('personnel:camps')
    return render(request,'personnel/form.html',{'form':form,'title':'Изменить ВЖГ' if pk else 'Новый ВЖГ'})


def roster_rows(company,camp,day,query='',position=''):
    stays=housing.present_stays(company,day).filter(camp=camp).select_related('person__employee__position__planning_mapping__brigade__group','person__schedule_slot__pattern','person__schedule_slot__group__board','recorded_by')
    if query:stays=stays.filter(Q(person__full_name__icontains=query)|Q(person__employee_number__icontains=query)|Q(person__employee__employee_number__icontains=query)|Q(person__employee__zup_full_name__icontains=query)|Q(person__employee__position__name__icontains=query))
    if position:stays=stays.filter(person__employee__position_id=position)
    days=[day.replace(day=n) for n in range(1,monthrange(day.year,day.month)[1]+1)]
    rows=[]
    for stay in stays.order_by('person__employee__position__name','person__full_name','pk'):
        person=stay.person;slot=housing.schedule_slot(person);calendar=calendar_for(slot,start=day.replace(day=1),end=day+timedelta(days=730)) if slot else {}
        planned_departure=next(iter(departure_days(calendar,stay.arrived_on)),'')
        departure_date=date.fromisoformat(planned_departure) if planned_departure else None
        until_departure=(departure_date-day).days if departure_date else None
        rows.append({'stay':stay,'person':person,'position':person.employee.position if person.employee_id else None,'slot':slot,
            'codes':[calendar.get(d.isoformat(),'') for d in days],'today_code':calendar.get(day.isoformat(),''),
            'planned_departure':planned_departure,'planned_departure_display':departure_date,'departure_days':until_departure,
            'departure_warning':('overdue' if until_departure<0 else 'soon' if until_departure<=14 else '') if until_departure is not None else '',
            'position_label':str(person.employee.position) if person.employee_id and person.employee.position_id else slot.group.position if slot else 'Должность не указана'})
    from apps.resources.group_order import position_order
    rows.sort(key=lambda row:(position_order(row['position'],row['position_label']),row['person'].full_name.casefold()))
    return rows,days


@login_required
@require_http_methods(['GET','POST'])
def workspace(request):
    access(request);day=selected_day(request)
    from .rotation import RotationFilter,rotation_rows
    raw=(request.GET if request.method=='GET' else request.POST).copy()
    raw.setdefault('date',str(day))
    if raw.get('camp') and not raw.getlist('camps'):raw.setlist('camps',[raw['camp']])
    selector=RotationFilter(raw,company=request.user.company)
    if not selector.is_valid():
        if raw.get('camp'):raise Http404('ВЖГ не найден.')
        return render(request,'personnel/housing.html',{'rotation_filter':selector,'day':day},status=400)
    selected_camps=selector.selected_camps(request.user.company)
    camps=scoped(request,Camp).order_by('name')
    camp_id=(request.POST if request.method=='POST' else request.GET).get('camp')
    camp=get_object_or_404(camps,pk=camp_id) if camp_id else selected_camps.first()
    if camp is None:return render(request,'personnel/housing.html',{'rotation_filter':selector,'camps':camps,'day':day,'no_camp_match':True})
    if request.method=='POST':
        try:require_fact_day(request.user,day)
        except ValidationError as exc:raise PermissionDenied('; '.join(exc.messages))
    kind=request.GET.get('action','ARRIVAL');kind=kind if kind in {'ARRIVAL','DEPARTURE','TRANSFER','DISMISSAL'} else 'ARRIVAL'
    selected_employee=request.GET.get('employee','')
    if not selected_employee.isdigit() and request.GET.get('person','').isdigit():
        selected_employee=str(Person.objects.filter(company=request.user.company,pk=request.GET['person']).values_list('employee_id',flat=True).first() or '')
    from django.core import signing
    from apps.resources.models import Employee
    personal_token=request.POST.get('personal_scope','') if request.method=='POST' else ''
    if personal_token:
        try:
            personal=signing.loads(personal_token,salt='personnel-single-movement',max_age=12*3600)
            if personal.get('user')!=request.user.pk or personal.get('company')!=request.user.company_id:raise signing.BadSignature
            selected_employee=str(personal['employee'])
        except (signing.BadSignature,KeyError,TypeError):raise Http404('Персональная форма перемещения устарела.')
    fixed_employee=int(selected_employee) if selected_employee.isdigit() else None
    if fixed_employee:
        get_object_or_404(Employee,company=request.user.company,pk=fixed_employee)
        personal_token=signing.dumps({'user':request.user.pk,'company':request.user.company_id,'employee':fixed_employee},salt='personnel-single-movement')
    form=MovementForm(request.POST if request.method=='POST' else None,company=request.user.company,camp=camp,day=day,fixed_employee=fixed_employee,
        initial={'personal_scope':personal_token,'kind':kind,'date':day,'camp':None if request.GET.get('choose_camp')=='1' and kind=='ARRIVAL' else camp.pk,'snapshot':housing.movement_snapshot(request.user),'employees':[selected_employee] if selected_employee.isdigit() else []})
    if request.method=='POST' and form.is_valid():
        try:
            selected=form.cleaned_data['employees'];arrival=form.cleaned_data['kind']=='ARRIVAL'
            people=[] if arrival else list(Person.objects.filter(company=request.user.company,employee__in=selected))
            count=housing.move(request.user,form.cleaned_data['kind'],camp,form.cleaned_data['date'],people,
                employees=selected if arrival else (),target=form.cleaned_data.get('target'),comment=form.cleaned_data['comment'],snapshot=form.cleaned_data['snapshot'],departure_destination='DISMISSAL' if form.cleaned_data.get('destination')=='DISMISSAL' else 'REST')
        except (ValidationError,IntegrityError) as exc:form.add_error(None,'; '.join(exc.messages) if isinstance(exc,ValidationError) else 'Размещение изменилось одновременно с вашим вводом. Обновите список.')
        else:
            messages.success(request,f'Перемещение сохранено. Сотрудников: {count}.')
            from .rotation_calendar import cycle_definition
            missing=0
            for person in Person.objects.filter(company=request.user.company,employee__in=selected).select_related('schedule_slot__pattern'):
                assigned=housing.schedule_slot(person)
                if not (assigned and assigned.pattern_id and cycle_definition(assigned.pattern)):missing+=1
            if missing:messages.warning(request,f'У {missing} сотрудников цикл вахта/отдых не задан. Укажите его через «Изменить график»; прогноз для них пока не пересчитан.')
            return redirect(camp_url(camp,form.cleaned_data['date'])+'#roster')
    values=selector.cleaned_data
    rows,days,rotation_stats=rotation_rows(request.user.company,list(selected_camps),day,query=values.get('q',''),state=values.get('state',''),include_unassigned=not values.get('construction_object') and not values.get('camps'))
    from .rotation import positions_in_rows,filter_position,rotation_stats as count_rotation
    available=positions_in_rows(request.user.company,rows)
    selector.fields['position'].queryset=available
    if values.get('position') and not available.filter(pk=values['position'].pk).exists():
        values['position']=None;selector.data=selector.data.copy();selector.data['position']=''
        messages.info(request,'Фильтр должности сброшен: её нет в текущей выборке.')
    rows=filter_position(rows,values.get('position'));rotation_stats=count_rotation(rows)
    from .rotation import status_indicator
    for row in rows:row['status_tone'],row['status_caption'],row['status_help']=status_indicator(row)
    # The calendar is a forecast through the end of the following year.
    calendar_start=day.replace(day=1);calendar_end=date(day.year+1,12,31)
    calendar_dates=[calendar_start+timedelta(days=n) for n in range((calendar_end-calendar_start).days+1)]
    calendar_months=[]
    for current in calendar_dates:
        if not calendar_months or calendar_months[-1]['start'].month!=current.month:
            calendar_months.append({'start':current,'span':0})
        calendar_months[-1]['span']+=1
    from .rotation_calendar import stays_by_person
    history=stays_by_person(request.user.company)
    for row in rows:
        values_calendar=calendar_for(row['slot'],start=calendar_start,end=calendar_end,stays=history.get(row['person'].pk,[]) if row.get('person') else []) if row.get('slot') else {}
        row['calendar_codes']=[values_calendar.get(str(current),'') for current in calendar_dates]
    events=HousingEvent.objects.filter(company=request.user.company).filter(Q(from_camp__in=selected_camps)|Q(to_camp__in=selected_camps)).select_related('person','from_camp','to_camp','recorded_by')[:50]
    from .housing_plan import monthly_plan,position_headers
    plan_context=monthly_plan(request.user,camp,day,camps=selected_camps)
    position_groups=position_headers(request.user.company,rows,plan_context,query=values.get('q',''),position=values.get('position'),state=values.get('state',''))
    from .housing_plan import roster_sections
    roster_view='groups' if raw.get('roster_view')=='groups' else 'positions'
    view_links=[]
    for value,label in [('positions','По должностям'),('groups','По группам должностей')]:
        params=request.GET.copy();params['roster_view']=value
        view_links.append({'value':value,'label':label,'url':'?'+params.urlencode()+'#roster'})
    from apps.projects.models import ConstructionObject
    objects=ConstructionObject.objects.filter(company=request.user.company,housing_camps__in=selected_camps).distinct()
    from .lodging_views import summary_context
    plan_context.update(summary_context(request.user,camp,day,camps=selected_camps))
    return render(request,'personnel/housing.html',{**plan_context,'rotation_filter':selector,'rotation_stats':rotation_stats,'position_groups':position_groups,'selected_camps':selected_camps,
        'roster_view':roster_view,'roster_view_links':view_links,'roster_sections':roster_sections(position_groups,roster_view=='groups',request.user.company),
        'camp':camp,'camps':camps,'day':day,'rows':rows,'dates':days,'form':form,'query':values.get('q',''),
        'calendar_dates':calendar_dates,'calendar_months':calendar_months,'calendar_start':calendar_start,'calendar_end':calendar_end,'events':events,'total':rotation_stats['people'],'objects':objects,
        'unknown_count':Person.objects.filter(company=request.user.company,stays__isnull=True).count(),
        'unlinked_count':Person.objects.filter(company=request.user.company,employee=None).count(),**window_context(request.user,day)})



@login_required
def history(request,pk):
    person=get_object_or_404(scoped(request,Person),pk=pk)
    stays=person.stays.select_related('camp','recorded_by','cancelled_by')
    for stay in stays:stay.can_cancel=can_edit_fact_day(request.user,stay.arrived_on)
    events=person.housing_events.select_related('from_camp','to_camp','recorded_by')
    return render(request,'personnel/housing_history.html',{'person':person,'stays':stays,'events':events})


@login_required
@require_http_methods(['GET','POST'])
def cancel(request,pk):
    stay=get_object_or_404(scoped(request,CampStay).select_related('person','camp'),pk=pk)
    try:require_fact_day(request.user,stay.arrived_on)
    except ValidationError as exc:raise PermissionDenied('; '.join(exc.messages))
    class CancelForm(forms.Form):
        reason=forms.CharField(label='Причина отмены',max_length=500,widget=forms.Textarea(attrs={'class':'form-control','rows':3}))
        version=forms.CharField(widget=forms.HiddenInput)
    form=CancelForm(request.POST if request.method=='POST' else None,initial={'version':stay.updated_at.isoformat()})
    if request.method=='POST' and form.is_valid():
        try:housing.cancel_stay(request.user,stay,form.cleaned_data['reason'],form.cleaned_data['version'])
        except ValidationError as exc:form.add_error(None,'; '.join(exc.messages))
        else:messages.success(request,'Размещение отменено. История сохранена.');return redirect('personnel:housing_history',pk=stay.person_id)
    return render(request,'personnel/form.html',{'form':form,'title':f'Отменить размещение: {stay}'})


@login_required
@require_http_methods(['GET','POST'])
def schedule(request,pk):
    person=get_object_or_404(scoped(request,Person).select_related('employee__position','schedule_slot'),pk=pk)
    form=PersonScheduleForm(request.POST if request.method=='POST' else None,person=person)
    if request.method=='GET' and request.GET.get('slot','').isdigit():
        selected=person.assignments.filter(company=person.company,pk=request.GET['slot']).first()
        if selected:form.initial.update(slot=selected.pk,pattern=selected.pattern_id)
    if request.method=='POST' and form.is_valid():
        with transaction.atomic():
            from apps.accounts.models import Company
            from django.db.models import F
            Company.objects.filter(pk=request.user.company_id).update(name=F('name'))
            person=Person.objects.select_for_update().get(pk=person.pk,company=request.user.company)
            slot=form.cleaned_data.get('slot');pattern=form.cleaned_data.get('pattern')
            if slot:
                slot=StaffingSlot.objects.select_for_update().get(pk=slot.pk,person=person,company=person.company)
                if pattern and pattern.board_id!=slot.group.board_id:
                    form.add_error('pattern','График должен принадлежать таблице выбранной строки. Выберите другую строку или очистите выбор строки.')
            if not form.errors:
                try:housing.check_schedule(person,form.cleaned_data['version'],slot)
                except ValidationError as exc:form.add_error(None,'; '.join(exc.messages))
            if not form.errors:
                if slot and slot.group.board.import_digest:
                    # Personal changes must not rewrite the imported opening balance.
                    from .models import ShiftPattern
                    board=Board.objects.filter(company=person.company,name='Графики сотрудников ВЖГ',import_digest='').order_by('pk').first()
                    if board is None:board=Board.objects.create(company=person.company,name='Графики сотрудников ВЖГ')
                    position=str(person.employee.position) if person.employee_id and person.employee.position_id else slot.group.position
                    group=board.groups.filter(position=position).first() or StaffingGroup.objects.create(company=person.company,board=board,position=position)
                    original=slot
                    if pattern and not form.cleaned_data.get('on_days'):pattern=ShiftPattern.objects.create(company=person.company,board=board,name=pattern.name,kind=pattern.kind,calendar=pattern.calendar,on_days=pattern.on_days,off_days=pattern.off_days,cycle_codes=pattern.cycle_codes)
                    slot=StaffingSlot(company=person.company,group=group,person=person,housing_camp=original.housing_camp,camp=original.camp,status=original.status,arrival_date=original.arrival_date,departure_date=original.departure_date,calendar=dict(original.calendar),comment='Личный график; начальные данные сохранены в исходной таблице.')
                if not slot:
                    board=pattern.board if pattern else Board.objects.filter(company=person.company,name='Графики сотрудников ВЖГ',import_digest='').order_by('pk').first()
                    if board is None:board=Board.objects.create(company=person.company,name='Графики сотрудников ВЖГ')
                    position=str(person.employee.position) if person.employee_id and person.employee.position_id else 'Без должности'
                    group=board.groups.filter(position=position).first() or StaffingGroup.objects.create(company=person.company,board=board,position=position)
                    slot=StaffingSlot(company=person.company,group=group,person=person)
                if form.cleaned_data.get('on_days'):
                    from .models import ShiftPattern
                    from .rotation_calendar import cycle_definition,work
                    from collections import Counter
                    cycle=cycle_definition(pattern)
                    codes=cycle['codes'] if cycle else list(pattern.calendar.values()) if pattern else []
                    defaults=Counter(str(code) for code in codes if work(code)).most_common(1)
                    hours=str(form.cleaned_data['hours']) if form.cleaned_data.get('hours') is not None else defaults[0][0] if defaults else 'МВ'
                    on,off=form.cleaned_data['on_days'],form.cleaned_data['off_days']
                    pattern=ShiftPattern.objects.create(company=person.company,board=slot.group.board,name=f'Индивидуальный ({on}/{off})',kind=pattern.kind if pattern else '',on_days=on,off_days=off,cycle_codes=[hours]*on+['МО']*off)
                slot.pattern=pattern
                start,end=form.cleaned_data.get('start'),form.cleaned_data.get('end')
                if start:
                    calendar=dict(slot.calendar);code=form.cleaned_data.get('code','');d=start
                    while d<=end:
                        if code:calendar[d.isoformat()]=code
                        else:calendar.pop(d.isoformat(),None)
                        d+=timedelta(days=1)
                    slot.calendar=calendar
                slot.save();person.schedule_slot=slot;person.save(update_fields=['schedule_slot','updated_at'])
                messages.success(request,'График сотрудника сохранён. Размещение в ВЖГ не изменено.');return redirect('personnel:housing_history',pk=person.pk)
    return render(request,'personnel/form.html',{'form':form,'title':'График: '+person.full_name})


@login_required
@require_http_methods(['POST'])
def link_people(request):
    access(request);count=housing.link_people(request.user.company)
    messages.success(request,f'Связано карточек: {count}. Неоднозначные совпадения оставлены для ручного уточнения.');return redirect('personnel:people')


def workbook_response(filename,title,rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    book=Workbook();sheet=book.active;sheet.title=title
    for row in rows:sheet.append(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in row])
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='167A94')
    for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=26
    stream=BytesIO();book.save(stream);response=HttpResponse(stream.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet');response['Content-Disposition']=f'attachment; filename="{filename}"';return response


@login_required
def roster_export(request):
    from .rotation import RotationFilter,rotation_rows
    raw=request.GET.copy();raw.setdefault('date',str(timezone.localdate()))
    if raw.get('camp') and not raw.getlist('camps'):raw.setlist('camps',[raw['camp']])
    selector=RotationFilter(raw,company=request.user.company)
    if not selector.is_valid():raise Http404('Некорректные фильтры перевахты.')
    camps=selector.selected_camps(request.user.company);data=selector.cleaned_data;day=data['date']
    rows,days,_=rotation_rows(request.user.company,list(camps),day,query=data.get('q',''),position=data.get('position'),state=data.get('state',''),include_unassigned=not data.get('construction_object') and not data.get('camps'))
    values=[['ВЖГ','Дата','Сотрудник / вакансия','Табельный номер','Должность','Вахта / отдых','Заезд','Выезд','Выезд по графику','Возвращение по графику','Дней на вахте / МО','Начало текущей вахты / МО','Источник начала периода','График',*[d.strftime('%d.%m.%Y') for d in days]]]
    for row in rows:
        person=row['person'];fmt=lambda value:value.isoformat() if value else ''
        values.append([str(row['camp'] or ''),str(day),person.display_name if person else 'Вакансия',person.display_employee_number if person else '',row['position_label'],row['state_label'],fmt(row.get('arrived_on')),fmt(row.get('last_departure')),fmt(row.get('planned_departure_display')),fmt(row.get('planned_return')),row.get('state_days'),fmt(row.get('state_since')),row.get('state_days_source',''),row.get('schedule_label',''),*row['codes']])
    return workbook_response('rotation.xlsx','Перевахта',values)



@login_required
@require_http_methods(['GET'])
def individual_calendar(request,pk):
    from datetime import timedelta
    from django.http import JsonResponse
    from .services import is_work
    person=get_object_or_404(scoped(request,Person).select_related('employee','schedule_slot'),pk=pk)
    raw=request.GET.get('month',timezone.localdate().strftime('%Y-%m'))
    try:month=date.fromisoformat(raw+'-01')
    except ValueError:return JsonResponse({'error':'Укажите месяц в формате ГГГГ-ММ.'},status=400)
    slot=housing.schedule_slot(person);calendar=calendar_for(slot,start=month,end=month+timedelta(days=730)) if slot else {}
    last=monthrange(month.year,month.month)[1]
    days=[month.replace(day=n) for n in range(1,last+1)]
    previous=(month-timedelta(days=1)).replace(day=1);following=(month+timedelta(days=32)).replace(day=1)
    from .rotation_calendar import description
    from .rotation import employee_schedule_label
    return JsonResponse({'name':person.display_name,'number':person.display_employee_number,'month':month.strftime('%Y-%m'),
        'schedule_info':description(slot,min(max(timezone.localdate(),month),days[-1])),
        'pattern':employee_schedule_label(person.employee if person.employee_id else None),
        'offset':month.weekday(),'previous':previous.strftime('%Y-%m'),'next':following.strftime('%Y-%m'),
        'edit_url':reverse('personnel:person_schedule',args=[person.pk]),
        'days':[{'day':d.day,'date':str(d),'code':calendar.get(str(d),''),'work':is_work(calendar.get(str(d),'')),'today':d==timezone.localdate()} for d in days]})
